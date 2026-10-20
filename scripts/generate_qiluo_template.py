"""生成琪珞国际货代发票模板 qiluo.xlsx（出厂默认，随包预置）。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "backend" / "templates" / "forwarders" / "qiluo.xlsx"

HEADERS = [
    "箱序号*",
    "FBAID*",
    "POAnumber*",
    "物品中文名*",
    "物品英文名*",
    "单箱内数量*",
    "申报单价 (USD)*",
    "材质(中文)*",
    "用途(中文)*",
    "品牌*",
    "型号*",
    "海关编码*",
    "单件实重*",
    "长*",
    "宽*",
    "高*",
]

SAMPLE_ROWS: list[list] = []

COL_WIDTHS = [10, 16, 14, 14, 14, 12, 14, 14, 12, 12, 10, 14, 10, 8, 8, 8]

FILL_HEADER_YELLOW = PatternFill("solid", fgColor="FFF2CC")
FILL_HEADER_GRAY = PatternFill("solid", fgColor="D9D9D9")
FILL_DATA_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_DATA_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_DATA_YELLOW = PatternFill("solid", fgColor="FFFF00")

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_HEADER = Font(name="微软雅黑", size=10, bold=True)
FONT_DATA = Font(name="微软雅黑", size=10)
FONT_PO = Font(name="微软雅黑", size=10, color="FF0000")
ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_cell(cell, *, fill, font=None):
    cell.fill = fill
    cell.font = font or FONT_DATA
    cell.border = BORDER
    cell.alignment = ALIGN


def build_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "发票"

    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, title)
        fill = FILL_HEADER_GRAY if col >= 13 else FILL_HEADER_YELLOW
        _style_cell(cell, fill=fill, font=FONT_HEADER)

    for row_offset, row_vals in enumerate(SAMPLE_ROWS, start=2):
        for col, val in enumerate(row_vals, start=1):
            cell = ws.cell(row_offset, col, val)
            if col <= 2:
                fill = FILL_DATA_GREEN
                font = FONT_DATA
            elif col == 3:
                fill = FILL_DATA_WHITE
                font = FONT_PO
            else:
                fill = FILL_DATA_YELLOW
                font = FONT_DATA
            _style_cell(cell, fill=fill, font=font)

    for col, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 28

    return wb


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
