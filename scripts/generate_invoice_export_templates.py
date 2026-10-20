"""S8 D4：由真实样例生成琪珞（T5/T6）、德速（T3）发票导出用的空白模板（出厂默认，随包预置）。

只清掉“这一票”的数据（货件号、仓库、日期、件数、明细行），保留货代给的表头、下拉、说明页和
琪珞「！」仓库地址表（T6 收件人地址靠它 VLOOKUP）。来源：
- T3 德速 Commercial Invoice 9810：tests/golden/samples/invoice_templates/T3_第9批第1票_FBA19GQXZW0N.xlsx
- T6 琪珞 FBA 专线导入模版：      tests/golden/samples/invoice_templates/T6_第7批第1票_VGT2-29箱.xlsx
- T5 琪珞 单行汇总（=T7 空白版）：  tests/golden/samples/invoice_templates/T7_义乌宁波运单下单模板_空白.xlsx
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "backend" / "tests" / "golden" / "samples" / "invoice_templates"
OUT = ROOT / "backend" / "templates" / "forwarders"


def _clear(ws, cells):
    for c in cells:
        ws[c].value = None


def _clear_rows(ws, first: int):
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for c in row:
            c.value = None


def make_t3() -> Path:
    wb = openpyxl.load_workbook(SRC / "T3_第9批第1票_FBA19GQXZW0N.xlsx")
    ws = wb["发票填写模板"]
    _clear(ws, ["B2", "B4", "B6", "B7", "B8", "B15", "B17", "G3", "G4"])
    _clear_rows(ws, 21)
    p = OUT / "deshu_t3.xlsx"
    wb.save(p)
    return p


def make_t6() -> Path:
    wb = openpyxl.load_workbook(SRC / "T6_第7批第1票_VGT2-29箱.xlsx")
    ws = wb["发票"]
    _clear(ws, ["B3", "E3", "B6", "B10", "B11"])
    _clear_rows(ws, 14)
    p = OUT / "qiluo_t6.xlsx"
    wb.save(p)
    return p


def make_t5() -> Path:
    wb = openpyxl.load_workbook(SRC / "T7_义乌宁波运单下单模板_空白.xlsx")
    p = OUT / "qiluo_t5.xlsx"
    wb.save(p)
    return p


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for f in (make_t3, make_t6, make_t5):
        print("生成", f().relative_to(ROOT))


if __name__ == "__main__":
    main()
