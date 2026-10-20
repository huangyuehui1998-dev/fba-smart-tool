"""一次性：xlrd 读旧版 .xls → openpyxl 写 .xlsx，保留合并区、列宽与字体。"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from xlrd.formatting import XF

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = ROOT / "dist" / "app_data" / "templates"
DEFAULT_OUT = ROOT / "backend" / "templates" / "customs" / "报关资料模板.xlsx"


def _find_xls(src_dir: Path) -> Path:
    if not src_dir.is_dir():
        raise FileNotFoundError(f"源目录不存在：{src_dir}")
    candidates = sorted(src_dir.glob("*.xls"))
    if not candidates:
        candidates = sorted(src_dir.rglob("*.xls"))
    customs = [p for p in candidates if "9810" in p.name or "报关" in p.name]
    if customs:
        return customs[0]
    if candidates:
        return candidates[0]
    raise FileNotFoundError(f"未在 {src_dir} 找到 .xls 模板")


def _rgb_from_xf(book: xlrd.book.Book, xf: XF) -> str | None:
    try:
        bg = book.xf_list[xf.background.pattern_colour_index]
        fg = book.colour_map[bg]
        if fg:
            return f"{fg.red:02X}{fg.green:02X}{fg.blue:02X}"
    except Exception:
        pass
    return None


def _font_from_xf(book: xlrd.book.Book, xf: XF) -> Font:
    fnt = book.font_list[xf.font_index]
    color = None
    try:
        c = book.colour_map[fnt.colour_index]
        if c:
            color = f"{c.red:02X}{c.green:02X}{c.blue:02X}"
    except Exception:
        pass
    return Font(
        name=fnt.name,
        size=fnt.height / 20 if fnt.height else 11,
        bold=fnt.bold,
        italic=fnt.italic,
        underline="single" if fnt.underline_type else None,
        strike=fnt.struck_out,
        color=color,
    )


def _alignment_from_xf(xf: XF) -> Alignment:
    h = "general"
    if xf.alignment.hor_align == 1:
        h = "left"
    elif xf.alignment.hor_align == 2:
        h = "center"
    elif xf.alignment.hor_align == 3:
        h = "right"
    v = "bottom"
    if xf.alignment.vert_align == 1:
        v = "top"
    elif xf.alignment.vert_align == 2:
        v = "center"
    return Alignment(
        horizontal=h,
        vertical=v,
        wrap_text=bool(xf.alignment.text_wrapped),
        shrink_to_fit=bool(xf.alignment.shrink_to_fit),
    )


def _copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)


def convert_xls_to_xlsx(src: Path, dest: Path) -> None:
    book = xlrd.open_workbook(str(src), formatting_info=True)
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_idx, sh in enumerate(book.sheets()):
        ws = wb.create_sheet(title=sh.name[:31], index=sheet_idx)

        for col_idx in range(sh.ncols):
            try:
                colinfo = sh.colinfo_map.get(col_idx)
                if colinfo:
                    letter = get_column_letter(col_idx + 1)
                    ws.column_dimensions[letter].width = max(colinfo.width / 256, 8)
            except Exception:
                pass

        for row_idx in range(sh.nrows):
            for col_idx in range(sh.ncols):
                cell = sh.cell(row_idx, col_idx)
                ocell = ws.cell(row=row_idx + 1, column=col_idx + 1)
                ctype = cell.ctype
                if ctype == xlrd.XL_CELL_EMPTY:
                    continue
                if ctype == xlrd.XL_CELL_TEXT:
                    ocell.value = cell.value
                elif ctype == xlrd.XL_CELL_NUMBER:
                    ocell.value = cell.value
                elif ctype == xlrd.XL_CELL_DATE:
                    ocell.value = cell.value
                    ocell.number_format = "yyyy-mm-dd"
                elif ctype == xlrd.XL_CELL_BOOLEAN:
                    ocell.value = bool(cell.value)
                else:
                    ocell.value = cell.value

                try:
                    xf = book.xf_list[cell.xf_index]
                    ocell.font = _font_from_xf(book, xf)
                    ocell.alignment = _alignment_from_xf(xf)
                    rgb = _rgb_from_xf(book, xf)
                    if rgb:
                        ocell.fill = PatternFill("solid", fgColor=rgb)
                except Exception:
                    pass

        for rlo, rhi, clo, chi in sh.merged_cells:
            ws.merge_cells(
                start_row=rlo + 1,
                end_row=rhi,
                start_column=clo + 1,
                end_column=chi,
            )

        for row_idx in range(sh.nrows):
            try:
                rowinfo = sh.rowinfo_map.get(row_idx)
                if rowinfo and rowinfo.height:
                    ws.row_dimensions[row_idx + 1].height = rowinfo.height / 20
            except Exception:
                pass

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(f"已转换：{src.name} -> {dest}")


def main(argv: list[str] | None = None) -> int:
    argv = argv or sys.argv[1:]
    src = Path(argv[0]) if argv else _find_xls(DEFAULT_SRC)
    dest = Path(argv[1]) if len(argv) > 1 else DEFAULT_OUT
    if not src.is_file() and argv:
        print(f"源文件不存在：{src}", file=sys.stderr)
        return 1
    if not argv:
        src = _find_xls(DEFAULT_SRC)
    convert_xls_to_xlsx(src, dest)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
