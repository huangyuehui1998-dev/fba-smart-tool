"""S9-step5：从 backend/ 里的原件生成安装包用的 DEMO / 空白版模板和出厂配置。

    python packaging/build_release_templates.py          # 生成到 packaging/release_templates/

* 原件不动（测试和 golden 继续用原件）；输出写到 packaging/release_templates/{templates,config}/。
* 本脚本不含任何真实值：要替换的值都是从原件里按「位置 / 格式」现场学出来的（公司名、统一社会信用代码、
  订单号 / 货件号、日期、金额数量、品牌、品名、货源地），只在内存里用，不打印、不写日志。
* 只输出计数（替换了几格、删了几张图），不输出任何值。
兼容 Python 3.9。
"""
from __future__ import annotations

import argparse
import datetime as _dt
import io
import json
import re
import sys
from pathlib import Path
from typing import Dict, Iterable, List, Set

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from release_manifest import CONFIG_WHITELIST, RELEASE, TEMPLATE_WHITELIST, release_dir  # noqa: E402

COMPANY_RE = re.compile(r"[\u4e00-\u9fff（）()]{2,40}(?:有限责任公司|有限公司|股份公司|集团公司)")
CREDIT_RE = re.compile(r"(?<![0-9A-Z])[0-9A-HJ-NPQRTUWXY]{2}\d{6}[0-9A-HJ-NPQRTUWXY]{10}(?![0-9A-Z])")
ORDER_RE = re.compile(r"(?<![0-9A-Za-z])[A-Z]{2,4}\d[A-Z0-9]{2,}\d{2,}(?![0-9A-Za-z])")
SHIPMENT_RE = re.compile(r"(?<![0-9A-Za-z])FBA[0-9A-Z]{6,}(?![0-9A-Za-z])")
WAREHOUSE_RE = re.compile(r"(?<![0-9A-Za-z])[A-Z]{3}\d(?![0-9A-Za-z])")
DATE_TEXT_RE = re.compile(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日")
MIXED_TOKEN_RE = re.compile(r"(?<![0-9A-Za-z])(?=[0-9A-Za-z]*[A-Za-z])(?=(?:[A-Za-z]*\d){2})[0-9A-Za-z]{6,}(?![0-9A-Za-z])")

DEMO_COMPANY = "示例公司（请在出口商档案里填写）"
DEMO_PRODUCT = "示例商品"
DEMO_BRAND = "DEMO"
DEMO_CITY = "示例城市"
REDACTED = "<已脱敏>"


class Counter(dict):
    def add(self, key: str, n: int = 1) -> None:
        self[key] = self.get(key, 0) + n


# ---------------------------------------------------------------- 报关资料模板

def _cell_below(ws, label: str):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().startswith(label):
                return ws.cell(c.row + 1, c.column)
    return None


def _column_values(ws, header: str, header_row_max: int = 30) -> List[str]:
    """表头为 header 的那一列下面的文字值（只看文字）。"""
    for row in ws.iter_rows(max_row=header_row_max):
        for c in row:
            if isinstance(c.value, str) and c.value.strip() == header:
                vals = []
                for r in range(c.row + 1, ws.max_row + 1):
                    v = ws.cell(r, c.column).value
                    if isinstance(v, str) and v.strip():
                        vals.append(v.strip())
                return vals
    return []


def _learn_customs(wb) -> Dict[str, str]:
    """从原件学出要替换的值 → 替换成什么。只在内存里。"""
    repl: Dict[str, str] = {}
    sheets = {ws.title: ws for ws in wb.worksheets}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    for m in COMPANY_RE.finditer(c.value):
                        repl[m.group(0)] = DEMO_COMPANY
    el = sheets.get("申报要素")
    if el is not None:
        for v in _column_values(el, "品牌"):
            if v not in ("无", "无品牌"):
                repl[v] = DEMO_BRAND
        for v in _column_values(el, "品名"):
            repl[v] = DEMO_PRODUCT
    decl = sheets.get("报关单")
    if decl is not None:
        for v in _column_values(decl, "境内货源地"):
            repl[v] = DEMO_CITY
        for v in _column_values(decl, "商品名称及规格型号"):
            if v != "总计":
                repl.setdefault(v, DEMO_PRODUCT)
    contract = sheets.get("合同")
    if contract is not None:
        for row in contract.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() == "签约地点":
                    for dc in (2, 1):
                        v = contract.cell(c.row, c.column + dc).value
                        if isinstance(v, str) and v.strip() and v.strip()[0] >= "\u4e00":
                            repl.setdefault(v.strip(), DEMO_CITY)
                            break
    ent = sheets.get("委托书")
    if ent is not None:
        for row in ent.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ("原产地/货源地", "主要货物名称"):
                    v = ent.cell(c.row, c.column + 1).value
                    if isinstance(v, str) and v.strip():
                        repl.setdefault(v.strip(), DEMO_CITY if "货源" in c.value else DEMO_PRODUCT)
    # 只替换有实际意义的值（太短的不替换，免得误伤普通字）
    return {k: v for k, v in repl.items() if len(k) >= 2}


def _scrub_text(s: str, repl: Dict[str, str], cnt: Counter) -> str:
    out = s
    for k in sorted(repl, key=len, reverse=True):
        if k in out:
            out = out.replace(k, repl[k])
            cnt.add("learned_value")
    for rx, to, key in ((CREDIT_RE, "", "credit_code"), (SHIPMENT_RE, "FBA-DEMO", "shipment_id"),
                        (ORDER_RE, "DEMO-ORDER", "order_no"), (WAREHOUSE_RE, "DEMO", "warehouse_code"),
                        (DATE_TEXT_RE, "    年  月  日", "date_text")):
        out, n = rx.subn(to, out)
        if n:
            cnt.add(key, n)
    return out


def build_customs(src: Path, dest: Path) -> Counter:
    cnt = Counter()
    wb = load_workbook(src)
    repl = _learn_customs(wb)
    for idx, ws in enumerate(wb.worksheets):
        rules_sheet = idx == 1  # 填制规范：海关公开的填报规范原文，只做公司 / 代码类替换
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("="):
                    continue  # 公式保留
                if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
                    c.value = None
                    cnt.add("datetime")
                elif isinstance(v, bool):
                    continue
                elif isinstance(v, (int, float)):
                    if not rules_sheet:
                        c.value = None
                        cnt.add("number")
                elif isinstance(v, str):
                    if rules_sheet:
                        new = CREDIT_RE.sub("", COMPANY_RE.sub(DEMO_COMPANY, v))
                    else:
                        new = _scrub_text(v, repl, cnt)
                    if new != v:
                        c.value = new
                        cnt.add("text_cell")
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return cnt


# ---------------------------------------------------------------- 货代发票模板（琪珞 T6 / 德速 T3）

def build_forwarder(src: Path, dest: Path) -> Counter:
    """第一张表（发票主表）上的图片是真实产品图 → 全部去掉；「填写示例」等表上货代自带的示例图保留。
    文字里出现的公司名 / 信用代码清空。"""
    cnt = Counter()
    wb = load_workbook(src)
    first = wb.worksheets[0]
    cnt.add("image_removed", len(getattr(first, "_images", [])))
    first._images = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith("="):
                    new = CREDIT_RE.sub("", COMPANY_RE.sub("", v)).strip() if (COMPANY_RE.search(v) or CREDIT_RE.search(v)) else v
                    if new != v:
                        c.value = new or None
                        cnt.add("text_cell")
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return cnt


# ---------------------------------------------------------------- 出厂配置

def _redact_strings(obj, cnt: Counter):
    if isinstance(obj, dict):
        return {k: _redact_strings(v, cnt) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_redact_strings(v, cnt) for v in obj]
    if isinstance(obj, str):
        out = obj
        for rx in (SHIPMENT_RE, MIXED_TOKEN_RE, WAREHOUSE_RE):
            out, n = rx.subn(REDACTED, out)
            if n:
                cnt.add("token_redacted", n)
        out, n = COMPANY_RE.subn(REDACTED, out)
        if n:
            cnt.add("company_redacted", n)
        return out
    return obj


def _redact_source_fields(obj, cnt: Counter):
    if isinstance(obj, dict):
        out = {}
        for k, v in obj.items():
            if k == "source" and isinstance(v, str):
                new, n = SHIPMENT_RE.subn(REDACTED, v)
                if n:
                    cnt.add("token_redacted", n)
                out[k] = new
            else:
                out[k] = _redact_source_fields(v, cnt)
        return out
    if isinstance(obj, list):
        return [_redact_source_fields(v, cnt) for v in obj]
    return obj


def build_config(name: str, src: Path, dest: Path) -> Counter:
    cnt = Counter()
    if name in ("amazon_sku_metadata.default.json", "fnsku_mapping.default.json", "sku_mapping.default.json"):
        data = {}  # 新用户从空映射开始；SKU / FNSKU / ASIN 由用户自己的 FBA 表和产品库补
        cnt.add("emptied")
    elif name == "exporter_profiles.default.json":
        data = {"active_id": "", "profiles": []}  # 出口商档案由用户自己在设置里添加
        cnt.add("emptied")
    elif name == "invoice_template_map.json":
        raw = json.loads(src.read_text(encoding="utf-8"))
        data = dict(raw)
        tpls = raw.get("templates") or {}
        new_tpls = {}
        for key, tpl in tpls.items():
            tpl = dict(tpl)
            if isinstance(tpl.get("decision"), str):
                tpl["decision"] = _redact_strings(tpl["decision"], cnt)
            new_tpls[key] = tpl
        data["templates"] = new_tpls
    elif name == "forwarders.default.json":
        raw = json.loads(src.read_text(encoding="utf-8"))
        data = _redact_source_fields(raw, cnt)  # 只改「source」说明文字里的货件号，其余配置原样
    else:
        raise ValueError("不是 release 配置：%s" % name)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes((json.dumps(data, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))  # LF；Python 3.9 的 write_text 没有 newline 参数
    return cnt


def build_all(root: Path = None, templates_src: Path = None, config_src: Path = None, out: Path = None) -> Dict[str, Counter]:
    root = Path(root).resolve() if root else HERE.parent
    templates_src = Path(templates_src) if templates_src else root / "backend" / "templates"
    config_src = Path(config_src) if config_src else root / "backend" / "config"
    out = Path(out) if out else release_dir(root)
    report: Dict[str, Counter] = {}
    for rel, origin in sorted(TEMPLATE_WHITELIST.items()):
        if origin != RELEASE:
            continue
        src, dest = templates_src / rel, out / "templates" / rel
        if rel.startswith("customs/"):
            report["templates/" + rel] = build_customs(src, dest)
        elif rel.startswith("forwarders/"):
            report["templates/" + rel] = build_forwarder(src, dest)
        else:
            raise ValueError("未知 release 模板：%s" % rel)
    for name, origin in sorted(CONFIG_WHITELIST.items()):
        if origin == RELEASE:
            report["config/" + name] = build_config(name, config_src / name, out / "config" / name)
    return report


def main(argv: Iterable[str] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--root", default=None)
    ap.add_argument("--templates-src", default=None)
    ap.add_argument("--config-src", default=None)
    ap.add_argument("--out", default=None)
    a = ap.parse_args(list(argv) if argv is not None else None)
    rep = build_all(a.root, a.templates_src, a.config_src, a.out)
    for k in sorted(rep):
        print(k, " ".join("%s=%d" % kv for kv in sorted(rep[k].items())))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
