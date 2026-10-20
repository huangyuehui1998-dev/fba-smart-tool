"""端到端工作流模拟：Step 1 → 2 → 3 → 4 → 4.5（TestClient + 真实 templates 夹具）。"""
from __future__ import annotations

import base64
import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from fba_tool.fba_parser import simplify_sku
from fba_tool.input_template import ensure_template_on_disk
from fba_tool.shipment_parser import parse_csv_shipment
from main import app

ROOT = Path(__file__).resolve().parent
TPL = ROOT / "templates"


def _invoice_template_bytes() -> bytes:
    """优先使用与主测试货件匹配的真实发票模板。"""
    for pattern in ("*RDU2*发票*.xlsx", "*发票*.xlsx", "待填写的发票模板.xlsx"):
        matches = sorted(TPL.glob(pattern))
        if matches:
            return matches[0].read_bytes()
    return _minimal_invoice_template()


def _minimal_invoice_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "Reference"
    ws["B1"] = "SKU"
    for i in range(2, 12):
        ws.cell(i, 1, f"REF-{i}")
        ws.cell(i, 2, f"SKU-{i}")
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _packing_csvs() -> list[tuple[str, bytes]]:
    out: list[tuple[str, bytes]] = []
    for path in sorted(TPL.glob("*.csv")):
        if "产品信息" in path.name:
            continue
        out.append((path.name, path.read_bytes()))
    return out


def _packing_pdfs() -> list[tuple[str, bytes]]:
    return [(p.name, p.read_bytes()) for p in sorted(TPL.glob("*.pdf"))]


PRIMARY_CSV = "亚马逊FBA表1.csv"
PRIMARY_PDF = "箱唛1-RDU2.pdf"


def _payload_from_packing_csv() -> dict:
    """从主测试 CSV 聚合 SKU，构造与 Step 4 校验一致的 Step 1 录入 payload。"""
    csv_path = TPL / PRIMARY_CSV
    if not csv_path.exists():
        csv_path = next(p for p in sorted(TPL.glob("*.csv")) if "产品信息" not in p.name)
    ship = parse_csv_shipment(csv_path.read_bytes())
    sku_qty: dict[str, int] = {}
    for box in ship.box_data.values():
        for sku, qty in box.items():
            # S8 D13：用 CSV 里的完整 Merchant SKU。以前用 simplify_sku 简称，新库没有 SKU 映射时
            # Step 4 对不上 Merchant SKU，判成漏发/多发，整个脚本跑不过。
            key = str(sku).strip()
            sku_qty[key] = sku_qty.get(key, 0) + int(qty)
    rows = [
        {
            "SKU简称": sku,
            "FNSKU": "",
            "发货总数": total,
            "单箱装量": max(1, min(6, total)),
            "箱长(cm)": 48.5,
            "箱宽(cm)": 30.0,
            "箱高(cm)": 38.0,
            "单个重量(kg)": 2.5,
            "空纸箱重(kg)": 0.5,
            "单箱毛重(kg)": None,
        }
        for sku, total in sorted(sku_qty.items())
    ]
    return {"rows": rows, "fnsku_overrides": {}}


def _step1_checks(client: TestClient, payload: dict) -> None:
    """Step 1 API：预览、校验、样表下载；Excel 解析为附加检查。"""
    r = client.post("/api/compute_preview", json=payload)
    assert r.status_code == 200, r.text
    assert r.json().get("rows"), "compute_preview 无行"

    r = client.post("/api/validate_input", json=payload)
    assert r.status_code == 200, r.text

    r = client.get("/api/input_template")
    assert r.status_code == 200 and len(r.content) > 200

    inp = TPL / "数据录入样表.xlsx"
    if not inp.exists():
        inp = ensure_template_on_disk()
    files = {
        "file": (
            inp.name,
            inp.read_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post("/api/parse_excel", files=files)
    assert r.status_code == 200, r.text
    assert r.json().get("rows"), "parse_excel 无行"


def _step2_boxing(client: TestClient, payload: dict) -> dict:
    r = client.post("/api/calculate_boxes", json=payload)
    assert r.status_code == 200, r.text
    boxes = r.json()
    assert boxes.get("box_editor"), "calculate_boxes 无 editor"
    assert "fba_compliance" in boxes

    body = {**payload, "box_editor": boxes["box_editor"]}
    r = client.post("/api/box_metrics", json=body)
    assert r.status_code == 200, r.text

    r = client.post("/api/tail_pool", json=body)
    assert r.status_code == 200, r.text
    tail = r.json().get("tail_pool") or []
    if tail:
        row = tail[0]
        remain = int(row.get("待排尾数") or row.get("剩余数量") or 0)
        sku = row["SKU简称"]
        if remain > 0:
            qty = min(remain, max(1, remain // 2))
            assign_body = {**body, "sku": sku, "qty": qty, "box_no": 1}
            r = client.post("/api/box/assign_tail", json=assign_body)
            assert r.status_code == 200, r.text
            body["box_editor"] = r.json()["box_editor"]

    return body


def _step3_export(client: TestClient, body: dict) -> None:
    r = client.post("/api/export/step3/preview", json=body)
    assert r.status_code == 200, r.text
    preview = r.json()
    assert preview.get("matrix"), "step3 matrix 为空"
    assert preview.get("box_count", 0) > 0

    r = client.post("/api/export/step3/internal", json=body)
    assert r.status_code == 200, r.text
    assert len(r.content) > 500

    r = client.post("/api/export/step3/amazon", json=body)
    assert r.status_code == 200, r.text
    assert len(r.content) > 500


def _step4_validate(client: TestClient, body: dict) -> dict:
    csv_path = TPL / PRIMARY_CSV
    pdf_path = TPL / PRIMARY_PDF
    if not csv_path.exists() or not pdf_path.exists():
        csv_path = next(p for p in sorted(TPL.glob("*.csv")) if "产品信息" not in p.name)
        pdf_path = next(iter(sorted(TPL.glob("*.pdf"))), None)
    assert csv_path.exists(), "templates 下无装箱 CSV"
    assert pdf_path and pdf_path.exists(), "templates 下无箱唛 PDF"

    csv_payload = [(csv_path.name, csv_path.read_bytes())]
    pdf_payload = [(pdf_path.name, pdf_path.read_bytes())]

    workflow = {
        "rows": body["rows"],
        "box_editor": body["box_editor"],
        "packing_mode": "amazon_led",
    }
    files: list[tuple[str, tuple]] = []
    for name, data in csv_payload:
        files.append(("csv_files", (name, data, "text/csv")))
    for name, data in pdf_payload:
        files.append(("pdf_files", (name, data, "application/pdf")))

    r = client.post(
        "/api/shipments/validate",
        files=files,
        data={"throw_ratio": "6000", "workflow_json": json.dumps(workflow, ensure_ascii=False)},
    )
    assert r.status_code == 200, r.text
    result = r.json()
    assert result.get("ok"), f"Step4 校验未通过: {result.get('errors', [])[:5]}"
    assert result.get("shipments"), "无 shipments 返回"
    assert result.get("worker_packing_preview") is not None
    assert result.get("forwarder") is not None

    # 工人装箱单导出
    r = client.post(
        "/api/export/worker_packing_list",
        files=files,
        data={"workflow_json": json.dumps(workflow, ensure_ascii=False)},
    )
    assert r.status_code == 200, r.text
    assert len(r.content) > 500

    return result


def _step45_invoice(client: TestClient, step4: dict) -> None:
    shipment = step4["shipments"][0]
    template_bytes = _invoice_template_bytes()
    template_b64 = base64.b64encode(template_bytes).decode("ascii")

    r = client.post("/api/invoice/preview", json={"template_base64": template_b64})
    assert r.status_code == 200, r.text
    preview = r.json()
    assert preview.get("template_hash"), "invoice preview 无 template_hash"

    gen_body = {
        "shipment_id": shipment["shipment_id"],
        "csv_base64": shipment["csv_base64"],
        "pdf_base64": shipment["pdf_base64"],
        "template_base64": template_b64,
    }
    r = client.post("/api/invoice/generate", json=gen_body)
    assert r.status_code == 200, r.text
    draft = r.json()
    assert draft.get("cells"), "invoice generate 无 cells"
    template_hash = draft.get("template_hash") or preview["template_hash"]

    export_body = {
        "template_hash": template_hash,
        "shipment_id": draft["shipment_id"],
        "warehouse_code": draft.get("warehouse_code") or shipment["warehouse"],
        "template_sheet": draft.get("template_sheet") or "Invoice",
        "fields": draft["cells"],
    }
    r = client.post("/api/export_invoice", json=export_body)
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"PK")
    assert len(r.content) > 200


def run_unit_tests() -> None:
    scripts = sorted(ROOT.glob("test_*.py"))
    scripts = [s for s in scripts if s.name != Path(__file__).name]
    failed: list[str] = []
    for script in scripts:
        proc = subprocess.run(
            [sys.executable, str(script.name)],
            cwd=str(ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if proc.returncode != 0:
            failed.append(f"{script.name}: {proc.stdout[-200:]}{proc.stderr[-200:]}")
    if failed:
        raise AssertionError("单元测试失败:\n" + "\n".join(failed))


def _run_workflow() -> list[str]:
    client = TestClient(app)
    steps: list[str] = []
    payload = _payload_from_packing_csv()
    _step1_checks(client, payload)
    steps.append(f"Step 1  预览+校验+样表+Excel解析 ({len(payload['rows'])} SKU)")
    box_body = _step2_boxing(client, payload)
    steps.append(f"Step 2  排箱+指标+尾数分配 ({len(box_body['box_editor'])} 行)")
    _step3_export(client, box_body)
    steps.append("Step 3  矩阵预览+内部表/Amazon 导出")
    step4 = _step4_validate(client, box_body)
    steps.append(f"Step 4  货件校验+工人装箱单 ({len(step4['shipments'])} 货件)")
    _step45_invoice(client, step4)
    steps.append("Step 4.5 发票预览+自动填写+导出")
    return steps


def test_e2e_workflow_step1_to_45():
    """S8 D13：改成 pytest 用例（以前没有 test_ 函数，pytest 不收集）。素材是 backend/templates 下的货件 CSV + 箱唛 PDF。"""
    import pytest

    if not (TPL / PRIMARY_CSV).exists() or not list(TPL.glob("*.pdf")):
        pytest.skip("backend/templates 下没有货件 CSV / 箱唛 PDF 素材（不受 git 管理），跳过端到端用例")
    steps = _run_workflow()
    assert len(steps) == 5


def main() -> None:
    print("=" * 60)
    print("1/2 运行既有单元测试 …")
    run_unit_tests()
    print("   11 个 test_*.py 全部通过")

    print("2/2 模拟完整工作流 Step 1 → 4.5 …")
    client = TestClient(app)
    steps: list[str] = []

    payload = _payload_from_packing_csv()
    _step1_checks(client, payload)
    steps.append(f"Step 1  预览+校验+样表+Excel解析 ({len(payload['rows'])} SKU)")

    box_body = _step2_boxing(client, payload)
    steps.append(f"Step 2  排箱+指标+尾数分配 ({len(box_body['box_editor'])} 行)")

    _step3_export(client, box_body)
    steps.append("Step 3  矩阵预览+内部表/Amazon 导出")

    step4 = _step4_validate(client, box_body)
    steps.append(
        f"Step 4  货件校验+工人装箱单 ({len(step4['shipments'])} 货件, "
        f"cross_check={'OK' if step4.get('cross_check_ok') else '—'})"
    )

    _step45_invoice(client, step4)
    steps.append("Step 4.5 发票预览+自动填写+导出")

    print("\n模拟结果:")
    for line in steps:
        print(f"  [OK] {line}")
    print("\nE2E SIMULATION: ALL PASSED")


if __name__ == "__main__":
    main()
