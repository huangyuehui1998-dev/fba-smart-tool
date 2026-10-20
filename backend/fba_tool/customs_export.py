"""Step 5 报关资料 Excel 导出 — 基于 openpyxl 模板基底写入动态单元格。"""
from __future__ import annotations

import io
from copy import copy
from datetime import date, datetime
from typing import Any, List, Tuple

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fba_tool.customs_package import CustomsLineItem, CustomsPackageData


def _parse_date(val: str | date | datetime | None) -> date | None:
    """解析为 date 对象供 openpyxl 写入；无法解析则返回 None（不写字符串）。"""
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    text = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _top_left_cell(ws: Worksheet, addr: str) -> str:
    col = "".join(c for c in addr if c.isalpha())
    row = int("".join(c for c in addr if c.isdigit()))
    col_idx = column_index_from_string(col)
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col_idx <= merged.max_col:
            return f"{get_column_letter(merged.min_col)}{merged.min_row}"
    return addr


DATE_FORMAT = "yyyy-mm-dd"


def _set(ws: Worksheet, addr: str, value: Any) -> None:
    if value is None or value == "":
        return
    target = _top_left_cell(ws, addr)
    # 直接赋 value，覆盖模板中原有 SUM 等公式（insert_rows 后公式引用易错乱）
    if isinstance(value, date) and not isinstance(value, datetime):
        value = datetime(value.year, value.month, value.day)
    ws[target].value = value
    if isinstance(value, datetime):
        ws[target].number_format = DATE_FORMAT  # Excel 日期序列号 + 日期格式（v8 §3.6 验收 2）


def _set_date_or_clear(ws: Worksheet, addr: str, val: Any) -> None:
    d = _parse_date(val)
    if d:
        _set(ws, addr, d)
    else:
        _clear(ws, addr)


def _clear_rows(ws: Worksheet, first: int, last: int, max_col: int) -> None:
    """清掉模板里没用到的数据行（模板残留的 0 / 序号）。"""
    for r in range(first, last + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if type(cell).__name__ == "MergedCell":
                continue
            cell.value = None


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, min_col: int = 1, max_col: int = 25) -> None:
    for col in range(min_col, max_col + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
            dst.alignment = copy(src.alignment)


def _ensure_data_rows(
    ws: Worksheet,
    *,
    first_row: int,
    template_row: int,
    count: int,
    total_row: int,
    max_col: int = 25,
) -> int:
    """在 total_row 之前插入不足的数据行，返回（可能下移后的）total_row。"""
    if count <= 0:
        return total_row
    available = max(0, total_row - first_row)
    extra = max(0, count - available)
    if extra > 0:
        ws.insert_rows(total_row, amount=extra)
        for i in range(extra):
            _copy_row_style(ws, template_row, total_row - available + i, max_col=max_col)
        total_row += extra
    return total_row


def _clear(ws: Worksheet, addr: str) -> None:
    ws[_top_left_cell(ws, addr)].value = None


def _write_declaration_sheet(ws: Worksheet, data: CustomsPackageData) -> None:
    exp = data.exporter
    _set(ws, "A4", exp.get("name"))
    _set(ws, "C3", exp.get("credit_code"))
    _set(ws, "A8", exp.get("name"))
    _set(ws, "C7", exp.get("credit_code"))
    _set(ws, "A6", data.buyer)
    _set(ws, "A10", data.order_no)
    # v6 Q16：出口日期默认留空（模板里的旧日期要清掉）；用户填了才写
    _set_date_or_clear(ws, "H3", data.dates.get("export_date"))   # 出口日期
    _set_date_or_clear(ws, "L3", data.dates.get("declare_date"))  # 申报日期（K3 标签，L3:N3）
    _set(ws, "E12", data.total_pieces)
    _set(ws, "F12", data.total_gross_kg)  # 毛重（千克）
    _set(ws, "G12", data.total_net_kg)    # 净重（千克）
    _clear(ws, "K12")                     # v6 Q11：运费留空
    _set(ws, "E6", data.transport_mode)
    _set(ws, "B16", data.remarks)

    first_row = 20
    template_row = 20
    total_row = _ensure_data_rows(
        ws, first_row=first_row, template_row=template_row, count=len(data.lines), total_row=27, max_col=25
    )

    for i, line in enumerate(data.lines):
        r = first_row + i
        _set(ws, f"A{r}", line.seq)
        _set(ws, f"B{r}", line.hs_code)
        _set(ws, f"D{r}", line.chinese_name)
        _set(ws, f"G{r}", line.qty)
        _set(ws, f"H{r}", line.unit)
        _set(ws, f"I{r}", line.unit_price)
        _set(ws, f"J{r}", line.amount)
        _set(ws, f"K{r}", line.origin_country)
        _set(ws, f"M{r}", line.dest_country)
        _set(ws, f"P{r}", line.source_place)
        _set(ws, f"S{r}", line.tax_exemption)
    _clear_rows(ws, first_row + len(data.lines), total_row - 1, 25)

    _set(ws, f"G{total_row}", data.total_qty)
    _set(ws, f"J{total_row}", data.total_amount)


def _write_elements_sheet(ws: Worksheet, lines: List[CustomsLineItem]) -> None:
    first_row = 10
    template_row = 10
    end_row = _ensure_data_rows(ws, first_row=first_row, template_row=template_row, count=len(lines), total_row=16, max_col=9)
    for i, line in enumerate(lines):
        r = first_row + i
        _set(ws, f"A{r}", line.seq)
        _set(ws, f"B{r}", line.chinese_name)
        _set(ws, f"C{r}", line.hs_code)
        _set(ws, f"D{r}", line.quarantine_code)
        _set(ws, f"E{r}", line.brand)
        _set(ws, f"F{r}", line.model)
        _set(ws, f"G{r}", line.brand_type)
        _set(ws, f"H{r}", line.export_preferential)
        _set(ws, f"I{r}", line.declare_elements)
    _clear_rows(ws, first_row + len(lines), end_row - 1, 9)


def _write_invoice_sheet(ws: Worksheet, data: CustomsPackageData) -> None:
    _set(ws, "G2", data.order_no)
    _set(ws, "C3", data.buyer)
    _set(ws, "G3", _parse_date(data.dates.get("invoice_date")))
    _set(ws, "G6", data.trade_term)

    first_row = 8
    total_row = _ensure_data_rows(ws, first_row=first_row, template_row=8, count=len(data.lines), total_row=15, max_col=7)
    for i, line in enumerate(data.lines):
        r = first_row + i
        _set(ws, f"A{r}", line.seq)
        _set(ws, f"C{r}", line.chinese_name)
        _set(ws, f"D{r}", line.qty)
        _set(ws, f"E{r}", line.unit)
        _set(ws, f"F{r}", line.unit_price)
        _set(ws, f"G{r}", line.amount)
    _clear_rows(ws, first_row + len(data.lines), total_row - 1, 7)
    _set(ws, f"G{total_row}", data.total_amount)


def _write_packing_sheet(ws: Worksheet, data: CustomsPackageData) -> None:
    _set(ws, "G2", _parse_date(data.dates.get("packing_date")))
    _set(ws, "G3", data.order_no)
    _set(ws, "B4", data.buyer)
    _set(ws, "G4", data.order_no)
    _set(ws, "D6", data.dest_country)

    first_row = 9
    total_row = _ensure_data_rows(ws, first_row=first_row, template_row=9, count=len(data.lines), total_row=15, max_col=7)
    for i, line in enumerate(data.lines):
        r = first_row + i
        _set(ws, f"B{r}", line.chinese_name)
        _set(ws, f"C{r}", line.box_count)
        _set(ws, f"D{r}", line.qty)
        _set(ws, f"E{r}", line.unit)
        _set(ws, f"F{r}", line.gross_kg)
        _set(ws, f"G{r}", line.net_kg)
    _clear_rows(ws, first_row + len(data.lines), total_row - 1, 7)
    _set(ws, f"C{total_row}", data.total_pieces)
    _set(ws, f"D{total_row}", data.total_qty)
    _set(ws, f"F{total_row}", data.total_gross_kg)
    _set(ws, f"G{total_row}", data.total_net_kg)


def _write_contract_sheet(ws: Worksheet, data: CustomsPackageData) -> None:
    exp = data.exporter
    _set(ws, "B2", exp.get("name"))
    _set(ws, "F6", data.order_no)
    _set(ws, "F8", _parse_date(data.dates.get("contract_date")))
    _set(ws, "F10", exp.get("signed_at") or "武汉")
    _set(ws, "F11", data.trade_term)
    _set(ws, "B8", data.buyer)
    _set(ws, "B10", data.remarks)  # 合同地址栏的订单备注，与报关单 B16 一致

    first_row = 18
    total_row = _ensure_data_rows(ws, first_row=first_row, template_row=18, count=len(data.lines), total_row=24, max_col=7)
    for i, line in enumerate(data.lines):
        r = first_row + i
        _set(ws, f"A{r}", line.chinese_name)
        _set(ws, f"C{r}", line.qty)
        _set(ws, f"D{r}", line.unit)
        _set(ws, f"E{r}", line.unit_price)
        _set(ws, f"F{r}", line.amount)
        _set(ws, f"G{r}", data.currency)
    _clear_rows(ws, first_row + len(data.lines), total_row - 1, 7)
    _set(ws, f"C{total_row}", data.total_qty)
    _set(ws, f"F{total_row}", data.total_amount)
    shift = total_row - 24
    _set(ws, f"A{28 + shift}", data.amount_in_words)
    _set(ws, f"B{32 + shift}", _parse_date(data.dates.get("shipment_deadline")))
    _set(ws, f"C{33 + shift}", data.dest_country)
    _set(ws, f"A{33 + shift}", f"(9)装运口岸和目的地           {data.loading_port}---")


def _write_entrust_sheet(ws: Worksheet, data: CustomsPackageData) -> None:
    exp = data.exporter
    _set(ws, "H5", data.order_no)
    _set(ws, "C21", exp.get("name"))
    _set(ws, "C22", data.main_goods_name)
    _set(ws, "C23", data.hs_codes_summary)
    _set(ws, "C24", data.total_amount)
    _set(ws, "C25", _parse_date(data.dates.get("entrust_import_export_date")))
    _set(ws, "C28", exp.get("domestic_source") or exp.get("signed_at") or "武汉")
    _set(ws, "H23", _parse_date(data.dates.get("entrust_doc_received_date")))
    _set(ws, "H16", _parse_date(data.dates.get("entrust_sign_date")))  # 签署日期（H24 是“合同□√”勾选格，不能写）
    _set(ws, "E12", data.entrust_valid_until)


# 各写入函数依赖的模板锚点。报关行换模板版本时，行号一旦错位就会静默填错格，
# 因此导出前先按 sheet 名 + 表头/合计行关键字体检一遍。
CUSTOMS_SHEET_ANCHORS: Tuple[dict, ...] = (
    {
        "index": 0,
        "name": "报关单",
        "header_row": 19,
        "header_keywords": ("项号", "商品编号"),
        "total_row": 27,
        "total_keyword": "总计",
    },
    {
        "index": 2,
        "name": "申报要素",
        "header_row": 9,
        "header_keywords": ("序号", "品名"),
        "total_row": 16,
        "total_keyword": None,
    },
    {
        "index": 3,
        "name": "发票",
        "header_row": 7,
        "header_keywords": ("數量", "單價"),
        "total_row": 15,
        "total_keyword": "TOTAL",
    },
    {
        "index": 4,
        "name": "箱单",
        "header_row": 8,
        "header_keywords": ("箱号", "毛重"),
        "total_row": 15,
        "total_keyword": "合计",
    },
    {
        "index": 5,
        "name": "合同",
        "header_row": 17,
        "header_keywords": ("Quantity", "Unit Price"),
        "total_row": 24,
        "total_keyword": "合计",
    },
    {
        "index": 6,
        "name": "委托书",
        "header_row": None,
        "header_keywords": (),
        "total_row": None,
        "total_keyword": None,
    },
)

CUSTOMS_REQUIRED_SHEET_COUNT = 7


def _row_text(ws: Worksheet, row: int) -> str:
    if row > ws.max_row:
        return ""
    parts: List[str] = []
    for col in range(1, min(ws.max_column, 26) + 1):
        val = ws.cell(row, col).value
        if val not in (None, ""):
            parts.append(str(val).replace("\n", " ").strip())
    return " | ".join(parts)


def validate_customs_template(wb) -> List[str]:
    """返回模板结构问题清单；空列表表示可安全按锚点写入。"""
    problems: List[str] = []
    sheets = wb.worksheets
    if len(sheets) < CUSTOMS_REQUIRED_SHEET_COUNT:
        problems.append(
            f"sheet 数量为 {len(sheets)}，需要 {CUSTOMS_REQUIRED_SHEET_COUNT} 个"
            f"（报关单 / 填制规范 / 申报要素 / 发票 / 箱单 / 合同 / 委托书）"
        )
        return problems

    for anchor in CUSTOMS_SHEET_ANCHORS:
        ws = sheets[anchor["index"]]
        title = str(ws.title).strip()
        if anchor["name"] not in title:
            problems.append(
                f"第 {anchor['index'] + 1} 个 sheet 名为「{title}」，预期包含「{anchor['name']}」"
            )
            continue

        header_row = anchor["header_row"]
        if header_row:
            text = _row_text(ws, header_row)
            missing = [kw for kw in anchor["header_keywords"] if kw not in text]
            if missing:
                problems.append(
                    f"「{title}」第 {header_row} 行应为数据区表头，未找到 {'、'.join(missing)}"
                    f"（实际：{text[:60] or '空行'}）"
                )

        total_row = anchor["total_row"]
        if total_row:
            if ws.max_row < total_row:
                problems.append(f"「{title}」只有 {ws.max_row} 行，缺少第 {total_row} 行合计行")
            elif anchor["total_keyword"] and anchor["total_keyword"] not in _row_text(ws, total_row):
                problems.append(
                    f"「{title}」第 {total_row} 行应为合计行，未找到「{anchor['total_keyword']}」"
                    f"（实际：{_row_text(ws, total_row)[:60] or '空行'}）"
                )

    return problems


class CustomsExportBlocked(ValueError):
    """导出前检查不通过（例如缺申报价），由接口返回 422。"""


def export_blockers(data: CustomsPackageData, exempt_codes: set | None = None) -> List[str]:
    """exempt_codes 仅供测试（golden B 级没有重量数据），正式导出不传。"""
    """按当前行数据检查（用户在预览里改过的值也算）。"""
    missing = [
        f"第{line.seq}项「{line.chinese_name}」"
        for line in data.lines
        if float(line.qty or 0) > 0 and float(line.unit_price or 0) <= 0
    ]
    out: List[str] = []
    cfg = [f"第{line.seq}项「{line.chinese_name}」{line.config_missing}" for line in data.lines if getattr(line, "config_missing", "")]
    if cfg:
        out.append("缺报关配置，不能导出：" + "、".join(cfg) + "。钓具等非宠物品类需先在 HS 配置里录入申报要素。")
    if missing:
        out.append("缺申报价（单价为 0 或空），不能导出：" + "、".join(missing) + "。请在产品库导入领星售价或手填申报价。")
    from fba_tool.customs_package import consistency_checks

    out.extend("一致性校验不通过：" + i["message"] for i in consistency_checks(data) if i["code"] not in (exempt_codes or set()))
    out.extend(b["message"] for b in (data.blocking or []) if b.get("code") == "weight_precision")
    return out


def export_customs_package(data: CustomsPackageData | dict, template_bytes: bytes, fmt: str = "xlsx",
                           *, _test_exempt_codes: set | None = None) -> bytes:
    """导出报关资料；fmt = xlsx（默认）| xls（v5 Q15，xlwt 按 7 张表重建）。"""
    fmt = str(fmt or "xlsx").lower().lstrip(".")
    if fmt not in ("xlsx", "xls"):
        raise ValueError(f"导出格式只能是 xlsx 或 xls：{fmt}")
    if isinstance(data, dict):
        lines = [CustomsLineItem(**ln) for ln in data.get("lines") or []]
        payload = dict(data)
        payload["lines"] = lines
        data = CustomsPackageData(**{k: v for k, v in payload.items() if k in CustomsPackageData.__dataclass_fields__})
        data.lines = lines

    blockers = export_blockers(data, _test_exempt_codes)
    if blockers:
        raise CustomsExportBlocked("；".join(blockers))

    wb = load_workbook(io.BytesIO(template_bytes))
    problems = validate_customs_template(wb)
    if problems:
        raise ValueError(
            "报关资料模板结构与预期不符，已中止导出以免把数据写进错误单元格：\n- "
            + "\n- ".join(problems)
            + "\n请确认使用的是随包发布的《报关资料模板.xlsx》；"
            "若报关行确实更新了模板，需同步调整 customs_export.py 中的锚点定义。"
        )

    sheets = wb.worksheets
    _write_declaration_sheet(sheets[0], data)
    _write_elements_sheet(sheets[2], data.lines)
    _write_invoice_sheet(sheets[3], data)
    _write_packing_sheet(sheets[4], data)
    _write_contract_sheet(sheets[5], data)
    _write_entrust_sheet(sheets[6], data)

    if fmt == "xls":
        from fba_tool.xls_writer import workbook_to_xls

        return workbook_to_xls(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
