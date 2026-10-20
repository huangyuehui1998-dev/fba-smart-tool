"""Send to Amazon (STA) — 纯二维数组构建，像素级对齐官方空白模板粘贴区。

列坐标（1-based Excel）：
  A~J (1~10)  SKU 基础列
  K   (11)    死亡空列 — 必须留空
  L+  (12+)   包装箱 n 数量 / P1-Bn / 箱规参数
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List

import pandas as pd
from openpyxl import Workbook

from fba_tool.models import SkuRow
from fba_tool.product_catalog import lookup_sta_product_name
from fba_tool.user_config import label, render_formula

# 0-based 列索引
IDX_DEATH_COL = 10  # K 列
IDX_FIRST_BOX = 11  # L 列
BASE_COL_COUNT = 11  # A~K

STA_LB_PER_KG = 2.2046
STA_CM_PER_INCH = 2.54

INSTRUCTION_TEXT = "请在下方提供此装箱组的包装箱详情。如果您有任何疑问，请参阅说明表。"
PACK_GROUP_TEXT = "装箱组：1"

HEADER_BASE: List[str] = [
    "SKU",
    "商品名称 ",
    "ASIN",
    "FNSKU",
    "状况",
    "预处理类型",
    "商品预处理方是谁？",
    "商品贴标方是谁？",
    "预计数量",
    "装箱数量",
    "",
]


def _matrix_box_col_name(box_no: int) -> str:
    return render_formula(label("internal_box_header", "第{box_no}箱"), box_no=box_no)


def _resolve_box_count(editor_df: pd.DataFrame, max_boxes: int | None) -> int:
    """实际箱数。仅在调用方显式给出上限时才截断——绝不用 amazon_max_boxes 静默丢箱。"""
    if editor_df.empty:
        return 0
    actual = int(editor_df["箱号"].astype(int).max())
    if max_boxes is None:
        return actual
    return min(actual, max_boxes)


def _matrix_max_box(box_matrix: pd.DataFrame) -> int:
    if box_matrix.empty:
        return 0
    return len([c for c in box_matrix.columns if c != "SKU简称"])


def _kg_to_lb(kg: float) -> float:
    return round(float(kg) * STA_LB_PER_KG, 2)


def _cm_to_in(cm: float) -> float:
    return round(float(cm) / STA_CM_PER_INCH, 2)


def _cell_qty(val) -> Any:
    """0 或空 → 空字符串，供粘贴对齐。"""
    if val in (None, "", 0) or (isinstance(val, float) and pd.isna(val)):
        return ""
    return int(val)


def _cell_metric(val) -> Any:
    if val in (None, "", 0) or (isinstance(val, float) and pd.isna(val)):
        return ""
    return round(float(val), 2)


def _footer_prefix(label_text: str) -> List[Any]:
    return [label_text] + [""] * IDX_DEATH_COL


def _per_box_metrics_imperial(
    editor_df: pd.DataFrame,
    sku_map: Dict[str, SkuRow],
    max_box: int,
) -> Dict[int, dict]:
    result: Dict[int, dict] = {}
    for box_no in range(1, max_box + 1):
        result[box_no] = {
            "weight_lb": "",
            "width_in": "",
            "length_in": "",
            "height_in": "",
        }
    if editor_df.empty:
        return result
    for box_no, group in editor_df.groupby("箱号"):
        box_no = int(box_no)
        if box_no < 1 or box_no > max_box:
            continue
        from fba_tool.box_weight import box_gross_kg

        max_l = max_w = max_h = 0.0
        has_dim = False
        lines = []
        for _, row in group.iterrows():
            sku = sku_map.get(str(row["SKU简称"]))
            if not sku:
                continue
            qty = int(row["数量"])
            lines.append((sku.sku, qty))
            if sku.length_cm or sku.width_cm or sku.height_cm:
                has_dim = True
            max_l = max(max_l, sku.length_cm)
            max_w = max(max_w, sku.width_cm)
            max_h = max(max_h, sku.height_cm)
        gross_kg = round(box_gross_kg(lines, sku_map), 2)
        result[box_no] = {
            "weight_lb": _kg_to_lb(gross_kg) if gross_kg else "",
            "width_in": _cm_to_in(max_w) if has_dim and max_w else "",
            "length_in": _cm_to_in(max_l) if has_dim and max_l else "",
            "height_in": _cm_to_in(max_h) if has_dim and max_h else "",
        }
    return result


def build_sta_grid(
    *,
    sku_rows: List[SkuRow],
    box_matrix: pd.DataFrame,
    editor_df: pd.DataFrame,
    amazon_meta: Dict[str, dict],
    fnsku_map: Dict[str, str],
    product_names: Dict[str, str],
    max_box: int,
) -> List[List[Any]]:
    """按 STA 粘贴坐标构建二维数组（List of Lists）。"""
    sku_map = {s.sku: s for s in sku_rows}
    sku_count = len(sku_rows)
    total_units = sum(int(s.total_qty) for s in sku_rows)
    box_metrics = _per_box_metrics_imperial(editor_df, sku_map, max_box)

    if not box_matrix.empty:
        ordered_skus = box_matrix["SKU简称"].astype(str).tolist()
    else:
        ordered_skus = [s.sku for s in sku_rows]

    rows: List[List[Any]] = []

    # ── Row 1~4 顶部 ──
    rows.append([INSTRUCTION_TEXT])
    rows.append([PACK_GROUP_TEXT])
    row3: List[Any] = [f"SKU 总数：{sku_count}（{total_units} 件商品）"]
    row3.extend([""] * (IDX_DEATH_COL - 1))
    row3.extend(["包装箱总数：", max_box])
    rows.append(row3)
    rows.append([])

    # ── Row 5 表头 ──
    header = list(HEADER_BASE) + [f"包装箱 {i} 数量" for i in range(1, max_box + 1)]
    rows.append(header)

    # ── Row 6+ SKU 明细 ──
    for sku in ordered_skus:
        src = sku_map.get(sku)
        if not src:
            continue
        meta = amazon_meta.get(sku, {})
        fnsku = fnsku_map.get(sku) or meta.get("fnsku") or ""
        mat_row = box_matrix[box_matrix["SKU简称"].astype(str) == sku]
        box_vals: List[Any] = []
        packed_total = 0
        for box_no in range(1, max_box + 1):
            qty_raw = None
            if not mat_row.empty:
                col_name = _matrix_box_col_name(box_no)
                if col_name in mat_row.columns:
                    qty_raw = mat_row.iloc[0][col_name]
            q = _cell_qty(qty_raw)
            if q != "":
                packed_total += int(q)
            box_vals.append(q)

        line: List[Any] = [
            sku,
            product_names.get(sku) or lookup_sta_product_name(sku, amazon_meta=amazon_meta),
            meta.get("asin") or "",
            fnsku,
            meta.get("condition") or "NewItem",
            meta.get("prep_owner") or "贴标",
            "",
            meta.get("label_owner") or "由卖家提供",
            int(src.total_qty),
            packed_total if packed_total > 0 else "",
        ]
        line.append("")
        line.extend(box_vals)
        rows.append(line)

    # ── Footer ──
    rows.append([])
    rows.append(_footer_prefix("包装箱名称") + [f"P1 - B{i}" for i in range(1, max_box + 1)])
    rows.append(
        _footer_prefix("包装箱重量（磅）：")
        + [_cell_metric(box_metrics[i].get("weight_lb")) for i in range(1, max_box + 1)]
    )
    rows.append(
        _footer_prefix("包装箱宽度（英寸）：")
        + [_cell_metric(box_metrics[i].get("width_in")) for i in range(1, max_box + 1)]
    )
    rows.append(
        _footer_prefix("包装箱长度（英寸）：")
        + [_cell_metric(box_metrics[i].get("length_in")) for i in range(1, max_box + 1)]
    )
    rows.append(
        _footer_prefix("包装箱高度（英寸）：")
        + [_cell_metric(box_metrics[i].get("height_in")) for i in range(1, max_box + 1)]
    )

    return rows


def _write_grid_to_xlsx(rows: List[List[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SendToAmazon"
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            if val == "" or val is None:
                continue
            ws.cell(row=r_idx, column=c_idx, value=val)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def pick_sta_packing_sheet(wb) -> Any:
    """返回 STA 数据工作表（纯网格导出时为 active sheet）。"""
    for ws in wb.worksheets:
        a1 = ws.cell(1, 1).value
        if a1 and "请在下方提供" in str(a1):
            return ws
    return wb.active


def export_sta_template_bytes(
    editor_df: pd.DataFrame,
    sku_rows: List[SkuRow],
    *,
    max_boxes: int | None = None,
    fnsku_overrides: Dict[str, str] | None = None,
    amazon_meta: Dict[str, dict] | None = None,
    fnsku_map: Dict[str, str] | None = None,
    box_matrix: pd.DataFrame | None = None,
    product_names: Dict[str, str] | None = None,
) -> bytes:
    meta = amazon_meta or {}
    catalog = None
    matrix = box_matrix if box_matrix is not None else pd.DataFrame()
    max_box = _matrix_max_box(matrix) if not matrix.empty else _resolve_box_count(editor_df, max_boxes)
    if max_box <= 0:
        max_box = _resolve_box_count(editor_df, max_boxes)

    names: Dict[str, str] = dict(product_names or {})
    for s in sku_rows:
        if s.sku not in names:
            names[s.sku] = lookup_sta_product_name(s.sku, amazon_meta=meta, catalog=catalog)

    grid = build_sta_grid(
        sku_rows=sku_rows,
        box_matrix=matrix,
        editor_df=editor_df,
        amazon_meta=meta,
        fnsku_map=fnsku_map or {},
        product_names=names,
        max_box=max_box,
    )
    return _write_grid_to_xlsx(grid)
