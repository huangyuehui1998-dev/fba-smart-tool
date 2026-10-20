"""箱规库导入。沿用箱规模板，并多读箱规名称、空箱重、备注。"""
from __future__ import annotations

from typing import Any, Dict, List, Tuple

from fba_tool.product_import import _issue, parse_carton_rows


def parse_carton_spec_rows(data: bytes, file_name: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    records, errors, _first = parse_carton_rows(data, file_name)
    extra = _extra_columns(data, file_name)
    out = []
    for record in records:
        row_no = int(record["row"])
        values = record["values"]
        more = extra.get(row_no, {})
        out.append({
            "row": row_no,
            "sku": record["sku"],
            "spec_name": more.get("spec_name") or "",
            "length_cm": values.get("carton_length"),
            "width_cm": values.get("carton_width"),
            "height_cm": values.get("carton_height"),
            "qty_per_carton": values.get("qty_per_box"),
            "gross_kg": values.get("carton_gross_weight"),
            "tare_kg": more.get("tare_kg"),
            "note": more.get("note") or "",
            "source": "import_xlsx",
        })
    return out, errors


def _extra_columns(data: bytes, file_name: str) -> Dict[int, Dict[str, Any]]:
    import io

    from openpyxl import load_workbook

    if not file_name.lower().endswith((".xlsx", ".xlsm")):
        return {}
    wb = load_workbook(io.BytesIO(data), data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return {}
    headers = [str(cell or "").strip() for cell in rows[0]]
    name_at = _index(headers, ["箱规名称", "规格名称", "spec_name"])
    tare_at = _index(headers, ["空箱重", "空箱重(kg)", "空纸箱重(kg)"])
    note_at = _index(headers, ["备注", "note"])
    found: Dict[int, Dict[str, Any]] = {}
    for offset, row in enumerate(rows[1:], start=2):
        item: Dict[str, Any] = {}
        if name_at is not None and name_at < len(row) and row[name_at] not in (None, ""):
            item["spec_name"] = str(row[name_at]).strip()
        if tare_at is not None and tare_at < len(row) and row[tare_at] not in (None, ""):
            item["tare_kg"] = row[tare_at]
        if note_at is not None and note_at < len(row) and row[note_at] not in (None, ""):
            item["note"] = str(row[note_at]).strip()
        if item:
            found[offset] = item
    return found


def _index(headers: List[str], names: List[str]) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None
