from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any, Dict, List

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fba_tool.converters import cm_to_inch, kg_to_lb
from fba_tool.fba_parser import FACTORY_WAREHOUSE_ORDER
from fba_tool.forwarder_quote import (
    DIM_FIELDS,
    FORWARDER_COLUMNS,
    SUBTOTAL_CBM_FLAG,
    SUBTOTAL_FLAG,
    append_warehouse_subtotals,
    compute_forwarder_total_cbm,
    format_subtotal_cbm_label,
    is_forwarder_subtotal_row,
)
from fba_tool.models import SkuRow
from fba_tool.shipping_sheet import (
    SHIPPING_SHEET_COLUMNS,
    SHIPPING_SHEET_HEADERS,
    SHIPPING_SHEET_LEGEND_COLUMN,
    SHIPPING_SHEET_LEGEND_COLUMNS,
    SHIPPING_SHEET_LEGEND_SKU_SPAN,
)
from fba_tool.sta_template_export import export_sta_template_bytes
from fba_tool.config_cache import load_json_cached
from fba_tool.template_layout import (
    InternalLayout,
    clear_range,
    detect_factory_warehouse_columns,
    detect_internal_layout,
    ensure_sku_capacity,
    safe_write,
)
from fba_tool.template_loader import find_template, template_path
from fba_tool.user_config import (
    factory_warehouse_columns,
    formula,
    forwarder_column_map,
    label,
    param,
    render_formula,
)

def _box_col_name(box_no: int) -> str:
    return render_formula(label("internal_box_header", "第{box_no}箱"), box_no=box_no)


def export_internal_archive(
    editor_df: pd.DataFrame,
    sku_rows: List[SkuRow],
    max_boxes: int | None = None,
    fnsku_overrides: Dict[str, str] | None = None,
) -> bytes:
    template = template_path("内部留档表.xlsx") or find_template("内部", "留档")
    sku_map = {s.sku: s for s in sku_rows}
    max_box = _resolve_box_count(editor_df, max_boxes)
    matrix = build_box_matrix(editor_df, max_boxes=max_box)
    fnsku_map = _load_fnsku_mapping(sku_rows, fnsku_overrides)

    if template:
        wb = load_workbook(template)
        ws = wb.active
        layout = detect_internal_layout(ws)
        _prepare_internal_structure(ws, layout, len(sku_rows), max_box)
        _fill_internal_sheet(ws, layout, sku_rows, matrix, sku_map, fnsku_map, editor_df, max_box)
        return _to_bytes(wb)

    return _write_internal_from_scratch(sku_rows, matrix, sku_map, fnsku_map, editor_df, max_box)


def export_amazon_template(
    editor_df: pd.DataFrame,
    sku_rows: List[SkuRow],
    max_boxes: int | None = None,
    fnsku_overrides: Dict[str, str] | None = None,
) -> bytes:
    from fba_tool.product_catalog import lookup_sta_product_name

    amazon_meta = _load_amazon_sku_metadata()
    catalog = None
    matrix = build_box_matrix(editor_df, max_boxes=max_boxes)
    product_names = {
        s.sku: lookup_sta_product_name(s.sku, amazon_meta=amazon_meta, catalog=catalog)
        for s in sku_rows
    }
    return export_sta_template_bytes(
        editor_df,
        sku_rows,
        max_boxes=max_boxes,
        fnsku_overrides=fnsku_overrides,
        amazon_meta=amazon_meta,
        fnsku_map=_load_fnsku_mapping(sku_rows, fnsku_overrides),
        box_matrix=matrix,
        product_names=product_names,
    )


def export_factory_sheet(factory_long_df: pd.DataFrame) -> bytes:
    if not factory_long_df.empty and "本地箱号" in factory_long_df.columns:
        cols = ["本地箱号", "产品名称", "数量", "FBA箱号", "目的仓", "货件编号", "贴标说明"]
        existing = [c for c in cols if c in factory_long_df.columns]
        return _write_simple_df(factory_long_df[existing], title="工厂贴标表")
    template = template_path("工厂表.xlsx")
    if template:
        wb = load_workbook(template)
        ws = wb.active
        if not factory_long_df.empty:
            _fill_factory_template(ws, factory_long_df)
        return _to_bytes(wb)
    return _write_simple_df(build_factory_layout(factory_long_df), title="工厂表")


def export_forwarder_quote(forwarder_df: pd.DataFrame) -> bytes:
    """导出货代表：与页面预览列一致，不使用模板。"""
    df = append_warehouse_subtotals(forwarder_df)
    return _write_forwarder_simple(df)


def build_box_matrix(editor_df: pd.DataFrame, max_boxes: int | None = None) -> pd.DataFrame:
    max_box = _resolve_box_count(editor_df, max_boxes)
    if editor_df.empty or max_box <= 0:
        return pd.DataFrame(columns=["SKU简称"])
    work = editor_df.copy()
    work["箱号"] = work["箱号"].astype(int)
    work["SKU简称"] = work["SKU简称"].astype(str)
    grouped = work.groupby(["SKU简称", "箱号"], as_index=False)["数量"].sum()
    pivot = grouped.pivot(index="SKU简称", columns="箱号", values="数量")
    skus = sorted(pivot.index.tolist())
    cols = ["SKU简称"] + [_box_col_name(i) for i in range(1, max_box + 1)]
    rows = []
    for sku in skus:
        row = {"SKU简称": sku}
        for i in range(1, max_box + 1):
            val = pivot.at[sku, i] if i in pivot.columns else None
            row[_box_col_name(i)] = int(val) if val not in (None, 0) and not pd.isna(val) else None
        rows.append(row)
    return pd.DataFrame(rows, columns=cols)


def build_factory_horizontal(box_items_df: pd.DataFrame) -> pd.DataFrame:
    return build_factory_layout(box_items_df)


def build_factory_layout(box_items_df: pd.DataFrame) -> pd.DataFrame:
    if box_items_df.empty:
        return pd.DataFrame()
    if "本地箱号" in box_items_df.columns:
        cols = ["本地箱号", "产品名称", "数量", "FBA箱号", "目的仓", "贴标说明"]
        existing = [c for c in cols if c in box_items_df.columns]
        sub = box_items_df[existing].copy()
        sort_col = "本地箱号" if "本地箱号" in sub.columns else existing[0]
        return sub.sort_values(sort_col).reset_index(drop=True)
    data_whs = set(box_items_df["仓库代码"].astype(str))
    config_whs = list(factory_warehouse_columns().values())
    warehouses = [w for w in config_whs if w in data_whs]
    for w in FACTORY_WAREHOUSE_ORDER:
        if w in data_whs and w not in warehouses:
            warehouses.append(w)
    for w in sorted(data_whs):
        if w not in warehouses:
            warehouses.append(w)
    columns: Dict[str, list] = {}
    max_rows = 0
    for wh in warehouses:
        sub = box_items_df[box_items_df["仓库代码"].astype(str) == wh].copy()
        box_col = "箱号" if "箱号" in sub.columns else "纸箱编号"
        sub = sub.sort_values([box_col, "产品名称"])
        columns[f"{wh}_纸箱编号"] = [f"{param('factory_box_label_prefix', '纸箱')}{int(x)}" for x in sub[box_col].astype(int).tolist()]
        columns[f"{wh}_产品名称"] = sub["产品名称"].astype(str).tolist()
        max_rows = max(max_rows, len(sub))
    for values in columns.values():
        while len(values) < max_rows:
            values.append("")
    return pd.DataFrame(columns)


def _prepare_internal_structure(ws, layout: InternalLayout, sku_count: int, max_box: int) -> None:
    ensure_sku_capacity(ws, layout, sku_count, layout.sku_start_row)
    last_col = layout.first_box_col + max(max_box, layout.template_max_boxes) - 1
    clear_range(ws, layout.sku_start_row, layout.footer_start_row - 1, 1, last_col)
    hdr = layout.sku_header_row
    # 仅扩展超出模板原有箱数的列，避免破坏合并表头
    for i in range(layout.template_max_boxes + 1, max_box + 1):
        safe_write(ws, hdr, layout.first_box_col + i - 1, _box_col_name(i))


def _fill_internal_sheet(
    ws,
    layout: InternalLayout,
    sku_rows: List[SkuRow],
    matrix: pd.DataFrame,
    sku_map: Dict[str, SkuRow],
    fnsku_map: Dict[str, str],
    editor_df: pd.DataFrame,
    max_box: int,
) -> None:
    cols = layout.sku_columns
    c_prod = cols.get("product", 1)
    c_fnsku = cols.get("fnsku", 2)
    c_total = cols.get("total", 3)
    c_carton = cols.get("carton", 4)
    c_ct_cnt = cols.get("carton_count", 5)
    c_check = cols.get("check", 6)
    last_col_letter = get_column_letter(layout.first_box_col + max_box - 1)
    first_col_letter = get_column_letter(layout.first_box_col)
    t_let_tpl = get_column_letter(c_total)
    c_let_tpl = get_column_letter(c_carton)
    vol_div = param("volume_divisor", 6000)

    for idx, src in enumerate(sku_rows):
        row_idx = layout.sku_start_row + idx
        sku = src.sku
        safe_write(ws, row_idx, c_prod, sku)
        safe_write(ws, row_idx, c_fnsku, fnsku_map.get(sku, ""))
        safe_write(ws, row_idx, c_total, src.total_qty)
        safe_write(ws, row_idx, c_carton, src.carton_qty)
        safe_write(
            ws,
            row_idx,
            c_ct_cnt,
            render_formula(
                formula("internal", "carton_count", f"={t_let_tpl}{{row}}/{{carton_col}}{{row}}"),
                row=row_idx,
                total_col=t_let_tpl,
                carton_col=c_let_tpl,
            ),
        )
        safe_write(
            ws,
            row_idx,
            c_check,
            render_formula(
                formula("internal", "check_qty", f"=SUM({first_col_letter}{{row}}:{last_col_letter}{{row}})"),
                row=row_idx,
                first_box_col=first_col_letter,
                last_box_col=last_col_letter,
            ),
        )
        mat_row = matrix[matrix["SKU简称"] == sku]
        if not mat_row.empty:
            mat = mat_row.iloc[0]
            for i in range(1, max_box + 1):
                val = mat.get(_box_col_name(i))
                col_idx = layout.first_box_col + i - 1
                if val not in (None, "", 0) and not pd.isna(val):
                    safe_write(ws, row_idx, col_idx, int(val))

    box_metrics = _per_box_metrics(editor_df, sku_map)
    fr = layout.footer_rows
    lr, wr, hr = fr.get("length"), fr.get("width"), fr.get("height")
    wt, vol = fr.get("weight"), fr.get("volume")
    glb = fr.get("gross_lb")
    win, lin, hin = fr.get("width_in"), fr.get("length_in"), fr.get("height_in")
    clear_cols = max(max_box, layout.template_max_boxes)

    for i in range(1, clear_cols + 1):
        col_idx = layout.first_box_col + i - 1
        col = get_column_letter(col_idx)
        m = box_metrics.get(i, {}) if i <= max_box else {}
        if lr:
            safe_write(ws, lr, col_idx, m.get("长(cm)", "") if i <= max_box else None)
        if wr:
            safe_write(ws, wr, col_idx, m.get("宽(cm)", "") if i <= max_box else None)
        if hr:
            safe_write(ws, hr, col_idx, m.get("高(cm)", "") if i <= max_box else None)
        if wt:
            safe_write(ws, wt, col_idx, m.get("实重(kg)", "") if i <= max_box else None)
        if vol and lr and wr and hr:
            if i <= max_box:
                safe_write(
                    ws,
                    vol,
                    col_idx,
                    render_formula(
                        formula("internal", "volume", f"={col}{{length_row}}*{col}{{width_row}}*{col}{{height_row}}/{{volume_divisor}}"),
                        col=col,
                        length_row=lr,
                        width_row=wr,
                        height_row=hr,
                        volume_divisor=vol_div,
                    ),
                )
            else:
                safe_write(ws, vol, col_idx, None)
        if glb and wt:
            if i <= max_box:
                safe_write(
                    ws,
                    glb,
                    col_idx,
                    render_formula(
                        formula("internal", "gross_lb", f"={col}{{weight_row}}*{{internal_gross_lb_factor}}"),
                        col=col,
                        weight_row=wt,
                    ),
                )
            else:
                safe_write(ws, glb, col_idx, None)
        if win and wr:
            if i <= max_box:
                safe_write(
                    ws,
                    win,
                    col_idx,
                    render_formula(
                        formula("internal", "width_in", f"={col}{{width_row}}/{{inch_to_cm_factor}}"),
                        col=col,
                        width_row=wr,
                    ),
                )
            else:
                safe_write(ws, win, col_idx, None)
        if lin and lr:
            if i <= max_box:
                safe_write(
                    ws,
                    lin,
                    col_idx,
                    render_formula(
                        formula("internal", "length_in", f"={col}{{length_row}}/{{inch_to_cm_factor}}"),
                        col=col,
                        length_row=lr,
                    ),
                )
            else:
                safe_write(ws, lin, col_idx, None)
        if hin and hr:
            if i <= max_box:
                safe_write(
                    ws,
                    hin,
                    col_idx,
                    render_formula(
                        formula("internal", "height_in", f"={col}{{height_row}}/{{inch_to_cm_factor}}"),
                        col=col,
                        height_row=hr,
                    ),
                )
            else:
                safe_write(ws, hin, col_idx, None)

    if wt:
        safe_write(
            ws,
            wt,
            2,
            render_formula(
                formula("internal", "sum_weight_row", f"=SUM({first_col_letter}{{weight_row}}:{last_col_letter}{{weight_row}})"),
                weight_row=wt,
                first_box_col=first_col_letter,
                last_box_col=last_col_letter,
            ),
        )
    if vol:
        safe_write(
            ws,
            vol,
            2,
            render_formula(
                formula("internal", "sum_volume_row", f"=SUM({first_col_letter}{{volume_row}}:{last_col_letter}{{volume_row}})"),
                volume_row=vol,
                first_box_col=first_col_letter,
                last_box_col=last_col_letter,
            ),
        )
    tq = fr.get("total_qty")
    if tq:
        safe_write(ws, tq, 3, sum(s.total_qty for s in sku_rows))


def _resolve_box_count(editor_df: pd.DataFrame, max_boxes: int | None) -> int:
    """实际箱数。仅在调用方显式给出上限时才截断——绝不用 amazon_max_boxes 静默丢箱。"""
    if editor_df.empty:
        return 0
    actual = int(editor_df["箱号"].astype(int).max())
    if max_boxes is None:
        return actual
    return min(actual, max_boxes)


def _per_box_metrics(editor_df: pd.DataFrame, sku_map: Dict[str, SkuRow], imperial: bool = False) -> Dict[int, dict]:
    result: Dict[int, dict] = {}
    if editor_df.empty:
        return result
    for box_no, group in editor_df.groupby("箱号"):
        box_no = int(box_no)
        from fba_tool.box_weight import box_gross_kg, outer_dims_cm

        lines = []
        for _, row in group.iterrows():
            sku = sku_map.get(str(row["SKU简称"]))
            if not sku:
                continue
            lines.append((sku.sku, int(row["数量"])))
        max_l, max_w, max_h = outer_dims_cm(lines, sku_map)
        gross_kg = round(box_gross_kg(lines, sku_map), 2)
        entry = {"长(cm)": max_l, "宽(cm)": max_w, "高(cm)": max_h, "实重(kg)": gross_kg}
        if imperial:
            entry.update(
                {
                    "weight_lb": kg_to_lb(gross_kg),
                    "width_in": cm_to_inch(max_w),
                    "length_in": cm_to_inch(max_l),
                    "height_in": cm_to_inch(max_h),
                }
            )
        result[box_no] = entry
    return result


def _fill_factory_template(ws, factory_long_df: pd.DataFrame):
    wh_cols = factory_warehouse_columns()
    if not wh_cols:
        wh_cols = detect_factory_warehouse_columns(ws, header_row=2)
    if not wh_cols:
        wh_cols = {2: "FWA4", 5: "ORF2", 9: "GYR2", 13: "SMF3", 17: "RDU2"}

    start_row = int(param("factory_data_start_row", 4))
    prefix = str(param("factory_box_label_prefix", "纸箱"))
    max_data_row = max(start_row, ws.max_row)
    for col in wh_cols:
        clear_range(ws, start_row, max_data_row, col, col + 1)

    for col, wh in sorted(wh_cols.items()):
        sub = factory_long_df[factory_long_df["仓库代码"].astype(str) == wh].copy()
        if sub.empty:
            continue
        sub = sub.sort_values(["纸箱编号", "产品名称"])
        for i, (_, row) in enumerate(sub.iterrows()):
            r = start_row + i
            safe_write(ws, r, col, f"{prefix}{int(row['纸箱编号'])}")
            safe_write(ws, r, col + 1, str(row["产品名称"]))


def _detect_forwarder_columns(ws) -> Dict[str, int]:
    keywords = {
        "仓库代码": ("仓库",),
        "SKU": ("sku",),
        "箱序号": ("箱序号", "箱号"),
        "单件实重(kg)": ("单件", "单重"),
        "箱数": ("箱数",),
        "箱子长(cm)": ("长", "length"),
        "箱子宽(cm)": ("宽", "width"),
        "箱子高(cm)": ("高", "height"),
        "总实重(kg)": ("总实重", "实重", "重量", "weight"),
        "总体积重(kg)": ("体积", "抛重"),
    }
    found: Dict[str, int] = {}
    for r in range(1, min(ws.max_row, 5) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            val = str(ws.cell(r, c).value or "").strip().lower()
            if not val:
                continue
            for field, kws in keywords.items():
                if field in found:
                    continue
                if any(kw.lower() in val for kw in kws):
                    found[field] = c
    return found


def _forwarder_subtotal_styles():
    return (
        PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid"),
        Font(bold=True),
    )


def _style_forwarder_subtotal_row(ws, row_idx: int, max_col: int) -> None:
    fill, font = _forwarder_subtotal_styles()
    for col in range(1, max_col + 1):
        cell = ws.cell(row_idx, col)
        cell.fill = fill
        cell.font = font


def _write_forwarder_global_summary(ws, total_cbm: float, start_row: int) -> None:
    summary_row = max(1, start_row - 1)
    cell = ws.cell(summary_row, 2)
    cell.value = f"总方数: {total_cbm:.3f} CBM"
    cell.font = Font(bold=True, size=12)
    cell.fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")


def _write_forwarder_subtotal_cbm(ws, row_idx: int, col_map: Dict[str, int], cbm_value: float) -> None:
    l_col = col_map.get("箱子长(cm)")
    w_col = col_map.get("箱子宽(cm)")
    h_col = col_map.get("箱子高(cm)")
    text = format_subtotal_cbm_label(float(cbm_value or 0))
    if l_col and w_col and h_col:
        safe_write(ws, row_idx, l_col, text)
        safe_write(ws, row_idx, w_col, "")
        safe_write(ws, row_idx, h_col, "")
        ws.merge_cells(
            start_row=row_idx,
            start_column=l_col,
            end_row=row_idx,
            end_column=h_col,
        )
    elif l_col:
        safe_write(ws, row_idx, l_col, text)


def _forwarder_subtotal_cbm(row) -> float:
    if isinstance(row, pd.Series):
        if SUBTOTAL_CBM_FLAG in row.index and pd.notna(row.get(SUBTOTAL_CBM_FLAG)):
            return round(float(row[SUBTOTAL_CBM_FLAG]), 3)
        text = str(row.get("箱子长(cm)", ""))
    else:
        val = row.get(SUBTOTAL_CBM_FLAG)
        if val not in (None, ""):
            return round(float(val), 3)
        text = str(row.get("箱子长(cm)", ""))
    if "方数:" in text and "CBM" in text.upper():
        match = re.search(r"方数:\s*([\d.]+)\s*CBM", text, flags=re.I)
        if match:
            return round(float(match.group(1)), 3)
    return 0.0


def _fill_forwarder_template(ws, forwarder_df: pd.DataFrame):
    col_map = forwarder_column_map()
    if not col_map:
        col_map = _detect_forwarder_columns(ws)
    if not col_map:
        col_map = {
            "仓库代码": 8,
            "箱数": 9,
            "箱子长(cm)": 10,
            "箱子宽(cm)": 11,
            "箱子高(cm)": 12,
            "总实重(kg)": 13,
            "总体积重(kg)": 14,
        }

    start_row = int(param("forwarder_data_start_row", 3))
    country = str(param("forwarder_country_default", "美国"))
    total_cbm = compute_forwarder_total_cbm(forwarder_df)
    _write_forwarder_global_summary(ws, total_cbm, start_row)
    max_data_row = max(start_row, ws.max_row)
    used_cols = set(col_map.values()) | {1}
    for c in used_cols:
        clear_range(ws, start_row, max_data_row, c, c)

    for i, (_, row) in enumerate(forwarder_df.iterrows()):
        r = start_row + i
        safe_write(ws, r, 1, country)
        if is_forwarder_subtotal_row(row):
            cbm_value = _forwarder_subtotal_cbm(row)
            for field, col in col_map.items():
                if field in DIM_FIELDS:
                    continue
                safe_write(ws, r, col, row.get(field, ""))
            _write_forwarder_subtotal_cbm(ws, r, col_map, cbm_value)
            max_col = max(max(col_map.values()), 1)
            _style_forwarder_subtotal_row(ws, r, max_col)
        else:
            for field, col in col_map.items():
                safe_write(ws, r, col, row.get(field, ""))


def _forwarder_display_columns(df: pd.DataFrame) -> List[str]:
    from fba_tool.forwarder_quote import DETAIL_CBM_FLAG

    hidden = {SUBTOTAL_FLAG, SUBTOTAL_CBM_FLAG, DETAIL_CBM_FLAG}
    extras = [c for c in df.columns if c not in FORWARDER_COLUMNS and c not in hidden]
    return list(FORWARDER_COLUMNS) + extras


def _format_forwarder_export_cell(col: str, row: pd.Series) -> Any:
    subtotal = is_forwarder_subtotal_row(row)
    raw = row.get(col, "")

    if subtotal and col in DIM_FIELDS:
        if col == "箱子长(cm)":
            if isinstance(raw, str) and "方数:" in raw:
                return raw
            return format_subtotal_cbm_label(_forwarder_subtotal_cbm(row))
        return ""

    if raw is None or (isinstance(raw, float) and pd.isna(raw)) or raw == "":
        return ""

    if subtotal:
        if col == "箱数":
            return int(round(float(raw)))
        if col == "总实重(kg)":
            return round(float(raw), 2)
        if col == "总体积重(kg)":
            return round(float(raw), 3)
        return raw

    if col == "箱数":
        return int(round(float(raw)))
    if col == "单件实重(kg)":
        return round(float(raw), 3)
    if col in ("总实重(kg)", "单箱体积重(kg)", "单箱实重(kg)"):
        return round(float(raw), 2)
    if col == "总体积重(kg)":
        return round(float(raw), 3)
    return raw


def _write_forwarder_simple(df: pd.DataFrame) -> bytes:
    display_cols = _forwarder_display_columns(df)
    wb = Workbook()
    ws = wb.active
    ws.title = "货代表"

    header_font = Font(bold=True)
    for ci, col in enumerate(display_cols, start=1):
        cell = ws.cell(1, ci, col)
        cell.font = header_font

    dim_cols = {name: idx + 1 for idx, name in enumerate(display_cols) if name in DIM_FIELDS}
    subtotal_rows: List[int] = []
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for ci, col in enumerate(display_cols, start=1):
            if is_forwarder_subtotal_row(row) and col in ("箱子宽(cm)", "箱子高(cm)"):
                continue
            ws.cell(row_idx, ci, _format_forwarder_export_cell(col, row))

        if is_forwarder_subtotal_row(row):
            subtotal_rows.append(row_idx)
            cbm_value = _forwarder_subtotal_cbm(row)
            l_idx = dim_cols.get("箱子长(cm)")
            w_idx = dim_cols.get("箱子宽(cm)")
            h_idx = dim_cols.get("箱子高(cm)")
            text = format_subtotal_cbm_label(cbm_value)
            if l_idx and w_idx and h_idx:
                ws.cell(row_idx, l_idx).value = text
                ws.cell(row_idx, w_idx).value = ""
                ws.cell(row_idx, h_idx).value = ""
                ws.merge_cells(
                    start_row=row_idx,
                    start_column=l_idx,
                    end_row=row_idx,
                    end_column=h_idx,
                )
            elif l_idx:
                ws.cell(row_idx, l_idx).value = text

    for r in subtotal_rows:
        _style_forwarder_subtotal_row(ws, r, len(display_cols))
    return _to_bytes(wb)


def _write_internal_from_scratch(sku_rows, matrix, sku_map, fnsku_map, editor_df, max_box):
    wb = Workbook()
    ws = wb.active
    ws.title = "亚马逊发货明细"
    headers = ["产品", "FNSKU", "总数量", "单箱装量", "箱数", "校验数量"] + [
        _box_col_name(i) for i in range(1, max_box + 1)
    ]
    ws.append(headers)
    layout = InternalLayout(
        sku_header_row=1,
        sku_start_row=2,
        first_box_col=7,
        sku_columns={"product": 1, "fnsku": 2, "total": 3, "carton": 4, "carton_count": 5, "check": 6},
        footer_rows={
            "length": 2 + len(sku_rows) + 1,
            "width": 2 + len(sku_rows) + 2,
            "height": 2 + len(sku_rows) + 3,
            "weight": 2 + len(sku_rows) + 4,
            "volume": 2 + len(sku_rows) + 5,
        },
    )
    _fill_internal_sheet(ws, layout, sku_rows, matrix, sku_map, fnsku_map, editor_df, max_box)
    return _to_bytes(wb)


def _write_simple_df(df: pd.DataFrame, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append(list(row.values))
    return _to_bytes(wb)


def _shipping_num(val: Any) -> float | None:
    if val in ("", None):
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def export_shipping_sheet(
    shipping_df: pd.DataFrame,
    legend_rows: List[Dict[str, Any]] | None = None,
) -> bytes:
    """导出发货装箱单（亚马逊 FBA Packing List 样式，无模板）。"""
    from openpyxl.styles import Border, Side

    shipping_df = shipping_df.rename(
        columns={"单箱数量": "总数量", "单箱毛重(KG)": "总毛重(KG)"}
    )
    legend_rows = legend_rows or []
    cols = [c for c in SHIPPING_SHEET_COLUMNS if c in shipping_df.columns] or list(
        SHIPPING_SHEET_COLUMNS
    )
    ncols = len(cols)

    navy = PatternFill(start_color="2E3A4E", end_color="2E3A4E", fill_type="solid")
    legend_fill = PatternFill(start_color="DDE5F0", end_color="DDE5F0", fill_type="solid")
    total_fill = PatternFill(start_color="EEF2F8", end_color="EEF2F8", fill_type="solid")
    white_bold = Font(bold=True, color="FFFFFF")
    legend_font = Font(bold=True, color="1F3354", size=9)
    thin = Side(style="thin", color="B6C2D4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "发货装箱单"

    # Row 1 — 标题
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(1, 1, "亚马逊 FBA 发货装箱单 (Amazon FBA Packing List)")
    title.fill = navy
    title.font = Font(bold=True, color="FFFFFF", size=14)
    title.alignment = center
    for c in range(1, ncols + 1):
        ws.cell(1, c).fill = navy
    ws.row_dimensions[1].height = 30

    # Row 2 — 浅色字段带表头
    for c in range(1, ncols + 1):
        cell = ws.cell(2, c)
        cell.fill = legend_fill
    for key in SHIPPING_SHEET_LEGEND_COLUMNS:
        col = SHIPPING_SHEET_LEGEND_COLUMN[key]
        if col > ncols:
            continue
        label = f"{key}*"
        if key == "SKU简称":
            end_col = min(col + SHIPPING_SHEET_LEGEND_SKU_SPAN - 1, ncols)
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=end_col)
            cell = ws.cell(2, col, label)
        else:
            cell = ws.cell(2, col, label)
        cell.fill = legend_fill
        cell.font = legend_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    # Row 3+ — Step 1 数据行（若有）
    legend_data_font = Font(color="1F3354", size=9)
    for r_off, leg_row in enumerate(legend_rows):
        r = 3 + r_off
        for c in range(1, ncols + 1):
            ws.cell(r, c).fill = legend_fill
        for key in SHIPPING_SHEET_LEGEND_COLUMNS:
            col = SHIPPING_SHEET_LEGEND_COLUMN[key]
            if col > ncols:
                continue
            val = leg_row.get(key, "")
            if key in ("发货总数", "单箱装量") and val not in ("", None):
                out: Any = int(round(float(val)))
            elif key in ("箱长(cm)", "箱宽(cm)", "箱高(cm)") and val not in ("", None):
                out = round(float(val), 1)
            else:
                out = val
            if key == "SKU简称":
                end_col = min(col + SHIPPING_SHEET_LEGEND_SKU_SPAN - 1, ncols)
                ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=end_col)
                cell = ws.cell(r, col, out)
            else:
                cell = ws.cell(r, col, out)
            cell.fill = legend_fill
            cell.font = legend_data_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 18

    header_row = 3 + len(legend_rows)

    # 深蓝双语表头
    for ci, key in enumerate(cols, start=1):
        cn, en = SHIPPING_SHEET_HEADERS.get(key, (key, ""))
        cell = ws.cell(header_row, ci, f"{cn}\n({en})" if en else cn)
        cell.fill = navy
        cell.font = white_bold
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[header_row].height = 36

    # 数据行
    start = header_row + 1
    for r_off, (_, row) in enumerate(shipping_df.iterrows()):
        r = start + r_off
        for ci, key in enumerate(cols, start=1):
            val = row.get(key, "")
            if key == "总数量" and val not in ("", None):
                out: Any = int(round(float(val)))
            elif key == "总毛重(KG)" and val not in ("", None):
                out = round(float(val), 2)
            elif key == "单箱体积(CBM)" and val not in ("", None):
                out = round(float(val), 3)
            elif key in ("外箱长(CM)", "外箱宽(CM)", "外箱高(CM)") and val not in ("", None):
                out = round(float(val), 1)
            else:
                out = val
            cell = ws.cell(r, ci, out)
            cell.alignment = center
            cell.border = border

    # 合计行
    total_r = start + len(shipping_df)
    qty_sum = int(sum(v for v in (_shipping_num(x) for x in shipping_df.get("总数量", [])) if v))
    gw_sum = round(sum(v for v in (_shipping_num(x) for x in shipping_df.get("总毛重(KG)", [])) if v), 2)
    cbm_sum = round(sum(v for v in (_shipping_num(x) for x in shipping_df.get("单箱体积(CBM)", [])) if v), 3)

    col_index = {key: i + 1 for i, key in enumerate(cols)}
    for c in range(1, ncols + 1):
        cell = ws.cell(total_r, c)
        cell.fill = total_fill
        cell.border = border
        cell.font = Font(bold=True)
        cell.alignment = center
    label_col = col_index.get("中文品名", 1)
    ws.cell(total_r, label_col, "合计 (Total)")
    if "总数量" in col_index:
        ws.cell(total_r, col_index["总数量"], qty_sum)
    if "总毛重(KG)" in col_index:
        ws.cell(total_r, col_index["总毛重(KG)"], gw_sum)
    if "单箱体积(CBM)" in col_index:
        ws.cell(total_r, col_index["单箱体积(CBM)"], cbm_sum)

    # 列宽
    widths = {
        "票数": 10,
        "仓库代码": 10,
        "FBA货件编号": 22,
        "箱号": 12,
        "亚马逊条码(FNSKU)": 15,
        "中文品名": 22,
        "总数量": 11,
        "总毛重(KG)": 13,
        "外箱长(CM)": 11,
        "外箱宽(CM)": 11,
        "外箱高(CM)": 11,
        "单箱体积(CBM)": 14,
    }
    for ci, key in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(key, 14)

    # Step 1 浅色字段带：加宽 SKU 区（A-C）
    if ncols >= SHIPPING_SHEET_LEGEND_SKU_SPAN:
        sku_col_width = 22
        for ci in range(1, SHIPPING_SHEET_LEGEND_SKU_SPAN + 1):
            ws.column_dimensions[get_column_letter(ci)].width = sku_col_width

    return _to_bytes(wb)


def _load_fnsku_mapping(
    sku_rows: List[SkuRow],
    fnsku_overrides: Dict[str, str] | None = None,
) -> Dict[str, str]:
    """与发货表/工人单同源的 FNSKU 解析，避免 STA 与贴标单印出不同条码。"""
    from fba_tool.fnsku_resolve import build_fnsku_map

    return build_fnsku_map(
        [s.sku for s in sku_rows],
        step1_map=fnsku_overrides or {},
    )


def _load_amazon_sku_metadata() -> Dict[str, dict]:
    from fba_tool.master_db import amazon_sku_metadata, init_master_db

    init_master_db()
    base = load_json_cached("amazon_sku_metadata.json", {})
    if not isinstance(base, dict):
        base = {}
    master = amazon_sku_metadata()
    return {**base, **master}


def _to_bytes(wb: Workbook) -> bytes:
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
