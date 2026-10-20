from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Dict

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fba_tool.user_config import formula, label, param, render_formula

BOX_HEADER_RE = re.compile(r"(?:第|单箱|包装箱)\s*(\d+)")
FOOTER_LABELS_INTERNAL = {
    "length": ("长/cm", "长（cm）"),
    "width": ("宽/cm", "宽（cm）"),
    "height": ("高/cm", "高（cm）"),
    "weight": ("实重",),
    "volume": ("体积", "抛重"),
    "gross_lb": ("包装箱重量",),
    "width_in": ("包装箱宽度",),
    "length_in": ("包装箱长度",),
    "height_in": ("包装箱高度",),
    "total_qty": ("商品总数",),
}
FOOTER_LABELS_STA = {
    "box_type": ("包装箱类型", "包装箱名称"),
    "weight_lb": ("包装箱重量",),
    "width_in": ("包装箱宽度",),
    "length_in": ("包装箱长度",),
    "height_in": ("包装箱高度",),
}


@dataclass
class InternalLayout:
    sku_header_row: int = 1
    sku_start_row: int = 2
    first_box_col: int = 7
    footer_rows: Dict[str, int] = field(default_factory=dict)
    template_max_boxes: int = 130
    sku_columns: Dict[str, int] = field(default_factory=dict)

    @property
    def footer_start_row(self) -> int:
        if self.footer_rows:
            metric = [self.footer_rows[k] for k in ("length", "width", "height", "weight", "volume") if k in self.footer_rows]
            if metric:
                return min(metric)
            return min(v for k, v in self.footer_rows.items() if k != "total_qty")
        return self.sku_start_row + 6


@dataclass
class StaLayout:
    sku_header_row: int = 5
    sku_start_row: int = 6
    first_box_col: int = 13  # M — 包装箱 1 数量
    box_count_row: int = 3
    box_count_col: int = 13  # M3
    footer_rows: Dict[str, int] = field(default_factory=dict)
    template_max_boxes: int = 130
    prep_col: int = 7
    label_owner_col: int = 9
    expected_qty_col: int = 10
    packed_qty_col: int = 11

    @property
    def footer_start_row(self) -> int:
        if self.footer_rows:
            return min(self.footer_rows.values())
        return self.sku_start_row + 7

    @property
    def box_count_ref(self) -> str:
        return f"${get_column_letter(self.box_count_col)}${self.box_count_row}"


def detect_internal_layout(ws: Worksheet) -> InternalLayout:
    layout = InternalLayout()
    for r in range(1, min(ws.max_row, 15) + 1):
        for c in range(1, min(ws.max_column, 200) + 1):
            val = _s(ws.cell(r, c).value)
            m = BOX_HEADER_RE.search(val)
            if m and int(m.group(1)) == 1:
                layout.first_box_col = c
                layout.sku_header_row = r
                layout.sku_start_row = r + 1
                layout.sku_columns = _map_internal_headers(ws, r)
                break
        if layout.sku_columns:
            break

    if not layout.sku_columns:
        for r in range(1, 10):
            if _s(ws.cell(r, 1).value) in {"产品", "SKU", "SKU简称"}:
                layout.sku_header_row = r
                layout.sku_start_row = r + 1
                layout.sku_columns = _map_internal_headers(ws, r)
                break

    layout.template_max_boxes = _count_box_cols(ws, layout.sku_header_row, layout.first_box_col)

    scan_from = max(layout.sku_start_row + 1, layout.sku_header_row + 1)
    for r in range(scan_from, min(ws.max_row, 60) + 1):
        val = _s(ws.cell(r, 3).value)
        if not val:
            continue
        for key, labels in FOOTER_LABELS_INTERNAL.items():
            if key in layout.footer_rows:
                continue
            if _footer_label_match(val, labels):
                layout.footer_rows[key] = r
    return layout


def _footer_label_match(val: str, labels: tuple[str, ...]) -> bool:
    for lb in labels:
        if val == lb or val.startswith(lb):
            return True
    return False


def detect_sta_layout(ws: Worksheet) -> StaLayout:
    layout = StaLayout()
    for r in range(1, min(ws.max_row, 30) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            val = _s(ws.cell(r, c).value)
            if "包装箱总数" in val or "包装箱数量" in val:
                layout.box_count_row = r
                for cc in range(c + 1, min(c + 6, ws.max_column + 1)):
                    v = ws.cell(r, cc).value
                    if isinstance(v, (int, float)) and v > 0:
                        layout.box_count_col = cc
                        layout.template_max_boxes = int(v)
                        break
            if val == "SKU" and c == 1:
                layout.sku_header_row = r
                layout.sku_start_row = r + 1
                for cc in range(1, ws.max_column + 1):
                    hv = _s(ws.cell(r, cc).value)
                    m = BOX_HEADER_RE.search(hv)
                    if m and int(m.group(1)) == 1:
                        layout.first_box_col = cc
                layout.template_max_boxes = max(
                    layout.template_max_boxes,
                    _count_box_cols(ws, r, layout.first_box_col),
                )

    for r in range(1, min(ws.max_row, 80) + 1):
        val = _s(ws.cell(r, 1).value)
        for key, labels in FOOTER_LABELS_STA.items():
            if val and any(lb in val for lb in labels):
                layout.footer_rows[key] = r

    v = ws.cell(layout.box_count_row, layout.box_count_col).value
    if isinstance(v, (int, float)) and v > 0:
        layout.template_max_boxes = max(layout.template_max_boxes, int(v))
    return layout


def ensure_sku_capacity(ws: Worksheet, layout: InternalLayout | StaLayout, sku_count: int, sample_row: int) -> None:
    footer = layout.footer_start_row
    available = max(0, footer - layout.sku_start_row)
    if sku_count <= available:
        return
    need = sku_count - available
    ws.insert_rows(footer, need)
    layout.footer_rows = {k: v + need for k, v in layout.footer_rows.items()}
    for i in range(need):
        _copy_row_formulas(ws, sample_row, footer + i, layout.first_box_col)


def ensure_box_columns_sta(ws: Worksheet, layout: StaLayout, max_box: int) -> None:
    header = layout.sku_header_row
    ref = layout.box_count_ref
    fc = get_column_letter(layout.first_box_col)
    for i in range(1, max_box + 1):
        col = layout.first_box_col + i - 1
        safe_write(
            ws,
            header,
            col,
            render_formula(
                formula("sta", "box_header", f'=IF({ref}>={i},"单箱 {i} 装箱","")'),
                box_count_ref=ref,
                box_no=i,
                first_box_col=fc,
            ),
        )
    bt = layout.footer_rows.get("box_type")
    if bt:
        for i in range(1, max_box + 1):
            safe_write(
                ws,
                bt,
                layout.first_box_col + i - 1,
                render_formula(
                    formula("sta", "box_type", f'=IF({ref}>={i},"P1 - B{i}","")'),
                    box_count_ref=ref,
                    box_no=i,
                ),
            )


def detect_factory_warehouse_columns(ws: Worksheet, header_row: int = 2) -> dict[int, str]:
    """从工厂表模板第2行自动识别仓库代码及起始列。"""
    found: dict[int, str] = {}
    pat = re.compile(r"^[A-Z0-9]{4}$")
    for c in range(1, min(ws.max_column, 80) + 1):
        val = _s(ws.cell(header_row, c).value).upper()
        if pat.match(val):
            found[c] = val
    return found


def safe_write(ws: Worksheet, row: int, col: int, value) -> None:
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            try:
                ws.unmerge_cells(str(merged))
            except KeyError:
                pass
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                try:
                    ws.unmerge_cells(str(merged))
                except KeyError:
                    pass
                break
        cell = ws.cell(row=row, column=col)
    cell.value = value


def clear_range(ws: Worksheet, row_start: int, row_end: int, col_start: int, col_end: int) -> None:
    if row_end < row_start:
        return
    for r in range(row_start, row_end + 1):
        for c in range(col_start, col_end + 1):
            safe_write(ws, r, c, None)


def _map_internal_headers(ws: Worksheet, header_row: int) -> Dict[str, int]:
    mapping = {}
    names = {
        "product": ("产品", "sku简称"),
        "fnsku": ("fnsku",),
        "total": ("总数量", "总数", "发货总数"),
        "carton": ("单箱装量", "单箱"),
        "carton_count": ("箱数",),
        "check": ("校验",),
    }
    for c in range(1, min(ws.max_column, 20) + 1):
        val = _s(ws.cell(header_row, c).value).lower()
        if not val:
            continue
        if val == "fnsku":
            mapping["fnsku"] = c
        elif val in {"产品", "sku简称"}:
            mapping["product"] = c
        elif any(lb in val for lb in names["total"]):
            mapping["total"] = c
        elif any(lb in val for lb in names["carton"]):
            mapping["carton"] = c
        elif any(lb in val for lb in names["carton_count"]):
            mapping["carton_count"] = c
        elif any(lb in val for lb in names["check"]):
            mapping["check"] = c
    if "product" not in mapping:
        mapping["product"] = 1
    if "fnsku" not in mapping:
        mapping["fnsku"] = 2
    if "total" not in mapping:
        mapping["total"] = 3
    if "carton" not in mapping:
        mapping["carton"] = 4
    if "carton_count" not in mapping:
        mapping["carton_count"] = 5
    if "check" not in mapping:
        mapping["check"] = 6
    return mapping


def _copy_row_formulas(ws: Worksheet, src_row: int, dst_row: int, first_box_col: int) -> None:
    for c in range(1, first_box_col):
        cell = ws.cell(src_row, c)
        if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
            ws.cell(dst_row, c).value = cell.value.replace(str(src_row), str(dst_row))


def _count_box_cols(ws: Worksheet, header_row: int, first_col: int) -> int:
    count = 0
    for c in range(first_col, min(ws.max_column, first_col + 200) + 1):
        val = _s(ws.cell(header_row, c).value)
        raw = ws.cell(header_row, c).value
        if BOX_HEADER_RE.search(val) or (isinstance(raw, str) and "IF(" in raw.upper()):
            count += 1
        elif count > 0:
            break
    return max(count, 1)


def _s(v) -> str:
    if v is None:
        return ""
    return str(v).strip()
