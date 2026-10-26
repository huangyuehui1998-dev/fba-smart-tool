"""【备用通道】Listing 以领星 MCP erp_listing 拉取为主（lingxing_preview），本文件导入只在 MCP 不可用时用。
S2 产品库批量导入：领星 Listing 导出（Excel/CSV）、箱规模板。

流程统一为：读表 → 逐行校验（错误按行号）→ 预览差异（不写库）→ 确认后写入（同一文件再导一次应为 0 差异）。
写入走 product_fields.set_product_fields，记录字段来源；手填（manual）字段不被覆盖。
"""
from __future__ import annotations

import csv
import io
import unicodedata
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from fba_tool.lingxing_adapter import load_field_map, parse_price_reason
from fba_tool.master_db import batch_transaction, get_products_by_skus, sku_conflicts_for
from fba_tool.lingxing_adapter import normalize_currency
from fba_tool.product_fields import NEVER_LOCKED_FOR_LINGXING, SOURCE_LINGXING, SOURCE_MANUAL, set_product_fields
from fba_tool.shipment_parser import decode_csv_bytes
from fba_tool.sku_keys import sku_text

Row = Tuple[int, Dict[str, Any]]  # (Excel/CSV 行号, {表头: 值})


def _norm_header(h: Any) -> str:
    return unicodedata.normalize("NFKC", str(h or "")).strip().replace("*", "").replace(" ", "").lower()


def read_table(data: bytes, file_name: str, required_aliases: List[List[str]]) -> Tuple[List[str], List[Row]]:
    """读第一个工作表（xlsx）或 CSV；表头行 = 前 10 行里第一个包含全部必需列的行。"""
    name = (file_name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
        raw_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        raw_rows = list(csv.reader(io.StringIO(decode_csv_bytes(data))))
    else:
        raise ValueError("只支持 .xlsx 或 .csv 文件")
    want = [{_norm_header(a) for a in aliases} for aliases in required_aliases]
    for i, r in enumerate(raw_rows[:10]):
        heads = [_norm_header(c) for c in r]
        if all(any(h in w for h in heads) for w in want):
            header = [str(c or "").strip() for c in r]
            body: List[Row] = []
            for j, rr in enumerate(raw_rows[i + 1:], start=i + 2):
                if not any(str(c or "").strip() for c in rr):
                    continue
                body.append((j, {header[k]: rr[k] if k < len(rr) else None for k in range(len(header))}))
            return header, body
    names = "、".join(aliases[0] for aliases in required_aliases)
    raise ValueError(f"没有找到表头：前 10 行里需要同时有这些列：{names}")


def _pick(row: Dict[str, Any], aliases: List[str]) -> Any:
    wanted = {_norm_header(a) for a in aliases}
    for k, v in row.items():
        if _norm_header(k) in wanted and v not in (None, ""):
            return v
    return None


def _header_present(row: Dict[str, Any], aliases: List[str]) -> bool:
    wanted = {_norm_header(a) for a in aliases}
    return any(_norm_header(k) in wanted for k in row)


def _issue(row: int, field: str, value: Any, message: str) -> Dict[str, Any]:
    return {"row": row, "field": field, "value": "" if value is None else str(value), "message": message}


class ImportRowsError(ValueError):
    """S9-step2-fix2 N2：导入有非法行时报 422，带行号和原始值（errors）。"""

    def __init__(self, message: str, errors: List[Dict[str, Any]]) -> None:
        super().__init__(message)
        self.errors = errors


def error_rows_text(errors: List[Dict[str, Any]]) -> str:
    """S9-step4-fix3：出错的行号（去重、排序），如“第 2、3 行”。"""
    rows = sorted({int(e["row"]) for e in errors if str(e.get("row", "")).isdigit()})
    if not rows:
        return ""
    shown = "、".join(str(r) for r in rows[:20]) + ("…" if len(rows) > 20 else "")
    return f"第 {shown} 行"


def split_case_conflicts(records: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """S9-final N5：产品库里有多条只差大小写的记录、这一行又对不上原样写法的 → 冲突（带行号），不进差异、不写入。"""
    conflicts = sku_conflicts_for([r["sku"] for r in records])
    if not conflicts:
        return records, []
    ok: List[Dict[str, Any]] = []
    errs: List[Dict[str, Any]] = []
    for r in records:
        cands = conflicts.get(sku_text(r["sku"]))
        if cands:
            errs.append(_issue(r["row"], "sku", r["sku"],
                               f"第{r['row']}行 SKU“{r['sku']}”在产品库里有 {len(cands)} 条只差大小写的记录（{'、'.join(cands)}），"
                               "无法确定是哪一条，请先在产品库里合并或改名"))
        else:
            ok.append(r)
    return ok, errs


def _raise_conflicts(conflicts: List[Dict[str, Any]]) -> None:
    if conflicts:
        raise ImportRowsError(f"有 SKU 大小写冲突，整个文件未写入（出错：{error_rows_text(conflicts)}）：" + "；".join(e["message"] for e in conflicts[:5]), conflicts)


# ---------------- 领星 Listing 导出 ----------------

def _listing_aliases(field: str) -> List[str]:
    fm = load_field_map()
    return list(fm.get("excel_headers", {}).get(field, [])) + list(fm["listing"].get(field, []))


def parse_listing_rows(data: bytes, file_name: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    _h, body = read_table(data, file_name, [_listing_aliases("sku")])
    records: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    for row_no, row in body:
        sku = str(_pick(row, _listing_aliases("sku")) or "").strip()
        if not sku:
            errors.append(_issue(row_no, "sku", "", f"第{row_no}行缺少 MSKU"))
            continue
        values: Dict[str, Any] = {}
        raw_price = _pick(row, _listing_aliases("listing_price"))
        if raw_price not in (None, ""):
            price, reason = parse_price_reason(raw_price)  # N1–N3：与领星同步同一套解析
            if price is None:
                errors.append(_issue(row_no, "listing_price", raw_price,
                                     f"第{row_no}行价格“{raw_price}”无效：{reason or '不是大于 0 的数字'}，这一行已跳过"))
                continue
            currency_aliases = _listing_aliases("listing_price_currency")
            cur = _pick(row, currency_aliases)
            try:
                code = normalize_currency(cur)
            except ValueError as exc:
                errors.append(_issue(row_no, "listing_price_currency", cur, f"第{row_no}行{exc}"))
                continue
            # 币种列在但单元格为空：与 import_retail_rows 一样跳过，不默认 USD。
            # 整张表没有币种列的旧 Listing 导入（只有 MSKU/价格）仍记 USD，那不是零售站点表。
            if _header_present(row, currency_aliases) and not code:
                errors.append(_issue(row_no, "listing_price_currency", cur, f"第{row_no}行币种为空，未写入"))
                continue
            # 这张表没有国家列，不能判断是不是美国站，也不能在这里换汇。非 USD 不写入售价。
            if code and code != "USD":
                errors.append(_issue(row_no, "listing_price_currency", code, f"第{row_no}行价格币种是 {code}，不是 USD，未写入"))
                continue
            values["listing_price"] = price
            values["listing_price_basis"] = "listing_price"
            values["listing_price_currency"] = code or "USD"
        for f in ("fnsku", "asin", "local_name"):
            v = _pick(row, _listing_aliases(f))
            if v not in (None, ""):
                values[f] = str(v).strip()
        from fba_tool.policy import map_category_name

        mapped_category = map_category_name(
            _pick(row, ["品类", "分类", "category", "category_name", "Category"])
        )
        if mapped_category:
            values["category"] = mapped_category
        records.append({"row": row_no, "sku": sku, "values": values})
    return records, errors


def _diff(records: List[Dict[str, Any]], source: str, priority_over: Tuple[str, ...] = ()) -> Dict[str, Any]:
    """priority_over：这些来源的现有值也不覆盖（例如箱规导入时领星值优先）。"""
    masters = get_products_by_skus([r["sku"] for r in records])
    items: List[Dict[str, Any]] = []
    field_summary = {"new": 0, "changed": 0, "unchanged": 0, "locked": 0}
    sku_status: Dict[str, set] = {}
    for rec in records:
        m = masters.get(rec["sku"])
        sources = (m or {}).get("field_sources") or {}
        for field, new in rec["values"].items():
            old = (m or {}).get(field)
            src = sources.get(field)
            never_locked = source == SOURCE_LINGXING and field in NEVER_LOCKED_FOR_LINGXING
            if not never_locked and (src == SOURCE_MANUAL and source != SOURCE_MANUAL or src in priority_over):
                status = "locked"
            elif m is None or old in (None, ""):
                status = "new"
            elif _same(old, new):
                status = "unchanged"
            else:
                status = "changed"
            field_summary[status] += 1
            sku_status.setdefault(rec["sku"], set()).add(status)
            if m is None:
                sku_status[rec["sku"]].add("_missing")
            items.append({"row": rec["row"], "sku": rec["sku"], "field": field, "old": old, "new": new,
                          "status": status, "old_source": src or ""})
    # 按 SKU 汇总：new=库里没有；changed=有字段新增或变化；locked=有字段被锁定（手填/优先来源）
    summary = {"new": 0, "changed": 0, "unchanged": 0, "locked": 0, "skus": len(sku_status)}
    for st in sku_status.values():
        if "_missing" in st:
            summary["new"] += 1
        elif st & {"new", "changed"}:
            summary["changed"] += 1
        else:
            summary["unchanged"] += 1
        if "locked" in st:
            summary["locked"] += 1
    return {"items": items, "summary": summary, "field_summary": field_summary}


def _same(a: Any, b: Any) -> bool:
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def _apply(records: List[Dict[str, Any]], diff: Dict[str, Any], source: str) -> Dict[str, Any]:
    todo: Dict[str, Dict[str, Any]] = {}
    for it in diff["items"]:
        if it["status"] in ("new", "changed"):
            todo.setdefault(it["sku"], {})[it["field"]] = it["new"]
    written = 0
    with batch_transaction():  # S9-final N5：一个事务里写完；中途任何异常整体回滚，不会只写一半
        for sku, vals in todo.items():
            set_product_fields(sku, vals, source)
            written += 1
    return {"written": written, "fields": sum(len(v) for v in todo.values()), "summary": diff["summary"]}


def _table_headers(data: bytes, file_name: str) -> List[str]:
    name = (file_name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
        first = next(ws.iter_rows(values_only=True), None) or ()
        return [str(cell or "").strip() for cell in first]
    text = decode_csv_bytes(data)
    return next(csv.reader(io.StringIO(text)), [])


def _is_price_chain(data: bytes, file_name: str) -> bool:
    headers = set(_table_headers(data, file_name))
    return "国家" in headers and "币种" in headers and ("价格" in headers or "售价(总价)" in headers)


def preview_listing_import(data: bytes, file_name: str) -> Dict[str, Any]:
    if _is_price_chain(data, file_name):
        from fba_tool.declared_value import parse_listing_xlsx, parse_listing_table, import_retail_rows
        from fba_tool.master_db import get_products_by_skus

        if (file_name or "").lower().endswith((".xlsx", ".xlsm")):
            rows = parse_listing_xlsx(data)
        else:
            text = decode_csv_bytes(data)
            table = list(csv.reader(io.StringIO(text)))
            rows = parse_listing_table(table[0], table[1:])
        products = get_products_by_skus([str(row.get("msku") or "") for row in rows])
        return import_retail_rows(rows, products, source="lingxing_xlsx")
    records, errors = parse_listing_rows(data, file_name)
    records, conflicts = split_case_conflicts(records)
    d = _diff(records, SOURCE_LINGXING)
    return {"errors": errors + conflicts, "conflict_rows": sorted({e["row"] for e in conflicts}), **d}


def apply_listing_import(data: bytes, file_name: str) -> Dict[str, Any]:
    if _is_price_chain(data, file_name):
        from fba_tool.declared_value import parse_listing_table, parse_listing_xlsx, persist_retail_import

        if (file_name or "").lower().endswith((".xlsx", ".xlsm")):
            rows = parse_listing_xlsx(data)
        else:
            text = decode_csv_bytes(data)
            table = list(csv.reader(io.StringIO(text)))
            rows = parse_listing_table(table[0], table[1:])
        return persist_retail_import(rows, source="lingxing_xlsx")
    records, errors = parse_listing_rows(data, file_name)
    records, conflicts = split_case_conflicts(records)
    _raise_conflicts(conflicts)
    if errors and not records:
        raise ImportRowsError(f"没有可写入的行，整个文件未写入（出错：{error_rows_text(errors)}）：" + "；".join(e["message"] for e in errors[:5]), errors)
    out = _apply(records, _diff(records, SOURCE_LINGXING), SOURCE_LINGXING)
    out["skipped"] = [dict(e) for e in errors]
    out["skipped_rows"] = sorted({int(e["row"]) for e in errors if str(e.get("row", "")).isdigit()})
    return out


# ---------------- 箱规批量导入（v8 E3） ----------------

CARTON_TEMPLATE_HEADERS = ["SKU", "箱长(cm)", "箱宽(cm)", "箱高(cm)", "单箱数量", "单箱重(kg)"]
_CARTON_FIELDS = [
    ("SKU", None, ["SKU", "SKU简称", "MSKU", "sku"]),
    ("箱长(cm)", "carton_length", ["箱长(cm)", "箱长", "外箱长(cm)"]),
    ("箱宽(cm)", "carton_width", ["箱宽(cm)", "箱宽", "外箱宽(cm)"]),
    ("箱高(cm)", "carton_height", ["箱高(cm)", "箱高", "外箱高(cm)"]),
    ("单箱数量", "qty_per_box", ["单箱数量", "单箱装量", "每箱数量", "装箱数"]),
    ("单箱重(kg)", "carton_gross_weight", ["单箱重(kg)", "单箱重", "单箱毛重(kg)", "单箱毛重"]),
]


def build_carton_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "箱规"
    ws.append(CARTON_TEMPLATE_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEF", (36, 11, 11, 11, 10, 12)):
        ws.column_dimensions[col].width = w
    note = wb.create_sheet("说明")
    for line in (
        "每行一个 SKU，6 列都必填。",
        "箱长/宽/高单位 cm，单箱重单位 kg，为整箱（含纸箱）称重的毛重；单箱数量为整数。",
        "导入时先显示差异，确认后才写入产品库；同一文件导入两次，第二次应为 0 差异。",
        "产品库里来自领星的值优先，不会被本表覆盖；手填（界面里改过）的值也不会被覆盖。",
    ):
        note.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _carton_number(raw: Any, label: str, row_no: int, errors: List[Dict[str, Any]], integer: bool) -> Optional[float]:
    if raw in (None, ""):
        errors.append(_issue(row_no, label, "", f"第{row_no}行缺少「{label}」"))
        return None
    text = unicodedata.normalize("NFKC", str(raw)).strip().replace(",", "")
    try:
        v = float(text)
    except ValueError:
        errors.append(_issue(row_no, label, raw, f"第{row_no}行「{label}」的值“{raw}”不是有效数字"))
        return None
    if v <= 0:
        errors.append(_issue(row_no, label, raw, f"第{row_no}行「{label}」必须大于 0，现在是“{raw}”"))
        return None
    if integer and v != int(v):
        errors.append(_issue(row_no, label, raw, f"第{row_no}行「{label}」必须是整数，现在是“{raw}”"))
        return None
    return float(int(v)) if integer else v


def parse_carton_rows(data: bytes, file_name: str):
    try:
        _h, body = read_table(data, file_name, [aliases for _l, _f, aliases in _CARTON_FIELDS])
    except ValueError:
        # 找出缺哪几列，给明确提示
        _h, _b = None, None
        missing = []
        try:
            ws_rows = read_table(data, file_name, [_CARTON_FIELDS[0][2]])
            heads = {_norm_header(h) for h in ws_rows[0]}
            missing = [label for label, _f, aliases in _CARTON_FIELDS if not heads & {_norm_header(a) for a in aliases}]
        except ValueError:
            missing = ["SKU"]
        raise ValueError("箱规表缺少列：" + "、".join(missing) + "（请用 箱规导入模板.xlsx）")
    records: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    first_row = body[0][0] if body else 2
    for row_no, row in body:
        sku = str(_pick(row, _CARTON_FIELDS[0][2]) or "").strip()
        before = len(errors)
        if not sku:
            errors.append(_issue(row_no, "SKU", "", f"第{row_no}行缺少 SKU"))
        values: Dict[str, Any] = {}
        for label, field, aliases in _CARTON_FIELDS[1:]:
            v = _carton_number(_pick(row, aliases), label, row_no, errors, integer=(field == "qty_per_box"))
            if v is not None:
                values[field] = v
        if len(errors) == before:
            records.append({"row": row_no, "sku": sku, "values": values})
    return records, errors, first_row


def preview_carton_import(data: bytes, file_name: str) -> Dict[str, Any]:
    from fba_tool.product_fields import SOURCE_CARTON_IMPORT

    records, errors, first_row = parse_carton_rows(data, file_name)
    records, conflicts = split_case_conflicts(records)
    d = _diff(records, SOURCE_CARTON_IMPORT, priority_over=(SOURCE_LINGXING,))
    return {"errors": errors + conflicts, "conflict_rows": sorted({e["row"] for e in conflicts}),
            "first_data_row": first_row, **d}


def apply_carton_import(data: bytes, file_name: str) -> Dict[str, Any]:
    from fba_tool.product_fields import SOURCE_CARTON_IMPORT

    records, errors, _first = parse_carton_rows(data, file_name)
    records, conflicts = split_case_conflicts(records)
    _raise_conflicts(conflicts)
    if errors:  # S9-step4-fix3：带行号（errors 列表 + 文字里的行号），接口 422
        raise ImportRowsError(f"箱规表有错误，整个文件未写入（出错：{error_rows_text(errors)}）：" + "；".join(e["message"] for e in errors[:5]), errors)
    d = _diff(records, SOURCE_CARTON_IMPORT, priority_over=(SOURCE_LINGXING,))
    return _apply(records, d, SOURCE_CARTON_IMPORT)
