"""把 Univer 快照里改过的值写回原始 xlsx。

只改不相等的普通单元格。公式格保持原公式；快照把公式换成纯值时拒绝。
合并单元格的从属格拒绝写入。不改数据验证、图片、条件格式和隐藏表。
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class ApplyRejected(Exception):
    def __init__(self, code: str, cells: List[str], message: str) -> None:
        super().__init__(message)
        self.code = code
        self.cells = cells
        self.message = message


def apply_cell_writes(data: bytes, writes: List[Dict[str, Any]], sheet_name: Optional[str] = None) -> bytes:
    """把导出字段写回原始 xlsx。不删工作表、图片、数据验证或隐藏表。

    合并从属格和公式格改成纯值时拒绝，与 apply_snapshot_values 同一套规则。
    自动填写要直接写主格，不要把从属格坐标交给这里。
    """
    wb = openpyxl.load_workbook(io.BytesIO(data))
    formula_cells: List[str] = []
    slave_cells: List[str] = []
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        slaves = _slave_map(ws)
        for item in writes or []:
            ref = str(item.get("cell") or "").replace("$", "").strip().upper()
            if not ref or item.get("value") in (None, ""):
                continue
            from openpyxl.utils.cell import coordinate_to_tuple

            row, col = coordinate_to_tuple(ref)
            value = coerce_sheet_value(
                ws,
                row,
                col,
                item.get("value"),
                str(item.get("field_key") or item.get("key") or ""),
            )
            if (row, col) in slaves:
                master = ws.cell(*slaves[(row, col)])
                if not _values_equal(master.value, value):
                    slave_cells.append(ref)
                continue
            cell = ws.cell(row, col)
            if _is_formula(cell):
                if _norm_formula(value) != _norm_formula(cell.value):
                    formula_cells.append(ref)
                continue
            if _values_equal(cell.value, value) and _stored_type_matches(cell.value, value):
                continue
            _write(cell, {"v": value})
        _raise_conflicts(formula_cells, slave_cells)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def apply_snapshot_values(data: bytes, snapshot: Dict[str, Any]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    formula_cells: List[str] = []
    slave_cells: List[str] = []
    try:
        by_name = {ws.title: ws for ws in wb.worksheets}
        for sheet in _snapshot_sheets(snapshot):
            ws = by_name.get(sheet.get("name"))
            if ws is None:
                continue
            slaves = _slave_map(ws)
            for row, col, payload in _iter_cells(sheet):
                ref = f"{get_column_letter(col)}{row}"
                if (row, col) in slaves:
                    master = ws.cell(*slaves[(row, col)])
                    if not _values_equal(master.value, payload.get("v")):
                        slave_cells.append(ref)
                    continue
                cell = ws.cell(row, col)
                if _is_formula(cell):
                    snap_f = payload.get("f")
                    if snap_f and _norm_formula(snap_f) == _norm_formula(cell.value):
                        continue
                    if not snap_f and "v" in payload:
                        formula_cells.append(ref)
                    continue
                if "v" not in payload and "f" not in payload:
                    continue
                value = coerce_sheet_value(ws, row, col, payload.get("v"))
                if _values_equal(cell.value, value) or _values_equal(cell.value, payload.get("v")):
                    continue
                _write(cell, {**payload, "v": value})
        _raise_conflicts(formula_cells, slave_cells)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def _raise_conflicts(formula_cells: List[str], slave_cells: List[str]) -> None:
    if not formula_cells and not slave_cells:
        return
    code = "FORMULA_OVERWRITE" if formula_cells else "MERGED_SLAVE"
    cells = formula_cells or slave_cells
    label = "公式单元格不能改成纯值" if code == "FORMULA_OVERWRITE" else "不能写入合并单元格的从属格"
    raise ApplyRejected(code, cells, f"{label}：{', '.join(cells[:30])}")


def _snapshot_sheets(snapshot: Dict[str, Any]) -> List[Dict[str, Any]]:
    sheets = snapshot.get("sheets") or {}
    order = snapshot.get("sheetOrder") or list(sheets.keys())
    return [sheets[key] for key in order if key in sheets]


def _iter_cells(sheet: Dict[str, Any]):
    cell_data = sheet.get("cellData") or {}
    for row_key, cols in cell_data.items():
        if not isinstance(cols, dict):
            continue
        row = int(row_key) + 1
        for col_key, payload in cols.items():
            if not isinstance(payload, dict):
                continue
            yield row, int(col_key) + 1, payload


def _slave_map(ws: Worksheet) -> Dict[tuple, tuple]:
    slaves: Dict[tuple, tuple] = {}
    for mr in ws.merged_cells.ranges:
        if mr.max_row > 20000 or mr.max_col > 512:
            continue
        origin = (mr.min_row, mr.min_col)
        for row in range(mr.min_row, mr.max_row + 1):
            for col in range(mr.min_col, mr.max_col + 1):
                if (row, col) != origin:
                    slaves[(row, col)] = origin
    return slaves


def _is_formula(cell: Any) -> bool:
    if isinstance(cell, MergedCell):
        return False
    if getattr(cell, "data_type", None) == "f":
        return True
    return isinstance(cell.value, str) and cell.value.startswith("=")


def _norm_formula(value: Any) -> str:
    text = str(value or "").strip()
    if text.startswith("="):
        text = text[1:]
    return text


def _values_equal(orig: Any, snap: Any) -> bool:
    if isinstance(orig, bool) or isinstance(snap, bool):
        return bool(orig) is bool(snap) and type(orig) is type(snap) or (isinstance(orig, bool) and snap in (orig, str(orig)))
    if orig is None:
        return snap in (None, "")
    if type(orig).__name__ in {"CellRichText", "ArrayFormula"}:
        return str(orig) == "" or str(orig) == str(snap) or snap in (None, "")
    if isinstance(orig, datetime):
        text = orig.isoformat()
        snap_text = "" if snap is None else str(snap)
        return snap_text in {text, text.replace("T", " "), orig.strftime("%Y-%m-%d"), orig.strftime("%Y-%m-%d %H:%M:%S")}
    if isinstance(orig, date):
        return str(snap) in {orig.isoformat(), orig.strftime("%Y-%m-%d")}
    if isinstance(orig, (int, float)) and not isinstance(orig, bool):
        try:
            if snap in (None, ""):
                return False
            return float(orig) == float(snap)
        except (TypeError, ValueError):
            return False
    return str(orig) == str(snap)


def _keep_numeric(cell: Any, value: Any) -> Any:
    """已经是数值的格子，快照里的文本数字仍写回数值，避免把 4 盖成 \"4.0\"。"""
    if isinstance(value, bool) or not isinstance(value, str):
        return value
    current = cell.value
    if isinstance(current, bool) or not isinstance(current, (int, float)):
        return value
    from fba_tool.invoice_rules import to_number

    num = to_number(value.strip())
    if isinstance(num, bool) or not isinstance(num, (int, float)):
        return value
    if isinstance(current, int) and float(num) == int(num):
        return int(num)
    return int(num) if isinstance(num, int) else float(num)


def _cell_text(ws: Worksheet, row: int, col: int) -> str:
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        return ""
    value = cell.value
    if value in (None, ""):
        return ""
    if isinstance(value, str) and value.startswith("="):
        return ""
    return str(value).strip()


def _plain_number(text: str) -> bool:
    token = text.replace(",", "").strip()
    if not token:
        return False
    try:
        float(token)
    except ValueError:
        return False
    return True


def column_context_label(ws: Worksheet, row: int, col: int) -> str:
    """这一格正上方最近的表头；表头对不上时再看同一行左侧的字段名。

    只跳过空格和纯数字，避免把更上面另一块区域的「预报件数」错当成箱序号的表头。
    """
    from fba_tool.invoice_rules import header_number_kind

    header = ""
    for above in range(row - 1, max(0, row - 45), -1):
        text = _cell_text(ws, above, col)
        if not text or _plain_number(text):
            continue
        header = text
        break
    if header_number_kind(header):
        return header
    if col > 1:
        left = _cell_text(ws, row, col - 1)
        if left and not _plain_number(left) and header_number_kind(left):
            return left
    return ""


def coerce_sheet_value(ws: Worksheet, row: int, col: int, value: Any, field_key: str = "") -> Any:
    from fba_tool.invoice_rules import coerce_export_number

    return coerce_export_number(value, field_key=field_key, header=column_context_label(ws, row, col))


def _stored_type_matches(current: Any, value: Any) -> bool:
    """数量要落成 int，价格落成 float。文本格里的 "4.0" 不算已经写对。"""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return True
    if isinstance(current, bool) or not isinstance(current, (int, float)):
        return False
    if type(value) is int:
        return type(current) is int
    return type(current) in {int, float}


def _write(cell: Any, payload: Dict[str, Any]) -> None:
    value = _keep_numeric(cell, payload.get("v"))
    kind = payload.get("t")
    if kind == 3 or isinstance(value, bool):
        cell.value = bool(value)
        return
    if kind == 2 and isinstance(value, str):
        from fba_tool.invoice_rules import to_number

        num = to_number(value.strip())
        if isinstance(num, (int, float)) and not isinstance(num, bool):
            value = int(num) if isinstance(num, float) and num == int(num) else num
    if kind == 2 or isinstance(value, (int, float)):
        cell.value = value
        return
    cell.value = None if value is None else value
