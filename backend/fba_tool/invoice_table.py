"""货代发票表格模板：表头行 + 按 SKU 逐行填写。"""
from __future__ import annotations

import io
import csv
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl.utils import get_column_letter

from fba_tool.product_catalog import ProductCatalog, ProductRecord
from fba_tool.sku_keys import lookup_sku, sku_key as _norm_key, sku_text
from fba_tool.shipment_parser import CsvShipment
from fba_tool.master_db import get_product_by_sku, invoice_hints_from_master
from fba_tool.shipment_parser import decode_csv_bytes

TABLE_HEADER_SCAN_MAX_ROW = 40
INVOICE_TABLE_DATA_CLEAR_ROWS = 100

# 已废弃：请使用 forwarder_mapping.json 中的 global_table_column_aliases
TABLE_COLUMN_ALIASES: Dict[str, List[str]] = {}

PCS_IN_SKU_RE = re.compile(r"(\d+)\s*PCS", re.I)


@dataclass
class InvoiceTableLayout:
    header_row: int
    first_data_row: int
    columns: Dict[str, int]  # field_key -> 1-based column index


@dataclass
class InvoiceLineItem:
    values: Dict[str, Any] = field(default_factory=dict)
    sources: Dict[str, str] = field(default_factory=dict)


def sanitize_header_text(text: Any) -> str:
    """表头清洗：换行→空，去所有空白，去掉必填星号，casefold 便于比对。"""
    s = str(text or "").replace("\r", " ").replace("\n", " ")
    s = re.sub(r"\s+", "", s)
    s = s.rstrip("*?")
    return s.casefold()


def merge_table_column_aliases(
    cfg: dict | None = None,
    forwarder_id: str | None = None,
) -> Dict[str, List[str]]:
    from fba_tool.forwarder_mapping import merge_table_column_aliases as _merge_mapping

    return _merge_mapping(forwarder_id=forwarder_id, extra_cfg=cfg)


def header_matches_alias(header_raw: str, alias: str) -> bool:
    h = sanitize_header_text(header_raw)
    a = sanitize_header_text(alias)
    if not h or not a:
        return False
    if a == h:
        return True
    # 「PO」不能因为 purpose 里含有 po 就抢走用途列。短英文别名按单词比，不按子串。
    if a.isascii() and len(a) <= 3:
        tokens = re.findall(r"[a-z0-9]+", str(header_raw or "").casefold())
        return a in tokens
    return a in h or h in a


def _is_sku_header(text: str) -> bool:
    k = sanitize_header_text(text)
    return k == "sku"


def _is_shipment_id_header(text: str) -> bool:
    k = sanitize_header_text(text)
    if not k:
        return False
    if "fba单号" in k or "货件编号" in k:
        return True
    compact = k.replace("-", "").replace("_", "")
    return "shipmentid" in compact


def _is_box_seq_header(text: str) -> bool:
    k = sanitize_header_text(text)
    return k in ("箱序号", "箱号段", "箱号范围") or "箱序号" in k


def _is_fba_id_header(text: str) -> bool:
    k = sanitize_header_text(text)
    if not k:
        return False
    if k == "fbaid":
        return True
    return _is_shipment_id_header(text)


def _is_detail_companion_header(text: str) -> bool:
    """明细区同行：箱号或品名/型号/材质。不要求这一行再写货件号。"""
    k = sanitize_header_text(text)
    if not k:
        return False
    if _is_box_seq_header(text) or k.startswith("箱号") or "boxno" in k:
        return True
    return any(token in k for token in ("英文品名", "中文品名", "物品英文名", "物品中文名", "型号", "材质", "用途"))


def _row_has_table_markers(cell_texts: List[str]) -> bool:
    has_sku = any(_is_sku_header(t) for t in cell_texts if t)
    has_shipment = any(_is_shipment_id_header(t) for t in cell_texts if t)
    if has_sku and has_shipment:
        return True
    has_fba = any(_is_fba_id_header(t) for t in cell_texts if t)
    has_box_seq = any(_is_box_seq_header(t) for t in cell_texts if t)
    if has_fba and has_box_seq:
        return True
    # SEND 第 14 行是「箱号 / SKU / 品名 / 型号」，货件号在上面的合并主格，不在这一行。
    return has_sku and any(_is_detail_companion_header(t) for t in cell_texts if t)


def _merge_at(ws, row: int, col: int):
    for merged in ws.merged_cells.ranges:
        if merged.max_row > 20000 or merged.max_col > 512:
            continue
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return merged
    return None


def value_cell_right_of(ws, row: int, col: int) -> tuple[int, int]:
    """标签右侧的值写在合并主格上。

    A4:B4 是「FBA Shipment ID」，紧挨着的 B4 是从属格，值的主格是 C4（C4:H4）。
    标签本身没有合并、右侧已有别的字时，仍往后空一格，避免盖住旁边的标签。
    """
    label_merge = _merge_at(ws, row, col)
    next_col = label_merge.max_col + 1 if label_merge else col + 1
    host = _merge_at(ws, row, next_col)
    if host and (host.min_row, host.min_col) != (row, next_col):
        origin = (host.min_row, host.min_col)
        if ws.cell(*origin).value in (None, ""):
            return origin
        after = host.max_col + 1
        host_after = _merge_at(ws, row, after)
        if host_after:
            return host_after.min_row, host_after.min_col
        return row, after
    if host and (host.min_row, host.min_col) == (row, next_col):
        return host.min_row, host.min_col
    right = ws.cell(row, next_col)
    if label_merge is None and right.value not in (None, ""):
        return row, next_col + 1
    return row, next_col


def _map_columns_from_headers(
    headers: Dict[int, str],
    column_aliases: Dict[str, List[str]],
) -> Dict[str, int]:
    columns: Dict[str, int] = {}
    for col_idx, header in headers.items():
        for field_key, aliases in column_aliases.items():
            if field_key in columns:
                continue
            if any(header_matches_alias(header, alias) for alias in aliases):
                columns[field_key] = col_idx
                break
    return columns


def detect_invoice_table_layout(
    wb,
    max_scan_row: int | None = None,
    cfg: dict | None = None,
    forwarder_id: str | None = None,
) -> Optional[InvoiceTableLayout]:
    """第 1~40 行全域扫描：同行出现 SKU + FBA单号/Shipment ID 即锁定表头行。"""
    ws = wb.active
    scan_to = max_scan_row or int((cfg or {}).get("table_header_scan_max_row", TABLE_HEADER_SCAN_MAX_ROW))
    scan_to = min(scan_to, ws.max_row or scan_to)
    column_aliases = merge_table_column_aliases(cfg, forwarder_id=forwarder_id)

    for row_idx in range(1, scan_to + 1):
        headers: Dict[int, str] = {}
        raw_texts: List[str] = []
        for col_idx in range(1, min(ws.max_column or 1, 60) + 1):
            raw = str(ws.cell(row_idx, col_idx).value or "").strip()
            if raw:
                headers[col_idx] = raw
                raw_texts.append(raw)
        if not headers:
            continue
        if not _row_has_table_markers(raw_texts):
            continue

        columns = _map_columns_from_headers(headers, column_aliases)
        if "sku" not in columns and not ("fba_id" in columns and "box_range" in columns):
            continue

        return InvoiceTableLayout(
            header_row=row_idx,
            first_data_row=row_idx + 1,
            columns=columns,
        )
    return None


def extract_csv_sku_meta(csv_bytes: bytes) -> Dict[str, Dict[str, str]]:
    text = decode_csv_bytes(csv_bytes)
    rows = list(csv.reader(io.StringIO(text)))
    header_idx = -1
    for i, row in enumerate(rows):
        if row and _norm_key(row[0]) == "sku":
            header_idx = i
            break
    if header_idx < 0:
        return {}

    header = [str(c).strip() for c in rows[header_idx]]
    col = {name: idx for idx, name in enumerate(header)}

    def pick(row: List[str], *names: str) -> str:
        for name in names:
            if name in col and col[name] < len(row):
                val = str(row[col[name]]).strip()
                if val:
                    return val
        return ""

    meta: Dict[str, Dict[str, str]] = {}
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        sku = str(row[0]).strip()
        if not sku or re.search(r"(length|width|height|weight|包装)", sku, flags=re.I):
            break
        key = sku_text(sku)  # S9-step4-fix3：原样写法作键，查找一律 lookup_sku（原样优先、唯一规范命中、多个不任选）
        entry = meta[key] = {
            "title": pick(row, "商品名称", "Product Name", "Title"),
            "asin": pick(row, "ASIN"),
            "fnsku": pick(row, "FNSKU"),
        }
        master = get_product_by_sku(sku)
        if master:
            if not entry["asin"] and master.get("asin"):
                entry["asin"] = master["asin"]
            if not entry["fnsku"] and master.get("fnsku"):
                entry["fnsku"] = master["fnsku"]
    return meta


def overlay_csv_asin(products: Dict[str, Dict[str, Any]], csv_bytes: bytes | None) -> Dict[str, Dict[str, Any]]:
    """产品库缺 ASIN 时，用这份 CSV 的 ASIN 填到副本上。不写回产品库。"""
    meta = extract_csv_sku_meta(csv_bytes or b"")
    by_fnsku: Dict[str, str] = {}
    for entry in meta.values():
        fnsku = str(entry.get("fnsku") or "").strip().upper()
        asin = str(entry.get("asin") or "").strip()
        if fnsku and asin:
            by_fnsku[fnsku] = asin
    filled: Dict[str, Dict[str, Any]] = {}
    for key, product in (products or {}).items():
        item = dict(product or {})
        if str(item.get("asin") or "").strip():
            filled[key] = item
            continue
        fnsku = str(item.get("fnsku") or key or "").strip().upper()
        asin = by_fnsku.get(fnsku, "")
        if not asin:
            hit = lookup_sku(meta, item.get("sku"))
            if isinstance(hit, dict):
                asin = str(hit.get("asin") or "").strip()
        if asin:
            item["asin"] = asin
        filled[key] = item
    return filled


def _pcs_from_sku(sku: str) -> Optional[int]:
    m = PCS_IN_SKU_RE.search(sku)
    if m:
        return int(m.group(1))
    return None


def _box_range(boxes: List[int]) -> str:
    if not boxes:
        return ""
    boxes = sorted(set(boxes))
    return f"{boxes[0]}~{boxes[-1]}"


def _source_for(field: str, from_csv: bool, from_catalog: bool, from_ai: bool) -> str:
    if from_csv:
        return "green"
    if from_catalog:
        return "blue"
    if from_ai:
        return "yellow"
    if field in ("reference_id", "sku_image"):
        return "red"
    return "red"


def _declared_unit_price(purchase_price: Optional[float]) -> float:
    if not purchase_price:
        return 0.0
    return round(float(purchase_price), 2)


def _catalog_record(catalog: Optional[ProductCatalog], sku: str) -> Optional[ProductRecord]:
    if not catalog:
        return None
    return catalog.lookup(sku)


def _combine_material_purpose(material: str, purpose: str) -> str:
    mat = str(material or "").strip()
    pur = str(purpose or "").strip()
    if mat and pur and pur not in mat:
        return f"{mat}/{pur}"
    return mat or pur


def build_invoice_line_items(
    ship: CsvShipment,
    csv_bytes: bytes,
    catalog: Optional[ProductCatalog],
    customs_by_sku: Optional[Dict[str, Dict[str, Any]]] = None,
    customs_ai_keys: Optional[Dict[str, set]] = None,
    reference_id: str = "",
) -> List[InvoiceLineItem]:
    """按 SKU 聚合货件，四层降级：CSV 🟩 → 产品库 🟦 → AI 🟨 → 留空 🟥。"""
    if ship.box_items.empty:
        return []

    sku_meta = extract_csv_sku_meta(csv_bytes)
    customs_by_sku = customs_by_sku or {}
    customs_ai_keys = customs_ai_keys or {}
    items: List[InvoiceLineItem] = []

    grouped = ship.box_items.groupby("SKU", sort=True)
    for sku, grp in grouped:
        sku_str = str(sku)
        boxes = sorted(int(b) for b in grp["箱号"].tolist())
        total_qty = int(grp["数量"].sum())
        box_count = len(boxes)

        spec_rows = (
            ship.box_specs[ship.box_specs["箱号"].isin(boxes)]
            if not ship.box_specs.empty
            else pd.DataFrame()
        )
        length = width = height = weight = 0.0
        if not spec_rows.empty:
            length = float(spec_rows["长(cm)"].median()) if "长(cm)" in spec_rows else 0
            width = float(spec_rows["宽(cm)"].median()) if "宽(cm)" in spec_rows else 0
            height = float(spec_rows["高(cm)"].median()) if "高(cm)" in spec_rows else 0
            weight = float(spec_rows["实重(kg)"].median()) if "实重(kg)" in spec_rows else 0

        csv_row = lookup_sku(sku_meta, sku_str, {})  # S9-step3：统一大小写 / 全角规则
        rec = _catalog_record(catalog, sku_str)
        master = get_product_by_sku(sku_str)
        customs = lookup_sku(customs_by_sku, sku_str, {})
        ai_keys = lookup_sku(customs_ai_keys, sku_str, set())
        master_hints = invoice_hints_from_master(master)

        english = (
            csv_row.get("title")
            or master_hints.get("英文品名")
            or customs.get("英文品名")
            or ""
        )
        chinese = master_hints.get("中文品名") or customs.get("中文品名") or ""
        material = master_hints.get("材料") or customs.get("材料") or ""
        purpose = master_hints.get("用途") or customs.get("用途") or "宠物用品"
        brand = master_hints.get("品牌") or customs.get("品牌") or ""
        hs_code = master_hints.get("海关编码") or customs.get("海关编码") or ""
        purchase = rec.unit_cost if rec and rec.unit_cost else None
        declared = _declared_unit_price(purchase)
        if master_hints.get("单个货值") not in (None, ""):
            try:
                declared = round(float(master_hints["单个货值"]), 2)
            except (TypeError, ValueError):
                pass
        elif customs.get("单个货值") not in (None, ""):
            try:
                declared = round(float(customs["单个货值"]), 2)
            except (TypeError, ValueError):
                pass
        asin = csv_row.get("asin") or (master.get("asin") if master else "") or (rec.asin if rec else "")
        pcs = _pcs_from_sku(sku_str)

        total_weight_kg = round(weight * box_count, 2) if weight and box_count else ""
        unit_weight_kg = ""
        if total_weight_kg and total_qty > 0:
            unit_weight_kg = round(float(total_weight_kg) / total_qty, 3)

        total_declared_value_usd = ""
        if declared not in (None, "") and total_qty > 0:
            total_declared_value_usd = round(float(declared) * total_qty, 2)

        material_display = _combine_material_purpose(material, purpose)

        values = {
            "shipment_id": ship.shipment_id,
            "reference_id": reference_id,
            "box_range": _box_range(boxes),
            "box_count": box_count,
            "sku": sku_str,
            "english_name": english,
            "chinese_name": chinese,
            "brand": brand,
            "material": material_display,
            "purpose": purpose,
            "hs_code": hs_code,
            "asin": asin,
            "battery": "否",
            "magnet": "否",
            "model": str((master or {}).get("model_name") or "").strip(),
            "total_qty": total_qty,
            "unit": "套",
            "pcs_per_set": pcs if pcs else "",
            "purchase_price": purchase if purchase else "",
            "declared_price": declared if declared else "",
            "total_declared_value_usd": total_declared_value_usd,
            "unit_weight_kg": unit_weight_kg,
            "sku_image": "",
            "shipping_mark": ship.warehouse_code or "",
            "box_weight": round(weight, 2) if weight else "",
            "length": length if length else "",
            "width": width if width else "",
            "height": height if height else "",
        }

        sources: Dict[str, str] = {}
        field_layer: Dict[str, str] = {
            "english_name": (
                "green"
                if csv_row.get("title")
                else "blue"
                if master_hints.get("英文品名")
                else "yellow"
                if english and "英文品名" in ai_keys
                else "red"
            ),
            "chinese_name": (
                "blue" if master_hints.get("中文品名") else "yellow" if chinese and "中文品名" in ai_keys else "red"
            ),
            "material": (
                "blue" if master_hints.get("材料") else "yellow" if material and "材料" in ai_keys else "red"
            ),
            "purpose": (
                "blue"
                if master_hints.get("用途")
                else "yellow"
                if purpose != "宠物用品" and "用途" in ai_keys
                else "green"
                if purpose == "宠物用品"
                else "red"
            ),
            "brand": "blue" if master_hints.get("品牌") else "yellow" if brand and "品牌" in ai_keys else "red",
            "hs_code": (
                "blue" if master_hints.get("海关编码") else "yellow" if hs_code and "海关编码" in ai_keys else "red"
            ),
            "declared_price": (
                "blue"
                if rec and rec.unit_cost
                else "blue"
                if master_hints.get("单个货值")
                else "yellow"
                if declared and customs.get("单个货值") and "单个货值" in ai_keys
                else "blue"
                if declared and customs.get("单个货值")
                else "red"
            ),
        }
        for key, val in values.items():
            if val in (None, ""):
                sources[key] = "red" if key in ("reference_id", "sku_image") else field_layer.get(key, "red")
            elif key in field_layer:
                sources[key] = field_layer[key]
            elif key in ("shipment_id", "box_range", "box_count", "total_qty", "sku", "asin"):
                sources[key] = "green"
            elif key in ("length", "width", "height", "box_weight", "total_declared_value_usd", "unit_weight_kg"):
                sources[key] = "green"
            elif key in ("reference_id",):
                sources[key] = "green" if reference_id else "red"
            elif key in ("purchase_price",) and purchase:
                sources[key] = "blue"
            elif key in ("battery", "magnet", "unit", "model", "shipping_mark", "pcs_per_set"):
                sources[key] = "green"
            elif key == "purpose" and purpose == "宠物用品" and "用途" not in ai_keys and not master_hints.get("用途"):
                sources[key] = "green"
            else:
                sources[key] = "green" if val else "red"

        items.append(InvoiceLineItem(values=values, sources=sources))

    fallback_hs = next((it.values.get("hs_code") for it in items if it.values.get("hs_code")), "")
    if fallback_hs:
        for item in items:
            if not item.values.get("hs_code"):
                item.values["hs_code"] = fallback_hs
                item.sources["hs_code"] = "blue"

    return items


def clear_invoice_table_data_zone(
    ws,
    layout: InvoiceTableLayout,
    *,
    max_data_rows: int | None = None,
) -> None:
    """清空表头下方数据区（防止模板脏数据残留）。"""
    rows_to_clear = max_data_rows or INVOICE_TABLE_DATA_CLEAR_ROWS
    if not layout.columns:
        return
    end_row = layout.first_data_row + rows_to_clear - 1
    max_col = max(layout.columns.values())
    for row in range(layout.first_data_row, end_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None


def clear_cells_for_table_zone(
    layout: InvoiceTableLayout,
    *,
    max_data_rows: int | None = None,
) -> List[Dict[str, Any]]:
    """生成空值 cells，供 FortuneSheet 覆盖旧数据。"""
    rows_to_clear = max_data_rows or INVOICE_TABLE_DATA_CLEAR_ROWS
    cells: List[Dict[str, Any]] = []
    end_row = layout.first_data_row + rows_to_clear - 1
    for row in range(layout.first_data_row, end_row + 1):
        for field_key, col in layout.columns.items():
            cells.append(
                {
                    "cell": f"{get_column_letter(col)}{row}",
                    "value": "",
                    "source": "模板解析",
                    "field_key": field_key,
                    "required": False,
                    "clear": True,
                }
            )
    return cells


def table_cells_from_line_items(
    layout: InvoiceTableLayout,
    line_items: List[InvoiceLineItem],
    source_labels: Dict[str, str],
) -> List[Dict[str, Any]]:
    """映射发票行到单元格。清空脏数据由 export 侧处理，此处不批量下发 clear cells（避免 FortuneSheet 崩溃）。"""
    cells: List[Dict[str, Any]] = []
    always_include = {"reference_id", "sku_image"}
    for row_offset, item in enumerate(line_items):
        row = layout.first_data_row + row_offset
        for field_key, col in layout.columns.items():
            val = item.values.get(field_key, "")
            if val in (None, "") and field_key not in always_include:
                continue
            src_kind = item.sources.get(field_key, "green")
            cells.append(
                {
                    "cell": f"{get_column_letter(col)}{row}",
                    "value": val if val not in (None, "") else "",
                    "source": source_labels.get(src_kind, "精准提取"),
                    "field_key": field_key,
                    "required": field_key == "reference_id",
                }
            )
    return cells


def preview_table_cells(layout: InvoiceTableLayout) -> List[Dict[str, Any]]:
    """预览态：标记表头下方首行各列位置（空值）。"""
    cells: List[Dict[str, Any]] = []
    row = layout.first_data_row
    for field_key, col in layout.columns.items():
        cells.append(
            {
                "cell": f"{get_column_letter(col)}{row}",
                "value": "",
                "source": "模板解析",
                "field_key": field_key,
                "required": field_key == "reference_id",
            }
        )
    return cells
