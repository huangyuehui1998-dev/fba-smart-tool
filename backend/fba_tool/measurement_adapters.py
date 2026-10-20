"""实测文件适配器。新增来源只加一个 sniff/parse，接口不变。"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Protocol

from fba_tool.carton_math import inches_to_cm, lb_to_kg, q


class MeasurementAdapter(Protocol):
    name: str

    def sniff(self, filename: str, head: str) -> bool: ...

    def parse(self, data: bytes) -> List[Dict[str, Any]]: ...


class ParseError(ValueError):
    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


def _xlsx_headers(data: bytes) -> List[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        first = next(ws.iter_rows(max_row=1, values_only=True), ())
        return [str(cell or "").strip() for cell in first]
    finally:
        wb.close()


def _xlsx_rows(data: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(cell or "").strip() for cell in rows[0]]
    out = []
    for index, row in enumerate(rows[1:], start=2):
        if not any(cell not in (None, "") for cell in row):
            continue
        out.append({"_row": index, **{headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}})
    return out


def _pick(row: Dict[str, Any], names: List[str]) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def _fnum(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(q(value, 3))


class JiufangBoxAdapter:
    name = "jiufang_box_xlsx"

    def sniff(self, filename: str, head: str) -> bool:
        name = filename.lower()
        return "九方" in filename or "jiufang" in name or ("外箱号" in head and "实重" in head and "入库计费重" not in head)

    def parse(self, data: bytes) -> List[Dict[str, Any]]:
        rows = []
        for raw in _xlsx_rows(data):
            box_id = str(_pick(raw, ["外箱号", "FBA箱号", "箱号"]) or "").strip()
            rows.append({
                "kind": "box",
                "fba_id": str(_pick(raw, ["货件编号", "FBA号"]) or ""),
                "fba_box_id": box_id,
                "sku": str(_pick(raw, ["SKU", "SKU简称"]) or ""),
                "length_cm": _fnum(_pick(raw, ["长", "长(cm)", "长度"])),
                "width_cm": _fnum(_pick(raw, ["宽", "宽(cm)", "宽度"])),
                "height_cm": _fnum(_pick(raw, ["高", "高(cm)", "高度"])),
                "weight_kg": _fnum(_pick(raw, ["实重", "实重(kg)", "重量"])),
                "chargeable_kg": _fnum(_pick(raw, ["计费重", "单箱计费重", "材重"])),
                "source_row": raw["_row"],
                "raw": {k: raw[k] for k in raw if k != "_row"},
                "raw_unit": "cm_kg",
            })
        return rows


class SendBoxAdapter:
    name = "send_box_xlsx"

    def sniff(self, filename: str, head: str) -> bool:
        name = filename.lower()
        return "send" in name or ("外箱号" in head and "SEND" in head.upper())

    def parse(self, data: bytes) -> List[Dict[str, Any]]:
        return JiufangBoxAdapter().parse(data)


class ManualBoxAdapter:
    name = "manual_xlsx"

    def sniff(self, filename: str, head: str) -> bool:
        return "手工" in filename or ("外箱号" in head and "重量" in head and "实重" not in head)

    def parse(self, data: bytes) -> List[Dict[str, Any]]:
        return JiufangBoxAdapter().parse(data)


class ForwarderBillAdapter:
    name = "forwarder_bill_xlsx"

    def sniff(self, filename: str, head: str) -> bool:
        return "入库计费重" in head or "货代表" in filename

    def parse(self, data: bytes) -> List[Dict[str, Any]]:
        rows = []
        seen = set()
        for raw in _xlsx_rows(data):
            fba_id = str(_pick(raw, ["货件编号", "FBA货件号", "FBA号"]) or "").strip()
            bill = _pick(raw, ["入库计费重"])
            box_id = _pick(raw, ["外箱号", "FBA箱号"])
            if box_id not in (None, ""):
                rows.append({
                    "kind": "box",
                    "fba_id": fba_id,
                    "fba_box_id": str(box_id).strip(),
                    "length_cm": _fnum(_pick(raw, ["长", "长(cm)"])),
                    "width_cm": _fnum(_pick(raw, ["宽", "宽(cm)"])),
                    "height_cm": _fnum(_pick(raw, ["高", "高(cm)"])),
                    "weight_kg": _fnum(_pick(raw, ["实重", "实重(kg)", "重量"])),
                    "chargeable_kg": _fnum(_pick(raw, ["单箱计费重", "计费重"])),
                    "source_row": raw["_row"],
                    "raw": {},
                    "raw_unit": "cm_kg",
                })
            if fba_id and bill not in (None, "") and fba_id not in seen:
                seen.add(fba_id)
                rows.append({
                    "kind": "forwarder_bill",
                    "fba_id": fba_id,
                    "bill_chargeable_kg": float(q(bill, 2)),
                    "channel_min_kg": _fnum(_pick(raw, ["渠道最低"])),
                    "forwarder": str(_pick(raw, ["货代"]) or ""),
                })
        return rows


class FeePreviewAdapter:
    name = "fee_preview_txt"

    def sniff(self, filename: str, head: str) -> bool:
        lowered = head.lower()
        return filename.lower().endswith((".txt", ".tsv")) or "product-size-tier" in lowered or "item-package-weight" in lowered

    def parse(self, data: bytes) -> List[Dict[str, Any]]:
        text = data.decode("utf-8-sig")
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            raise ParseError("Fee Preview 文件是空的")
        headers = [cell.strip() for cell in lines[0].split("\t")]
        lowered = [cell.lower() for cell in headers]
        weight_names = ["item-package-weight", "weight", "item_package_weight"]
        if not any(name in lowered for name in weight_names):
            raise ParseError("缺少列：item-package-weight")
        def col(*names: str) -> int | None:
            for name in names:
                if name in lowered:
                    return lowered.index(name)
            return None
        sku_i = col("sku")
        asin_i = col("asin")
        long_i = col("longest-side")
        med_i = col("median-side")
        short_i = col("shortest-side")
        weight_i = col("item-package-weight", "weight")
        dim_unit_i = col("unit-of-dimension")
        weight_unit_i = col("unit-of-weight")
        tier_i = col("product-size-tier", "product-size-weight-band")
        store_i = col("amazon-store", "marketplace")
        rows = []
        for line_no, line in enumerate(lines[1:], start=2):
            cells = line.split("\t")
            def take(index: int | None) -> str:
                if index is None or index >= len(cells):
                    return ""
                return cells[index].strip()
            dim_unit = take(dim_unit_i).lower()
            weight_unit = take(weight_unit_i).lower()
            def length(index: int | None):
                raw = take(index)
                if not raw:
                    return None
                number = q(raw, 2)
                if dim_unit in {"inches", "inch", "in"}:
                    return float(inches_to_cm(number))
                return float(number)
            weight_raw = take(weight_i)
            if not weight_raw:
                raise ParseError(f"第{line_no}行缺少列：item-package-weight")
            weight = q(weight_raw, 3)
            if weight_unit in {"pounds", "pound", "lb", "lbs"}:
                weight = lb_to_kg(weight)
            store = take(store_i).upper()
            if store in {"AMAZON.COM", "US", "ATVPDKIKX0DER"}:
                marketplace = "US"
            elif store in {"AMAZON.CA", "CA", "A2EUQ1WTGCTBG2"}:
                marketplace = "CA"
            else:
                marketplace = store or "US"
            rows.append({
                "kind": "item",
                "sku": take(sku_i),
                "asin": take(asin_i),
                "marketplace": marketplace,
                "longest_cm": length(long_i),
                "median_cm": length(med_i),
                "shortest_cm": length(short_i),
                "weight_kg": float(q(weight, 3)),
                "size_tier": take(tier_i),
                "source_row": line_no,
                "raw": {"line": line},
            })
        return rows


ADAPTERS: List[MeasurementAdapter] = [
    FeePreviewAdapter(),
    ForwarderBillAdapter(),
    JiufangBoxAdapter(),
    SendBoxAdapter(),
    ManualBoxAdapter(),
]


def head_text(filename: str, data: bytes) -> str:
    if filename.lower().endswith((".txt", ".tsv", ".csv")) or not data.startswith(b"PK"):
        return data[:4000].decode("utf-8", errors="ignore")
    try:
        return " ".join(_xlsx_headers(data))
    except Exception:
        return ""


def detect_adapter(filename: str, data: bytes) -> MeasurementAdapter:
    head = head_text(filename, data)
    for adapter in ADAPTERS:
        if adapter.sniff(filename, head):
            return adapter
    raise ParseError("认不出实测文件格式")
