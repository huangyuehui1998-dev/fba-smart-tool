"""票级数据导入器（方案 v8 §3.2 P14）：腾讯表格 / 运费测算表「FBA运费-记录」/ 九方单箱数据 → shipment_ticket。

- 列名映射在 config/ticket_import_map.json；先预览差异，确认再写。
- 按 FBA 号关联；每个字段记录来源（field_source_json），排名：manual > tencent > freight = jiufang，
  排名低的来源不覆盖排名高的值。
- 总运费只作参考，不写入报关单（v6 Q11）；出口/申报日期不从这里取（v6 Q16）。
"""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from fba_tool.app_paths import bundled_config_dir

TEXT_FIELDS = ("batch", "ticket_no", "order_no", "declare_mode", "warehouse_code", "region", "forwarder", "channel", "lcl_no")
INT_FIELDS = ("box_count", "total_qty")
FLOAT_FIELDS = ("gross_kg", "net_kg", "chargeable_kg", "unit_freight", "total_freight")
DATE_FIELDS = (
    "etd", "eta_expected", "delivered_at", "amazon_received_at",
    "export_date", "declare_date", "invoice_date", "contract_date", "shipment_deadline", "entrust_date",
)
TICKET_FIELDS = TEXT_FIELDS + INT_FIELDS + FLOAT_FIELDS + DATE_FIELDS + ("boxes_json",)
_FBA_RE = re.compile(r"^FBA[0-9A-Z]{6,}$")
_ORDER_RE = re.compile(r"^J[A-Z0-9]{4,}$")
_JF_BOX_RE = re.compile(r"^(FBA[0-9A-Z]+?)U\d{6}$")


def _load_map() -> Dict[str, Any]:
    p = bundled_config_dir() / "ticket_import_map.json"
    return json.loads(p.read_text(encoding="utf-8"))


def _rank(source: str) -> int:
    return int(_load_map()["priority"].get(source, 0))


def _norm(text: Any) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        v = v.replace("\xa0", " ").strip()
        return v or None
    return v


def _workbook(data: bytes, file_name: str):
    import openpyxl

    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{file_name or '文件'}（{exc.__class__.__name__}）") from exc


def _sheet_rows(ws) -> List[Tuple[Any, ...]]:
    return [tuple(_cell(c) for c in r) for r in ws.iter_rows(values_only=True)]


def detect_ticket_source(data: bytes, file_name: str = "") -> str:
    wb = _workbook(data, file_name)
    if "FBA运费-记录" in wb.sheetnames:
        return "freight"
    rows = _sheet_rows(wb.worksheets[0])[:10]
    heads = {_norm(c) for r in rows for c in r if c is not None}
    if "外箱号" in heads and "单箱实重" in heads:
        return "jiufang"
    if "FBA编码" in heads and "毛重" in heads:
        return "tencent"
    raise ValueError("认不出是哪种票级数据表：需要腾讯表格（含 FBA编码、毛重）、运费测算表（含「FBA运费-记录」）或九方单箱数据（含 外箱号、单箱实重）")


def _columns(header: Tuple[Any, ...], mapping: Dict[str, List[str]]) -> Dict[str, List[int]]:
    heads = [_norm(h) for h in header]
    out: Dict[str, List[int]] = {}
    for field, cands in mapping.items():
        if field.startswith("_"):
            continue
        cn = [_norm(c) for c in cands]
        exact = [i for i, h in enumerate(heads) if h in cn]
        loose = [i for i, h in enumerate(heads) if h and any(h.startswith(c) for c in cn)]
        cols = exact or loose
        if cols:
            out[field] = cols
    return out


def _find_header(rows: List[Tuple[Any, ...]], mapping: Dict[str, List[str]], key: str) -> int:
    want = [_norm(c) for c in mapping[key]]
    for i, r in enumerate(rows[:15]):
        if any(_norm(c) in want for c in r if c is not None):
            return i
    raise ValueError(f"没找到表头「{mapping[key][0]}」")


def _pick(row: Tuple[Any, ...], cols: List[int]) -> Any:
    for i in reversed(cols):  # 同名列（如两列“总运费”）取靠右第一个有值的
        if i < len(row) and row[i] not in (None, ""):
            return row[i]
    return None


def _num(v: Any) -> Optional[float]:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _date(v: Any) -> Optional[str]:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", v.strip())
        if m:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    return None


def _issue(row: int, field: str, value: Any, msg: str) -> Dict[str, Any]:
    return {"row": row, "field": field, "value": "" if value is None else str(value), "message": msg}


def _typed(row_no: int, field: str, raw: Any, issues: List[Dict[str, Any]]) -> Any:
    if raw in (None, ""):
        return None
    if field in FLOAT_FIELDS:
        v = _num(raw)
        if v is None:
            issues.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」不是数字：{raw}"))
        return v
    if field in INT_FIELDS:
        v = _num(raw)
        if v is None or v != int(v):
            issues.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」不是整数：{raw}"))
            return None
        return int(v)
    if field in DATE_FIELDS:
        v = _date(raw)
        if v is None:
            issues.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」不是日期：{raw}（未导入）"))
        return v
    text = str(raw).strip()
    if field == "order_no" and not _ORDER_RE.match(text):
        issues.append(_issue(row_no, field, raw, f"第{row_no}行订单编号格式不对：{text}（未导入）"))
        return None
    return text


def _parse_table(rows, mapping, source) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    hi = _find_header(rows, mapping, "fba_id")
    cols = _columns(rows[hi], mapping)
    out: Dict[str, Dict[str, Any]] = {}
    issues: List[Dict[str, Any]] = []
    for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
        fba = str(_pick(r, cols.get("fba_id", [])) or "").strip().upper()
        if not _FBA_RE.match(fba):
            continue
        rec: Dict[str, Any] = {"fba_id": fba, "_source": source, "_row": ri}
        for field, cidx in cols.items():
            if field in ("fba_id", "plan_name"):
                continue
            val = _typed(ri, field, _pick(r, cidx), issues)
            if val is not None:
                rec[field] = val
        plan = str(_pick(r, cols.get("plan_name", [])) or "")
        m = re.search(r"(\d+)\s*TH", plan, re.IGNORECASE)
        if m:
            rec["batch"] = m.group(1)
        if rec.get("warehouse_code"):
            rec["warehouse_code"] = str(rec["warehouse_code"]).upper()
        out[fba] = rec
    return list(out.values()), issues


def _parse_jiufang(rows, mapping) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    hi = _find_header(rows, mapping, "box_no")
    cols = _columns(rows[hi], mapping)
    issues: List[Dict[str, Any]] = []
    tickets: Dict[str, Dict[str, Any]] = {}
    for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
        box_no = str(_pick(r, cols.get("box_no", [])) or "").strip().upper()
        m = _JF_BOX_RE.match(box_no)
        if not m:
            if box_no:
                issues.append(_issue(ri, "box_no", box_no, f"第{ri}行外箱号认不出 FBA 号：{box_no}"))
            continue
        weight = _num(_pick(r, cols.get("box_weight", [])))
        if weight is None or weight <= 0:
            issues.append(_issue(ri, "box_weight", _pick(r, cols.get("box_weight", [])), f"第{ri}行单箱实重无效"))
            continue
        t = tickets.setdefault(m.group(1), {"fba_id": m.group(1), "_source": "jiufang", "_row": ri, "boxes": []})
        lcl = str(_pick(r, cols.get("lcl_no", [])) or "").strip()
        if lcl:
            t["lcl_no"] = lcl
        t["boxes"].append({
            "box_no": box_no,
            "weight": weight,
            "length": _num(_pick(r, cols.get("length", []))),
            "width": _num(_pick(r, cols.get("width", []))),
            "height": _num(_pick(r, cols.get("height", []))),
            "chargeable_kg": _num(_pick(r, cols.get("chargeable_kg", []))),
        })
    out = []
    from decimal import ROUND_HALF_UP, Decimal

    for t in tickets.values():
        boxes = sorted(t.pop("boxes"), key=lambda b: b["box_no"])
        t["boxes_json"] = json.dumps(boxes, ensure_ascii=False)
        t["box_count"] = len(boxes)
        t["gross_kg"] = float(sum(Decimal(str(b["weight"])) for b in boxes).quantize(Decimal("0.01"), ROUND_HALF_UP))
        out.append(t)
    return out, issues


def parse_ticket_file(data: bytes, file_name: str = "") -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    source = detect_ticket_source(data, file_name)
    mp = _load_map()[source]
    wb = _workbook(data, file_name)
    if source == "freight":
        rows = _sheet_rows(wb["FBA运费-记录"])
        return _parse_table(rows, mp, source)
    rows = _sheet_rows(wb.worksheets[0])
    if source == "jiufang":
        return _parse_jiufang(rows, mp)
    return _parse_table(rows, mp, source)


# ---------------- 存取 ----------------
def _conn():
    from fba_tool.master_db import _connect

    return _connect()


def _row_dict(cur, row) -> Dict[str, Any]:
    d = dict(zip([c[0] for c in cur.description], row))
    try:
        d["field_sources"] = json.loads(d.get("field_source_json") or "{}")
    except ValueError:
        d["field_sources"] = {}
    try:
        d["boxes"] = json.loads(d.get("boxes_json") or "[]")
    except ValueError:
        d["boxes"] = []
    return d


def get_ticket(fba_id: str) -> Optional[Dict[str, Any]]:
    with _conn() as conn:
        cur = conn.execute("SELECT * FROM shipment_ticket WHERE fba_id = ?", (str(fba_id or "").strip().upper(),))
        row = cur.fetchone()
        return _row_dict(cur, row) if row else None


def list_tickets(batch: str = "") -> List[Dict[str, Any]]:
    with _conn() as conn:
        if batch:
            cur = conn.execute("SELECT * FROM shipment_ticket WHERE batch = ? ORDER BY order_no, fba_id", (str(batch),))
        else:
            cur = conn.execute("SELECT * FROM shipment_ticket ORDER BY batch, order_no, fba_id")
        return [_row_dict(cur, r) for r in cur.fetchall()]


def _plan(existing: Optional[Dict[str, Any]], values: Dict[str, Any], source: str) -> Dict[str, Any]:
    """返回会被写入的字段（考虑来源排名）。"""
    sources = (existing or {}).get("field_sources") or {}
    my = _rank(source)
    out: Dict[str, Any] = {}
    for f, v in values.items():
        if f not in TICKET_FIELDS or v in (None, ""):
            continue
        cur_src = sources.get(f)
        if cur_src and _rank(cur_src) > my:
            continue
        out[f] = v
    return out


def set_ticket_fields(fba_id: str, values: Dict[str, Any], source: str) -> Dict[str, Any]:
    if source not in _load_map()["priority"]:
        raise ValueError(f"未知来源：{source}")
    key = str(fba_id or "").strip().upper()
    if not _FBA_RE.match(key):
        raise ValueError(f"FBA 号格式不对：{fba_id}")
    existing = get_ticket(key)
    writes = _plan(existing, values, source)
    sources = dict((existing or {}).get("field_sources") or {})
    for f in writes:
        sources[f] = source
    now = datetime.now(timezone.utc).isoformat()
    with _conn() as conn:
        if not existing:
            conn.execute("INSERT INTO shipment_ticket (fba_id, updated_at) VALUES (?, ?)", (key, now))
        if writes:
            sets = dict(writes)
            sets["field_source_json"] = json.dumps(sources, ensure_ascii=False, sort_keys=True)
            sets["updated_at"] = now
            conn.execute(
                f"UPDATE shipment_ticket SET {', '.join(f'{k} = ?' for k in sets)} WHERE fba_id = ?",
                (*sets.values(), key),
            )
        conn.commit()
    return {"written": sorted(writes)}


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, (int, float)) or isinstance(b, (int, float)):
        try:
            return abs(float(a) - float(b)) < 1e-9
        except (TypeError, ValueError):
            return False
    return str(a or "") == str(b or "")


def preview_ticket_import(data: bytes, file_name: str = "") -> Dict[str, Any]:
    source = detect_ticket_source(data, file_name)
    rows, issues = parse_ticket_file(data, file_name)
    summary = {"new": 0, "changed": 0, "unchanged": 0, "total": len(rows)}
    items = []
    for rec in rows:
        ex = get_ticket(rec["fba_id"])
        vals = {k: v for k, v in rec.items() if not k.startswith("_") and k != "fba_id"}
        writes = _plan(ex, vals, source)
        changed = {f: {"old": (ex or {}).get(f), "new": v} for f, v in writes.items() if not ex or not _same(ex.get(f), v)}
        status = "new" if not ex else ("changed" if changed else "unchanged")
        summary[status] += 1
        items.append({"fba_id": rec["fba_id"], "status": status, "fields": changed})
    return {"source": source, "summary": summary, "items": items, "issues": issues}


def apply_ticket_import(data: bytes, file_name: str = "") -> Dict[str, Any]:
    source = detect_ticket_source(data, file_name)
    rows, issues = parse_ticket_file(data, file_name)
    written = 0
    for rec in rows:
        vals = {k: v for k, v in rec.items() if not k.startswith("_") and k != "fba_id"}
        res = set_ticket_fields(rec["fba_id"], vals, source)
        written += 1 if res["written"] else 0
    return {"ok": True, "source": source, "tickets": len(rows), "tickets_written": written, "issues": issues}
