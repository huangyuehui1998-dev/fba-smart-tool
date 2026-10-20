"""工人装箱贴标实操单 — 按箱唛 PDF 上传顺序横向并排导出。"""
from __future__ import annotations

import re
from collections import defaultdict
from io import BytesIO
from typing import Any, Dict, List, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from fba_tool.box_label_filename import column_label_for_pair, sort_pairs_by_box_label_pdf
from fba_tool.fnsku_resolve import (
    fnsku_map_from_pairs,
    fnsku_map_from_step1_rows,
    resolve_fnsku,
)
from fba_tool.master_db import lookup_model_name
from fba_tool.shipment_parser import normalize_fba_box_id
from fba_tool.template_loader import load_sku_mapping
from fba_tool.user_config import param

_THIN = Side(style="thin", color="000000")
CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
# 便携版/其他电脑打开 Excel 时避免中文乱码
WORKER_FONT = Font(name="Microsoft YaHei", size=11)
WORKER_FONT_BOLD = Font(name="Microsoft YaHei", bold=True, size=11)
TITLE_FONT = Font(name="Microsoft YaHei", bold=True, size=14)
SUB_FONT = Font(name="Microsoft YaHei", bold=True, size=11)
SUMMARY_FONT = Font(name="Microsoft YaHei", bold=True, size=11)

_DIGITS = "零一二三四五六七八九"


def _chinese_num(n: int) -> str:
    if n <= 0:
        return str(n)
    if n < 10:
        return _DIGITS[n]
    if n < 20:
        return "十" + (_DIGITS[n % 10] if n % 10 else "")
    if n < 100:
        tens, ones = divmod(n, 10)
        text = _DIGITS[tens] + "十"
        if ones:
            text += _DIGITS[ones]
        return text
    return str(n)


def ticket_label(ticket_index: int, warehouse: str) -> str:
    """例：5 + AVP1 -> 第五票 AVP1"""
    wh = str(warehouse or "UNKN").strip()
    return f"第{_chinese_num(ticket_index)}票 {wh}"


def warehouse_labels_from_pairs(pairs: List[Any]) -> List[str]:
    ordered = sort_pairs_by_box_label_pdf(pairs)
    return [column_label_for_pair(p, i) for i, p in enumerate(ordered, 1)]


def _is_worker_summary_row(box_val: str, prod_val: str) -> bool:
    box = str(box_val or "").strip()
    prod = str(prod_val or "").strip()
    if box == "【汇总】":
        return True
    return not box and "共" in prod and prod.endswith("箱")


def _fba_tail(fba_id: str) -> str:
    m = re.search(r"U0*(\d+)\s*$", str(fba_id or "").upper())
    return f"U{int(m.group(1)):03d}" if m else ""


def _carton_label(seq: int, fba_id: str) -> str:
    prefix = str(param("factory_box_label_prefix", "纸箱"))
    base = f"{prefix}{seq}"
    tail = _fba_tail(fba_id)
    return f"{base} (尾号{tail})" if tail else base


def _worker_product_name(sku: str, catalog=None) -> str:
    """工人实操单产品名称：主库型号名 / 报关中文名 / 英文名，最后 SKU。"""
    key = str(sku or "").strip()
    if not key:
        return ""
    from fba_tool.master_db import get_product_by_sku, lookup_customs_name

    name = lookup_model_name(key) or lookup_customs_name(key)
    if name:
        return name
    rec = get_product_by_sku(key)
    if rec and rec.get("english_name"):
        return str(rec["english_name"]).strip()
    mapping = load_sku_mapping()
    if key in mapping:
        return mapping[key]
    return key


def _product_label(sku: str, qty: int, catalog=None) -> str:
    """保留兼容：产品名称不再拼接数量（数量单独成列）。"""
    return _worker_product_name(sku, catalog=catalog)


def _fnsku_for_sku(
    sku: str,
    *,
    catalog=None,
    csv_fnsku: str = "",
    step1_fnsku: str = "",
) -> str:
    return resolve_fnsku(
        sku,
        csv_fnsku=csv_fnsku,
        step1_fnsku=step1_fnsku,
        catalog=catalog,
    )


def _boxes_from_pair(pair: Any) -> List[Tuple[str, str, Dict[str, int]]]:
    """单货件 (sort_key, fba_id, skus) 列表。"""
    boxes: List[Tuple[str, str, Dict[str, int]]] = []
    csv = pair.csv
    fba_map = dict(csv.box_fba_ids)
    for box_no, skus in csv.box_data.items():
        fba_id = str(fba_map.get(int(box_no)) or "").strip()
        if not fba_id:
            continue
        boxes.append((normalize_fba_box_id(fba_id), fba_id, dict(skus)))
    boxes.sort(key=lambda x: x[0])
    return boxes


def _blocks_from_pairs(pairs: List[Any]) -> List[Tuple[str, List[Tuple[str, str, Dict[str, int]]]]]:
    """按箱唛 PDF 票序返回 (列标题, 箱列表)。"""
    blocks: List[Tuple[str, List[Tuple[str, str, Dict[str, int]]]]] = []
    ordered = sort_pairs_by_box_label_pdf(pairs)
    for i, pair in enumerate(ordered, start=1):
        label = column_label_for_pair(pair, i)
        blocks.append((label, _boxes_from_pair(pair)))
    return blocks


def _warehouse_product_summary_rows(
    boxes: List[Tuple[str, str, Dict[str, int]]],
    catalog=None,
    csv_fnsku_map: Dict[str, str] | None = None,
    step1_map: Dict[str, str] | None = None,
) -> List[Dict[str, str]]:
    """单仓产品汇总：相同 SKU 合并显示名称、FNSKU、总件数与总箱数。"""
    box_count_by_sku: Dict[str, int] = defaultdict(int)
    total_qty_by_sku: Dict[str, int] = defaultdict(int)
    for _, _, skus in boxes:
        seen_in_box: set[str] = set()
        for sku, qty in skus.items():
            key = str(sku or "").strip()
            if int(qty) <= 0 or not key or key in seen_in_box:
                continue
            seen_in_box.add(key)
            box_count_by_sku[key] += 1
            total_qty_by_sku[key] += int(qty)

    if not box_count_by_sku:
        return []

    csv_map = csv_fnsku_map or {}
    s1_map = step1_map or {}

    blank = {"纸箱编号": "", "产品名称": "", "FNSKU": "", "数量": ""}
    rows: List[Dict[str, str]] = [dict(blank)]
    rows.append({"纸箱编号": "【汇总】", "产品名称": "", "FNSKU": "", "数量": ""})
    for sku in sorted(box_count_by_sku.keys()):
        name = _worker_product_name(sku, catalog=catalog)
        rows.append(
            {
                "纸箱编号": "",
                "产品名称": f"{name}  共{box_count_by_sku[sku]}箱",
                "FNSKU": _fnsku_for_sku(
                    sku,
                    catalog=catalog,
                    csv_fnsku=csv_map.get(sku, ""),
                    step1_fnsku=s1_map.get(sku, ""),
                ),
                "数量": str(total_qty_by_sku[sku]),
            }
        )
    return rows


WORKER_WAREHOUSE_COLUMNS = ["纸箱编号", "产品名称", "FNSKU", "数量"]


def build_warehouse_worker_df(
    boxes: List[Tuple[str, str, Dict[str, int]]],
    *,
    catalog=None,
    csv_fnsku_map: Dict[str, str] | None = None,
    step1_map: Dict[str, str] | None = None,
) -> pd.DataFrame:
    """单仓四列 DataFrame：纸箱编号、产品名称、FNSKU、数量（混装允许多行同箱号）。"""
    if catalog is None:
        catalog = None
    csv_map = csv_fnsku_map or {}
    s1_map = step1_map or {}
    rows: List[Dict[str, str]] = []
    for seq, (_, fba_id, skus) in enumerate(boxes, start=1):
        label = _carton_label(seq, fba_id)
        items = [(sku, int(qty)) for sku, qty in sorted(skus.items()) if int(qty) > 0]
        if not items:
            rows.append({"纸箱编号": label, "产品名称": "", "FNSKU": "", "数量": ""})
            continue
        for sku, qty in items:
            rows.append(
                {
                    "纸箱编号": label,
                    "产品名称": _worker_product_name(sku, catalog=catalog),
                    "FNSKU": _fnsku_for_sku(
                        sku,
                        catalog=catalog,
                        csv_fnsku=csv_map.get(sku, ""),
                        step1_fnsku=s1_map.get(sku, ""),
                    ),
                    "数量": str(qty),
                }
            )
    rows.extend(
        _warehouse_product_summary_rows(
            boxes, catalog=catalog, csv_fnsku_map=csv_map, step1_map=s1_map
        )
    )
    if not rows:
        return pd.DataFrame(columns=WORKER_WAREHOUSE_COLUMNS)
    return pd.DataFrame(rows, columns=WORKER_WAREHOUSE_COLUMNS)


def build_worker_packing_horizontal_df(
    pairs: List[Any],
    step1_rows: List[Dict[str, Any]] | None = None,
) -> pd.DataFrame:
    """多票横向 concat(axis=1)，列名前缀为「第N票 仓库」。"""
    blocks = _blocks_from_pairs(pairs)
    if not blocks:
        return pd.DataFrame()

    catalog = None
    csv_map = fnsku_map_from_pairs(pairs)
    s1_map = fnsku_map_from_step1_rows(step1_rows)
    frames: List[pd.DataFrame] = []
    for col_label, boxes in blocks:
        block = build_warehouse_worker_df(
            boxes, catalog=catalog, csv_fnsku_map=csv_map, step1_map=s1_map
        )
        block = block.rename(
            columns={
                "纸箱编号": f"{col_label}_纸箱编号",
                "产品名称": f"{col_label}_产品名称",
                "FNSKU": f"{col_label}_FNSKU",
                "数量": f"{col_label}_数量",
            }
        )
        frames.append(block.reset_index(drop=True))

    combined = pd.concat(frames, axis=1).fillna("")
    for col in combined.columns:
        combined[col] = combined[col].astype(str).replace("nan", "")
    return combined


def _labels_from_df_columns(df: pd.DataFrame) -> List[str]:
    labels: List[str] = []
    seen: set[str] = set()
    for col in df.columns:
        m = re.match(r"^(.+?)_(?:纸箱编号|产品名称|FNSKU|数量)$", str(col))
        if m:
            label = m.group(1)
            if label not in seen:
                seen.add(label)
                labels.append(label)
    return labels


def export_worker_packing_preview(
    preview_rows: List[Dict[str, Any]],
    warehouse_labels: List[str] | None = None,
) -> bytes:
    """从已校验的预览数据快速导出（无需重传 CSV/PDF）。"""
    df = pd.DataFrame(preview_rows)
    if df.empty:
        wb = Workbook()
        ws = wb.active
        ws.title = "工人装箱实操单"
        ws.cell(1, 1, "无可用装箱数据").font = WORKER_FONT
        return _to_bytes(wb)
    labels = warehouse_labels or _labels_from_df_columns(df)
    return _export_worker_workbook(df, labels)


def export_worker_packing_list(
    pairs: List[Any],
    step1_rows: List[Dict[str, Any]] | None = None,
) -> bytes:
    """导出横向分票并排的工人装箱贴标实操单。"""
    df = build_worker_packing_horizontal_df(pairs, step1_rows=step1_rows)
    if df.empty:
        wb = Workbook()
        ws = wb.active
        ws.title = "工人装箱实操单"
        ws.cell(1, 1, "无可用装箱数据").font = WORKER_FONT
        return _to_bytes(wb)
    return _export_worker_workbook(df, _labels_from_df_columns(df))


def _export_worker_workbook(df: pd.DataFrame, warehouse_labels: List[str]) -> bytes:
    max_data_rows = len(df)
    wb = Workbook()
    ws = wb.active
    ws.title = "工人装箱实操单"

    col_start = 1
    for wh_label in warehouse_labels:
        col_box = col_start
        col_prod = col_start + 1
        col_fnsku = col_start + 2
        col_qty = col_start + 3
        box_key = f"{wh_label}_纸箱编号"
        prod_key = f"{wh_label}_产品名称"
        fnsku_key = f"{wh_label}_FNSKU"
        qty_key = f"{wh_label}_数量"

        ws.merge_cells(start_row=1, start_column=col_box, end_row=1, end_column=col_qty)
        title_cell = ws.cell(1, col_box, wh_label)
        title_cell.font = TITLE_FONT
        title_cell.alignment = CENTER
        title_cell.border = CELL_BORDER
        for c in (col_prod, col_fnsku, col_qty):
            ws.cell(1, c).border = CELL_BORDER

        for hdr_text, c in (
            ("纸箱编号", col_box),
            ("产品名称", col_prod),
            ("FNSKU", col_fnsku),
            ("数量", col_qty),
        ):
            cell = ws.cell(2, c, hdr_text)
            cell.font = SUB_FONT
            cell.alignment = CENTER
            cell.border = CELL_BORDER

        for i in range(max_data_rows):
            r = 3 + i
            box_val = str(df.iloc[i].get(box_key, "")) if box_key in df.columns else ""
            prod_val = str(df.iloc[i].get(prod_key, "")) if prod_key in df.columns else ""
            fnsku_val = str(df.iloc[i].get(fnsku_key, "")) if fnsku_key in df.columns else ""
            qty_val = str(df.iloc[i].get(qty_key, "")) if qty_key in df.columns else ""
            if box_val in ("nan", "None"):
                box_val = ""
            if prod_val in ("nan", "None"):
                prod_val = ""
            if fnsku_val in ("nan", "None"):
                fnsku_val = ""
            if qty_val in ("nan", "None"):
                qty_val = ""
            summary = _is_worker_summary_row(box_val, prod_val)
            for cell_val, c in (
                (box_val, col_box),
                (prod_val, col_prod),
                (fnsku_val, col_fnsku),
                (qty_val, col_qty),
            ):
                cell = ws.cell(r, c, cell_val)
                cell.alignment = CENTER
                cell.border = CELL_BORDER
                cell.font = SUMMARY_FONT if summary else WORKER_FONT

        ws.column_dimensions[get_column_letter(col_box)].width = 22
        ws.column_dimensions[get_column_letter(col_prod)].width = 40
        ws.column_dimensions[get_column_letter(col_fnsku)].width = 18
        ws.column_dimensions[get_column_letter(col_qty)].width = 10
        col_start += 4

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22
    return _to_bytes(wb)


def _to_bytes(wb: Workbook) -> bytes:
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
