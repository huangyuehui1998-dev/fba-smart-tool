"""产品库本地批量编辑（D-2026-10-16-03）：导出 Excel → 本地改 → 导回预览 → 确认写入。

写入一律走 patch_product_field（来源记 manual 并锁定，MCP 拉取不覆盖）；
申报价例外：listing_price 等不锁定，申报价仍按 D-2026-10-16-01 领星优先。
"""
from __future__ import annotations

import io
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook

from fba_tool.master_db import PATCHABLE_FIELDS, get_product_by_sku, list_products, patch_product_field, resolve_stored_sku
from fba_tool.sku_keys import SkuCaseConflict, sku_key
from fba_tool.product_fields import OWN_FIELDS, _FLOAT_FIELDS

EXPORT_FIELDS = (
    "sku", "fnsku", "asin", "local_sku", "local_name", "english_name", "model_name", "chinese_name",
    "hs_code", "category", "brand", "material", "purpose", "declare_unit", "declared_value",
    "listing_price", "listing_price_currency",
    "unit_length", "unit_width", "unit_height", "unit_weight",
    "qty_per_box", "carton_length", "carton_width", "carton_height", "carton_weight", "carton_gross_weight",
)
SHEET = "产品库"


def editable_fields() -> List[str]:
    allowed = set(PATCHABLE_FIELDS) | set(OWN_FIELDS)
    return [f for f in EXPORT_FIELDS if f != "sku" and f in allowed]


def _num_fields() -> set:
    from fba_tool.master_db import CARTON_FIELDS, INVOICE_FLOAT_FIELDS, UNIT_FIELDS

    return set(_FLOAT_FIELDS) | set(CARTON_FIELDS) | set(UNIT_FIELDS) | set(INVOICE_FLOAT_FIELDS)


def _all_products() -> List[Dict[str, Any]]:
    res = list_products()
    items = res.get("items", res) if isinstance(res, dict) else res
    return [get_product_by_sku(p["sku"]) or p for p in items]


def export_products_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(list(EXPORT_FIELDS))
    for p in _all_products():
        ws.append([p.get(f) for f in EXPORT_FIELDS])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _norm(field: str, v: Any, nums: set):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if field in nums:
        return float(str(v).replace(",", "").strip())
    return str(v).strip()


def preview_products_edit(content: bytes, filename: str = "") -> Dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"无法读取 Excel：{exc}") from exc
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("表格为空")
    head = [str(h or "").strip() for h in rows[0]]
    if "sku" not in head:
        raise ValueError("表头缺少 sku 列（请用“导出产品库”得到的表格）")
    fields = [f for f in head if f in editable_fields()]
    nums = _num_fields()
    changes: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    # S8 D11：同一 SKU 出现多行 → 这些行都不写，提示合并（以前最后一行悄悄生效）
    seen: Dict[str, List[int]] = {}
    for i, row in enumerate(rows[1:], start=2):
        k = sku_key(dict(zip(head, row)).get("sku"))  # S9-step3：统一大小写 / 全角规则
        if k:
            seen.setdefault(k, []).append(i)
    dups = {k: v for k, v in seen.items() if len(v) > 1}
    for i, row in enumerate(rows[1:], start=2):
        vals = dict(zip(head, row))
        sku = str(vals.get("sku") or "").strip()
        if not sku:
            continue
        if sku_key(sku) in dups:
            lines = dups[sku_key(sku)]
            errors.append({"row": i, "sku": sku, "field": "sku", "value": sku,
                           "message": f"SKU 重复出现在第 {'、'.join(map(str, lines))} 行，这几行都没有写入；请合并成一行后重新导入"})
            continue
        try:  # S9-step4-fix3：只差大小写的多条 → 报冲突（之前误报“产品库里没有”）
            resolve_stored_sku(sku)
        except SkuCaseConflict as exc:
            errors.append({"row": i, "sku": sku, "field": "sku", "value": sku, "message": str(exc)})
            continue
        cur = get_product_by_sku(sku)
        if not cur:
            errors.append({"row": i, "sku": sku, "field": "sku", "value": sku, "message": "产品库里没有这个 SKU，不新增"})
            continue
        for f in fields:
            try:
                new = _norm(f, vals.get(f), nums)
            except ValueError:
                errors.append({"row": i, "sku": sku, "field": f, "value": vals.get(f), "message": "不是数字"})
                continue
            if new is not None and f in ("hs_code", "dest_hs_code"):
                from fba_tool.hs_code import HsCodeError, normalize_hs_code

                try:
                    new = normalize_hs_code(new)
                except HsCodeError as exc:
                    errors.append({"row": i, "sku": sku, "field": f, "value": vals.get(f), "message": str(exc)})
                    continue
            if new is not None and f in nums and new < 0:
                errors.append({"row": i, "sku": sku, "field": f, "value": vals.get(f), "message": "不能为负数"})
                continue
            try:
                old = _norm(f, cur.get(f), nums)
            except ValueError:
                old = cur.get(f)
            if new is None:
                continue  # 留空 = 不改（不清空库内值）
            same = (abs(new - old) < 1e-9) if (f in nums and isinstance(old, float)) else (new == old)
            if not same:
                changes.append({"row": i, "sku": sku, "field": f, "old": cur.get(f), "new": new,
                                "source_before": (cur.get("field_sources") or {}).get(f, "")})
    return {"ok": True, "filename": filename, "changes": changes, "errors": errors,
            "note": "确认后写入，改动字段来源记 manual 并锁定（申报价仍领星优先）。留空单元格不改库内值。"}


def apply_products_edit(content: bytes, filename: str = "") -> Dict[str, Any]:
    prev = preview_products_edit(content, filename)
    if prev["errors"] and not prev["changes"]:  # S9-step4-fix3：全是坏行 → 422，列出行号
        from fba_tool.product_import import ImportRowsError, error_rows_text

        raise ImportRowsError(f"没有可写入的改动，整个文件未写入（出错：{error_rows_text(prev['errors'])}）：" + "；".join(
            str(e.get("message") or "") for e in prev["errors"][:5]), prev["errors"])
    for c in prev["changes"]:
        patch_product_field(c["sku"], c["field"], c["new"])
    return {**prev, "written": len(prev["changes"])}
