"""Step 4.5 生成发票与 /api/invoice-export/export 使用同一套货代适配器。

有导出适配器的货代不再按通用表头猜格子。生成时渲染正式发票并记住结果，
导出时直接交这本工作簿；表格里的手改再叠加上去。
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from fba_tool.invoice_automation import InvoiceDraft
from fba_tool.univer_apply import (
    ApplyRejected,
    _is_formula,
    _iter_cells,
    _norm_formula,
    _snapshot_sheets,
    _values_equal,
    _write,
)

_CACHE: Dict[str, bytes] = {}


def _key(forwarder_id: str, shipment_id: str) -> str:
    return f"{forwarder_id}\n{shipment_id}"


def remember_render(forwarder_id: str, shipment_id: str, data: bytes) -> None:
    if forwarder_id and shipment_id and data:
        _CACHE[_key(forwarder_id, shipment_id)] = data


def recall_render(forwarder_id: str | None, shipment_id: str | None) -> Optional[bytes]:
    if not forwarder_id or not shipment_id:
        return None
    return _CACHE.get(_key(forwarder_id, shipment_id))


def _adapter_id(forwarder_id: str | None) -> str:
    if not forwarder_id:
        return ""
    from fba_tool.invoice_export import export_forwarders

    return str((export_forwarders().get(forwarder_id) or {}).get("export_adapter") or "")


def render_adapter_invoice(
    forwarder_id: str,
    csv_bytes: bytes,
    parsed_shipment: dict | None = None,
) -> Dict[str, Any]:
    """和 /api/invoice-export/export 同一套解析、产品、仓库和适配器。"""
    from fba_tool.invoice_export import export_invoice
    from fba_tool.master_db import lookup_by_fnsku
    from fba_tool.shipment_parser import parse_csv_shipment
    from fba_tool.shipment_serialize import resolve_csv_shipment
    from fba_tool.warehouse_db import get_warehouse
    from routers.invoice_export import resolve_export_warehouse

    if csv_bytes:
        ship = parse_csv_shipment(csv_bytes, file_name=f"{forwarder_id}.csv", allow_unassigned_fba=True)
    else:
        ship = resolve_csv_shipment(parsed_shipment=parsed_shipment, csv_bytes=None)
    products: Dict[str, Any] = {}
    fnskus = set()
    if not ship.box_items.empty and "FNSKU" in ship.box_items.columns:
        fnskus = {str(x).upper() for x in ship.box_items["FNSKU"] if str(x or "").strip()}
    for fnsku in fnskus:
        product = lookup_by_fnsku(fnsku)
        if product:
            products[fnsku] = product
    code, _source = resolve_export_warehouse(forwarder_id, "", ship.ship_to, ship.warehouse_code)
    result = export_invoice(
        forwarder_id,
        shipment=ship,
        products=products,
        reference_id="",
        warehouse=code,
        warehouse_info=get_warehouse(code) if code else None,
        order_no="",
        channel="",
        invoice_date=None,
    )
    result["shipment"] = ship
    return result


def _json_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _slave_coords(ws) -> set:
    slaves = set()
    for merged in ws.merged_cells.ranges:
        if merged.max_row > 20000 or merged.max_col > 512:
            continue
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if (row, col) != (merged.min_row, merged.min_col):
                    slaves.add((row, col))
    return slaves


def preview_cells(template: bytes, rendered: bytes, sheet_name: str) -> List[Dict[str, Any]]:
    """只列出这一张表上和模板不同的主格，供预览铺格。导出仍用渲染好的整本。"""
    blank = openpyxl.load_workbook(io.BytesIO(template))
    filled = openpyxl.load_workbook(io.BytesIO(rendered))
    try:
        if sheet_name not in filled.sheetnames:
            return []
        ws = filled[sheet_name]
        src = blank[sheet_name] if sheet_name in blank.sheetnames else None
        slaves = _slave_coords(ws)
        cells: List[Dict[str, Any]] = []
        max_row = min(ws.max_row or 1, 80)
        max_col = min(ws.max_column or 1, 40)
        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            for cell in row:
                if isinstance(cell, MergedCell) or (cell.row, cell.column) in slaves:
                    continue
                if _is_formula(cell):
                    continue
                old = src[cell.coordinate].value if src is not None else None
                if _values_equal(old, cell.value):
                    continue
                cells.append(
                    {
                        "cell": cell.coordinate,
                        "value": _json_cell(cell.value),
                        "source": "精准提取",
                        "field_key": cell.coordinate,
                        "required": False,
                    }
                )
        return cells
    finally:
        blank.close()
        filled.close()


def try_adapter_draft(
    forwarder_id: str | None,
    csv_bytes: bytes,
    template_bytes: bytes,
    parsed_shipment: dict | None,
) -> Optional[InvoiceDraft]:
    if not _adapter_id(forwarder_id):
        return None
    from fba_tool.xlsx_io import load_workbook_fast

    rendered = render_adapter_invoice(forwarder_id or "", csv_bytes, parsed_shipment)
    remember_render(forwarder_id or "", rendered["shipment"].shipment_id, rendered["bytes"])
    wb = load_workbook_fast(template_bytes)
    sheet = wb.active.title
    wb.close()
    warnings = []
    for item in rendered.get("warnings") or []:
        if isinstance(item, dict):
            warnings.append(str(item.get("message") or item.get("code") or ""))
        elif item:
            warnings.append(str(item))
    return InvoiceDraft(
        shipment_id=rendered["shipment"].shipment_id,
        warehouse_code=rendered["shipment"].warehouse_code or "",
        fields=[],
        template_sheet=sheet,
        table_cells=preview_cells(template_bytes, rendered["bytes"], sheet),
        is_table_mode=True,
        warnings=[w for w in warnings if w],
    )


def export_rendered_workbook(
    forwarder_id: str | None,
    shipment_id: str | None,
    template_bytes: bytes,
    snapshot: Optional[Dict[str, Any]],
) -> Optional[bytes]:
    """有生成结果时导出那本正式发票。快照里多出来的手改写回去，模板残留不覆盖。"""
    rendered = recall_render(forwarder_id, shipment_id)
    if rendered is None:
        return None
    if not snapshot:
        return rendered
    return _apply_user_edits(rendered, template_bytes, snapshot)


def _apply_user_edits(rendered: bytes, template: bytes, snapshot: Dict[str, Any]) -> bytes:
    filled = openpyxl.load_workbook(io.BytesIO(rendered))
    blank = openpyxl.load_workbook(io.BytesIO(template))
    formula_cells: List[str] = []
    slave_cells: List[str] = []
    try:
        by_name = {ws.title: ws for ws in filled.worksheets}
        blank_by_name = {ws.title: ws for ws in blank.worksheets}
        for sheet in _snapshot_sheets(snapshot):
            ws = by_name.get(sheet.get("name"))
            src = blank_by_name.get(sheet.get("name"))
            if ws is None:
                continue
            slaves = _slave_coords(ws)
            for row, col, payload in _iter_cells(sheet):
                ref = f"{get_column_letter(col)}{row}"
                snap = payload.get("v")
                if (row, col) in slaves:
                    master_coord = None
                    for merged in ws.merged_cells.ranges:
                        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                            master_coord = (merged.min_row, merged.min_col)
                            break
                    master = ws.cell(*master_coord) if master_coord else ws.cell(row, col)
                    template_master = src.cell(*master_coord) if src is not None and master_coord else None
                    template_value = template_master.value if template_master is not None else None
                    if _values_equal(master.value, snap) or _values_equal(template_value, snap):
                        continue
                    slave_cells.append(ref)
                    continue
                cell = ws.cell(row, col)
                template_value = src.cell(row, col).value if src is not None else None
                if _is_formula(cell):
                    snap_f = payload.get("f")
                    if snap_f and _norm_formula(snap_f) == _norm_formula(cell.value):
                        continue
                    if _values_equal(template_value, snap) or (snap_f and _norm_formula(snap_f) == _norm_formula(template_value)):
                        continue
                    if not snap_f and "v" in payload and not _values_equal(cell.value, snap):
                        formula_cells.append(ref)
                    continue
                if "v" not in payload and "f" not in payload:
                    continue
                if _values_equal(cell.value, snap) or _values_equal(template_value, snap):
                    continue
                _write(cell, payload)
        if formula_cells or slave_cells:
            code = "FORMULA_OVERWRITE" if formula_cells else "MERGED_SLAVE"
            cells = formula_cells or slave_cells
            label = "公式单元格不能改成纯值" if code == "FORMULA_OVERWRITE" else "不能写入合并单元格的从属格"
            raise ApplyRejected(code, cells, f"{label}：{', '.join(cells[:30])}")
        buf = io.BytesIO()
        filled.save(buf)
        return buf.getvalue()
    finally:
        filled.close()
        blank.close()
