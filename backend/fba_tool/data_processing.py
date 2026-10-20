from __future__ import annotations

import io
import re
import unicodedata
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook

REQUIRED_COLUMNS = [
    "SKU简称",
    "发货总数",
    "单箱装量",
    "箱长(cm)",
    "箱宽(cm)",
    "箱高(cm)",
    "单个重量(kg)",
    "空纸箱重(kg)",
    "单箱毛重(kg)",
]

OPTIONAL_COLUMNS = ["FNSKU"]
INPUT_COLUMNS = ["SKU简称", "FNSKU"] + REQUIRED_COLUMNS[1:]

COLUMN_ALIASES: Dict[str, List[str]] = {
    "SKU简称": ["sku简称", "sku", "产品", "product", "seller sku", "msku"],
    "FNSKU": ["fnsku"],
    "发货总数": ["发货总数", "总数", "总数量", "qty", "quantity"],
    "单箱装量": ["单箱装量", "单箱", "装箱数", "每箱数量", "carton qty"],
    "箱长(cm)": ["箱长(cm)", "箱长", "长(cm)", "长", "length"],
    "箱宽(cm)": ["箱宽(cm)", "箱宽", "宽(cm)", "宽", "width"],
    "箱高(cm)": ["箱高(cm)", "箱高", "高(cm)", "高", "height"],
    "单个重量(kg)": ["单个重量(kg)", "单个重量", "单重", "净重", "unit weight"],
    "空纸箱重(kg)": ["空纸箱重(kg)", "空纸箱重", "纸箱重", "皮重", "tare"],
    "单箱毛重(kg)": ["单箱毛重(kg)", "单箱毛重", "毛重", "gross weight"],
}


def empty_input_df() -> pd.DataFrame:
    return pd.DataFrame(columns=INPUT_COLUMNS)


# ---- 文本录入解析（S1c，方案 3.1）-------------------------------------------
# 数值列：(列名, 是否整数, 是否必填且必须 > 0)
_TEXT_NUMERIC_FIELDS: Dict[str, Tuple[bool, bool]] = {
    "发货总数": (True, True),
    "单箱装量": (True, True),
    "箱长(cm)": (False, True),
    "箱宽(cm)": (False, True),
    "箱高(cm)": (False, True),
    "单个重量(kg)": (False, False),
    "空纸箱重(kg)": (False, False),
    "单箱毛重(kg)": (False, False),
}
_TEXT_REQUIRED = ["SKU简称", "发货总数", "单箱装量", "箱长(cm)", "箱宽(cm)", "箱高(cm)"]
# 表头里必须有的列（S2：单箱装量、箱规可以整列不给，由产品库补齐）
_HEADER_REQUIRED = ["SKU简称", "发货总数"]
_CARTON_DIMS = ("箱长(cm)", "箱宽(cm)", "箱高(cm)")
_POSITIONAL = [c for c in INPUT_COLUMNS if c != "FNSKU"]
_FNSKU_RE = re.compile(r"^(X0|B0)[0-9A-Z]{8}$")
_NUM_RE = re.compile(r"^-?\d+(?:\.\d+)?$")
_NUM_PREFIX_RE = re.compile(r"^(-?\d+(?:\.\d+)?)\s*([^\d\s.,-][^\d]*)$")


def _split_text_line(line: str) -> List[str]:
    """Tab 优先（保留空单元格，Excel 复制粘贴不会错列）；其次逗号；最后连续空白。"""
    if "\t" in line:
        return [p.strip() for p in line.split("\t")]
    if "," in line or "，" in line:
        return [p.strip() for p in re.split(r"[,，]", line)]
    return [p for p in re.split(r"\s+", line.strip()) if p != ""]


_THOUSANDS_IN_COMMA_LINE = re.compile(r"(?:^|[,，])\s*\d{1,3}[,，]\d{3}\s*(?=[,，]|$)")
_THOUSANDS_HINT = "这一行用逗号分隔，又出现了“1,000”这类千分位写法，千分位的逗号会被当成分隔符导致错列；请去掉数字里的逗号，或改用 Tab 分隔（从 Excel 直接复制粘贴就是 Tab）"


def _thousands_hint(line: str, line_no: int, errors: list, start: int, warnings: list) -> None:
    """逗号分隔且疑似含千分位时：给本行已有错误补一句说明；没有错误也加一条提示。"""
    if "\t" in line or not _THOUSANDS_IN_COMMA_LINE.search(line):
        return
    new_errors = errors[start:]
    for e in new_errors:
        if "千分位" not in e["message"]:
            e["message"] = f"{e['message']}。可能原因：{_THOUSANDS_HINT}"
    if not new_errors:
        warnings.append(_issue(line_no, "整行", line.strip(), f"第{line_no}行疑似千分位：{_THOUSANDS_HINT}"))


def _header_mapping(parts: List[str]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for idx, raw in enumerate(parts):
        norm = _normalize_header(unicodedata.normalize("NFKC", raw))
        if not norm:
            continue
        for field, aliases in COLUMN_ALIASES.items():
            if field not in mapping and (norm in aliases or norm == field.lower()):
                mapping[field] = idx
                break
    return mapping


def _clean_number(raw: str, field: str, row_no: int, errors: list, warnings: list):
    """返回数值或 None（空）；非法时写 errors 并返回 _INVALID。"""
    text = unicodedata.normalize("NFKC", str(raw)).strip()
    if text == "":
        return None
    candidate = text.replace(",", "")
    if not _NUM_RE.match(candidate):
        m = _NUM_PREFIX_RE.match(candidate)
        if not m:
            errors.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」的值“{raw}”不是有效数字"))
            return _INVALID
        candidate = m.group(1)
        warnings.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」的值“{raw}”带了单位，已按 {candidate} 读取，请核对"))
    is_int, positive = _TEXT_NUMERIC_FIELDS[field]
    value = float(candidate)
    if is_int:
        if value != int(value):
            errors.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」必须是整数，现在是“{raw}”"))
            return _INVALID
        value = int(value)
    if positive and value <= 0:
        errors.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」必须大于 0，现在是“{raw}”"))
        return _INVALID
    if not positive and value < 0:
        errors.append(_issue(row_no, field, raw, f"第{row_no}行「{field}」不能是负数，现在是“{raw}”"))
        return _INVALID
    return value


_INVALID = object()


_MASTER_FILLABLE = ("单箱装量", "箱长(cm)", "箱宽(cm)", "箱高(cm)")


def _validate_cells(
    cells: Dict[str, Any],
    line_no: int,
    errors: list,
    warnings: list,
    master_fillable: bool = False,
) -> Optional[dict]:
    """文本录入与 Excel 导入共用的逐行校验；有错误返回 None（错误写进 errors）。

    master_fillable=True（Excel 样表）：单箱装量、箱规留空时不报错，给提示并留空，
    由产品库补齐（apply_master_to_input_rows）；产品库也没有时 Step1 校验会拦下。
    填了非法值（abc、负数、0、小数件数）一律报错。
    """
    row = _blank_row()
    before = len(errors)
    sku = unicodedata.normalize("NFKC", str(cells.get("SKU简称", "") or "")).strip()
    if not sku:
        errors.append(_issue(line_no, "SKU简称", "", f"第{line_no}行缺少 SKU简称"))
    row["SKU简称"] = sku
    row["FNSKU"] = unicodedata.normalize("NFKC", str(cells.get("FNSKU", "") or "")).strip()
    for field in _TEXT_NUMERIC_FIELDS:
        raw = cells.get(field, "")
        value = _clean_number("" if raw is None else raw, field, line_no, errors, warnings)
        if value is _INVALID:
            continue
        if value is None:
            if master_fillable and field in _MASTER_FILLABLE:
                warnings.append(_issue(line_no, field, "", f"第{line_no}行「{field}」为空，将从产品库补齐"))
            elif field in _TEXT_REQUIRED:
                errors.append(_issue(line_no, field, "", f"第{line_no}行缺少必填的「{field}」"))
            continue
        row[field] = value
    # 箱长宽高要么都填、要么都空（都空才由产品库补；只填一部分会和产品库的值混在一起）
    filled = [f for f in _CARTON_DIMS if str(cells.get(f) if cells.get(f) is not None else "").strip() != ""]
    if filled and len(filled) < len(_CARTON_DIMS):
        missing = [f for f in _CARTON_DIMS if f not in filled]
        errors.append(_issue(line_no, "箱规", "", f"第{line_no}行箱长/宽/高只填了一部分，缺「{'、'.join(missing)}」；要么都填，要么都留空由产品库补齐"))
        warnings[:] = [w for w in warnings if not (w["row"] == line_no and w["field"] in _CARTON_DIMS)]
    return row if len(errors) == before else None


def _issue(row_no: int, field: str, value, message: str) -> Dict[str, Any]:
    return {"row": int(row_no), "field": field, "value": "" if value is None else str(value), "message": message}


def parse_smart_text_detailed(raw: str) -> Tuple[pd.DataFrame, List[Dict[str, Any]], List[Dict[str, Any]]]:
    """解析粘贴的文本，返回 (数据表, 错误列表, 提示列表)。

    - 第一行如果是表头（能认出 SKU简称 等列名），按表头定位列，顺序随意，FNSKU 列可有可无；
    - 没有表头时按样表列序：SKU简称, [FNSKU], 发货总数, 单箱装量, 箱长, 箱宽, 箱高, 单个重量, 空纸箱重, 单箱毛重；
      第 2 列形如 X0xxxxxxxx / B0xxxxxxxx 时视为 FNSKU；
    - 分隔符：Tab、逗号（含全角）、连续空格；
    - 任何解析失败都写进错误列表 {row, field, value, message}，**不再静默填 0**；有错误的行不进入结果。
    """
    rows: List[dict] = []
    errors: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    mapping: Optional[Dict[str, int]] = None
    lines = (raw or "").replace("\ufeff", "").splitlines()
    first_content = True
    for line_no, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        parts = _split_text_line(line)
        if first_content:
            first_content = False
            header = _header_mapping(parts)
            if "SKU简称" in header and len(header) >= 2:
                missing = [c for c in _HEADER_REQUIRED if c not in header]
                if missing:
                    errors.append(_issue(line_no, "表头", line.strip(), f"第{line_no}行表头缺少必填列：{'、'.join(missing)}"))
                    return empty_input_df(), errors, warnings
                mapping = header
                continue
        if mapping is not None:
            cells = {field: (parts[idx] if idx < len(parts) else "") for field, idx in mapping.items()}
        else:
            values = list(parts)
            has_fnsku = len(values) >= 2 and bool(_FNSKU_RE.match(unicodedata.normalize("NFKC", values[1]).upper()))
            cols = (["SKU简称", "FNSKU"] + _POSITIONAL[1:]) if has_fnsku else _POSITIONAL
            need = 3 if has_fnsku else 2  # S2：单箱装量、箱规可留空，由产品库补齐
            if len(values) < need:
                errors.append(_issue(line_no, "整行", line.strip(),
                    f"第{line_no}行只有 {len(values)} 列，至少需要 {need} 列（SKU简称、{'FNSKU、' if has_fnsku else ''}发货总数）"))
                _thousands_hint(line, line_no, errors, len(errors) - 1, warnings)
                continue
            if len(values) > len(cols):
                errors.append(_issue(line_no, "整行", line.strip(), f"第{line_no}行有 {len(values)} 列，超过 {len(cols)} 列，无法确定每列含义，请加表头"))
                _thousands_hint(line, line_no, errors, len(errors) - 1, warnings)
                continue
            cells = dict(zip(cols, values))
        errors_before_row = len(errors)
        row = _validate_cells(cells, line_no, errors, warnings, master_fillable=True)
        if row is not None:
            rows.append(row)
        _thousands_hint(line, line_no, errors, errors_before_row, warnings)
    if not rows:
        return empty_input_df(), errors, warnings
    return pd.DataFrame(rows, columns=INPUT_COLUMNS), errors, warnings


def parse_smart_text(raw: str) -> pd.DataFrame:
    """兼容旧调用：只返回解析成功的行（坏行不再以 0 混进结果）。需要错误明细请用 parse_smart_text_detailed。"""
    df, _errors, _warnings = parse_smart_text_detailed(raw)
    return df


def load_from_excel_detailed(
    file_bytes: bytes,
) -> Tuple[pd.DataFrame, Dict[str, str], List[Dict[str, Any]], List[Dict[str, Any]]]:
    """解析发货录入样表，返回 (数据表, SKU->FNSKU, 错误, 提示)。

    与文本录入同一口径（_validate_cells）：非数字、负数、0、小数件数、缺必填都按 Excel 行号报错，
    有错误的行不进入结果，**不再静默填 0**（D3）。
    """
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    header_row, col_map = _detect_header_map(ws)
    if not col_map or "SKU简称" not in col_map:
        df, fn = _load_from_excel_fallback(file_bytes)
        return df, fn, [], []

    rows: List[dict] = []
    fnsku_map: Dict[str, str] = {}
    errors: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        cells = {field: ws.cell(row=r, column=col_idx).value for field, col_idx in col_map.items()}
        sku = str(cells.get("SKU简称") or "").strip()
        if not sku:
            continue
        if sku.startswith("填写") or sku.startswith("示例"):
            continue
        row = _validate_cells(cells, r, errors, warnings, master_fillable=True)
        if row is None:
            continue
        if row.get("FNSKU"):
            fnsku_map[row["SKU简称"]] = row["FNSKU"]
        rows.append(row)

    if not rows:
        return empty_input_df(), fnsku_map, errors, warnings
    return pd.DataFrame(rows, columns=INPUT_COLUMNS), fnsku_map, errors, warnings


def load_from_excel(file_bytes: bytes) -> Tuple[pd.DataFrame, Dict[str, str]]:
    """兼容旧调用：有错误时抛 ValueError（含行号），不再静默填 0。"""
    df, fnsku_map, errors, _warnings = load_from_excel_detailed(file_bytes)
    if errors:
        raise ValueError("；".join(e["message"] for e in errors[:20]))
    return df, fnsku_map


def with_computed_columns(df: pd.DataFrame) -> pd.DataFrame:
    work = df.copy()
    for col in INPUT_COLUMNS:
        if col not in work.columns:
            work[col] = "" if col == "FNSKU" else (None if col == "单箱毛重(kg)" else 0)
    work["发货总数"] = work["发货总数"].apply(_to_int)
    work["单箱装量"] = work["单箱装量"].apply(_to_int)
    work["单个重量(kg)"] = work["单个重量(kg)"].apply(_to_float)
    work["空纸箱重(kg)"] = work["空纸箱重(kg)"].apply(_to_float)
    work["单箱毛重(kg)"] = work["单箱毛重(kg)"].apply(_to_optional_float)
    if work.empty:
        work["预计整箱数"] = pd.Series(dtype=int)
        work["尾数"] = pd.Series(dtype=int)
        work["预计总净重(kg)"] = pd.Series(dtype=float)
        work["预计总毛重(kg)"] = pd.Series(dtype=float)
        return work
    qty = work["单箱装量"].astype(int)
    total = work["发货总数"].astype(int)
    work["预计整箱数"] = np.where(qty > 0, total // qty, 0)
    work["尾数"] = np.where(qty > 0, total % qty, 0)
    work["预计总净重(kg)"] = (work["发货总数"] * work["单个重量(kg)"]).round(3)
    from fba_tool.box_weight import carton_gross_kg

    totals = []
    for _, row in work.iterrows():
        full = int(row["预计整箱数"])
        rem = int(row["尾数"])
        per = int(row["单箱装量"])
        unit = float(row["单个重量(kg)"] or 0)
        tare = float(row["空纸箱重(kg)"] or 0)
        declared = row["单箱毛重(kg)"]
        gross = 0.0
        if full > 0 and per > 0:
            gross += carton_gross_kg(per, unit, tare, declared, per) * full
        if rem > 0:
            gross += carton_gross_kg(rem, unit, tare, None, per)
        totals.append(round(gross, 3))
    work["预计总毛重(kg)"] = totals
    return work


def compute_step1_logistics_totals(
    df: pd.DataFrame,
    *,
    throw_ratio_val: float | None = None,
) -> dict[str, float]:
    """Step1 汇总：预计总方数 (CBM) 与预计总体积重 (kg)。"""
    from fba_tool.user_config import throw_ratio as config_throw_ratio
    from fba_tool.volume_weight import volume_weight_kg

    work = with_computed_columns(df)
    ratio = float(throw_ratio_val if throw_ratio_val is not None else config_throw_ratio())

    total_net = float(work["预计总净重(kg)"].sum()) if not work.empty else 0.0
    total_gross = float(work["预计总毛重(kg)"].sum()) if not work.empty else 0.0

    total_cbm = 0.0
    total_volumetric_weight = 0.0

    for _, row in work.iterrows():
        if not str(row.get("SKU简称", "")).strip():
            continue
        full = int(row["预计整箱数"])
        rem = int(row["尾数"])
        box_count = full + (1 if rem > 0 else 0)
        if box_count <= 0:
            continue

        l = _to_float(row["箱长(cm)"])
        w = _to_float(row["箱宽(cm)"])
        h = _to_float(row["箱高(cm)"])
        if l <= 0 or w <= 0 or h <= 0:
            continue

        per_cbm = l * w * h / 1_000_000
        per_vol = volume_weight_kg(l, w, h, throw_ratio=ratio)
        total_cbm += round(per_cbm * box_count, 3)
        total_volumetric_weight += round(per_vol * box_count, 3)

    from fba_tool.logistics_metrics import compute_logistics_metrics

    metrics = compute_logistics_metrics(
        total_net_kg=total_net,
        total_gross_kg=total_gross,
        total_cbm=round(total_cbm, 3),
        total_volumetric_weight_kg=round(total_volumetric_weight, 3),
    )
    return metrics


def validate_step1(df: pd.DataFrame) -> List[Tuple[int, str]]:
    errors: List[Tuple[int, str]] = []
    required = [
        "SKU简称",
        "发货总数",
        "单箱装量",
        "箱长(cm)",
        "箱宽(cm)",
        "箱高(cm)",
    ]
    for idx, row in df.iterrows():
        if not str(row.get("SKU简称", "")).strip():
            continue
        for col in required:
            val = row.get(col, "")
            if col == "单箱装量":
                if not _is_positive_int(val):
                    errors.append((idx, col))
            elif col == "发货总数":
                if _to_int(val) <= 0:
                    errors.append((idx, col))
            elif col in {"箱长(cm)", "箱宽(cm)", "箱高(cm)"}:
                if _to_float(val) <= 0:
                    errors.append((idx, col))
            elif str(val).strip() in {"", "nan", "None"}:
                errors.append((idx, col))
    return errors


# 录入毛重与理论毛重的可接受偏差；超过即提示核对（软告警，不阻断）
GROSS_WEIGHT_TOLERANCE = 0.08


def _unit_weight_missing(value: Any) -> bool:
    """单重为空或 0 视同缺失。"""
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    if isinstance(value, str) and not value.strip():
        return True
    return _to_float(value) <= 0


def _is_positive_int(value: Any) -> bool:
    """单箱装量必须是正整数。2.5、0、空都不是。"""
    if isinstance(value, bool) or value is None:
        return False
    try:
        if pd.isna(value):
            return False
    except (TypeError, ValueError):
        pass
    if isinstance(value, int):
        return value > 0
    if isinstance(value, float):
        return value > 0 and value == int(value)
    text = str(value).strip()
    if not text:
        return False
    try:
        number = float(text)
    except ValueError:
        return False
    return number > 0 and number == int(number)


def check_gross_weight_consistency(
    df: pd.DataFrame,
    tolerance: float = GROSS_WEIGHT_TOLERANCE,
) -> List[Dict[str, Any]]:
    """单箱毛重三档：

    1a 单重缺失且录了毛重：要确认，不给建议值（空箱重不是建议毛重）。
    1b 单重有值但毛重空着：只警告。
    规则 2 两边都有且不一致：要确认，并给出计算值。
    tolerance 保留给旧调用方，不再把小偏差静默掉。
    """
    del tolerance  # 不一致就要确认，不再按 8% 静默
    warnings: List[Dict[str, Any]] = []
    if df.empty:
        return warnings

    for idx, row in df.iterrows():
        sku = str(row.get("SKU简称", "")).strip()
        if not sku:
            continue
        declared = _to_optional_float(row.get("单箱毛重(kg)"))
        unit_missing = _unit_weight_missing(row.get("单个重量(kg)"))
        per_box = _to_int(row.get("单箱装量"))
        unit_kg = _to_float(row.get("单个重量(kg)"))
        tare_kg = _to_float(row.get("空纸箱重(kg)"))
        expected = round(per_box * unit_kg + tare_kg, 3)

        if unit_missing and declared is not None and declared > 0:
            warnings.append(
                {
                    "row_index": int(idx),
                    "sku": sku,
                    "kind": "unit_weight_missing",
                    "requires_confirm": True,
                    "suggested_gross_kg": None,
                    "declared_kg": round(declared, 3),
                    "expected_kg": expected,
                    "diff_kg": None,
                    "deviation_pct": None,
                    "message": (
                        f"SKU {sku} 录了单箱毛重 {declared:g} kg，但单个重量为空。"
                        f"请补单重，或确认沿用录入毛重。不要把空箱重当成建议毛重"
                    ),
                }
            )
            continue

        if not unit_missing and (declared is None or declared <= 0):
            warnings.append(
                {
                    "row_index": int(idx),
                    "sku": sku,
                    "kind": "gross_missing",
                    "requires_confirm": False,
                    "suggested_gross_kg": expected,
                    "declared_kg": None,
                    "expected_kg": expected,
                    "diff_kg": None,
                    "deviation_pct": None,
                    "message": (
                        f"SKU {sku} 没有填单箱毛重。按「单箱装量 × 单个重量 + 空纸箱重」"
                        f"可算出 {expected:g} kg"
                    ),
                }
            )
            continue

        if declared is None or expected <= 0:
            continue
        if abs(declared - expected) <= 0.001:
            continue
        deviation = abs(declared - expected) / expected
        warnings.append(
            {
                "row_index": int(idx),
                "sku": sku,
                "kind": "gross_mismatch",
                "requires_confirm": True,
                "suggested_gross_kg": expected,
                "declared_kg": round(declared, 3),
                "expected_kg": expected,
                "diff_kg": round(declared - expected, 3),
                "deviation_pct": round(deviation * 100, 1),
                "message": (
                    f"SKU {sku} 录入单箱毛重 {declared:g} kg，按「单箱装量 × 单个重量 + 空纸箱重」"
                    f"应为 {expected:g} kg，偏差 {deviation * 100:.1f}%。请确认用哪一个"
                ),
            }
        )
    return warnings


def _detect_header_map(ws) -> Tuple[int, Dict[str, int]]:
    for r in range(1, min(ws.max_row, 15) + 1):
        mapping: Dict[str, int] = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            raw = ws.cell(row=r, column=c).value
            if raw is None:
                continue
            norm = _normalize_header(str(raw))
            for field, aliases in COLUMN_ALIASES.items():
                if norm in aliases and field not in mapping:
                    mapping[field] = c
        if "SKU简称" in mapping and len(mapping) >= 4:
            return r, mapping
    return -1, {}


def _normalize_header(text: str) -> str:
    text = text.strip().lower()
    text = text.replace("*", "").replace("（", "(").replace("）", ")")
    text = re.sub(r"\s+", "", text)
    return text


def _blank_row() -> dict:
    return {col: (None if col == "单箱毛重(kg)" else ("" if col == "FNSKU" else 0)) for col in INPUT_COLUMNS}


def _load_from_excel_fallback(file_bytes: bytes) -> Tuple[pd.DataFrame, Dict[str, str]]:
    df = pd.read_excel(io.BytesIO(file_bytes))
    renamed = df.rename(
        columns={
            "SKU": "SKU简称",
            "产品": "SKU简称",
            "总数": "发货总数",
            "总数量": "发货总数",
            "单箱": "单箱装量",
            "长": "箱长(cm)",
            "宽": "箱宽(cm)",
            "高": "箱高(cm)",
        }
    )
    for col in INPUT_COLUMNS:
        if col not in renamed.columns:
            renamed[col] = "" if col == "FNSKU" else (None if col == "单箱毛重(kg)" else "")
    out = renamed[INPUT_COLUMNS]
    fnsku_map = {
        str(r["SKU简称"]): str(r["FNSKU"])
        for _, r in out.iterrows()
        if str(r.get("SKU简称", "")).strip() and str(r.get("FNSKU", "")).strip()
    }
    return out, fnsku_map


def _to_int(value) -> int:
    try:
        if pd.isna(value):
            return 0
        return int(float(value))
    except Exception:
        return 0


def _to_float(value) -> float:
    try:
        if pd.isna(value) or value == "":
            return 0.0
        return float(value)
    except Exception:
        return 0.0


def _to_optional_float(value):
    try:
        if pd.isna(value) or value == "":
            return None
        return float(value)
    except Exception:
        return None
