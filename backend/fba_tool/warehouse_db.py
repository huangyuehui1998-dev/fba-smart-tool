"""fba_warehouse 仓库地址库（琪珞模板「！」表导入，先预览后写入）与 customs_rule_profile 报关规则。"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone
from typing import Optional, Any, Dict, List, Tuple

WAREHOUSE_FIELDS = ("name", "company", "address1", "address2", "address3", "city", "state", "zip", "country", "region", "phone")
_HEADERS = {
    "code": ("FBA仓码",),
    "name": ("收件人名称",),
    "company": ("收件人公司",),
    "zip": ("邮编",),
    "address1": ("地址1",),
    "address2": ("地址2",),
    "address3": ("地址3",),
    "country": ("国家",),
    "city": ("城市",),
    "state": ("州", "省/州", "State"),
    "phone": ("收件人电话",),
}
_WEST = {"CA", "WA", "OR", "NV", "AZ", "UT", "ID", "MT", "WY", "CO", "NM", "AK", "HI"}
_CENTRAL = {"TX", "OK", "KS", "NE", "SD", "ND", "MN", "IA", "MO", "AR", "LA", "WI", "IL", "IN", "MI", "OH", "KY", "TN", "MS", "AL"}
# 地址库收全部行（含 Walmart、加拿大仓、同仓备用地址如 MDT2-1）；进仓库白名单的只有亚马逊格式代码
WHITELIST_CODE_RE = re.compile(r"^(?:[A-Z]{3}\d{1,2}|[A-Z0-9]{4})$")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def region_for_state(state: str) -> str:
    s = (state or "").upper()
    if not s:
        return ""
    if s in _WEST:
        return "美西"
    if s in _CENTRAL:
        return "美中"
    return "美东"


_CORR_FIELDS = ("company", "address1", "address2", "address3", "city", "state", "zip")


def load_corrections() -> Dict[str, Dict[str, Any]]:
    """config/fba_warehouse_corrections.json：按公开资料核实过的纠错（带出处）。"""
    import json

    from fba_tool.app_paths import bundled_config_dir

    p = bundled_config_dir() / "fba_warehouse_corrections.json"
    if not p.is_file():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"配置文件格式错误：{p.name} 第 {exc.lineno} 行（{exc.msg}）") from exc
    return {str(k).upper(): v for k, v in (data.get("corrections") or {}).items()}


def apply_correction(rec: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    if not rec:
        return rec
    c = load_corrections().get(str(rec.get("code") or "").upper())
    if not c:
        return rec
    out = dict(rec)
    for k in _CORR_FIELDS:
        if k in c:
            out[k] = c[k]
    out["region"] = region_for_state(out.get("state") or "")
    out["corrected"] = {"was": c.get("was", ""), "reason": c.get("reason", ""), "source": c.get("source", "")}
    return out


def parse_warehouse_sheet(data: bytes, file_name: str = "") -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """读取含「FBA仓码」表头的工作表（琪珞模板「！」表）。返回 (rows, issues)。"""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{file_name or '文件'}（{exc.__class__.__name__}）") from exc
    sheets = [wb["！"]] if "！" in wb.sheetnames else list(wb.worksheets)
    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True))
        for hi, header in enumerate(rows[:20]):
            texts = [_clean(c) for c in header]
            col: Dict[str, int] = {}
            for key, names in _HEADERS.items():
                for i, t in enumerate(texts):
                    if any(t.startswith(n) for n in names):
                        col.setdefault(key, i)
            if "code" not in col:
                continue
            if "state" not in col and "city" in col:
                col["state"] = col["city"] + 1  # 琪珞表州列表头写成了“TX,”，按“城市”右边一列取
            out: Dict[str, Dict[str, Any]] = {}
            issues: List[Dict[str, Any]] = []
            for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
                def g(k: str) -> str:
                    i = col.get(k)
                    return _clean(r[i]) if i is not None and i < len(r) else ""

                code = g("code").upper()
                if not code:
                    continue
                if len(code) > 40:
                    issues.append({"row": ri, "field": "code", "value": code, "message": f"第{ri}行仓库代码太长：{code}"})
                    continue
                zip_code = g("zip")
                if zip_code.isdigit() and len(zip_code) < 5:
                    zip_code = zip_code.zfill(5)
                state = g("state").upper().rstrip(",")
                rec = {
                    "code": code,
                    "name": g("name"),
                    "company": g("company"),
                    "address1": g("address1"),
                    "address2": g("address2"),
                    "address3": g("address3"),
                    "city": g("city"),
                    "state": state,
                    "zip": zip_code,
                    "country": g("country").split("-")[0] or "US",
                    "region": region_for_state(state),
                    "phone": g("phone"),
                }
                if code in out:
                    issues.append({"row": ri, "field": "code", "value": code, "message": f"第{ri}行仓库代码重复：{code}（取后一行）"})
                fixed = apply_correction(rec)
                if fixed is not rec:
                    issues.append({"row": ri, "field": "address", "value": code, "level": "info",
                                   "message": f"第{ri}行 {code} 地址按公开资料纠正为 {fixed['address1']}, {fixed['city']}, {fixed['state']} {fixed['zip']}（{fixed['corrected']['source'][:60]}…）"})
                    fixed = {k: v for k, v in fixed.items() if k != "corrected"}
                out[code] = fixed
            return list(out.values()), issues
    raise ValueError("没找到含「FBA仓码」表头的工作表（应为琪珞发票模板的「！」表）")


def _conn():
    from fba_tool.master_db import _connect

    return _connect()


def list_warehouse_codes(whitelist_only: bool = False) -> List[str]:
    with _conn() as conn:
        codes = [r[0] for r in conn.execute("SELECT code FROM fba_warehouse ORDER BY code").fetchall()]
    return [c for c in codes if WHITELIST_CODE_RE.match(c)] if whitelist_only else codes


def get_warehouse(code: str) -> Dict[str, Any] | None:
    with _conn() as conn:
        cur = conn.execute("SELECT * FROM fba_warehouse WHERE code = ?", (str(code or "").strip().upper(),))
        row = cur.fetchone()
        return apply_correction(dict(zip([d[0] for d in cur.description], row))) if row else None


def list_warehouses(q: str = "", limit: int = 1000) -> List[Dict[str, Any]]:
    like = f"%{q.strip().upper()}%"
    with _conn() as conn:
        cur = conn.execute(
            "SELECT * FROM fba_warehouse WHERE UPPER(code) LIKE ? OR UPPER(city) LIKE ? OR UPPER(state) LIKE ? ORDER BY code LIMIT ?",
            (like, like, like, int(limit)),
        )
        names = [d[0] for d in cur.description]
        return [apply_correction(dict(zip(names, r))) for r in cur.fetchall()]


def _diff(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    with _conn() as conn:
        cur = conn.execute("SELECT * FROM fba_warehouse")
        names = [d[0] for d in cur.description]
        existing = {r[0]: dict(zip(names, r)) for r in cur.fetchall()}
    summary = {"new": 0, "changed": 0, "unchanged": 0, "total": len(rows)}
    changed: List[Dict[str, Any]] = []
    for rec in rows:
        old = existing.get(rec["code"])
        if not old:
            summary["new"] += 1
            continue
        diffs = {f: {"old": old.get(f) or "", "new": rec[f]} for f in WAREHOUSE_FIELDS if (old.get(f) or "") != rec[f]}
        if diffs:
            summary["changed"] += 1
            changed.append({"code": rec["code"], "fields": diffs})
        else:
            summary["unchanged"] += 1
    return {"summary": summary, "changed": changed[:200]}


def preview_warehouse_import(data: bytes, file_name: str = "") -> Dict[str, Any]:
    rows, issues = parse_warehouse_sheet(data, file_name)
    out = _diff(rows)
    out["issues"] = issues
    return out


def apply_warehouse_import(data: bytes, file_name: str = "") -> Dict[str, Any]:
    rows, issues = parse_warehouse_sheet(data, file_name)
    now = _now()
    cols = ("code",) + WAREHOUSE_FIELDS + ("source", "updated_at")
    with _conn() as conn:
        for rec in rows:
            vals = [rec["code"]] + [rec[f] for f in WAREHOUSE_FIELDS] + [file_name or "import", now]
            upd = ", ".join(f"{c} = excluded.{c}" for c in cols[1:])
            conn.execute(
                f"INSERT INTO fba_warehouse ({', '.join(cols)}) VALUES ({', '.join('?' for _ in cols)}) "
                f"ON CONFLICT(code) DO UPDATE SET {upd}",
                vals,
            )
        conn.commit()
    return {"ok": True, "written": len(rows), "issues": issues}


# ---------------- customs_rule_profile ----------------
_PROFILE_TEXT = (
    "net_weight_mode", "total_price_rule", "loading_port", "entrust_valid_until_rule",
    "exemption_nature", "transport_mode", "declared_value_basis", "trade_term", "tax_fx_basis",
)
_PROFILE_NUM = ("net_minus_per_box", "net_ratio", "forwarder_gross_tolerance_kg")
_PROFILE_BOOL = ("fill_freight", "fill_export_declare_date")


def get_default_rule_profile() -> Dict[str, Any]:
    with _conn() as conn:
        cur = conn.execute("SELECT * FROM customs_rule_profile ORDER BY is_default DESC, id LIMIT 1")
        row = cur.fetchone()
        return dict(zip([d[0] for d in cur.description], row)) if row else {}


NET_WEIGHT_MODES = ("tare_from_carton", "minus_per_box", "ratio")
DECLARED_VALUE_BASES = ("retail", "fob", "cost")


TAX_FX_LOCK_MESSAGE = "出口退税折合率确定后 12 个月内不得变更"


def _plus_twelve_months(day: date) -> date:
    try:
        return day.replace(year=day.year + 1)
    except ValueError:
        return day.replace(year=day.year + 1, day=28)


def update_rule_profile(values: Dict[str, Any]) -> Dict[str, Any]:
    prof = get_default_rule_profile()
    if not prof:
        raise ValueError("没有默认报关规则")
    incoming = dict(values or {})
    confirm = bool(incoming.pop("confirm", False))
    mode = str(incoming.get("net_weight_mode") or prof.get("net_weight_mode") or "tare_from_carton").strip()
    if mode not in NET_WEIGHT_MODES:
        raise ValueError(f"净重算法只能是 {' / '.join(NET_WEIGHT_MODES)}：{mode}")
    basis = str(incoming.get("declared_value_basis") or prof.get("declared_value_basis") or "retail").strip()
    if "declared_value_basis" in incoming and basis not in DECLARED_VALUE_BASES:
        raise ValueError(f"申报价口径只能是 {' / '.join(DECLARED_VALUE_BASES)}：{basis}")
    ratio = incoming.get("net_ratio", prof.get("net_ratio"))
    if mode == "ratio" and ratio in (None, ""):
        incoming["net_ratio"] = 0.95
        ratio = 0.95
    if "tax_fx_basis" in incoming:
        new_basis = str(incoming.get("tax_fx_basis") or "").strip()
        old_basis = str(prof.get("tax_fx_basis") or "month_first").strip()
        locked_until = str(prof.get("tax_fx_locked_until") or "").strip()[:10]
        if new_basis and new_basis != old_basis and locked_until and not confirm:
            try:
                until = date.fromisoformat(locked_until)
            except ValueError:
                until = None
            if until and date.today() < until:
                raise ValueError(TAX_FX_LOCK_MESSAGE)
    sets: Dict[str, Any] = {}
    for k, v in incoming.items():
        if k in _PROFILE_TEXT:
            text = str(v or "").strip()
            if not text:
                raise ValueError(f"「{k}」不能为空")
            sets[k] = text
        elif k in _PROFILE_NUM:
            if v in (None, ""):
                sets[k] = None
                continue
            try:
                num = float(v)
            except (TypeError, ValueError):
                raise ValueError(f"「{k}」不是数字：{v}")
            if num < 0 or (k == "net_ratio" and num > 1):
                raise ValueError(f"「{k}」超出范围：{v}")
            sets[k] = num
        elif k in _PROFILE_BOOL:
            sets[k] = 1 if v in (True, 1, "1", "true", "True") else 0
        else:
            raise ValueError(f"不允许修改的字段：{k}")
    if "tax_fx_basis" in incoming:
        old_basis = str(prof.get("tax_fx_basis") or "month_first").strip()
        new_basis = str(incoming.get("tax_fx_basis") or "").strip()
        if new_basis and new_basis != old_basis:
            sets["tax_fx_locked_until"] = _plus_twelve_months(date.today()).isoformat()
    if sets:
        sets["updated_at"] = _now()
        with _conn() as conn:
            conn.execute(
                f"UPDATE customs_rule_profile SET {', '.join(f'{k} = ?' for k in sets)} WHERE id = ?",
                (*sets.values(), prof["id"]),
            )
            if "net_weight_mode" in incoming and str(incoming.get("net_weight_mode") or "") != str(prof.get("net_weight_mode") or ""):
                conn.execute(
                    """CREATE TABLE IF NOT EXISTS customs_rule_profile_history (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        profile_name TEXT,
                        field TEXT,
                        old_value TEXT,
                        new_value TEXT,
                        reason TEXT,
                        created_at TEXT
                    )"""
                )
                conn.execute(
                    """INSERT INTO customs_rule_profile_history
                       (profile_name, field, old_value, new_value, reason, created_at)
                       VALUES (?,?,?,?,?,?)""",
                    (
                        prof.get("name") or "default",
                        "net_weight_mode",
                        str(prof.get("net_weight_mode") or ""),
                        str(incoming.get("net_weight_mode") or ""),
                        "user",
                        sets["updated_at"],
                    ),
                )
            conn.commit()
    return get_default_rule_profile()
