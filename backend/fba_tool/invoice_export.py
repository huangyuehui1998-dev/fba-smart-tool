"""S4 发票按货代模板导出（方案 v8 §3.5）。

适配器只负责“写到模板哪个格子”，规则都在 invoice_rules.py。
货代在 config/forwarders.default.json 和 config/invoice_export_forwarders.default.json 里登记 export_adapter 和 invoice_export 常量。
适配器共 5 个（见 ADAPTERS）：send、jiufang、qiluo_t6、qiluo_t5、deshu_t3（德速）。
"""
from __future__ import annotations

import io
import json
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from fba_tool.invoice_rules import (
    DECLARED_CURRENCY,
    InvoiceBlocked,
    InvoiceModel,
    build_model,
    cbm,
    round2,
    to_number,
)

SEND_SHEET = "Invoice\xa0Template"  # SEND 原表的工作表名里是不换行空格（方案 §5-7）


class ExportConfigError(ValueError):
    pass


# ---------------------------------------------------------------- 配置
def _read_json(path: Path) -> Dict[str, Any]:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ExportConfigError(f"配置文件格式错误：{Path(path).name} 第 {exc.lineno} 行（{exc.msg}），请修正或删除后重试") from exc


def export_forwarders() -> Dict[str, Dict[str, Any]]:
    """登记了 export_adapter 的货代：出厂默认 + 用户 forwarders.json（同 id 时用户的覆盖）。"""
    from fba_tool.app_paths import bundled_config_dir, resolve_config_read

    out: Dict[str, Dict[str, Any]] = {}
    paths = [bundled_config_dir() / "forwarders.default.json",
             bundled_config_dir() / "invoice_export_forwarders.default.json"]  # S8 D4：琪珞 T5/T6、德速 T3
    try:
        user = resolve_config_read("forwarders.json")
        if Path(user).is_file():
            paths.append(Path(user))
    except Exception:  # noqa: BLE001
        pass
    for p in paths:
        if not Path(p).is_file():
            continue
        for item in _read_json(p).get("forwarders", []):
            if item.get("export_adapter"):
                out[item["id"]] = {**out.get(item["id"], {}), **item}
    return out


def get_export_forwarder(forwarder_id: str) -> Dict[str, Any]:
    fw = export_forwarders()
    if forwarder_id not in fw:
        raise ExportConfigError(f"不支持该货代：{forwarder_id} 没有发票导出适配器（可选：{'、'.join(fw) or '无'}）")
    return fw[forwarder_id]


def _template_bytes(fw: Dict[str, Any]) -> bytes:
    from fba_tool.app_paths import bundled_forwarders_templates_dir, user_forwarders_templates_dir

    for d in (user_forwarders_templates_dir(), bundled_forwarders_templates_dir()):
        p = Path(d) / fw["template_file"]
        if p.is_file():
            return p.read_bytes()
    raise ExportConfigError(f"货代模板不存在：{fw['template_file']}")


# ---------------------------------------------------------------- 写格子工具
def _style_row(ws, src_row: int) -> List[Any]:
    return [copy(c._style) for c in ws[src_row]]


def _apply_style(ws, row: int, styles: List[Any]) -> None:
    for i, st in enumerate(styles, start=1):
        ws.cell(row=row, column=i)._style = copy(st)


def _fmt_dim(v: Any) -> str:
    if v is None:
        return ""
    f = round2(v)
    return str(int(f)) if f == int(f) else str(f)


def _sheet(wb, name: str):
    from fba_tool.invoice_templates import norm

    for ws in wb.worksheets:
        if norm(ws.title) == norm(name):
            return ws
    raise ExportConfigError(f"模板里找不到工作表「{name}」")


# ---------------------------------------------------------------- SEND
def write_send(wb, model: InvoiceModel, fw: Dict[str, Any], ctx: Dict[str, Any]) -> None:
    cfg = fw.get("invoice_export") or {}
    ws = _sheet(wb, SEND_SHEET)
    ws.title = SEND_SHEET  # 保持原表名（含 \xa0）
    ws["C4"] = model.shipment_id
    ws["C5"] = cfg.get("duty_included", "是 Yes")
    ws["C6"] = cfg.get("use_forwarder_ior", "否 No")
    styles = _style_row(ws, 15)
    lines = {l["fnsku"]: l for l in model.lines}
    for i, r in enumerate(model.ranges):
        row = 15 + i
        _apply_style(ws, row, styles)
        l = lines[r.fnsku]
        pcs = to_number(l["pcs_per_set"]) or 1
        vals = {
            "A": r.label,                                   # 箱号区间（文本）
            "B": "*".join(_fmt_dim(x) for x in r.dims) if all(x is not None for x in r.dims) else "",
            "C": round2(r.weight) if r.weight is not None else None,
            "D": model.reference_id,
            "E": r.fnsku,
            "F": r.qty,
            "G": l["english_name"], "H": l["chinese_name"], "I": l["brand"], "J": l["material_cn"], "K": l["purpose_cn"],
            "L": l["model"], "M": str(l["hs_code"]), "N": l["asin"], "O": l["invoice_unit"],
            "P": pcs if pcs and pcs != 1 else None,        # 模板说明：套装才填每套个数
            "Q": to_number(l["purchase_price"]),           # 采购价照抄（CNY）
            "R": to_number(l["price"]),                     # 申报单价（D-01）
            "S": to_number(l["listing_price"]),
            "T": DECLARED_CURRENCY,
        }
        for col, v in vals.items():
            ws[f"{col}{row}"] = v


# ---------------------------------------------------------------- 九方
def write_jiufang(wb, model: InvoiceModel, fw: Dict[str, Any], ctx: Dict[str, Any]) -> None:
    cfg = fw.get("invoice_export") or {}
    inv, pk = _sheet(wb, "Invoice"), _sheet(wb, "Packing list")
    wh = ctx.get("warehouse_info") or {}
    code = model.warehouse
    the_date = ctx["invoice_date"]
    inv["B2"] = ctx.get("channel") or cfg.get("channel", "")
    inv["P2"] = ctx.get("order_no", "")
    inv["P3"] = model.shipment_id
    inv["P4"] = the_date
    inv["P5"] = ctx.get("order_no", "")
    shipper = cfg.get("shipper") or {}
    inv["B4"], inv["B5"], inv["B7"], inv["B8"] = shipper.get("company", ""), shipper.get("address", ""), shipper.get("contact", ""), shipper.get("tel", "")
    inv["B12"] = code
    inv["B13"] = " ".join(x for x in (wh.get("address1"), wh.get("address2"), wh.get("address3")) if x)
    inv["B15"] = wh.get("zip", "")
    inv["B16"] = wh.get("city", "")
    inv["B17"] = wh.get("state", "")
    inv["B19"] = code
    inv["B22"] = code
    data_st, total_st = _style_row(inv, 25), _style_row(inv, 26)
    for c in inv[26]:
        c.value = None
    for i, l in enumerate(model.lines):
        row = 25 + i
        _apply_style(inv, row, data_st)
        vals = {"A": model.shipment_id, "B": model.reference_id, "C": l["fnsku"], "D": l["english_name"], "E": l["chinese_name"],
                "F": l["brand"], "G": l["model"], "H": l["material_en"], "I": l["has_battery"], "J": "", "K": "", "L": "",
                "M": l["purpose_en"], "N": str(l["hs_code"]), "O": l["qty"], "P": l["invoice_unit"],
                "Q": to_number(l["pcs_per_set"]), "R": to_number(l["price"]), "S": l["total"],
                "T": cfg.get("image_placeholder", "") if i == 0 else "", "U": l["asin"], "V": ""}
        for col, v in vals.items():
            inv[f"{col}{row}"] = v
    trow = 25 + len(model.lines)
    _apply_style(inv, trow, total_st)
    inv[f"M{trow}"] = "*Total（汇总）:"
    inv.merge_cells(f"M{trow}:N{trow}")
    inv[f"O{trow}"] = model.total_qty
    inv[f"Q{trow}"] = "Currency"
    inv[f"R{trow}"] = DECLARED_CURRENCY
    inv[f"S{trow}"] = model.total_value
    inv[f"T{trow}"] = "箱数（CTNS）"
    inv[f"U{trow}"] = len(model.boxes)

    pk["C2"] = inv["B2"].value
    pk["C3"] = the_date
    names = {l["fnsku"]: l for l in model.lines}
    p_data, p_total = _style_row(pk, 6), _style_row(pk, 7)
    for c in pk[7]:
        c.value = None
    row = 6
    total_cbm, total_w = 0.0, 0.0
    for b in model.boxes:
        for j, (sku, q) in enumerate(b.items):
            _apply_style(pk, row, p_data)
            first = j == 0
            box_cbm = cbm(*b.dims) if first and all(x is not None for x in b.dims) else None
            vals = {"A": model.shipment_id, "B": model.reference_id, "C": sku, "D": names[sku]["english_name"], "E": names[sku]["chinese_name"],
                    "F": b.fba_box_id, "G": 1 if first else None, "H": q,
                    "I": b.dims[0] if first else None, "J": b.dims[1] if first else None, "K": b.dims[2] if first else None,
                    "L": box_cbm, "M": round2(b.weight) if first and b.weight is not None else None, "N": ""}
            for col, v in vals.items():
                pk[f"{col}{row}"] = v
            if first:
                total_cbm += box_cbm or 0
                total_w += b.weight or 0
            row += 1
    _apply_style(pk, row, p_total)
    pk[f"F{row}"] = "*Total（汇总）:"
    pk[f"G{row}"] = len(model.boxes)
    pk[f"L{row}"] = round(total_cbm, 3)
    pk[f"M{row}"] = round2(total_w)


# ---------------------------------------------------------------- 德速 T3（Commercial Invoice 9810）
def _per_box_qty(r) -> Any:
    n = r.end - r.start + 1
    q = r.qty / n if n else r.qty
    return int(q) if q == int(q) else round(q, 2)


WEIGHT_BASES = ("per_box", "per_unit")
WEIGHT_BASIS_LABEL = {"per_box": "单箱重", "per_unit": "单件重"}


def _box_by_no(model: InvoiceModel, no: int):
    for b in model.boxes:
        if b.no == no:
            return b
    return None


def _weight_cell(model: InvoiceModel, r, ctx: Dict[str, Any]):
    """S8-fix1 N2：重量列按口径写。per_box = 单箱重；per_unit = 单箱重 ÷ 该箱件数（混装箱按整箱件数，只在首行写）。
    口径没定（None）时留空，由 _warn_weight_basis 提示——不再固定写单箱重。"""
    basis = ctx.get("weight_basis")
    if r.weight is None or basis not in WEIGHT_BASES:
        return None
    if basis == "per_box":
        return round2(r.weight)
    b = _box_by_no(model, r.start)
    units = sum(q for _, q in b.items) if b is not None else (r.qty / (r.end - r.start + 1))
    return round(r.weight / units, 4) if units else None


def _warn_weight_basis(model: InvoiceModel, ctx: Dict[str, Any], forwarder_name: str) -> None:
    basis, src = ctx.get("weight_basis"), ctx.get("weight_basis_source") or ""
    if basis in WEIGHT_BASES:
        model.warnings.append({"code": "weight_basis", "basis": basis, "source": src,
                               "message": f"重量列按{WEIGHT_BASIS_LABEL[basis]}填写（{src or '指定'}）"})
        return
    r = next((x for x in model.ranges if x.weight is not None), None)
    eg = ""
    if r is not None:
        b = _box_by_no(model, r.start)
        units = sum(q for _, q in b.items) if b is not None else 0
        eg = f"例如第 {r.label} 箱：单箱重 {round2(r.weight)} kg" + (f"，单件重 {round(r.weight / units, 4)} kg" if units else "")
    # 示例按货代给：德速 T3 第 2 批单箱重、第 9 批单件重；琪珞 T6 第 6 批单件重、第 7 批单箱重（S8 评审 v2）
    sample = "如德速第2批单箱重、第9批单件重" if "德速" in (forwarder_name or "") else "如琪珞第6批单件重、第7批单箱重"
    model.warnings.append({"code": "weight_basis_undecided",
                           "message": f"{forwarder_name}发票的重量列有的票填单箱重、有的填单件重（{sample}），"
                                      f"这票还没定口径，重量列已留空。请在面板选“单箱重/单件重”，或上传货代给的这票发票让程序按货件 CSV 核对判定。{eg}"})


def decide_weight_basis(reference: bytes, reference_name: str, csv_bytes: bytes, shipment_id: str = ""):
    """S8-fix1 N2：和琪珞模板解析同一做法——把货代原件里的重量按 单箱重/单件重 两种口径各算一遍，
    和货件 CSV 的 Σ单箱实重 比对，只有一种对得上才采用。返回 (basis 或 None, 说明)。"""
    from fba_tool.invoice_templates import parse_invoice_file, reconcile_with_csv

    res = reconcile_with_csv(parse_invoice_file(reference, reference_name or ""), [csv_bytes])
    tks = [t for t in res.get("tickets", []) if not shipment_id or str(t.get("shipment_id") or "").upper() == shipment_id.upper()]
    if not tks:
        return None, f"货代发票 {reference_name} 里没有货件 {shipment_id}，无法判定重量口径"
    tk = tks[0]
    codes = {i["code"]: i for i in tk.get("issues", [])}
    if "weight_meaning_proposed" in codes:
        basis = codes["weight_meaning_proposed"].get("proposed")
    elif "weight_meaning_ambiguous" in codes or "weight_mismatch" in codes or "weight_meaning_undecided" in codes:
        return None, f"货代发票 {reference_name} 的重量按两种口径和货件 CSV 都对不上或都对得上，无法自动判定"
    else:
        basis = tk.get("weight_meaning")
    if basis not in WEIGHT_BASES:
        return None, f"货代发票 {reference_name} 的重量口径是 {basis}，不是单箱重/单件重，无法套用"
    return basis, f"按货代发票 {reference_name} 和货件 CSV 的重量核对判定为{WEIGHT_BASIS_LABEL[basis]}"


def write_deshu_t3(wb, model: InvoiceModel, fw: Dict[str, Any], ctx: Dict[str, Any]) -> None:
    ws = _sheet(wb, "发票填写模板")
    ws["B2"] = model.shipment_id
    ws["G2"] = DECLARED_CURRENCY
    ws["G4"] = model.total_qty
    ws["B6"] = model.warehouse
    if ctx.get("channel"):
        ws["B7"] = ctx["channel"]
    if ctx.get("order_no"):
        ws["B15"] = ctx["order_no"]
    lines = {l["fnsku"]: l for l in model.lines}
    for i, r in enumerate(model.ranges):
        row = 21 + i
        l = lines[r.fnsku]
        vals = {"A": model.shipment_id, "B": model.reference_id, "C": f"{r.start}-{r.end}", "D": l["sku"],
                "E": l["english_name"], "F": l["chinese_name"], "G": str(l["hs_code"]), "H": l["brand"],
                "I": l["material_cn"], "J": l["purpose_cn"], "K": l["asin"], "L": l["has_battery"], "M": l["model"],
                "N": _per_box_qty(r), "O": l["invoice_unit"], "P": to_number(l["pcs_per_set"]),
                "Q": to_number(l["purchase_price"]), "R": to_number(l["price"]),
                "T": _weight_cell(model, r, ctx),   # S8-fix1：口径按 CSV 核对 / 用户选择（混装箱只在首行写）
                "U": r.dims[0], "V": r.dims[1], "W": r.dims[2]}
        for col, v in vals.items():
            ws[f"{col}{row}"] = v
    _warn_weight_basis(model, ctx, "德速")


# ---------------------------------------------------------------- 琪珞 T6（FBA 专线导入模版）
def write_qiluo_t6(wb, model: InvoiceModel, fw: Dict[str, Any], ctx: Dict[str, Any]) -> None:
    ws = _sheet(wb, "发票")
    ws["B3"] = model.shipment_id
    ws["E3"] = model.warehouse       # E4–E10 收件人信息由模板 VLOOKUP「！」表带出
    ws["B6"] = len(model.boxes)
    if ctx.get("channel"):
        ws["B4"] = ctx["channel"]
    try:
        bang = _sheet(wb, "！")
        codes = {str(c.value).strip().upper() for c in bang["W"] if c.value}
        if model.warehouse and model.warehouse.upper() not in codes:
            model.warnings.append({"code": "qiluo_warehouse_not_in_sheet",
                                   "message": f"仓库 {model.warehouse} 不在琪珞「！」地址表里，收件人地址公式会显示 #N/A，请在 E3 选琪珞表里的代码（如带 -1 后缀）"})
    except ExportConfigError:
        pass
    lines = {l["fnsku"]: l for l in model.lines}
    for i, r in enumerate(model.ranges):
        row = 14 + i
        l = lines[r.fnsku]
        vals = {"A": ",".join(str(n) for n in range(r.start, r.end + 1)), "B": model.shipment_id, "C": model.reference_id,
                "D": l["chinese_name"], "E": l["english_name"], "F": _per_box_qty(r), "G": to_number(l["price"]),
                "H": l["material_cn"], "I": l["purpose_cn"], "J": l["brand"], "K": l["model"], "L": str(l["hs_code"]),
                "M": _weight_cell(model, r, ctx),   # S8-fix1：T6“单件实重”实际口径因票而异（第6批单件重、第7批单箱重）
                "N": r.dims[0], "O": r.dims[1], "P": r.dims[2],
                "R": f"https://www.amazon.com/dp/{l['asin']}" if l["asin"] else ""}
        for col, v in vals.items():
            ws[f"{col}{row}"] = v
    _warn_weight_basis(model, ctx, "琪珞")


# ---------------------------------------------------------------- 琪珞 T5（下单模板 单行汇总；T7 是它的空白版）
def write_qiluo_t5(wb, model: InvoiceModel, fw: Dict[str, Any], ctx: Dict[str, Any]) -> None:
    ws = wb.worksheets[0]
    wh = ctx.get("warehouse_info") or {}
    ws["B3"] = model.shipment_id
    ws["D3"] = model.reference_id
    if ctx.get("channel"):
        ws["B2"] = ctx["channel"]
    ws["B7"] = model.warehouse
    ws["F7"] = wh.get("zip", "")
    ws["G7"] = "US"
    ws["H7"] = wh.get("state", "")
    ws["I7"] = wh.get("city", "")
    ws["J7"] = " ".join(x for x in (wh.get("address1"), wh.get("address2"), wh.get("address3")) if x)
    ls = model.lines
    join = lambda key: ",".join(str(l[key]) for l in ls)  # noqa: E731
    ws["A10"] = join("hs_code")
    ws["B10"] = join("chinese_name")
    ws["C10"] = join("english_name")
    ws["E10"] = join("material_cn")
    ws["F10"] = join("purpose_cn")
    ws["G10"] = join("brand")
    ws["H10"] = ",".join(str(l["qty"]) for l in ls) if len(ls) > 1 else ls[0]["qty"]
    ws["I10"] = ",".join(str(to_number(l["price"])) for l in ls) if len(ls) > 1 else to_number(ls[0]["price"])
    ws["J10"] = len(model.boxes)
    dims = {tuple(b.dims) for b in model.boxes}
    if len(dims) == 1 and all(x is not None for x in next(iter(dims))):
        ws["K10"], ws["L10"], ws["M10"] = next(iter(dims))
    else:
        model.warnings.append({"code": "qiluo_t5_mixed_dims",
                               "message": "这票箱子尺寸不止一种，琪珞单行汇总模板只有一组长宽高，已留空，请手工填写或改用 T6 模板"})
    weights = [b.weight for b in model.boxes]
    ws["N10"] = round2(sum(w for w in weights if w is not None)) if all(w is not None for w in weights) else None


ADAPTERS: Dict[str, Callable] = {"send": write_send, "jiufang": write_jiufang, "deshu_t3": write_deshu_t3,
                                 "qiluo_t6": write_qiluo_t6, "qiluo_t5": write_qiluo_t5}
# 需要派送地址（仓库地址库）才能填的适配器
ADDRESS_REQUIRED = {"jiufang", "qiluo_t5"}
# 各货代发票必填的产品字段（除通用必填外）
ADAPTER_REQUIRED = {
    "send": {"material_cn": "材质", "purpose_cn": "用途", "asin": "ASIN"},
    "jiufang": {"material_en": "英文材质", "purpose_en": "英文用途", "asin": "ASIN"},
    "deshu_t3": {"material_cn": "材质", "purpose_cn": "用途"},
    "qiluo_t6": {"material_cn": "材质(中文)", "purpose_cn": "用途(中文)"},
    "qiluo_t5": {"material_cn": "材料", "purpose_cn": "用途"},
}


# ---------------------------------------------------------------- 入口
def export_invoice(forwarder_id: str, *, shipment, products: Dict[str, Dict[str, Any]], reference_id: str = "",
                   warehouse: str = "", warehouse_info: Optional[Dict[str, Any]] = None, order_no: str = "",
                   channel: str = "", invoice_date: Optional[str] = None,
                   measured: Optional[Dict[str, Dict[str, Any]]] = None, hs_templates=None,
                   weight_basis: Optional[str] = None, weight_basis_source: str = "") -> Dict[str, Any]:
    """返回 {"bytes", "file_name", "warnings", "model"}；有拦截项时抛 InvoiceBlocked。"""
    import openpyxl
    from fba_tool.invoice_rules import REQUIRED_PRODUCT_FIELDS

    fw = get_export_forwarder(forwarder_id)
    adapter = ADAPTERS.get(fw["export_adapter"])
    if adapter is None:
        raise ExportConfigError(f"未知的导出适配器：{fw['export_adapter']}")
    required = {**REQUIRED_PRODUCT_FIELDS, **ADAPTER_REQUIRED.get(fw["export_adapter"], {})}
    model = build_model(forwarder=forwarder_id, shipment=shipment, products=products, reference_id=reference_id,
                        warehouse=warehouse or shipment.warehouse_code, measured=measured, required=required,
                        hs_templates=hs_templates)
    if fw["export_adapter"] in ADDRESS_REQUIRED and not warehouse_info:
        raise InvoiceBlocked([{"code": "warehouse_address_missing",
                               "message": f"仓库 {model.warehouse} 在 FBA 仓库地址库里没有地址，{fw.get('name') or forwarder_id}发票的派送地址填不了"}])
    no_mark = [b.no for b in model.boxes if not b.fba_box_id]
    if no_mark and fw["export_adapter"] == "jiufang":
        # 待分配货件没有 FBA 箱号（U 号），九方 Packing list 的箱唛列会是空的
        model.warnings.append({"code": "box_mark_missing", "boxes": no_mark,
                               "message": f"货件 {model.shipment_id} 有 {len(no_mark)} 箱没有 FBA 箱号（货件“待分配”时常见），"
                                          "九方 Packing list 的箱唛（Mark of CTN）会是空的；请等亚马逊分配货件号后重新下载 CSV 再导出，或导出后手工补箱唛"})
    wb = openpyxl.load_workbook(io.BytesIO(_template_bytes(fw)))
    ctx = {"order_no": order_no, "channel": channel, "invoice_date": invoice_date or date.today().isoformat(),
           "warehouse_info": warehouse_info or {},
           "weight_basis": weight_basis or (fw.get("invoice_export") or {}).get("weight_basis"),
           "weight_basis_source": weight_basis_source or ("" if weight_basis else "货代配置")}
    if weight_basis and weight_basis not in WEIGHT_BASES:
        raise ValueError(f"重量口径只能是 per_box（单箱重）或 per_unit（单件重）：{weight_basis}")
    adapter(wb, model, fw, ctx)
    buf = io.BytesIO()
    wb.save(buf)
    name_tpl = (fw.get("invoice_export") or {}).get("file_name", "{shipment_id}发票.xlsx")
    file_name = name_tpl.format(shipment_id=model.shipment_id, order_no=order_no or model.shipment_id)
    return {"bytes": buf.getvalue(), "file_name": file_name, "warnings": model.warnings, "model": model}
