"""xlsx → Univer IWorkbookData，只用 openpyxl。

浏览器不解析货代模板。整列数据验证（例如 F15:F1048576）默认钳到已用行 + 200，
避免 Univer 的 rowCount 变成 Excel 最大行。
"""
from __future__ import annotations

import hashlib
import io
import json
import time
from copy import deepcopy
from datetime import date, datetime
from typing import Any, Callable, Dict, List, Optional

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

EXCEL_MAX_ROW = 1048576
EXCEL_MAX_COL = 16384
CLAMP_PAD = 200
DATA_VALIDATION_PLUGIN = "SHEET_DATA_VALIDATION_PLUGIN"

_TYPE_MAP = {
    "list": "list",
    "listmultiple": "listMultiple",
    "whole": "whole",
    "decimal": "decimal",
    "date": "date",
    "time": "time",
    "textlength": "textLength",
    "custom": "custom",
    "checkbox": "checkbox",
    "any": "any",
    "none": "none",
}
_OP_MAP = {
    "between": "between",
    "notbetween": "notBetween",
    "equal": "equal",
    "notequal": "notEqual",
    "lessthan": "lessThan",
    "lessthanorequal": "lessThanOrEqual",
    "greaterthan": "greaterThan",
    "greaterthanorequal": "greaterThanOrEqual",
}

_CACHE: Dict[tuple, Dict[str, Any]] = {}


class SnapshotCancelled(Exception):
    pass


def estimate_used_rows(data: bytes) -> int:
    """读 dimension。货代模板的 max_row 是真实使用行，不是整列验证的 1048576。"""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        worst = 1
        for ws in wb.worksheets:
            worst = max(worst, int(ws.max_row or 1))
        return worst
    finally:
        wb.close()


def workbook_to_snapshot(
    data: bytes,
    *,
    clamp: bool = True,
    progress: Optional[Callable[[int], None]] = None,
    cancel_event: Any = None,
) -> Dict[str, Any]:
    key = (hashlib.sha256(data).hexdigest(), bool(clamp))
    if cancel_event is None and key in _CACHE:
        if progress:
            progress(100)
        return deepcopy(_CACHE[key])

    started = time.perf_counter()
    _check_cancel(cancel_event)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    try:
        styles: Dict[str, Any] = {}
        style_cache: Dict[Any, Optional[str]] = {}
        sheets: Dict[str, Any] = {}
        order: List[str] = []
        rules_by_sheet: Dict[str, List[Dict[str, Any]]] = {}
        used = 1
        worksheets = list(wb.worksheets)
        sheet_total = max(len(worksheets), 1)
        for index, ws in enumerate(worksheets):
            _check_cancel(cancel_event)

            def _on_fraction(fraction: float, index: int = index) -> None:
                if not progress:
                    return
                overall = int(100 * (index + max(0.0, min(fraction, 0.99))) / sheet_total)
                progress(min(99, max(1, overall)))

            if progress:
                _on_fraction(0.01)
            sheet_id = f"s{index}"
            order.append(sheet_id)
            sheet, rules, content_rows = _sheet_to_univer(
                ws, sheet_id, clamp=clamp, styles=styles, style_cache=style_cache,
                on_fraction=_on_fraction,
            )
            sheets[sheet_id] = sheet
            if rules:
                rules_by_sheet[sheet_id] = rules
            used = max(used, content_rows)
        snapshot: Dict[str, Any] = {
            "id": "wb",
            "name": "workbook",
            "appVersion": "1.1.0",
            "locale": "zhCN",
            "styles": styles,
            "sheetOrder": order,
            "sheets": sheets,
        }
        if rules_by_sheet:
            snapshot["resources"] = [{
                "name": DATA_VALIDATION_PLUGIN,
                "data": json.dumps(rules_by_sheet, ensure_ascii=False),
            }]
        result = {
            "snapshot": snapshot,
            "meta": {
                "sha256": key[0],
                "clamp": bool(clamp),
                "used_rows": used,
                "sheet_count": len(order),
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
            },
        }
    finally:
        wb.close()
    if progress:
        progress(100)
    if cancel_event is None:
        _CACHE[key] = deepcopy(result)
    return result


def _check_cancel(cancel_event: Any) -> None:
    if cancel_event is not None and cancel_event.is_set():
        raise SnapshotCancelled()


def _sheet_to_univer(
    ws: Worksheet,
    sheet_id: str,
    *,
    clamp: bool,
    styles: Dict[str, Any],
    style_cache: Dict[Any, Optional[str]],
    on_fraction: Optional[Callable[[float], None]] = None,
) -> tuple:
    content_row, content_col = _content_extent(ws)
    dv_rules, dv_row, dv_col = _data_validations(ws, clamp=clamp, content_row=content_row, content_col=content_col)
    freeze = _freeze(ws)
    freeze_row = (freeze.get("startRow") or 0) + 1 if freeze else 1
    freeze_col = (freeze.get("startColumn") or 0) + 1 if freeze else 1
    if clamp:
        row_count = max(content_row, dv_row, freeze_row, 1)
        col_count = max(content_col, dv_col, freeze_col, 1)
    else:
        row_count = max(content_row, freeze_row, 1)
        col_count = max(content_col, freeze_col, 1)
    row_count = min(row_count, 99999)
    col_count = min(col_count, 512)

    cell_data: Dict[str, Dict[str, Any]] = {}
    owned = [
        cell for cell in ws._cells.values()
        if not isinstance(cell, MergedCell) and cell.row <= row_count and cell.column <= col_count
    ]
    owned.sort(key=lambda cell: (cell.row, cell.column))
    steps = max(len(owned), 1)
    last_bucket = -1
    for index, cell in enumerate(owned, start=1):
        payload = _cell_payload(cell, styles, style_cache)
        if payload is not None:
            cell_data.setdefault(str(cell.row - 1), {})[str(cell.column - 1)] = payload
        if on_fraction:
            bucket = int(4 * index / steps)
            if bucket != last_bucket:
                last_bucket = bucket
                on_fraction(min(index / steps, 0.99))

    sheet: Dict[str, Any] = {
        "id": sheet_id,
        "name": ws.title,
        "hidden": 0 if ws.sheet_state == "visible" else 1,
        "rowCount": row_count,
        "columnCount": col_count,
        "cellData": cell_data,
        "mergeData": _merges(ws),
        "rowData": _row_data(ws, row_count),
        "columnData": _column_data(ws, col_count),
    }
    if freeze:
        sheet["freeze"] = freeze
    return sheet, dv_rules, content_row


def _content_extent(ws: Worksheet) -> tuple:
    max_row = 1
    max_col = 1
    for cell in ws._cells.values():
        if isinstance(cell, MergedCell):
            continue
        if cell.value is None:
            continue
        max_row = max(max_row, int(cell.row))
        max_col = max(max_col, int(cell.column))
    for mr in ws.merged_cells.ranges:
        if mr.max_row < EXCEL_MAX_ROW - 10:
            max_row = max(max_row, int(mr.max_row))
        if mr.max_col < EXCEL_MAX_COL - 10:
            max_col = max(max_col, int(mr.max_col))
    if ws.freeze_panes:
        row, col = _pane_row_col(ws.freeze_panes)
        max_row = max(max_row, row)
        max_col = max(max_col, col)
    return max_row, max_col


def _pane_row_col(pane: Any) -> tuple:
    if isinstance(pane, str):
        from openpyxl.utils import coordinate_to_tuple

        return coordinate_to_tuple(pane)
    return int(pane.row), int(pane.column)


def _freeze(ws: Worksheet) -> Optional[Dict[str, int]]:
    if not ws.freeze_panes:
        return None
    row, col = _pane_row_col(ws.freeze_panes)
    return {
        "xSplit": max(col - 1, 0),
        "ySplit": max(row - 1, 0),
        "startRow": max(row - 1, 0),
        "startColumn": max(col - 1, 0),
    }


def _merges(ws: Worksheet) -> List[Dict[str, int]]:
    out = []
    for mr in ws.merged_cells.ranges:
        if mr.max_row >= EXCEL_MAX_ROW - 10 or mr.max_col >= EXCEL_MAX_COL - 10:
            continue
        out.append({
            "startRow": mr.min_row - 1,
            "endRow": mr.max_row - 1,
            "startColumn": mr.min_col - 1,
            "endColumn": mr.max_col - 1,
        })
    return out


def _row_data(ws: Worksheet, row_count: int) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    for idx, dim in ws.row_dimensions.items():
        if not isinstance(idx, int) or idx < 1 or idx > row_count:
            continue
        item: Dict[str, Any] = {}
        if dim.height:
            item["h"] = float(dim.height) * 4 / 3
        if dim.hidden:
            item["hd"] = 1
        if item:
            out[str(idx - 1)] = item
    return out


def _column_data(ws: Worksheet, col_count: int) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    for key, dim in ws.column_dimensions.items():
        if not isinstance(key, str) or key == "customWidth":
            continue
        try:
            from openpyxl.utils import column_index_from_string

            idx = column_index_from_string(key)
        except ValueError:
            continue
        if idx < 1 or idx > col_count or not dim.width:
            continue
        width = float(dim.width)
        px = ((256 * width + int(128 / 7)) / 256) * 7
        item: Dict[str, Any] = {"w": px}
        if dim.hidden:
            item["hd"] = 1
        out[str(idx - 1)] = item
    return out


def _cell_payload(cell: Any, styles: Dict[str, Any], style_cache: Dict[Any, Optional[str]]) -> Optional[Dict[str, Any]]:
    if cell.value is None and getattr(cell, "data_type", None) != "f":
        return None
    payload: Dict[str, Any] = {}
    if getattr(cell, "data_type", None) == "f" or _looks_formula(cell.value):
        formula = str(cell.value)
        if not formula.startswith("="):
            formula = "=" + formula
        payload["f"] = formula
    elif getattr(cell, "data_type", None) == "b" or isinstance(cell.value, bool):
        payload["v"] = bool(cell.value)
        payload["t"] = 3
    elif getattr(cell, "data_type", None) == "n" or isinstance(cell.value, (int, float)):
        payload["v"] = cell.value
        payload["t"] = 2
    elif isinstance(cell.value, (datetime, date)):
        payload["v"] = cell.value.isoformat()
        payload["t"] = 1
    else:
        payload["v"] = str(cell.value)
        payload["t"] = 1
    style_id = _style_id(cell, styles, style_cache)
    if style_id:
        payload["s"] = style_id
    return payload


def _looks_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _style_id(cell: Any, styles: Dict[str, Any], cache: Dict[Any, Optional[str]]) -> Optional[str]:
    key = getattr(cell, "_style", None)
    if not key:
        return None
    if key in cache:
        return cache[key]
    style: Dict[str, Any] = {}
    font = cell.font
    if font is not None:
        if font.name:
            style["ff"] = str(font.name)
        if font.size:
            style["fs"] = float(font.size)
        if font.bold:
            style["bl"] = 1
        if font.italic:
            style["it"] = 1
        color = _rgb(getattr(font, "color", None))
        if color:
            style["cl"] = {"rgb": color}
    fill = cell.fill
    if fill is not None and getattr(fill, "patternType", None) not in (None, "none"):
        color = _rgb(getattr(fill, "fgColor", None))
        if color and color.lower() not in {"#000000", "#ffffff"}:
            style["bg"] = {"rgb": color}
    number_format = getattr(cell, "number_format", None)
    if number_format and number_format != "General":
        style["n"] = {"pattern": str(number_format)}
    alignment = cell.alignment
    if alignment is not None:
        horizontal = {"left": 1, "center": 2, "right": 3}.get(alignment.horizontal or "")
        vertical = {"top": 1, "center": 2, "bottom": 3}.get(alignment.vertical or "")
        if horizontal:
            style["ht"] = horizontal
        if vertical:
            style["vt"] = vertical
        if alignment.wrap_text:
            style["tb"] = 3
    if not style:
        cache[key] = None
        return None
    sid = f"st{len(styles) + 1}"
    styles[sid] = style
    cache[key] = sid
    return sid


def _rgb(color: Any) -> Optional[str]:
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    if not rgb or not isinstance(rgb, str):
        return None
    if rgb in {"00000000", "None"}:
        return None
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6:
        return None
    return "#" + rgb


def _data_validations(ws: Worksheet, *, clamp: bool, content_row: int, content_col: int) -> tuple:
    rules = []
    max_end_row = 1
    max_end_col = 1
    collection = getattr(ws, "data_validations", None)
    items = list(getattr(collection, "dataValidation", None) or [])
    for index, dv in enumerate(items):
        if dv.type in (None, "none") and dv.formula1 in (None, ""):
            continue
        ranges = []
        for cr in _iter_ranges(dv.sqref):
            end_row = int(cr.max_row)
            end_col = int(cr.max_col)
            if clamp:
                end_row = _clamp_end(end_row, content_row, EXCEL_MAX_ROW)
                end_col = _clamp_end(end_col, content_col, EXCEL_MAX_COL)
            start_row = min(int(cr.min_row), end_row)
            start_col = min(int(cr.min_col), end_col)
            ranges.append({
                "startRow": start_row - 1,
                "endRow": end_row - 1,
                "startColumn": start_col - 1,
                "endColumn": end_col - 1,
            })
            max_end_row = max(max_end_row, end_row)
            max_end_col = max(max_end_col, end_col)
        if not ranges:
            continue
        rule: Dict[str, Any] = {
            "uid": f"dv-{ws.title}-{index}",
            "type": _TYPE_MAP.get(str(dv.type or "").lower(), str(dv.type or "custom")),
            "ranges": ranges,
            "showDropDown": _univer_show_dropdown(dv),
        }
        if dv.formula1 not in (None, ""):
            rule["formula1"] = str(dv.formula1)
        if dv.formula2 not in (None, ""):
            rule["formula2"] = str(dv.formula2)
        op = _OP_MAP.get(str(dv.operator or "").lower())
        if op:
            rule["operator"] = op
        if dv.allow_blank is not None:
            rule["allowBlank"] = bool(dv.allow_blank)
        rules.append(rule)
    return rules, max_end_row, max_end_col


def _clamp_end(end: int, content: int, excel_max: int) -> int:
    if end >= excel_max - 10:
        return min(content + CLAMP_PAD, 99999)
    return end


def _iter_ranges(sqref: Any):
    if sqref is None:
        return []
    if hasattr(sqref, "ranges"):
        return list(sqref.ranges)
    from openpyxl.worksheet.cell_range import MultiCellRange

    return list(MultiCellRange(str(sqref)).ranges)


def _univer_show_dropdown(dv: Any) -> bool:
    """Excel 的 showDropDown=True 表示隐藏下拉；Univer 的同名字段表示显示。"""
    flag = getattr(dv, "showDropDown", None)
    if flag is None:
        return str(dv.type or "") == "list"
    return not bool(flag)


def a1(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"
