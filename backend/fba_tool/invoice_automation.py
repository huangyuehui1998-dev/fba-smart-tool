"""Step 4.5 货代发票：四层数据源降级兜底 + DeepSeek 补全 + 保格式导出。"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Literal, Tuple

import pdfplumber
from openpyxl.utils import get_column_letter

from fba_tool.invoice_deepseek import (
    infer_customs_fields,
    infer_header_options_with_ai,
    parse_address_with_ai,
)
from fba_tool.deepseek_billing import DeepseekUsageAccumulator
from fba_tool.logistics_metrics import logistics_metrics_from_box_specs
from fba_tool.user_config import throw_ratio as config_throw_ratio
from fba_tool.invoice_table import (
    build_invoice_line_items,
    detect_invoice_table_layout,
    preview_table_cells,
    sanitize_header_text,
    table_cells_from_line_items,
)
from fba_tool.invoice_workflow_fill import build_invoice_line_items_workflow
from fba_tool.product_catalog import extract_csv_skus
from fba_tool.sku_keys import case_conflict_message, case_conflicts, lookup_sku, sku_key as _norm_key, sku_text
from fba_tool.shipment_parser import CsvShipment, parse_csv_shipment
from fba_tool.shipment_serialize import resolve_csv_shipment
from fba_tool.app_paths import invoice_history_path, resolve_config_read
from fba_tool.xlsx_io import load_workbook_fast
from fba_tool.shipment_parser import decode_csv_bytes

FieldSource = Literal["green", "yellow", "red", "blue", "preview"]

SOURCE_LABELS = {
    "green": "精准提取",
    "blue": "主数据库",
    "yellow": "AI推理(需复核)",
    "red": "必填待补",
    "preview": "模板解析",
}


@dataclass
class InvoiceField:
    key: str
    label: str
    value: Any
    source: FieldSource
    cell: str = ""
    required: bool = False


@dataclass
class InvoiceDraft:
    shipment_id: str
    warehouse_code: str
    fields: List[InvoiceField] = field(default_factory=list)
    template_sheet: str = "Sheet1"
    table_cells: List[Dict[str, Any]] = field(default_factory=list)
    is_table_mode: bool = False
    warnings: List[str] = field(default_factory=list)

    def to_grid_df(self):
        import pandas as pd

        return pd.DataFrame(
            [
                {
                    "字段": f.label,
                    "值": f.value if f.value not in (None, "") else "",
                    "来源": SOURCE_LABELS[f.source],
                    "单元格": f.cell,
                    "必填": "是" if f.required else "",
                }
                for f in self.fields
            ]
        )

    def field_map(self) -> Dict[str, InvoiceField]:
        return {f.key: f for f in self.fields}

    def missing_required(self) -> List[str]:
        return [f.label for f in self.fields if f.required and not str(f.value or "").strip()]

    def to_cells_list(self) -> List[Dict[str, Any]]:
        """Return coordinate-based cell fills for WYSIWYG spreadsheet UI."""
        if self.table_cells:
            return self.table_cells
        cells: List[Dict[str, Any]] = []
        for f in self.fields:
            if not f.cell:
                continue
            cells.append(
                {
                    "cell": str(f.cell).strip().upper(),
                    "value": "" if f.value in (None, "") else f.value,
                    "source": SOURCE_LABELS[f.source],
                    "field_key": f.key,
                    "required": f.required,
                }
            )
        return cells


def _ordered_template_field_keys(cfg: dict) -> List[str]:
    keys: List[str] = []
    for group in (
        list(cfg.get("header_defaults", {}).keys()),
        cfg.get("deterministic_fields", []),
        cfg.get("address_fields", []),
        cfg.get("ai_only_fields", []),
        cfg.get("required_fields", []),
    ):
        for key in group:
            if key and key not in keys:
                keys.append(key)
    return keys


def _scan_max_above_header(header_row: int | None, fallback: int = 15) -> int:
    return (header_row - 1) if header_row and header_row > 1 else fallback


def build_invoice_preview(template_bytes: bytes, forwarder_id: str | None = None) -> InvoiceDraft:
    """解析模板结构，返回空值字段列表（预览态）。"""
    cfg = load_invoice_config()
    wb = load_workbook_fast(template_bytes)
    table_layout = detect_invoice_table_layout(wb, cfg=cfg, forwarder_id=forwarder_id)
    if table_layout:
        return InvoiceDraft(
            shipment_id="",
            warehouse_code="",
            fields=[],
            template_sheet=wb.active.title,
            table_cells=preview_table_cells(table_layout),
            is_table_mode=True,
        )

    cell_map = detect_template_cells(wb, cfg)
    required_set = set(cfg.get("required_fields", []))

    fields: List[InvoiceField] = []
    for key in _ordered_template_field_keys(cfg):
        fields.append(
            InvoiceField(
                key=key,
                label=key,
                value="",
                source="preview",
                cell=cell_map.get(key, ""),
                required=key in required_set,
            )
        )

    return InvoiceDraft(
        shipment_id="",
        warehouse_code="",
        fields=fields,
        template_sheet=wb.active.title,
    )


def load_invoice_config() -> dict:
    path = resolve_config_read("invoice_mapping.json")
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def load_invoice_history() -> dict:
    path = invoice_history_path()
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def extract_csv_product_names(csv_bytes: bytes) -> List[str]:
    text = decode_csv_bytes(csv_bytes)
    import csv as csvmod

    rows = list(csvmod.reader(io.StringIO(text)))
    header_idx = -1
    for i, row in enumerate(rows):
        if row and _norm_key(row[0]) == "sku":
            header_idx = i
            break
    names: List[str] = []
    if header_idx < 0:
        return names
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        sku = str(row[0]).strip()
        if not sku or re.search(r"(length|width|height|weight|包装)", sku, flags=re.I):
            break
        name = str(row[1]).strip() if len(row) > 1 else ""
        if name:
            names.append(name)
    return names


def extract_total_product_qty(csv_bytes: bytes) -> int:
    text = decode_csv_bytes(csv_bytes)
    import csv as csvmod

    rows = list(csvmod.reader(io.StringIO(text)))
    for row in rows[:20]:
        if row and str(row[0]).strip().strip('"') == "商品数量" and len(row) > 1:
            try:
                return int(float(str(row[1]).strip().strip('"')))
            except Exception:
                pass
    return 0


def extract_pdf_address_block(pdf_bytes: bytes) -> str:
    if not pdf_bytes or not pdf_bytes.strip():
        return ""
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            if not pdf.pages:
                return ""
            text = pdf.pages[0].extract_text() or ""
    except Exception:
        return ""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    block: List[str] = []
    capture = False
    for ln in lines:
        if re.search(r"(Ship\s*To|目的地|配送地址|Deliver)", ln, flags=re.I):
            capture = True
            continue
        if capture:
            if re.search(r"(Single SKU|SKU|FBA[A-Z0-9]+U\d|数量|QTY|P1\s*-)", ln, flags=re.I):
                break
            block.append(ln)
    if not block:
        for ln in lines:
            if re.search(r"\b[A-Z]{2}\s+\d{5}", ln) or "SMITHFIELD" in ln.upper():
                block.append(ln)
    return " ".join(block)


def _label_matches(cell_text: str, variant: str) -> bool:
    c = sanitize_header_text(cell_text)
    v = sanitize_header_text(variant)
    if not c or not v:
        return False
    return v == c or v in c or c in v


def detect_template_cells(wb, cfg: dict, max_row: int | None = None) -> Dict[str, str]:
    overrides = cfg.get("cell_overrides", {})
    aliases: Dict[str, List[str]] = cfg.get("label_aliases", {})
    found: Dict[str, str] = dict(overrides)

    ws = wb.active
    scan_max = max_row if max_row is not None else min(ws.max_row or 80, 80)
    max_col = min(ws.max_column or 30, 40)
    for row in ws.iter_rows(min_row=1, max_row=scan_max, max_col=max_col):
        for cell in row:
            text = str(cell.value or "").strip()
            if not text:
                continue
            for field_key, variants in aliases.items():
                if field_key in found:
                    continue
                if any(_label_matches(text, v) for v in variants):
                    target_col = cell.column + 1
                    target_row = cell.row
                    if cell.column < ws.max_column:
                        right = ws.cell(cell.row, cell.column + 1)
                        if right.value in (None, ""):
                            target_col = cell.column + 1
                        else:
                            target_col = cell.column + 2
                    found[field_key] = f"{get_column_letter(target_col)}{target_row}"
    return found


def _header_option_cells(
    wb,
    cfg: dict,
    ship: CsvShipment | None = None,
    product_names: List[str] | None = None,
    api_key: str | None = None,
    model: str | None = None,
    temperature: float | None = None,
    usage_acc: DeepseekUsageAccumulator | None = None,
    header_row: int | None = None,
) -> List[Dict[str, Any]]:
    """第 1 层规则默认 + 第 3 层 AI：发票顶部选项标签（包税/交货方式/保险等）。"""
    scan_max = _scan_max_above_header(header_row)
    cell_map = detect_template_cells(wb, cfg, max_row=scan_max)
    defaults: Dict[str, Any] = dict(cfg.get("header_defaults", {}))
    ai_field_keys: List[str] = list(cfg.get("header_ai_fields", []))
    cells: List[Dict[str, Any]] = []
    filled: Dict[str, Any] = {}

    for key, val in defaults.items():
        cell = cell_map.get(key)
        if not cell or val in (None, ""):
            continue
        filled[key] = val
        cells.append(
            {
                "cell": str(cell).strip().upper(),
                "value": val,
                "source": SOURCE_LABELS["green"],
                "field_key": key,
                "required": False,
            }
        )

    missing_ai = [k for k in ai_field_keys if k not in filled and cell_map.get(k)]
    if api_key and missing_ai and ship:
        context = {
            "shipment_id": ship.shipment_id,
            "warehouse_code": ship.warehouse_code,
            "box_count": ship.box_count,
            "product_names": (product_names or [])[:10],
        }
        try:
            ai_json = infer_header_options_with_ai(
                missing_ai,
                context,
                api_key,
                model=model,
                temperature=temperature,
                usage_acc=usage_acc,
            )
            for key in missing_ai:
                val = ai_json.get(key)
                cell = cell_map.get(key)
                if not cell or val in (None, ""):
                    continue
                cells.append(
                    {
                        "cell": str(cell).strip().upper(),
                        "value": val,
                        "source": SOURCE_LABELS["yellow"],
                        "field_key": key,
                        "required": False,
                    }
                )
        except Exception:
            pass
    return cells


def _fill_header_zone_cells(
    wb,
    cfg: dict,
    pdf_bytes: bytes,
    ship: CsvShipment,
    product_names: List[str],
    metrics: Dict[str, Any],
    api_key: str | None = None,
    model: str | None = None,
    temperature: float | None = None,
    usage_acc: DeepseekUsageAccumulator | None = None,
    header_row: int | None = None,
) -> List[Dict[str, Any]]:
    """表格/标签模式共用：顶部选项 + 地址 + 宏观指标。"""
    cells: List[Dict[str, Any]] = []
    cells.extend(
        _header_option_cells(
            wb,
            cfg,
            ship=ship,
            product_names=product_names,
            api_key=api_key,
            model=model,
            temperature=temperature,
            usage_acc=usage_acc,
            header_row=header_row,
        )
    )
    cells.extend(
        _address_block_cells(
            wb,
            cfg,
            pdf_bytes,
            api_key=api_key,
            model=model,
            temperature=temperature,
            usage_acc=usage_acc,
            header_row=header_row,
        )
    )
    cells.extend(
        _shipment_metric_cells(
            wb,
            cfg,
            metrics,
            max_row=_scan_max_above_header(header_row),
        )
    )
    return cells


def _address_block_cells(
    wb,
    cfg: dict,
    pdf_bytes: bytes,
    api_key: str | None = None,
    model: str | None = None,
    temperature: float | None = None,
    usage_acc: DeepseekUsageAccumulator | None = None,
    header_row: int | None = None,
) -> List[Dict[str, Any]]:
    """表格模式仍写入表头上方的地址标签区（Consignee / Country 等）。"""
    scan_max = _scan_max_above_header(header_row)
    cell_map = detect_template_cells(wb, cfg, max_row=scan_max)
    addr_text = extract_pdf_address_block(pdf_bytes)
    addr_json: Dict[str, Any] = {}
    if addr_text.strip() and api_key:
        try:
            addr_json = parse_address_with_ai(
                addr_text,
                api_key,
                model=model,
                temperature=temperature,
                usage_acc=usage_acc,
            )
        except Exception:
            addr_json = {}

    cells: List[Dict[str, Any]] = []
    keys_to_fill = list(cfg.get("address_fields", ["邮编", "国家代码", "省州代码", "城市", "地址"]))
    if "Consignee" not in keys_to_fill:
        keys_to_fill.append("Consignee")

    for key in keys_to_fill:
        cell = cell_map.get(key)
        if not cell:
            continue
        val = str(addr_json.get(key, "") or "").strip()
        if key == "Consignee" and not val:
            val = str(addr_json.get("地址", "") or "").strip()
        if key == "地址" and not val:
            val = str(addr_json.get("Consignee", "") or addr_text or "").strip()
        if not val:
            continue
        src = SOURCE_LABELS["yellow"] if api_key else SOURCE_LABELS["green"]
        cells.append(
            {
                "cell": str(cell).strip().upper(),
                "value": val,
                "source": src,
                "field_key": key,
                "required": False,
            }
        )
    return cells


def _shipment_metric_cells(
    wb,
    cfg: dict,
    metrics: Dict[str, Any],
    max_row: int | None = None,
) -> List[Dict[str, Any]]:
    """模板表头区的计费吨等宏观指标单元格（精准提取）。"""
    cell_map = detect_template_cells(wb, cfg, max_row=max_row)
    cells: List[Dict[str, Any]] = []
    revenue_ton = metrics.get("revenue_ton")
    if revenue_ton in (None, ""):
        return cells
    for field_key in ("计费吨", "RT"):
        cell = cell_map.get(field_key)
        if not cell:
            continue
        cells.append(
            {
                "cell": str(cell).strip().upper(),
                "value": revenue_ton,
                "source": SOURCE_LABELS["green"],
                "field_key": "revenue_ton",
                "required": False,
            }
        )
        break
    return cells


def _customs_by_sku(
    csv_bytes: bytes,
    catalog,
    api_key: str | None,
    model: str | None = None,
    temperature: float | None = None,
    usage_acc: DeepseekUsageAccumulator | None = None,
) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, set]]:
    """第 2 层产品库 + 第 3 层 AI；返回 (customs_map, ai_filled_keys_by_sku)。"""
    from fba_tool.invoice_table import extract_csv_sku_meta

    sku_meta = extract_csv_sku_meta(csv_bytes)
    skus = extract_csv_skus(csv_bytes)
    customs_keys = [
        "海关编码",
        "中文品名",
        "英文品名",
        "材料",
        "用途",
        "品牌",
        "单个货值",
    ]
    result: Dict[str, Dict[str, Any]] = {}
    ai_keys_by_sku: Dict[str, set] = {}

    for sku in skus:
        sku_key = sku_text(sku)  # S9-step4-fix3：键用原样写法，只差大小写的两个 SKU 各自一条，不互相覆盖
        if not sku_key:
            continue
        merged: Dict[str, Any] = {}
        from fba_tool.master_db import get_product_by_sku, invoice_hints_from_master
        master = get_product_by_sku(sku)
        if master:
            master_hints = invoice_hints_from_master(master)
            for key, val in master_hints.items():
                if val not in (None, ""):
                    merged[key] = val
        title = lookup_sku(sku_meta, sku, {}).get("title", "")  # S9-step3：统一大小写 / 全角规则
        missing = [k for k in customs_keys if not str(merged.get(k, "") or "").strip()]
        ai_filled: set = set()
        if api_key and missing and (title or sku):
            try:
                ai_json = infer_customs_fields(
                    [title or sku],
                    load_invoice_history(),
                    api_key,
                    missing_keys=missing,
                    known_values=merged,
                    model=model,
                    temperature=temperature,
                    usage_acc=usage_acc,
                )
                for k, v in ai_json.items():
                    if k in missing and v not in (None, "") and not str(merged.get(k, "") or "").strip():
                        merged[k] = v
                        ai_filled.add(k)
            except Exception:
                pass
        result[sku_key] = merged
        ai_keys_by_sku[sku_key] = ai_filled
    return result, ai_keys_by_sku


def _sku_case_warnings(ship) -> List[str]:
    """S9-step4-fix3：同一票里有只差大小写 / 全角半角的 SKU → 提示冲突（各行按原样分别取资料，不合并、不覆盖）。"""
    try:
        skus = [] if ship.box_items.empty else list(ship.box_items["SKU"])
    except Exception:  # noqa: BLE001
        return []
    groups = list(case_conflicts(skus).values())
    return [case_conflict_message(groups, "本票")] if groups else []


def build_invoice_draft(
    csv_bytes: bytes,
    pdf_bytes: bytes,
    template_bytes: bytes,
    api_key: str | None = None,
    model: str | None = None,
    temperature: float | None = None,
    workflow_json: dict | None = None,
    po_number: str = "",
    use_workflow_fill: bool = False,
    forwarder_id: str | None = None,
    parsed_shipment: dict | None = None,
) -> Tuple[InvoiceDraft, DeepseekUsageAccumulator | None]:
    usage_acc = DeepseekUsageAccumulator() if api_key else None
    if forwarder_id:
        from fba_tool.invoice_step45 import try_adapter_draft

        adapted = try_adapter_draft(forwarder_id, csv_bytes, template_bytes, parsed_shipment)
        if adapted is not None:
            return adapted, usage_acc
    cfg = load_invoice_config()
    ship = resolve_csv_shipment(parsed_shipment=parsed_shipment, csv_bytes=csv_bytes or None)
    ratio = float(config_throw_ratio())
    shipment_metrics = logistics_metrics_from_box_specs(ship.box_specs, throw_ratio_val=ratio)
    wb = load_workbook_fast(template_bytes)
    table_layout = detect_invoice_table_layout(wb, cfg=cfg, forwarder_id=forwarder_id)
    catalog = None
    case_warnings = _sku_case_warnings(ship)
    etl_warnings: List[str] = []
    from fba_tool.forwarder_mapping import get_forwarder_mapping
    from fba_tool.invoice_etl import run_invoice_etl

    fwd_mapping = get_forwarder_mapping(forwarder_id) if forwarder_id else None

    if table_layout:
        separate_mp = "material" in table_layout.columns and "purpose" in table_layout.columns
        if not separate_mp and fwd_mapping and fwd_mapping.get("separate_material_purpose") is not None:
            separate_mp = bool(fwd_mapping.get("separate_material_purpose"))
        if fwd_mapping:
            line_items, etl_warnings = run_invoice_etl(
                forwarder_id=forwarder_id or "",
                ship=ship,
                csv_bytes=csv_bytes,
                workflow_json=workflow_json,
                po_number=po_number,
                catalog=None,
                separate_material_purpose=separate_mp,
            )
        elif use_workflow_fill or workflow_json is not None:
            line_items = build_invoice_line_items_workflow(
                ship,
                csv_bytes,
                catalog,
                workflow_json,
                po_number=po_number,
                separate_material_purpose=separate_mp,
            )
        else:
            product_names = extract_csv_product_names(csv_bytes)
            customs_map, customs_ai_keys = _customs_by_sku(
                csv_bytes,
                catalog,
                api_key,
                model=model,
                temperature=temperature,
                usage_acc=usage_acc,
            )
            line_items = build_invoice_line_items(
                ship,
                csv_bytes,
                catalog,
                customs_by_sku=customs_map,
                customs_ai_keys=customs_ai_keys,
                reference_id="",
            )
        table_cells = table_cells_from_line_items(table_layout, line_items, SOURCE_LABELS)
        if not use_workflow_fill:
            product_names = extract_csv_product_names(csv_bytes)
            table_cells.extend(
                _fill_header_zone_cells(
                    wb,
                    cfg,
                    pdf_bytes,
                    ship,
                    product_names,
                    shipment_metrics,
                    api_key=api_key,
                    model=model,
                    temperature=temperature,
                    usage_acc=usage_acc,
                    header_row=table_layout.header_row,
                )
            )
        return InvoiceDraft(
            shipment_id=ship.shipment_id,
            warehouse_code=ship.warehouse_code,
            fields=[],
            template_sheet=wb.active.title,
            table_cells=table_cells,
            is_table_mode=True,
            warnings=case_warnings + list(etl_warnings or []),
        ), usage_acc

    cell_map = detect_template_cells(wb, cfg)

    product_names = extract_csv_product_names(csv_bytes)
    total_qty = extract_total_product_qty(csv_bytes)
    if total_qty <= 0 and not ship.box_items.empty:
        total_qty = int(ship.box_items["数量"].sum())

    specs = ship.box_specs
    if not specs.empty:
        length = float(specs["长(cm)"].mode().iloc[0]) if "长(cm)" in specs else 0
        width = float(specs["宽(cm)"].mode().iloc[0]) if "宽(cm)" in specs else 0
        height = float(specs["高(cm)"].mode().iloc[0]) if "高(cm)" in specs else 0
        single_weight = float(specs["实重(kg)"].mean()) if "实重(kg)" in specs else 0
        total_weight = round(single_weight * ship.box_count, 2)
    else:
        length = width = height = single_weight = total_weight = 0

    revenue_ton = shipment_metrics.get("revenue_ton")
    det_values = {
        "FBA单号": ship.shipment_id,
        "仓库代码": ship.warehouse_code,
        "件数": ship.box_count,
        "个数": total_qty,
        "单件长": length,
        "单件宽": width,
        "单件高": height,
        "总实重": total_weight,
        "计费吨": "" if revenue_ton in (None, "") else revenue_ton,
    }

    fields: List[InvoiceField] = []
    header_defaults = cfg.get("header_defaults", {})
    for key, val in header_defaults.items():
        fields.append(
            InvoiceField(
                key=key,
                label=key,
                value=val,
                source="green",
                cell=cell_map.get(key, ""),
                required=False,
            )
        )

    for key in cfg.get("deterministic_fields", list(det_values.keys())):
        fields.append(
            InvoiceField(
                key=key,
                label=key,
                value=det_values.get(key, ""),
                source="green",
                cell=cell_map.get(key, ""),
                required=False,
            )
        )

    addr_text = extract_pdf_address_block(pdf_bytes)
    addr_json: Dict[str, Any] = {}
    if addr_text.strip() and api_key:
        try:
            addr_json = parse_address_with_ai(
                addr_text,
                api_key,
                model=model,
                temperature=temperature,
                usage_acc=usage_acc,
            )
        except Exception:
            addr_json = {}
    for key in cfg.get("address_fields", ["邮编", "国家代码", "省州代码", "城市", "地址"]):
        val = addr_json.get(key, "")
        src: FieldSource = "yellow" if val else ("green" if not api_key else "yellow")
        if not val and not api_key:
            src = "red"
        fields.append(
            InvoiceField(key=key, label=key, value=val, source=src, cell=cell_map.get(key, ""), required=False)
        )

    catalog_json: Dict[str, Any] = {}
    from fba_tool.master_db import get_product_by_sku, invoice_hints_from_master

    skus = extract_csv_skus(csv_bytes)
    for sku in skus:
        hints = invoice_hints_from_master(get_product_by_sku(sku))
        for key, val in hints.items():
            if val not in (None, "") and not catalog_json.get(key):
                catalog_json[key] = val

    customs_keys = cfg.get("ai_only_fields", [])
    merged_customs: Dict[str, Any] = {}
    for key in customs_keys:
        cat_val = catalog_json.get(key, "")
        if cat_val not in (None, ""):
            merged_customs[key] = cat_val

    missing_keys = [k for k in customs_keys if not str(merged_customs.get(k, "") or "").strip()]
    ai_json: Dict[str, Any] = {}
    ai_filled_keys: set = set()
    if api_key and product_names and missing_keys:
        try:
            ai_json = infer_customs_fields(
                product_names,
                load_invoice_history(),
                api_key,
                missing_keys=missing_keys,
                known_values=merged_customs,
                model=model,
                temperature=temperature,
                usage_acc=usage_acc,
            )
            for k, v in ai_json.items():
                if k in missing_keys and v not in (None, ""):
                    ai_filled_keys.add(k)
        except Exception:
            ai_json = {}

    header_ai_keys = list(cfg.get("header_ai_fields", []))
    header_ai_json: Dict[str, Any] = {}
    if api_key and header_ai_keys:
        try:
            header_ai_json = infer_header_options_with_ai(
                header_ai_keys,
                {
                    "shipment_id": ship.shipment_id,
                    "warehouse_code": ship.warehouse_code,
                    "box_count": ship.box_count,
                    "product_names": product_names[:10],
                },
                api_key,
                model=model,
                temperature=temperature,
                usage_acc=usage_acc,
            )
        except Exception:
            header_ai_json = {}

    for key in header_ai_keys:
        if key in header_defaults:
            continue
        val = header_ai_json.get(key, "")
        if val in (None, ""):
            continue
        fields.append(
            InvoiceField(
                key=key,
                label=key,
                value=val,
                source="yellow",
                cell=cell_map.get(key, ""),
                required=False,
            )
        )

    for key in customs_keys:
        cat_val = merged_customs.get(key, "")
        ai_val = ai_json.get(key, "")
        if cat_val not in (None, ""):
            val, src = cat_val, "blue"
        elif ai_val not in (None, "") and key in ai_filled_keys:
            val, src = ai_val, "yellow"
        else:
            val, src = "", "red"
        fields.append(
            InvoiceField(
                key=key,
                label=key,
                value=val,
                source=src,
                cell=cell_map.get(key, ""),
                required=False,
            )
        )

    fields.append(
        InvoiceField(
            key="ID号",
            label="ID号",
            value="",
            source="red",
            cell=cell_map.get("ID号", ""),
            required=True,
        )
    )

    return InvoiceDraft(
        shipment_id=ship.shipment_id,
        warehouse_code=ship.warehouse_code,
        fields=fields,
        template_sheet=wb.active.title,
        warnings=case_warnings,
    ), usage_acc


def apply_grid_edits(draft: InvoiceDraft, edited_df) -> InvoiceDraft:
    fmap = draft.field_map()
    for _, row in edited_df.iterrows():
        label = str(row.get("字段", ""))
        for f in draft.fields:
            if f.label == label:
                f.value = row.get("值", "")
                if f.key == "ID号" and str(f.value or "").strip():
                    f.source = "green"
                    f.required = True
                break
    return draft


def export_invoice_xlsx(
    template_bytes: bytes,
    draft: InvoiceDraft,
    forwarder_id: str | None = None,
) -> bytes:
    wb = load_workbook_fast(template_bytes)
    ws = wb[draft.template_sheet] if draft.template_sheet in wb.sheetnames else wb.active

    cell_writes = draft.table_cells if draft.table_cells else []
    if not cell_writes:
        for f in draft.fields:
            if not f.cell or f.value in (None, ""):
                continue
            cell_writes.append({"cell": f.cell, "value": f.value})

    if draft.is_table_mode or draft.table_cells:
        from fba_tool.invoice_table import clear_invoice_table_data_zone, detect_invoice_table_layout

        cfg = load_invoice_config()
        layout = detect_invoice_table_layout(wb, cfg=cfg, forwarder_id=forwarder_id)
        if layout:
            clear_invoice_table_data_zone(ws, layout)

    # 表头下的脏数据先清掉，再整本写回。不删其它工作表、图片、数据验证或隐藏表。
    sheet_name = ws.title
    if draft.is_table_mode or draft.table_cells:
        bio = io.BytesIO()
        wb.save(bio)
        template_bytes = bio.getvalue()
    wb.close()

    from fba_tool.univer_apply import apply_cell_writes

    return apply_cell_writes(template_bytes, cell_writes, sheet_name)


def styled_grid(df):
    """返回带颜色高亮的 Styler。"""

    def _row_style(row):
        src = str(row.get("来源", ""))
        if "精准提取" in src:
            color = "#d4edda"
        elif "产品信息表" in src or "主数据库" in src:
            color = "#cce5ff"
        elif "AI推理" in src:
            color = "#fff3cd"
        elif "必填" in src or row.get("必填") == "是":
            color = "#f8d7da"
        else:
            color = ""
        return [f"background-color: {color}" if color else ""] * len(row)

    return df.style.apply(_row_style, axis=1)
