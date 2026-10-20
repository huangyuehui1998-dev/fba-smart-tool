"""S8 深度测评 V3.0 修复回归：D1–D15、P1–P3、D-12 促销申报、HS 校验与钓具预设。"""
from __future__ import annotations

import base64
import importlib
import io
import json
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from fastapi import HTTPException

ROOT = Path(__file__).resolve().parent


# ---------- D1 尾数分配 ----------
def _editor(rows):
    return pd.DataFrame(rows, columns=["箱号", "来源", "SKU简称", "数量"])


def test_d1_tail_assign_within_pool():
    from routers.boxing import tail_remaining, validate_tail_assign

    ed = _editor([[1, "整箱", "A", 10], [2, "整箱", "A", 10]])
    rows = [SimpleNamespace(sku="A", total_qty=23)]
    assert tail_remaining(ed, rows) == {"A": 3}
    validate_tail_assign(ed, rows, "A", 3, 3)  # 剩 3 件、新开第 3 箱：允许
    for sku, qty, box in [("A", 4, 1), ("B", 1, 1), ("A", 1, 5)]:
        with pytest.raises(HTTPException) as e:
            validate_tail_assign(ed, rows, sku, qty, box)
        assert e.value.status_code == 400
    with pytest.raises(HTTPException):
        validate_tail_assign(_editor([[1, "整箱", "A", 23]]), rows, "A", 1, 1)  # 已全部装箱


def test_d1_tail_pool_no_auto_fallback():
    """全部装箱后尾数池必须是空的（以前会退回“自动尾数”，显示已装箱的件数，导致超发）。"""
    from services.boxing_service import compute_tail_pool

    rows = [SimpleNamespace(sku="A", total_qty=20, unit_weight_kg=1.0), SimpleNamespace(sku="B", total_qty=7, unit_weight_kg=0.5)]
    full = _editor([[1, "整箱", "A", 10], [2, "整箱", "A", 10], [3, "整箱", "B", 7]])
    pool = compute_tail_pool(full, rows)
    assert pool.empty and list(pool.columns) == ["SKU简称", "待排尾数", "单个重量(kg)"]
    part = compute_tail_pool(_editor([[1, "整箱", "A", 10], [2, "整箱", "B", 5]]), rows)
    assert part.set_index("SKU简称")["待排尾数"].to_dict() == {"A": 10, "B": 2}


# ---------- D2 / D15 分仓 ----------
def test_d2_csv_warehouse_wins_over_filename():
    from routers.routing import resolve_routing_warehouse

    code, src, warn = resolve_routing_warehouse("SMF3", "第一票ONT8.csv")
    assert (code, src) == ("SMF3", "CSV") and "ONT8" in warn
    assert resolve_routing_warehouse("SMF3", "第一票SMF3.csv")[2] is None
    assert resolve_routing_warehouse("", "第一票ONT8.csv")[:2] == ("ONT8", "文件名")
    assert resolve_routing_warehouse("UNKN", "亚马逊FBA表1.csv")[0] == "UNKN"


def test_d15_empty_routing_export_400():
    from fastapi.testclient import TestClient
    from main import app

    r = TestClient(app).post("/api/routing/export", json={"forwarder": []})
    assert r.status_code == 400
    assert "数据为空" in r.json()["detail"]


# ---------- D3 模板识别 ----------
def test_d3_t4_detect_only_order_sheet():
    """真实样本识别：T4 原件 → T4；只剩“下单模板”以外表的 T4 副本不再被认成 T4；T7 空白 → T7，T6 → T6。"""
    from openpyxl import load_workbook
    from fba_tool.invoice_templates import detect_template

    S = ROOT / "tests/golden/samples/invoice_templates"
    t4 = S / "T4_第4批第一票_FWA4-27箱_去图片副本.xlsx"
    assert detect_template(t4.read_bytes(), t4.name)["template_id"] == "T4"
    wb = load_workbook(t4)
    for name in list(wb.sheetnames):
        if name == "下单模板":
            del wb[name]
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    buf = io.BytesIO()
    wb.save(buf)
    assert detect_template(buf.getvalue(), "x.xlsx").get("template_id") != "T4"
    for f, want in (("T7_义乌宁波运单下单模板_空白.xlsx", "T7"), ("T6_第7批第1票_VGT2-29箱.xlsx", "T6"),
                    ("T3_第9批第1票_FBA19GQXZW0N.xlsx", "T3")):
        assert detect_template((S / f).read_bytes(), f)["template_id"] == want


# ---------- D4 发票导出适配器 ----------
def test_d4_new_forwarders_registered():
    """注册 + 能真正导出（导出内容逐格比对见 test_s8fix1_invoice_adapters.py）。"""
    from fba_tool.invoice_export import ADAPTERS, ADDRESS_REQUIRED, export_forwarders

    assert set(ADAPTERS) == {"send", "jiufang", "deshu_t3", "qiluo_t6", "qiluo_t5"}
    assert ADDRESS_REQUIRED == {"jiufang", "qiluo_t5"}
    fw = export_forwarders()
    assert set(fw) == {"send", "jiufang", "deshu", "qiluo_t6", "qiluo_t5"}
    assert {fw[k]["export_adapter"] for k in ("deshu", "qiluo_t6", "qiluo_t5")} == {"deshu_t3", "qiluo_t6", "qiluo_t5"}
    for name in ("deshu_t3.xlsx", "qiluo_t6.xlsx", "qiluo_t5.xlsx"):
        assert (ROOT / "templates/forwarders" / name).is_file()


def test_d4_unknown_forwarder_message():
    from fba_tool.invoice_export import get_export_forwarder

    with pytest.raises(Exception) as e:
        get_export_forwarder("nope_xyz")
    assert "不支持该货代" in str(e.value)


# ---------- P1 仓库冲突 ----------
def test_p1_warehouse_conflict_warning():
    from routers.invoice_export import warehouse_conflict_warning

    w = warehouse_conflict_warning("ONT8", "界面指定", "SMF3", "")
    assert w["code"] == "warehouse_csv_conflict" and "SMF3" in w["message"]
    assert warehouse_conflict_warning("SMF3", "界面指定", "SMF3", "") is None
    assert warehouse_conflict_warning("SMF3", "配送地址", "UNKN", "") is None


# ---------- D5 / D6 HS ----------
@pytest.mark.parametrize("raw,want", [("9507.10.00 90", "9507100090"), ("95072000", "95072000"), ("9507-30-0000", "9507300000"), ("９５０７１０００９０", "9507100090")])
def test_d6_hs_normalize_ok(raw, want):
    from fba_tool.hs_code import normalize_hs_code

    assert normalize_hs_code(raw) == want


@pytest.mark.parametrize("raw", ["950710", "95071000901", "abc12345", "٩٥٠٧١٠٠٠٩٠"])
def test_d6_hs_normalize_bad(raw):
    from fba_tool.hs_code import HsCodeError, is_valid_hs_code, normalize_hs_code

    assert not is_valid_hs_code(raw)
    with pytest.raises(HsCodeError):
        normalize_hs_code(raw)


# ---------- HS 钓具预设 ----------
def test_hs_fishing_presets_verified_codes():
    from fba_tool.hs_code import is_valid_hs_code
    from fba_tool.hs_presets import hs_presets

    p = hs_presets("渔具")
    codes = {i["key"]: i["hs_code"] for i in p["items"]}
    assert codes == {"rod": "9507100090", "hook": "9507200000", "reel": "9507300000", "lure": "9507900000"}
    assert all(is_valid_hs_code(c) for c in codes.values())
    assert all(i["us_hts"] and all(h.startswith("9507.") for h in i["us_hts"]) for i in p["items"])
    line = p["not_preset"][0]
    assert line["key"] == "line" and "9507.90.20.00" in line["note"] and "零售包装" in line["note"]
    # HQ 954309 / HQ 956007：零售包装并标明鱼线 → 9507.90.20；只绕线轴 → 5404 等（S8 评审 v2 第 4 节核对裁定原文）
    note = line["note"]
    assert "954309" in note and "956007" in note and "9507.90.20" in note and "5404" in note
    assert "零售包装" in note and "只绕在线轴上" in note
    assert hs_presets("pet")["items"] == []


def test_hs_presets_endpoint_and_not_auto_seeded():
    from fastapi.testclient import TestClient
    from main import app

    c = TestClient(app)
    r = c.get("/api/master-data/hs-presets", params={"category": "fishing"})
    assert r.status_code == 200 and len(r.json()["items"]) == 4
    tpl = c.get("/api/master-data/hs-templates", params={"category": "fishing"}).json()
    items = tpl.get("items", tpl) if isinstance(tpl, dict) else tpl
    assert not any(str(t.get("source", "")).startswith("S8 预设") for t in items)  # 只读参考，不自动写库


# ---------- D7 PDF 工具 ----------
def test_d7_pdf_inputs_rejected():
    from routers.pdf_tools import _decode_pdf_b64, _ensure_pdf, _parse_mm

    assert _parse_mm("") is None and _parse_mm("100") == 100.0
    for bad in ("abc", "-5", "100000"):
        with pytest.raises(HTTPException) as e:
            _parse_mm(bad)
        assert e.value.status_code == 400
    with pytest.raises(HTTPException):
        _ensure_pdf("x.pdf", b"not a pdf")
    with pytest.raises(HTTPException):
        _ensure_pdf("y.pdf", b"%PDF-1.4 broken")
    with pytest.raises(HTTPException):
        _decode_pdf_b64("@@@", 1)
    with pytest.raises(HTTPException):
        _decode_pdf_b64(base64.b64encode(b"hello").decode(), 1)


# ---------- D8 / B 配置错误横幅 ----------
def test_d8_config_error_keyed_by_logical_name(monkeypatch, tmp_path):
    """坏的 user_settings.json → 有错误；删掉文件后重新读取 → 错误消失（不用重启）。"""
    from fba_tool import user_config

    cfg = tmp_path / "user_settings.json"
    cfg.write_text("{坏的 json", encoding="utf-8")
    monkeypatch.setattr(user_config, "_settings_path", lambda: cfg)
    user_config.CONFIG_ERRORS.clear()
    user_config.load_settings(force=True)
    assert "user_settings.json" in user_config.config_errors()
    cfg.unlink()
    user_config.load_settings(force=True)
    assert "user_settings.json" not in user_config.config_errors()
    user_config.load_settings(force=True)


# ---------- D9 run_app ----------
def test_d9_run_app_env(monkeypatch, tmp_path):
    sys.path.insert(0, str(ROOT.parent))
    ra = importlib.import_module("run_app")
    monkeypatch.setenv("FBA_APP_DATA_DIR", str(tmp_path))
    assert ra._data_dir() == tmp_path and ra._port_file().parent == tmp_path
    monkeypatch.setenv("FBA_NO_BROWSER", "1")
    opened = []
    monkeypatch.setattr(ra.webbrowser, "open", lambda u: opened.append(u))
    ra._open_url("http://x")
    assert opened == []
    monkeypatch.setenv("FBA_NO_BROWSER", "")
    ra._open_url("http://x")
    assert opened == ["http://x"]


# ---------- D10 领星新品需复核 ----------
def test_d10_review_reasons():
    from fba_tool.lingxing_preview import review_reasons

    assert review_reasons({"sku": "OK-1", "fnsku": "X0", "asin": "B0"}) == []
    r = review_reasons({"sku": "Ã©Ã¨-1", "fnsku": "", "asin": ""})
    assert len(r) == 3


# ---------- D11 产品库批量修改 ----------
def _xlsx(rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_d11_duplicate_sku_rows_rejected(monkeypatch):
    from fba_tool import product_edit

    monkeypatch.setattr(product_edit, "_all_products", lambda: [{"sku": "A", "hs_code": "6307909000"}])
    res = product_edit.preview_products_edit(_xlsx([["sku", "hs_code"], ["A", "6307909000"], ["A", "7326909000"]]))
    assert len(res["errors"]) >= 2 and not res.get("changes")
    res2 = product_edit.preview_products_edit(_xlsx([["sku", "hs_code"], ["A", "12AB"]]))
    assert res2["errors"]


# ---------- D12 重复上传 ----------
def test_d12_duplicate_upload_is_error():
    from fba_tool.shipment_validator import validate_shipments

    tpl = ROOT / "templates"
    a, b = tpl / "亚马逊FBA表.csv", tpl / "亚马逊FBA表1.csv"
    if not (a.exists() and b.exists()):
        pytest.skip("templates 缺样例")
    res = validate_shipments([(a.name, a.read_bytes()), (b.name, b.read_bytes())], [])
    assert not res.ok and any("重复上传" in e for e in res.errors)


# ---------- D-12 促销期按正常价申报 ----------
def _iso(days_ago=0):
    return (datetime.now(timezone.utc) - timedelta(days=days_ago)).isoformat()


def test_dec12_regular_price_used_in_promo():
    from fba_tool.product_fields import promo_info, resolve_declared_price

    m = {"listing_price": 15.99, "regular_price": 19.99}
    assert resolve_declared_price(m) == (19.99, "lingxing_regular")
    assert promo_info(m)["drop_pct"] == 20.0


def test_dec12_threshold_boundary():
    from fba_tool.product_fields import resolve_declared_price

    assert resolve_declared_price({"listing_price": 19.0, "regular_price": 20.0})[1] == "lingxing"   # 正好 5%：不算促销
    assert resolve_declared_price({"listing_price": 18.99, "regular_price": 20.0})[1] == "lingxing_regular"
    assert resolve_declared_price({"listing_price": 22.0, "regular_price": 20.0}) == (22.0, "lingxing")  # 涨价按实时价


def test_dec12_recent_nonpromo_fallback_and_expiry():
    from fba_tool.product_fields import resolve_declared_price

    m = {"listing_price": 12.0, "last_nonpromo_price": 16.0, "last_nonpromo_at": _iso(10)}
    assert resolve_declared_price(m) == (16.0, "lingxing_recent_nonpromo")
    m["last_nonpromo_at"] = _iso(120)  # 超过 90 天：不再拿来比
    assert resolve_declared_price(m) == (12.0, "lingxing")


def test_dec12_list_price_ignored_and_manual_unchanged():
    from fba_tool.product_fields import resolve_declared_price

    assert resolve_declared_price({"listing_price": 15.99, "list_price": 38.97}) == (15.99, "lingxing")
    assert resolve_declared_price({"declared_value": 9.0, "field_sources": {"declared_value": "manual"}}) == (9.0, "manual")


def test_dec12_track_nonpromo():
    from fba_tool.product_fields import _track_nonpromo

    s = {"listing_price": 19.5, "regular_price": 20.0}
    _track_nonpromo(s, {})
    assert s["last_nonpromo_price"] == 19.5
    s2 = {"listing_price": 15.0}
    _track_nonpromo(s2, {"regular_price": 20.0})
    assert "last_nonpromo_price" not in s2  # 促销价不记成“非促销价”


def test_dec12_price_rows_warn():
    """price_checks：带 promo_price 的行给 promo_price_not_used，写明促销价和实际申报价；没有促销的行不提示。"""
    from fba_tool.customs_package import price_checks

    rows = [{"sku": "P1", "unit_price": 19.99, "price_source": "lingxing_regular", "promo_price": 15.99, "raw_price": 19.99},
            {"sku": "P2", "unit_price": 12.0, "price_source": "lingxing", "promo_price": None, "raw_price": 12.0}]
    _, warnings, blocking = price_checks(rows)
    w = [x for x in warnings if x["code"] == "promo_price_not_used"]
    assert len(w) == 1 and "P1" in w[0]["message"] and "15.99" in w[0]["message"] and "19.99" in w[0]["message"]
    assert "P2" not in w[0]["message"] and not any(b["code"] == "promo_price_not_used" for b in blocking)