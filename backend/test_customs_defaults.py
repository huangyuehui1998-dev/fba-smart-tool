"""S1f：报关默认值（方案 v8 §3.6）。"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from fba_tool import customs_package as cp
from fba_tool.customs_export import export_customs_package
from fba_tool.customs_package import (
    CustomsPackageData,
    build_customs_data,
    money_half_up,
    net_weight_minus_per_box,
)
from test_customs_package import _sample_pair

TEMPLATE = Path(__file__).resolve().parent / "templates" / "customs" / "报关资料模板.xlsx"
PROFILE = {"id": "t", "name": "测试公司", "credit_code": "91420113TEST0000", "signed_at": "武汉", "domestic_source": "武汉"}


def test_dataclass_defaults():
    d = CustomsPackageData()
    assert d.exemption_nature == "101"
    assert d.transport_mode == "海路运输"
    assert d.loading_port == "宁波北仑"
    assert d.freight == ""
    assert d.entrust_valid_until == f"{date.today().year} 年 12 月 31日止。"


def test_default_dates_leave_export_and_declare_blank():
    dates = cp.default_dates()
    assert dates["export_date"] == ""
    assert dates["declare_date"] == ""
    assert dates["invoice_date"] == date.today().isoformat()


@pytest.mark.parametrize(
    "v,expect",
    [("2.345", "2.35"), ("2.335", "2.34"), ("0.125", "0.13"), (2.675, "2.68"), ("1.005", "1.01")],
)
def test_money_half_up(v, expect):
    assert money_half_up(v) == Decimal(expect)


def test_net_weight_minus_per_box():
    assert net_weight_minus_per_box("115.2", 12) == Decimal("112.80")
    assert net_weight_minus_per_box("661.4", 115) == Decimal("638.40")
    assert net_weight_minus_per_box(0, 3) == Decimal("0")


def test_merge_unit_price_half_up_and_total_is_price_times_qty():
    rows = [
        {"chinese_name": "宠物垫", "hs_code": "6307909000", "qty": 3, "unit": "个", "amount": 10.0,
         "box_count": 1, "gross_kg": 10.0, "net_kg": 0},
    ]
    lines = cp._merge_rows(rows, "武汉")
    # 10/3 = 3.333.. → 3.33；总价 = 3.33 × 3 = 9.99
    assert lines[0].unit_price == 3.33
    assert lines[0].amount == 9.99
    assert lines[0].unit == "个"
    assert lines[0].net_kg == 9.8


def test_unit_from_master_declare_unit(monkeypatch):
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {"中文品名": "宠物垫", "declare_unit": "个"})
    data = build_customs_data([_sample_pair()], [], PROFILE, source="shipment", order_no="T1")
    assert {l.unit for l in data.lines} == {"个"}


def test_totals_gross_and_net(monkeypatch):
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {"中文品名": "宠物垫"})
    data = build_customs_data([_sample_pair()], [], PROFILE, source="shipment", order_no="T1")
    # 2 箱：12.5 + 11.0；净重 = 23.5 − 0.2×2
    assert data.total_pieces == 2
    assert data.total_gross_kg == 23.5
    assert data.total_net_kg == 23.1


def _export(data):
    wb = openpyxl.load_workbook(io.BytesIO(export_customs_package(data, TEMPLATE.read_bytes())))
    return {ws.title: ws for ws in wb.worksheets}


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="无模板")
def test_export_cells(monkeypatch):
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {"中文品名": "宠物垫", "单个货值": 10.0, "申报价来源": "lingxing"})
    data = build_customs_data([_sample_pair()], [], PROFILE, source="shipment", order_no="T1")
    ws = _export(data)
    decl = ws["报关单"]
    assert decl["E6"].value == "海路运输"
    assert "101" in str(decl["H7"].value)
    assert decl["S20"].value == "照章征税"
    assert decl["F12"].value == 23.5          # 毛重
    assert decl["G12"].value == 23.1          # 净重
    assert decl["K12"].value in (None, "")    # 运费留空
    assert decl["H3"].value in (None, "")     # 出口日期留空
    assert "宁波北仑" in ws["合同"]["A33"].value
    assert "深圳" not in ws["合同"]["A33"].value
    assert ws["委托书"]["E12"].value == f"{date.today().year} 年 12 月 31日止。"


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="无模板")
def test_overrides_port_and_validity(monkeypatch):
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {"中文品名": "宠物垫", "单个货值": 10.0, "申报价来源": "lingxing"})
    data = build_customs_data(
        [_sample_pair()], [], PROFILE, source="shipment", order_no="T1",
        overrides={"loading_port": "上海", "entrust_valid_until": "2027 年 12 月 31日止。"},
    )
    ws = _export(data)
    assert "上海" in ws["合同"]["A33"].value
    assert ws["委托书"]["E12"].value == "2027 年 12 月 31日止。"
