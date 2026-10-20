"""发票导出体积与缓存测试。"""
from io import BytesIO

from openpyxl import Workbook

from fba_tool.invoice_automation import InvoiceDraft, export_invoice_xlsx
from fba_tool.invoice_template_cache import get_template, hash_template, store_template
from fba_tool.xlsx_io import recompress_xlsx, save_workbook_compact


def _minimal_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "Reference"
    ws["B1"] = "SKU"
    for i in range(2, 42):
        ws.cell(i, 1, "")
        ws.cell(i, 2, f"PLACEHOLDER-{i}")
    for r in range(43, 500):
        ws.cell(r, 1, None)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_template_cache_roundtrip():
    raw = _minimal_template()
    key = store_template(raw)
    assert key == hash_template(raw)
    assert get_template(key) == raw


def test_export_smaller_than_raw_template_bloat():
    raw = _minimal_template()
    draft = InvoiceDraft(
        shipment_id="FBA123",
        warehouse_code="FWA4",
        template_sheet="Invoice",
        table_cells=[
            {"cell": "A2", "value": "REF-001", "field_key": "reference_id"},
            {"cell": "B2", "value": "SKU-A", "field_key": "sku"},
        ],
        is_table_mode=True,
    )
    out = export_invoice_xlsx(raw, draft)
    assert out.startswith(b"PK")
    assert len(out) <= len(recompress_xlsx(raw)) + 2048


def test_save_workbook_compact():
    wb = Workbook()
    wb.active["A1"] = "x"
    data = save_workbook_compact(wb)
    assert data.startswith(b"PK")


if __name__ == "__main__":
    test_template_cache_roundtrip()
    test_export_smaller_than_raw_template_bloat()
    test_save_workbook_compact()
    print("ALL OK")
