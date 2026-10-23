"""报关资料 Excel 导出测试。"""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import openpyxl

from fba_tool.customs_export import export_customs_package, validate_customs_template
from fba_tool.customs_package import build_customs_data
from test_customs_package import _sample_pair


TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "customs" / "报关资料模板.xlsx"



import pytest as _pytest


@_pytest.fixture(autouse=True)
def _with_declared_price(monkeypatch):
    """S2-fix D2：无申报价会阻断导出；这里的导出用例给样例 SKU 一个领星价。"""
    from fba_tool import customs_package as _cp

    orig = _cp._catalog_hints
    monkeypatch.setattr(_cp, "_catalog_hints", lambda sku: {**orig(sku), "单个货值": 10.0, "申报价来源": "lingxing"})

def test_export_writes_key_cells():
    if not TEMPLATE.is_file():
        return
    profile = {
        "id": "test",
        "name": "示例公司（请在出口商档案里填写）",
        "credit_code": "DEMO-CREDIT-0001",
        "signed_at": "武汉",
        "domestic_source": "武汉",
    }
    data = build_customs_data([_sample_pair()], [], profile, source="shipment", order_no="DEMO-ORDER-1")
    data.dates["export_date"] = "2026-06-27"
    data.dates["invoice_date"] = "2026-06-27"
    out = export_customs_package(data, TEMPLATE.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(out))
    ws = wb.worksheets[0]
    assert ws["A10"].value == "DEMO-ORDER-1"
    assert ws["A4"].value == profile["name"]
    assert ws["C3"].value == profile["credit_code"]
    assert isinstance(ws["H3"].value, (date, datetime))
    cell_d = ws["H3"].value.date() if isinstance(ws["H3"].value, datetime) else ws["H3"].value
    assert cell_d == date(2026, 6, 27)


def test_export_total_cells_are_values_not_formulas():
    if not TEMPLATE.is_file():
        return
    profile = {
        "id": "test",
        "name": "测试公司",
        "credit_code": "91420113TEST0000",
        "signed_at": "武汉",
        "domestic_source": "武汉",
    }
    data = build_customs_data([_sample_pair()], [], profile, source="shipment", order_no="ORD001")
    out = export_customs_package(data, TEMPLATE.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(out))
    invoice = wb.worksheets[3]
    total_cell = invoice["G15"]
    assert total_cell.data_type != "f"
    assert total_cell.value == data.total_amount


def test_bundled_template_passes_anchor_validation():
    if not TEMPLATE.is_file():
        return
    wb = openpyxl.load_workbook(TEMPLATE)
    assert validate_customs_template(wb) == []


def test_renamed_sheet_is_rejected_instead_of_writing_wrong_cells():
    """报关行换模板导致 sheet 顺序/名称变化时，必须报错而不是静默填错格。"""
    if not TEMPLATE.is_file():
        return
    wb = openpyxl.load_workbook(TEMPLATE)
    wb.worksheets[4].title = "装箱明细"
    problems = validate_customs_template(wb)
    assert any("箱单" in p for p in problems)


def test_shifted_data_zone_is_detected():
    """表头行被上下挪动 → 数据会写到错误行，必须拦下。"""
    if not TEMPLATE.is_file():
        return
    wb = openpyxl.load_workbook(TEMPLATE)
    wb.worksheets[0].insert_rows(5, amount=2)
    problems = validate_customs_template(wb)
    assert any("报关单" in p and "表头" in p for p in problems)


def test_missing_sheets_rejected():
    wb = openpyxl.Workbook()
    problems = validate_customs_template(wb)
    assert len(problems) == 1
    assert "sheet 数量" in problems[0]


def test_export_aborts_on_bad_template():
    if not TEMPLATE.is_file():
        return
    import pytest

    profile = {"id": "t", "name": "测试公司", "credit_code": "X", "signed_at": "武汉"}
    data = build_customs_data([_sample_pair()], [], profile, source="shipment", order_no="ORD002")

    bad = openpyxl.load_workbook(TEMPLATE)
    bad.worksheets[3].title = "Invoice"
    buf = io.BytesIO()
    bad.save(buf)

    with pytest.raises(ValueError, match="模板结构与预期不符"):
        export_customs_package(data, buf.getvalue())