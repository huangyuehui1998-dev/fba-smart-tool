"""S3-2 报关单元格：日期留空/序列号带格式、委托书签署日期坐标、合同备注、模板残留 0 清理（方案 v8 §3.6 验收 2/4）。"""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from fba_tool import customs_package as cp
from fba_tool.customs_export import export_customs_package
from fba_tool.customs_package import build_customs_from_tickets
from fba_tool.warehouse_db import update_rule_profile

TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "customs" / "报关资料模板.xlsx"
PROFILE = {"id": "t", "name": "测试公司", "credit_code": "X", "signed_at": "武汉", "domestic_source": "武汉"}
PET = {"chinese_name": "宠物垫", "hs_code": "6307909000", "unit": "包", "brand": "DemoBrand", "category": "pet"}


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    ad = tmp_path / "ad"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    from fba_tool.master_db import init_master_db

    init_master_db()
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {})


def _data(**kw):
    t = {"shipment_id": "FBA1", "warehouse": "SBD1", "boxes": [10, 10], "lines": [dict(PET, sku="A", qty=8, unit_price=39.99)]}
    return build_customs_from_tickets([t], PROFILE, order_no="JUS1DD2", **kw)


def _wb(data):
    return openpyxl.load_workbook(io.BytesIO(export_customs_package(data, TEMPLATE.read_bytes())))


def test_export_and_declare_dates_blank_freight_blank():
    ws = _wb(_data())["报关单"]
    assert ws["H3"].value is None  # 出口日期
    assert ws["L3"].value is None  # 申报日期（K3 标签，L3:N3 合并）
    assert ws["K12"].value is None  # 运费


def test_ticket_freight_never_written():
    d = _data()
    d.freight = "USD/427.15/3"
    assert _wb(d)["报关单"]["K12"].value is None


def test_profile_switch_fills_export_declare_date_from_etd():
    update_rule_profile({"fill_export_declare_date": 1})
    ws = _wb(_data(dates={"etd": "2026-11-02"}))["报关单"]
    assert ws["H3"].value.date() == date(2026, 11, 2) and ws["L3"].value.date() == date(2026, 11, 2)
    assert ws["H3"].is_date and "yy" in ws["L3"].number_format


@pytest.mark.parametrize("sheet,cell,key", [
    ("发票", "G3", "invoice_date"), ("箱单", "G2", "packing_date"), ("合同", "F8", "contract_date"),
    ("委托书", "H16", "entrust_sign_date"), ("委托书", "H23", "entrust_doc_received_date"),
    ("委托书", "C25", "entrust_import_export_date"),
])
def test_dates_are_serial_with_date_format(sheet, cell, key):
    d = _data(dates={key: "2026-10-20"})
    c = _wb(d)[sheet][cell]
    assert isinstance(c.value, datetime) and c.value.date() == date(2026, 10, 20)
    assert c.is_date and "yy" in c.number_format.lower()


def test_entrust_checkbox_cell_not_overwritten():
    assert _wb(_data())["委托书"]["H24"].value == "合同□√"


def test_contract_b10_remarks_match_order():
    wb = _wb(_data())
    assert wb["合同"]["B10"].value == "订单编号：JUS1DD2。仓库编码：SBD1。FBA编码：FBA1。"
    assert wb["合同"]["F6"].value == "JUS1DD2"


def test_unused_template_rows_cleared():
    wb = _wb(_data())
    assert wb["发票"]["C9"].value is None and wb["发票"]["D9"].value is None
    assert wb["申报要素"]["B11"].value is None and wb["申报要素"]["I11"].value is None
    assert wb["合同"]["A19"].value is None and wb["合同"]["G19"].value is None
    assert wb["箱单"]["B10"].value is None
    assert wb["发票"]["A9"].value is None  # 序号也清掉


def test_port_and_validity_follow_profile():
    update_rule_profile({"loading_port": "上海"})
    wb = _wb(_data())
    assert "上海" in wb["合同"]["A33"].value
    assert wb["委托书"]["E12"].value == f"{date.today().year} 年 12 月 31日止。"
