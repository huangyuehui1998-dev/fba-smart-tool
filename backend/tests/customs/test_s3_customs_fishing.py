"""S3-6 钓具导出与边界（方案 v8 §3.6 验收 3、§5-7）：已配置钓具能导出、缺要素阻断、宠物+钓具混装输出 2 行各用各的要素；
单件装标准尺寸 23kg 箱按 v8 只提示（Team Lift）不拦截（用户未回复前维持 v8）。"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from fba_tool import customs_package as cp
from fba_tool.customs_export import CustomsExportBlocked, export_customs_package
from fba_tool.customs_package import build_customs_from_tickets
from fba_tool.hs_config import upsert_hs_template

TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "customs" / "报关资料模板.xlsx"
PROFILE = {"id": "t", "name": "测试公司", "credit_code": "X", "signed_at": "武汉", "domestic_source": "武汉"}
PET = {"sku": "P", "chinese_name": "宠物垫", "hs_code": "6307909000", "unit": "包", "brand": "DemoBrand", "category": "pet", "qty": 10, "unit_price": 39.99}
FISH = {"sku": "F", "chinese_name": "钓鱼竿", "hs_code": "9507100000", "unit": "支", "brand": "TESTFISH", "category": "fishing", "qty": 4, "unit_price": 25.0}


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    ad = tmp_path / "ad"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    from fba_tool.master_db import init_master_db

    init_master_db()
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {})


def _fish_template():
    # 测试数据，不是真实报关行确认的要素（外部核实项 E4）
    upsert_hs_template({"category": "fishing", "hs_code": "9507100000", "customs_name": "钓鱼竿",
                        "elements_pattern": "1|2|钓鱼用|碳素|品牌：{品牌}牌|型号：{型号}", "default_purpose": "钓鱼用"})


def _data(lines):
    return build_customs_from_tickets([{"shipment_id": "FBA1", "warehouse": "SBD1", "boxes": [10, 12], "lines": lines}], PROFILE, order_no="JUS1DD2")


def test_fishing_without_config_blocks_export():
    d = _data([FISH])
    with pytest.raises(CustomsExportBlocked) as e:
        export_customs_package(d, TEMPLATE.read_bytes())
    assert "缺报关配置" in str(e.value)


def test_fishing_with_config_exports_with_own_elements():
    _fish_template()
    d = _data([FISH])
    assert d.blocking == []
    wb = openpyxl.load_workbook(io.BytesIO(export_customs_package(d, TEMPLATE.read_bytes())))
    assert wb["申报要素"]["I10"].value == "1|2|钓鱼用|碳素|品牌：TESTFISH牌|型号：无型号"
    assert wb["报关单"]["H20"].value == "支"


@pytest.mark.parametrize("fmt", ["xlsx", "xls"])
def test_pet_and_fishing_mixed_two_lines(fmt):
    _fish_template()
    d = _data([PET, FISH])
    assert [l.hs_code for l in d.lines] == ["6307909000", "9507100000"]
    raw = export_customs_package(d, TEMPLATE.read_bytes(), fmt=fmt)
    if fmt == "xlsx":
        ws = openpyxl.load_workbook(io.BytesIO(raw))["申报要素"]
        vals = [ws["I10"].value, ws["I11"].value]
    else:
        import xlrd

        ws = xlrd.open_workbook(file_contents=raw).sheet_by_name("申报要素")
        vals = [ws.cell_value(9, 8), ws.cell_value(10, 8)]
    assert vals == ["1|2|宠物用|品牌：DemoBrand牌|型号：无型号", "1|2|钓鱼用|碳素|品牌：TESTFISH牌|型号：无型号"]


def test_fishing_never_uses_pet_template():
    d = _data([dict(FISH, hs_code="6307909000")])  # 钓具即使 HS 撞上宠物模板，也不能套用
    assert any(b["code"] == "customs_config_missing" for b in d.blocking)


def test_single_item_standard_23kg_only_notice_v8():
    """v8 §3.4：单件装且箱重 > 22.68kg 直接算大件 → 只提示 Team Lift，不拦截（用户未回复，维持 v8）。"""
    from test_fba_rules_s1e import _run
    from fba_tool.models import SkuRow

    sku = SkuRow("STD", 1, 1, 30, 20, 10, 23.0, 0.5)
    comp = _run([sku], [(1, "STD", 1)], {1: 23.0})
    assert comp["blocked"] is False
    assert [n["label"] for n in comp["lift_notices"]] == ["Team Lift"]
