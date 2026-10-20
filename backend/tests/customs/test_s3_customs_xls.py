"""S3-4 .xls 导出（v5 Q15）：同一票 .xlsx 与 .xls 的 7 张表关键字段一致；留空单元格在两种格式里都为空；日期是序列号带日期格式。"""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import openpyxl
import pytest
import xlrd

from fba_tool import customs_package as cp
from fba_tool.customs_export import export_customs_package
from fba_tool.customs_package import build_customs_from_tickets

TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "customs" / "报关资料模板.xlsx"
PROFILE = {"id": "t", "name": "测试公司", "credit_code": "X", "signed_at": "武汉", "domestic_source": "武汉"}
PET = {"chinese_name": "宠物垫", "hs_code": "6307909000", "unit": "包", "brand": "DemoBrand", "category": "pet"}
SHEETS = ("报关单", "填制规范", "申报要素", "发票", "箱单", "合同", "委托书")


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    ad = tmp_path / "ad"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    from fba_tool.master_db import init_master_db

    init_master_db()
    monkeypatch.setattr(cp, "_catalog_hints", lambda sku: {})


def _data():
    lines = [dict(PET, sku="A", qty=52, unit_price=39.99), dict(PET, sku="B", qty=32, unit_price=49.99),
             dict(sku="C", qty=5, unit_price=69.99, chinese_name="车载狗栏", hs_code="7326909000", unit="个", brand="DemoBrand", category="pet")]
    t = {"shipment_id": "FBA19P4NX8HH", "warehouse": "SBD1", "boxes": [7.5] * 13 + [9.5] * 8, "lines": lines}
    return build_customs_from_tickets([t], PROFILE, order_no="JUS13DD21", dates={"invoice_date": "2026-10-20"})


def _both():
    d = _data()
    x = openpyxl.load_workbook(io.BytesIO(export_customs_package(d, TEMPLATE.read_bytes(), fmt="xlsx")), data_only=True)
    b = xlrd.open_workbook(file_contents=export_customs_package(d, TEMPLATE.read_bytes(), fmt="xls"), formatting_info=True)
    return x, b


def _xls_value(b, sheet, ref):
    ws = b.sheet_by_name(sheet)
    col = openpyxl.utils.column_index_from_string("".join(c for c in ref if c.isalpha())) - 1
    row = int("".join(c for c in ref if c.isdigit())) - 1
    if row >= ws.nrows or col >= ws.ncols:
        return None, None
    c = ws.cell(row, col)
    return c, ws


def test_xls_is_real_biff_with_seven_sheets():
    raw = export_customs_package(_data(), TEMPLATE.read_bytes(), fmt="xls")
    assert raw[:8] == bytes.fromhex("D0CF11E0A1B11AE1")
    assert tuple(xlrd.open_workbook(file_contents=raw).sheet_names()) == SHEETS


def test_all_cells_match_between_formats():
    x, b = _both()
    diffs = []
    for name in SHEETS:
        wx = x[name]
        for row in wx.iter_rows():
            for c in row:
                if type(c).__name__ == "MergedCell":
                    continue
                cell, ws = _xls_value(b, name, c.coordinate)
                xv = c.value
                if hasattr(xv, "date"):
                    ok = cell is not None and cell.ctype == xlrd.XL_CELL_DATE and xlrd.xldate_as_datetime(cell.value, b.datemode) == xv
                elif xv in (None, ""):
                    ok = cell is None or cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
                elif isinstance(xv, (int, float)):
                    ok = cell is not None and cell.ctype == xlrd.XL_CELL_NUMBER and abs(cell.value - xv) < 1e-9
                else:
                    ok = cell is not None and cell.value == str(xv)
                if not ok:
                    diffs.append((name, c.coordinate, xv, cell and cell.value))
    assert diffs == []


@pytest.mark.parametrize("ref", ["报关单!K12", "报关单!H3", "报关单!L3"])
def test_blank_cells_blank_in_both(ref):
    x, b = _both()
    sheet, cell = ref.split("!")
    assert x[sheet][cell].value is None
    c, _ = _xls_value(b, sheet, cell)
    assert c is None or c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)


def test_xls_dates_are_date_serials():
    _, b = _both()
    c, _ = _xls_value(b, "发票", "G3")
    assert c.ctype == xlrd.XL_CELL_DATE
    assert xlrd.xldate_as_datetime(c.value, b.datemode).date() == date(2026, 10, 20)


def test_xls_merges_kept():
    x, b = _both()
    ws = b.sheet_by_name("报关单")
    xl = {(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col) for m in x["报关单"].merged_cells.ranges}
    assert xl == set(ws.merged_cells)


def test_two_lines_written_in_xls():
    _, b = _both()
    ws = b.sheet_by_name("报关单")
    assert ws.cell_value(19, 1) == "6307909000" and ws.cell_value(20, 1) == "7326909000"
    assert ws.cell_value(19, 8) == 43.8 and ws.cell_value(20, 8) == 69.99


def test_bad_format_rejected():
    with pytest.raises(ValueError):
        export_customs_package(_data(), TEMPLATE.read_bytes(), fmt="csv")


def test_api_export_format_and_filename(monkeypatch):
    from fastapi.testclient import TestClient
    from main import app

    d = _data()
    c = TestClient(app)
    r = c.post("/api/customs/export", json={"source": "shipment", "order_no": "JUS13DD21", "data": d.to_dict(), "export_format": "xls"})
    assert r.status_code == 200, r.text
    assert r.content[:4] == bytes.fromhex("D0CF11E0")
    assert "application/vnd.ms-excel" in r.headers["content-type"]
    assert "JUS13DD21-FBA19P4NX8HH-%E5%8D%95%E7%8B%AC" in r.headers["content-disposition"] or "单独报关" in r.headers["content-disposition"]
    r2 = c.post("/api/customs/export", json={"source": "shipment", "order_no": "JUS13DD21", "data": d.to_dict()})
    assert r2.status_code == 200 and r2.content[:2] == b"PK"
