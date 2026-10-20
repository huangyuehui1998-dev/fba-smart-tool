"""Tests for Send to Amazon (STA) paste-grid export (K 死亡空列 / L 起箱列)."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import load_workbook

from fba_tool.excel_export import build_box_matrix
from fba_tool.models import SkuRow
from fba_tool.sta_template_export import (
    IDX_DEATH_COL,
    IDX_FIRST_BOX,
    STA_CM_PER_INCH,
    build_sta_grid,
    export_sta_template_bytes,
)
from fba_tool.units import kg_to_lb

COL_K = IDX_DEATH_COL + 1
COL_L = IDX_FIRST_BOX + 1


def _sample_editor() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"箱号": 1, "SKU简称": "SKU-A", "数量": 6},
            {"箱号": 3, "SKU简称": "SKU-A", "数量": 6},
            {"箱号": 1, "SKU简称": "SKU-B", "数量": 5},
            {"箱号": 2, "SKU简称": "SKU-B", "数量": 8},
        ]
    )


def _sample_sku_rows() -> list[SkuRow]:
    return [
        SkuRow(
            sku="SKU-A",
            total_qty=12,
            carton_qty=6,
            length_cm=50.0,
            width_cm=40.0,
            height_cm=30.0,
            unit_weight_kg=1.0,
            carton_tare_kg=0.5,
            carton_gross_kg=10.0,
        ),
        SkuRow(
            sku="SKU-B",
            total_qty=13,
            carton_qty=5,
            length_cm=48.0,
            width_cm=38.0,
            height_cm=28.0,
            unit_weight_kg=1.2,
            carton_tare_kg=0.5,
            carton_gross_kg=12.5,
        ),
    ]


def _meta():
    return {
        "SKU-A": {"asin": "B001", "fnsku": "X001"},
        "SKU-B": {"asin": "B002", "fnsku": "X002"},
    }


def _load_ws():
    editor = _sample_editor()
    skus = _sample_sku_rows()
    matrix = build_box_matrix(editor, max_boxes=3)
    data = export_sta_template_bytes(
        editor,
        skus,
        max_boxes=3,
        box_matrix=matrix,
        amazon_meta=_meta(),
        fnsku_map={"SKU-A": "X001", "SKU-B": "X002"},
        product_names={"SKU-A": "Product A", "SKU-B": "Product B"},
    )
    return load_workbook(io.BytesIO(data)).active


def test_sta_row3_box_count_at_k_l():
    ws = _load_ws()
    assert "SKU 总数：2" in str(ws.cell(3, 1).value)
    assert ws.cell(3, COL_K).value == "包装箱总数："
    assert ws.cell(3, COL_L).value == 3


def test_sta_header_death_col_k():
    ws = _load_ws()
    assert ws.cell(5, 10).value == "装箱数量"
    assert ws.cell(5, COL_K).value in (None, "")
    assert ws.cell(5, COL_L).value == "包装箱 1 数量"
    assert ws.cell(5, COL_L + 1).value == "包装箱 2 数量"


def test_sta_sku_row_qty_and_zero_blank():
    ws = _load_ws()
    assert ws.cell(6, 1).value == "SKU-A"
    assert ws.cell(6, 2).value == "Product A"
    assert ws.cell(6, 3).value == "B001"
    assert ws.cell(6, COL_K).value in (None, "")
    assert ws.cell(6, COL_L).value == 6
    assert ws.cell(6, COL_L + 1).value in (None, "")
    assert ws.cell(6, COL_L + 2).value == 6
    assert ws.cell(6, 10).value == 12


def test_sta_footer_positions_and_p1b():
    ws = _load_ws()
    assert ws.cell(8, 1).value in (None, "")
    assert ws.cell(9, 1).value == "包装箱名称"
    assert ws.cell(9, COL_L).value == "P1 - B1"
    assert ws.cell(9, COL_L + 2).value == "P1 - B3"
    assert ws.cell(10, 1).value == "包装箱重量（磅）："
    box1_gross = round(6 * 1.0 + 5 * 1.2 + 0.5, 2)
    assert ws.cell(10, COL_L).value == kg_to_lb(box1_gross, 2)


def test_sta_metric_conversion():
    editor = _sample_editor()
    matrix = build_box_matrix(editor, max_boxes=1)
    grid = build_sta_grid(
        sku_rows=[_sample_sku_rows()[0]],
        box_matrix=matrix,
        editor_df=editor,
        amazon_meta=_meta(),
        fnsku_map={},
        product_names={"SKU-A": "A"},
        max_box=1,
    )
    width_row = grid[-3]
    assert width_row[IDX_FIRST_BOX] == round(40.0 / STA_CM_PER_INCH, 2)
