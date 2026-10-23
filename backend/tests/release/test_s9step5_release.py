"""S9-step5：安装包只打白名单（packaging/release_manifest.py），模板和出厂配置不含真实数据。

* 包内容扫描：真实 SKU / ASIN / FNSKU 从仓库里的真实产品表、原件出厂配置和原件模板现场读出来，只放在内存里；
  再加形状规则（ASIN / FNSKU / 货件号 / 公司名 / 信用代码）。失败信息只给文件名和计数，不给任何值。
* 老用户升级：出厂默认（空映射、空出口商档案、示例产品表）不会混进老用户已有的配置、模板和主库。
* release 版报关模板能通过锚点校验并正常导出；工厂表没有模板时走兜底版式。
全部用 tmp_path / 临时 FBA_APP_DATA_DIR，不碰真实 app_data。
"""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import io
import json
import re
import shutil
import sqlite3
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Dict, Iterable, List, Set

import openpyxl
import pandas as pd
import pytest

from fba_tool import app_paths

BACKEND = Path(__file__).resolve().parents[2]
ROOT = BACKEND.parent
PACKAGING = ROOT / "packaging"
REAL_CATALOG = ROOT / "archive" / "private_data" / "产品信息表.csv"


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


M = _load("release_manifest", PACKAGING / "release_manifest.py")
GEN = _load("build_release_templates_t", PACKAGING / "build_release_templates.py")

ASIN_RE = re.compile(r"(?<![0-9A-Za-z])B0[0-9A-Z]{8}(?![0-9A-Za-z])")
FNSKU_RE = re.compile(r"(?<![0-9A-Za-z])X00[0-9A-Z]{7}(?![0-9A-Za-z])")
SHIPMENT_RE = re.compile(r"(?<![0-9A-Za-z])FBA1[0-9A-Z]{7,}(?![0-9A-Za-z])")


# ---------------------------------------------------------------- 真实值参照集（只在内存）

def _csv_rows(path: Path) -> List[List[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(text)))


def _tokens(text: str) -> Set[str]:
    return set(re.findall(r"[0-9A-Za-z][0-9A-Za-z\-_.]{5,}", text))


def _json_strings(obj) -> Iterable[str]:
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield str(k)
            yield from _json_strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _json_strings(v)
    elif obj is not None:
        yield str(obj)


def build_real_refs() -> Dict[str, Set[str]]:
    refs: Dict[str, Set[str]] = {"sku": set(), "asin": set(), "fnsku": set(), "text": set()}
    if REAL_CATALOG.is_file():
        rows = [r for r in _csv_rows(REAL_CATALOG) if any(c.strip() for c in r)]
        head = [h.strip() for h in rows[0]]
        sku_cols = [i for i, h in enumerate(head) if h.split("\n")[0].strip().upper() in ("SKU", "MSKU", "SKU简称", "SELLER SKU", "SELLER_SKU")]
        for r in rows[1:]:
            for i, c in enumerate(r):
                c = c.strip()
                if ASIN_RE.fullmatch(c):
                    refs["asin"].add(c)
                elif FNSKU_RE.fullmatch(c):
                    refs["fnsku"].add(c)
                elif i in sku_cols and len(c) >= 6:
                    refs["sku"].add(c)
    cfg = BACKEND / "config"
    for name in ("amazon_sku_metadata.default.json", "fnsku_mapping.default.json", "sku_mapping.default.json"):
        p = cfg / name
        if p.is_file():
            for s in _json_strings(json.loads(p.read_text(encoding="utf-8"))):
                s = s.strip()
                if ASIN_RE.fullmatch(s):
                    refs["asin"].add(s)
                elif FNSKU_RE.fullmatch(s):
                    refs["fnsku"].add(s)
                elif re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9\-_.]{7,}", s) and "-" in s:
                    refs["sku"].add(s)
    for p in (BACKEND / "templates").glob("亚马逊FBA表*.csv"):
        for r in _csv_rows(p):
            for c in r:
                c = c.strip()
                if ASIN_RE.fullmatch(c):
                    refs["asin"].add(c)
                elif FNSKU_RE.fullmatch(c):
                    refs["fnsku"].add(c)
    # 原件报关模板 / 货代模板里学出来的公司名、品牌、品名、订单号等（生成器同一套规则）
    placeholders = {GEN.DEMO_BRAND, GEN.DEMO_COMPANY, GEN.DEMO_PRODUCT, GEN.DEMO_CITY, "DEMO-ORDER", "FBA-DEMO", "DEMO", "FBATEST0DEMO"}
    orig_customs = BACKEND / "templates" / "customs" / "报关资料模板.xlsx"
    if orig_customs.is_file():
        wb = openpyxl.load_workbook(orig_customs)
        # 只取公司名和品牌（品名、城市是普通词，不算真实标识）。已经是 DEMO 的不再当真实参照。
        refs["text"] |= {
            k for k, v in GEN._learn_customs(wb).items()
            if v in (GEN.DEMO_BRAND, GEN.DEMO_COMPANY) and len(k) >= 2 and k not in placeholders
        }
        public = "\n".join(str(c.value) for ws in wb.worksheets if ws.title == "填制规范" for row in ws.iter_rows() for c in row if c.value)
        for ws in wb.worksheets:
            if ws.title == "填制规范":
                continue
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        for rx in (GEN.ORDER_RE, GEN.SHIPMENT_RE, GEN.CREDIT_RE):
                            # 填制规范里也有的是海关公开的填报样式（如 AEO 编码示例），不算真实值。
                            # 已经换成 DEMO / FBA-DEMO 的格子也不再当真实参照。
                            refs["text"] |= {t for t in rx.findall(c.value) if t not in public and t not in placeholders}
    refs["text"] -= placeholders
    for rel in ("forwarders/qiluo_t6.xlsx", "forwarders/deshu_t3.xlsx"):
        p = BACKEND / "templates" / rel
        if p.is_file():
            for ws in openpyxl.load_workbook(p).worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str):
                            refs["text"] |= set(GEN.COMPANY_RE.findall(c.value)) | set(GEN.CREDIT_RE.findall(c.value))
    # D-OSS-01：三份出厂映射已与便携包一样是 {}。没有真实产品表时，参照集可以是空的。
    # 包内容仍按形状规则扫（ASIN / FNSKU / FBA1* 货件号 / 公司名 / 信用代码）。
    for name in ("amazon_sku_metadata.default.json", "fnsku_mapping.default.json", "sku_mapping.default.json"):
        assert json.loads((cfg / name).read_text(encoding="utf-8")) == {}
    if REAL_CATALOG.is_file():
        assert len(refs["sku"]) >= 6 and len(refs["asin"]) >= 6 and len(refs["fnsku"]) >= 6, {k: len(v) for k, v in refs.items()}
    return refs


@pytest.fixture(scope="module")
def real_refs() -> Dict[str, Set[str]]:
    return build_real_refs()


# ---------------------------------------------------------------- 取文件文字

def _xlsx_texts(path: Path) -> List[str]:
    out: List[str] = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".vml", ".rels")):
                out.append(z.read(n).decode("utf-8", "replace"))
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        out.append(ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out.append(str(c.value))
    return out


def _file_texts(path: Path) -> List[str]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _xlsx_texts(path)
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return [raw.decode(enc)]
        except UnicodeDecodeError:
            continue
    return [raw.decode("latin-1")]


def _hits(texts: List[str], refs: Dict[str, Set[str]], shapes: bool) -> Dict[str, int]:
    blob = "\n".join(texts)
    toks = _tokens(blob)
    cnt: Dict[str, int] = {}
    for key in ("sku", "asin", "fnsku"):
        n = len(refs[key] & toks) + sum(1 for v in refs[key] if v not in toks and v in blob)
        if n:
            cnt["real_" + key] = n
    n = sum(1 for v in refs["text"] if v in blob)
    if n:
        cnt["real_text"] = n
    if shapes:
        for key, rx in (("asin_shape", ASIN_RE), ("fnsku_shape", FNSKU_RE), ("shipment_shape", SHIPMENT_RE)):
            n = sum(1 for m in rx.findall(blob) if "DEMO" not in m)
            if n:
                cnt[key] = n
        n = sum(1 for m in GEN.COMPANY_RE.finditer(blob) if "示例公司" not in m.group(0))
        if n:
            cnt["company_shape"] = n
        n = len(GEN.CREDIT_RE.findall(blob))
        if n:
            cnt["credit_shape"] = n
    return cnt


# ---------------------------------------------------------------- 1. 白名单本身

def test_manifest_sources_exist_and_release_files_present():
    entries = M.entries(ROOT)
    missing = [d for s, d in entries if not s.is_file()]
    assert missing == []
    assert len(M.pyinstaller_datas(ROOT)) == len(entries)
    rel_root = str(M.release_dir(ROOT))
    for kind, table in (("templates", M.TEMPLATE_WHITELIST), ("config", M.CONFIG_WHITELIST)):
        for rel, origin in table.items():
            assert str(M.source_path(kind, rel, ROOT)).startswith(rel_root) == (origin == M.RELEASE), rel


def test_manifest_excludes_real_data_files():
    tpl = set(M.TEMPLATE_WHITELIST)
    bundled = {p.relative_to(BACKEND / "templates").as_posix() for p in (BACKEND / "templates").rglob("*") if p.is_file()}
    excluded = bundled - tpl
    for pat in (r"^亚马逊FBA表\d*\.csv$", r"^J-US.*发票\.xlsx$", r"^箱唛.*\.pdf$", r"^发货录入样表\.xlsx$",
                r"^已填写好的send to amazong\.xlsx$", r"^工厂表\.xlsx$", r"^货代表\.xlsx$", r"^Send to Amazon表格\.xlsx$"):
        assert not any(re.search(pat, r) for r in tpl), pat
    # Marco 2026-10-20：3 份真实数据模板原样留在仓库，不再要求“排除数 ≥ 19”。
    assert M.ALLOWED_REAL_TEMPLATES_DECISION_DATE == "2026-10-20"
    assert set(M.ALLOWED_REAL_TEMPLATES) == {"发货录入样表.xlsx", "工厂表.xlsx", "货代表.xlsx"}
    for name in M.ALLOWED_REAL_TEMPLATES:
        assert (BACKEND / "templates" / name).is_file(), name
        assert name not in tpl
    assert set(M.KNOWN_BLANK_TEMPLATES) == {"Send to Amazon表格.xlsx"}
    allowed_repo_binaries = set(tpl) | set(M.ALLOWED_REAL_TEMPLATES) | set(M.KNOWN_BLANK_TEMPLATES)
    assert excluded <= allowed_repo_binaries, sorted(excluded - allowed_repo_binaries)
    stray = Path("backend/not-allowed-real.xlsx")
    assert stray.as_posix() not in {f"templates/{n}" for n in allowed_repo_binaries}
    assert not any(app_paths.is_private_template(Path(r)) for r in tpl)
    cfg = set(M.CONFIG_WHITELIST)
    assert "invoice_history.default.json" not in cfg and "lingxing.json" not in cfg
    assert not any("sessions" in c for c in cfg)
    assert M.forbidden_real_hits(ROOT) == []
    assert not hasattr(M, "ALLOWED_LISTING_FIXTURES")
    assert not (ROOT / "review" / "lingxing_listing_export_2026-10-21.xlsx").exists()
    assert "lingxing_listing_prices.xlsx" in M.FORBIDDEN_REAL_NAMES


def test_oss_defaults_drop_real_catalog_and_shipment_ids():
    """D-OSS-01 / D-OSS-02：出厂映射是 {}；报关模板与便携包脱敏版单元格一致；样表和德速示例不再带原货号。"""
    banned = ("JoyBlooms", "J-PETWOODSTAIRS", "X0054", "B0H2Y", "JUS7HB115", "FBA19GV", "FBA22232AF13")
    paths = [
        BACKEND / "config" / "amazon_sku_metadata.default.json",
        BACKEND / "config" / "fnsku_mapping.default.json",
        BACKEND / "config" / "sku_mapping.default.json",
        BACKEND / "config" / "invoice_history.default.json",
        BACKEND / "templates" / "customs" / "报关资料模板.xlsx",
        BACKEND / "templates" / "发货录入样表.xlsx",
        BACKEND / "templates" / "forwarders" / "deshu_t3.xlsx",
        PACKAGING / "release_templates" / "templates" / "forwarders" / "deshu_t3.xlsx",
        PACKAGING / "release_templates" / "templates" / "customs" / "报关资料模板.xlsx",
        PACKAGING / "Step5报关资料_架构与工作流.txt",
    ]
    for path in paths:
        blob = "\n".join(_file_texts(path))
        found = [token for token in banned if token in blob]
        assert found == [], (path.relative_to(ROOT).as_posix(), found)
    src = openpyxl.load_workbook(BACKEND / "templates" / "customs" / "报关资料模板.xlsx")
    rel = openpyxl.load_workbook(M.source_path("templates", "customs/报关资料模板.xlsx", ROOT))
    for left, right in zip(src.worksheets, rel.worksheets):
        assert [[c.value for c in row] for row in left.iter_rows()] == [[c.value for c in row] for row in right.iter_rows()]
    assert "FBATEST0DEMO" in "\n".join(_file_texts(BACKEND / "templates" / "forwarders" / "deshu_t3.xlsx"))


def _gitignore_rules() -> List[str]:
    rules = []
    for raw in (ROOT / ".gitignore").read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line and not line.startswith("#"):
            rules.append(line)
    return rules


def _check_ignore_pattern(rel: str) -> str:
    """最后一条命中的 gitignore 规则。空字符串表示没有任何规则命中。以 ! 开头的是放行。"""
    proc = subprocess.run(
        ["git", "check-ignore", "-v", "--no-index", "--", rel],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode == 1:
        return ""
    assert proc.returncode == 0, (rel, proc.stderr)
    left = proc.stdout.strip().splitlines()[-1].split("\t", 1)[0]
    match = re.match(r"^.*?:\d+:(.*)$", left)
    assert match, proc.stdout
    return match.group(1)


def test_gitignore_matches_forbidden_real_rules():
    """运营 r1：.gitignore 与 FORBIDDEN_REAL_* 对拍。模板放行不能把 listing xlsx 放回来。"""
    rules = _gitignore_rules()
    for name in M.FORBIDDEN_REAL_NAMES:
        assert name in rules, name
    for pattern in M.FORBIDDEN_REAL_GLOBS:
        assert pattern in rules, pattern
    ignored = (
        "lingxing_listing_prices.xlsx",
        "review/lingxing_listing_export_2026-10-21.xlsx",
        "backend/templates/lingxing_listing_prices.xlsx",
        "backend/templates/foo-listing-bar.xlsx",
        "nested/xlistingx.xlsx",
    )
    for rel in ignored:
        pattern = _check_ignore_pattern(rel)
        assert pattern in set(M.FORBIDDEN_REAL_NAMES) | set(M.FORBIDDEN_REAL_GLOBS), (rel, pattern)
        assert not pattern.startswith("!")
    kept = (
        "backend/tests/fixtures/synthetic/lingxing_us_columns_82.xlsx",
        "backend/templates/forwarders/send.xlsx",
    )
    for rel in kept:
        pattern = _check_ignore_pattern(rel)
        assert pattern == "" or pattern.startswith("!"), (rel, pattern)


def test_spec_uses_manifest_only():
    text = (PACKAGING / "FBA_Smart_Tool.spec").read_text(encoding="utf-8")
    assert "from release_manifest import pyinstaller_datas" in text and "*_release_datas(ROOT)," in text
    assert "rglob" not in text and '"templates")' not in text


# ---------------------------------------------------------------- 2. 包内容扫描

def test_release_package_contents_have_no_real_values(real_refs):
    bad = {}
    for src, dest in M.entries(ROOT):
        cnt = _hits(_file_texts(src), real_refs, shapes=True)
        if cnt:
            bad[dest] = cnt
    assert bad == {}, bad  # 只有文件名和计数


def test_shipped_docs_have_no_real_values(real_refs):
    docs = [p for p in (PACKAGING / "使用说明.txt", ROOT / "使用说明.txt", ROOT / "FBA_Smart_Tool_工作流与运行逻辑.txt",
                        PACKAGING / "工作流与运行逻辑_V1.2.txt") if p.is_file()]
    assert docs
    bad = {p.name: c for p in docs for c in [_hits(_file_texts(p), real_refs, shapes=True)] if c}
    assert bad == {}, bad


def test_source_code_has_no_real_values(real_refs):
    files: List[Path] = []
    for base, pats in ((BACKEND / "fba_tool", ("*.py",)), (BACKEND / "routers", ("*.py",)), (BACKEND / "services", ("*.py",)),
                       (ROOT / "frontend" / "src", ("*.ts", "*.tsx"))):
        if base.is_dir():
            for pat in pats:
                files += list(base.rglob(pat))
    files += [p for p in (BACKEND / "main.py", BACKEND / "settings.py", BACKEND / "static_ui.py", ROOT / "run_app.py") if p.is_file()]
    assert len(files) > 50
    bad = {}
    for p in files:
        cnt = _hits(_file_texts(p), real_refs, shapes=False)
        if cnt:
            bad[p.relative_to(ROOT).as_posix()] = cnt
    assert bad == {}, bad


def _media_hashes(path: Path, sheet_index: int = None) -> Set[str]:
    out: Set[str] = set()
    wb = openpyxl.load_workbook(path)
    sheets = wb.worksheets if sheet_index is None else [wb.worksheets[sheet_index]]
    for ws in sheets:
        for img in getattr(ws, "_images", []):
            data = img._data()
            out.add(hashlib.sha256(data).hexdigest())
    return out


@pytest.mark.parametrize("rel", ["forwarders/qiluo_t6.xlsx", "forwarders/deshu_t3.xlsx"])
def test_forwarder_release_has_no_product_photos(rel):
    rel_path = M.source_path("templates", rel, ROOT)
    orig = BACKEND / "templates" / rel
    wb = openpyxl.load_workbook(rel_path)
    assert len(wb.worksheets[0]._images) == 0  # 发票主表上的真实产品图全部去掉
    assert _media_hashes(orig, 0) and not (_media_hashes(rel_path) & _media_hashes(orig, 0))
    assert len(_media_hashes(rel_path)) == 1  # 只留货代自带的示例截图
    owb = openpyxl.load_workbook(orig)
    assert wb.sheetnames == owb.sheetnames
    for a, b in zip(owb.worksheets, wb.worksheets):  # 版式不变：表头 / 合并单元格 / 下拉
        assert a.dimensions == b.dimensions and len(a.merged_cells.ranges) == len(b.merged_cells.ranges)
        assert len(a.data_validations.dataValidation) == len(b.data_validations.dataValidation)
        diff = sum(1 for row in a.iter_rows() for c in row if c.value != b[c.coordinate].value)
        assert diff <= 1


def test_release_configs_are_empty_or_redacted():
    for name in ("amazon_sku_metadata.default.json", "fnsku_mapping.default.json", "sku_mapping.default.json"):
        assert json.loads(M.source_path("config", name, ROOT).read_text(encoding="utf-8")) == {}
    assert json.loads(M.source_path("config", "exporter_profiles.default.json", ROOT).read_text(encoding="utf-8")) == {"active_id": "", "profiles": []}
    orig = json.loads((BACKEND / "config" / "invoice_template_map.json").read_text(encoding="utf-8"))
    rel = json.loads(M.source_path("config", "invoice_template_map.json", ROOT).read_text(encoding="utf-8"))
    assert set(rel["templates"]) == set(orig["templates"])
    for k, t in orig["templates"].items():  # 只改 decision 说明文字，其余字段原样
        assert {kk: vv for kk, vv in rel["templates"][k].items() if kk != "decision"} == {kk: vv for kk, vv in t.items() if kk != "decision"}
    assert {k: v for k, v in rel.items() if k != "templates"} == {k: v for k, v in orig.items() if k != "templates"}

    def strip_source(o):
        if isinstance(o, dict):
            return {k: (None if k == "source" else strip_source(v)) for k, v in o.items()}
        if isinstance(o, list):
            return [strip_source(v) for v in o]
        return o

    fo = json.loads((BACKEND / "config" / "forwarders.default.json").read_text(encoding="utf-8"))
    fr = json.loads(M.source_path("config", "forwarders.default.json", ROOT).read_text(encoding="utf-8"))
    assert strip_source(fr) == strip_source(fo)  # 货代配置只改 source 说明文字


def test_generator_is_reproducible(tmp_path):
    GEN.build_all(ROOT, out=tmp_path / "out")
    for kind in ("config",):
        for name, origin in M.CONFIG_WHITELIST.items():
            if origin == M.RELEASE:
                # 比内容不比字节：git 检出时 JSON 的换行可能被改成 CRLF
                got = json.loads((tmp_path / "out" / kind / name).read_text(encoding="utf-8"))
                assert got == json.loads(M.source_path(kind, name, ROOT).read_text(encoding="utf-8")), name
    for rel, origin in M.TEMPLATE_WHITELIST.items():
        if origin == M.RELEASE:
            a = openpyxl.load_workbook(tmp_path / "out" / "templates" / rel)
            b = openpyxl.load_workbook(M.source_path("templates", rel, ROOT))
            for wa, wb_ in zip(a.worksheets, b.worksheets):
                assert [[c.value for c in r] for r in wa.iter_rows()] == [[c.value for c in r] for r in wb_.iter_rows()], rel


# ---------------------------------------------------------------- 3. release 报关模板可用

@pytest.fixture()
def _declared_price(monkeypatch):
    from fba_tool import customs_package as _cp

    orig = _cp._catalog_hints
    monkeypatch.setattr(_cp, "_catalog_hints", lambda sku: {**orig(sku), "单个货值": 10.0, "申报价来源": "lingxing"})


def test_release_customs_template_validates_and_exports(_declared_price):
    from fba_tool.customs_export import export_customs_package, validate_customs_template
    from fba_tool.customs_package import build_customs_data
    from test_customs_package import _sample_pair

    path = M.source_path("templates", "customs/报关资料模板.xlsx", ROOT)
    assert validate_customs_template(openpyxl.load_workbook(path)) == []
    profile = {"id": "demo", "name": "DEMO 出口商有限公司", "credit_code": "DEMO-CREDIT-0001", "signed_at": "示例城市", "domestic_source": "示例城市"}
    data = build_customs_data([_sample_pair()], [], profile, source="shipment", order_no="DEMO-ORDER-1")
    data.dates["export_date"] = "2026-06-27"
    data.dates["invoice_date"] = "2026-06-27"
    out = openpyxl.load_workbook(io.BytesIO(export_customs_package(data, path.read_bytes())))
    ws = out.worksheets[0]
    assert ws["A10"].value == "DEMO-ORDER-1" and ws["A4"].value == profile["name"] and ws["C3"].value == profile["credit_code"]


def test_factory_sheet_falls_back_without_template(monkeypatch):
    from fba_tool import excel_export

    monkeypatch.setattr(excel_export, "template_path", lambda name: None)
    df = pd.DataFrame({"仓库代码": ["DEMO1", "DEMO1", "DEMO2"], "箱号": [1, 2, 1], "产品名称": ["DEMO-A", "DEMO-B", "DEMO-A"]})
    wb = openpyxl.load_workbook(io.BytesIO(excel_export.export_factory_sheet(df)))
    values = {str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None}
    assert {"DEMO1_纸箱编号", "DEMO1_产品名称", "DEMO2_纸箱编号", "DEMO-A", "DEMO-B"} <= values
    assert excel_export.export_factory_sheet(pd.DataFrame())  # 空表也能导出


# ---------------------------------------------------------------- 4. 老用户升级 / 新用户首次运行

def _package(tmp_path: Path) -> Path:
    """按白名单摆出一个「安装包里的 backend」目录。"""
    pkg = tmp_path / "pkg"
    for src, dest in M.entries(ROOT):
        target = pkg / dest
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, target)
    return pkg


def _use_app_data(monkeypatch, ad: Path, pkg: Path) -> None:
    from fba_tool import master_db, product_catalog, user_config

    ad.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("FBA_APP_DATA_DIR", str(ad))
    monkeypatch.setattr(app_paths, "app_data_dir", lambda: ad)
    monkeypatch.setattr(master_db, "_master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr(app_paths, "bundled_backend_dir", lambda: pkg)
    monkeypatch.setattr(product_catalog, "templates_dir", lambda: ad / "templates")  # 模块导入时算好的目录，这里改到本测试的 app_data
    if hasattr(user_config, "_cache"):
        monkeypatch.setattr(user_config, "_cache", {})


def _snapshot(ad: Path) -> Dict[str, bytes]:
    return {p.relative_to(ad).as_posix(): p.read_bytes() for p in ad.rglob("*")
            if p.is_file() and not p.name.startswith("fba_master.db")}


def _master_rows(ad: Path):
    con = sqlite3.connect(ad / "fba_master.db")
    try:
        return [tuple(r) for r in con.execute("SELECT * FROM product_master ORDER BY sku")]
    finally:
        con.close()


def _example_catalog_for_user() -> bytes:
    """用户自己的产品信息表（示例值，冒充老用户的真实表）：拿示例模板换成 DEMO-USER-* 的一行。"""
    rows = [r for r in _csv_rows(BACKEND / "templates" / "产品信息表_示例.csv") if any(c.strip() for c in r)]
    row = [c.replace("DEMO-SKU-001", "DEMO-USER-1").replace("B0DEMO0001", "B0DEMOUSR1").replace("X0DEMO0001", "X0DEMOUSR1") for c in rows[1]]
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows([rows[0], row])
    return buf.getvalue().encode("utf-8-sig")


def test_upgrade_keeps_user_data_and_ignores_examples(tmp_path, monkeypatch):
    from fba_tool import catalog_seed, exporter_profiles, master_db, product_catalog

    pkg = _package(tmp_path)
    ad = tmp_path / "old_user_app_data"
    _use_app_data(monkeypatch, ad, pkg)
    cfg = ad / "config"
    cfg.mkdir(parents=True)
    user_files = {
        "config/sku_mapping.json": json.dumps({"DEMO-USER-1": "DEMO-USER-1"}),
        "config/fnsku_mapping.json": json.dumps({"X0DEMOUSR1": "DEMO-USER-1"}),
        "config/amazon_sku_metadata.json": json.dumps({"DEMO-USER-1": {"asin": "B0DEMOUSR1", "fnsku": "X0DEMOUSR1"}}),
        "config/exporter_profiles.json": json.dumps({"active_id": "u1", "profiles": [{"id": "u1", "name": "DEMO 老用户出口商", "credit_code": "DEMO-USR"}]}, ensure_ascii=False),
    }
    for rel, text in user_files.items():
        (ad / rel).write_text(text, encoding="utf-8")
    (ad / "templates" / "customs").mkdir(parents=True)
    (ad / "templates" / "产品信息表.csv").write_bytes(_example_catalog_for_user())
    (ad / "templates" / "customs" / "报关资料模板.xlsx").write_bytes(b"user-own-customs-template")
    master_db.init_master_db()
    master_db.upsert_product({"sku": "DEMO-USER-1", "asin": "B0DEMOUSR1", "chinese_name": "老用户品名"})
    before_files = _snapshot(ad)
    before_rows = _master_rows(ad)

    app_paths.ensure_app_data_layout()
    master_db.init_master_db()
    catalog_seed.import_catalog_gaps_into_master()

    after = _snapshot(ad)
    assert {k: after.get(k) for k in before_files} == before_files  # 老用户已有的文件逐字节不变
    added = set(after) - set(before_files)
    allowed = {"config/" + n.replace(".default", "") for n in M.CONFIG_WHITELIST} | {"templates/" + r for r in M.TEMPLATE_WHITELIST}
    assert added <= allowed | {"config/forwarder_mapping.json"}, sorted(added - allowed)
    skus = [r[0] for r in _master_rows(ad)]
    assert skus == ["DEMO-USER-1"]  # 示例产品表 DEMO-SKU-* 没进主库
    assert product_catalog.load_product_catalog().source_path.name == "产品信息表.csv"
    assert exporter_profiles.get_active_profile()["id"] == "u1"
    rows_after = _master_rows(ad)
    assert len(rows_after) == len(before_rows)
    con = sqlite3.connect(ad / "fba_master.db")
    try:
        tables = {row[0] for row in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    finally:
        con.close()
    for name in (
        "carton_spec",
        "carton_spec_history",
        "shipment_box_snapshot",
        "carton_measurement",
        "item_measurement",
        "forwarder_bill",
    ):
        assert name in tables, name


def test_new_user_starts_empty(tmp_path, monkeypatch):
    from fastapi import HTTPException

    from fba_tool import catalog_seed, exporter_profiles, master_db

    pkg = _package(tmp_path)
    ad = tmp_path / "new_user_app_data"
    _use_app_data(monkeypatch, ad, pkg)
    app_paths.ensure_app_data_layout()
    master_db.init_master_db()
    catalog_seed.import_catalog_gaps_into_master()
    assert _master_rows(ad) == []
    for name in ("sku_mapping.json", "fnsku_mapping.json", "amazon_sku_metadata.json"):
        assert json.loads((ad / "config" / name).read_text(encoding="utf-8")) == {}
    assert not (ad / "templates" / "产品信息表.csv").exists()
    assert (ad / "templates" / "产品信息表_示例.csv").is_file()
    for rel in ("亚马逊FBA表.csv", "工厂表.xlsx", "货代表.xlsx", "发货录入样表.xlsx"):
        assert not (ad / "templates" / rel).exists(), rel
    assert exporter_profiles.list_profiles() == []
    with pytest.raises(HTTPException) as ei:
        exporter_profiles.get_active_profile()
    assert ei.value.status_code == 404
