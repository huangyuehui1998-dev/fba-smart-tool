"""S2-4（v8 E3）：箱规批量导入模板 + 导入（预览→写入，二次导入 0 差异；缺列/非数字/负数逐行报错；领星值优先）。"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from fba_tool.master_data_apply import apply_master_to_input_rows
from fba_tool.master_db import get_product_by_sku, init_master_db
from fba_tool.product_fields import set_product_fields
from fba_tool.product_import import (
    CARTON_TEMPLATE_HEADERS,
    apply_carton_import,
    build_carton_template,
    preview_carton_import,
)


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    ad = tmp_path / "app_data"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    init_master_db()


def _xlsx(rows, headers=None):
    wb = openpyxl.load_workbook(io.BytesIO(build_carton_template())) if headers is None else openpyxl.Workbook()
    ws = wb.worksheets[0]
    if headers is not None:
        ws.append(headers)
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_template_headers():
    ws = openpyxl.load_workbook(io.BytesIO(build_carton_template())).worksheets[0]
    heads = [c.value for c in next(ws.iter_rows(min_row=1, max_row=5)) if c.value]
    joined = "".join(str(v) for row in ws.iter_rows(min_row=1, max_row=3, values_only=True) for v in row if v)
    for h in CARTON_TEMPLATE_HEADERS:
        assert h in joined


def test_import_twice_zero_diff():
    data = _xlsx([["A", 60, 40, 35, 12, 18.5], ["B", 50.5, 40, 30, 6, 10.2]])
    pv = preview_carton_import(data, "box.xlsx")
    assert pv["errors"] == [] and pv["summary"]["new"] == 2
    apply_carton_import(data, "box.xlsx")
    a = get_product_by_sku("A")
    assert (a["carton_length"], a["carton_width"], a["carton_height"], a["qty_per_box"], a["carton_gross_weight"]) == (60, 40, 35, 12, 18.5)
    assert a["field_sources"]["carton_length"] == "carton_import"
    pv2 = preview_carton_import(data, "box.xlsx")
    assert pv2["summary"]["changed"] == 0 and pv2["summary"]["new"] == 0


def test_row_errors():
    data = _xlsx([["A", "abc", 40, 35, 12, 18], ["B", 60, -1, 35, 12, 18], ["C", 60, 40, 35, 2.5, 18], ["", 60, 40, 35, 12, 18]])
    pv = preview_carton_import(data, "box.xlsx")
    got = {(e["row"], e["field"]) for e in pv["errors"]}
    start = pv["first_data_row"]
    assert (start, "箱长(cm)") in got
    assert (start + 1, "箱宽(cm)") in got
    assert (start + 2, "单箱数量") in got
    assert (start + 3, "SKU") in got
    with pytest.raises(ValueError):
        apply_carton_import(data, "box.xlsx")


def test_missing_column():
    with pytest.raises(ValueError, match="单箱重"):
        preview_carton_import(_xlsx([["A", 1, 1, 1, 1]], headers=["SKU", "箱长(cm)", "箱宽(cm)", "箱高(cm)", "单箱数量"]), "box.xlsx")


def test_lingxing_values_take_priority():
    set_product_fields("A", {"carton_length": 61}, "lingxing")
    pv = preview_carton_import(_xlsx([["A", 60, 40, 35, 12, 18.5]]), "box.xlsx")
    locked = [i for i in pv["items"] if i["status"] == "locked"]
    assert [i["field"] for i in locked] == ["carton_length"]
    apply_carton_import(_xlsx([["A", 60, 40, 35, 12, 18.5]]), "box.xlsx")
    assert get_product_by_sku("A")["carton_length"] == 61


def test_step1_fills_blank_carton_from_master_including_gross():
    apply_carton_import(_xlsx([["A", 60, 40, 35, 12, 18.5]]), "box.xlsx")
    rows, _src = apply_master_to_input_rows([{"SKU简称": "A", "发货总数": 24, "单箱装量": None, "箱长(cm)": None,
                                               "箱宽(cm)": None, "箱高(cm)": None, "单箱毛重(kg)": None}])
    r = rows[0]
    assert (r["箱长(cm)"], r["单箱装量"], r["单箱毛重(kg)"]) == (60, 12, 18.5)


def test_api_template_and_import():
    from main import app
    c = TestClient(app)
    r = c.get("/api/master-data/import/carton/template")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    data = _xlsx([["A", 60, 40, 35, 12, 18.5]])
    r = c.post("/api/master-data/import/carton/preview", files={"file": ("b.xlsx", data, "application/octet-stream")})
    assert r.status_code == 200 and r.json()["summary"]["new"] == 1
    r = c.post("/api/master-data/import/carton/apply", files={"file": ("b.xlsx", data, "application/octet-stream")})
    assert r.status_code == 200
