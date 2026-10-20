"""S3 第 0 步：产品库本地编辑（D-2026-10-16-03）。网页编辑 + Excel 导出/导回预览确认；手改记 manual 锁定；申报价例外。"""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fba_tool.lingxing_preview import apply_lingxing_field_patch
from fba_tool.master_db import get_product_by_sku, init_master_db, patch_product_field, upsert_product
from fba_tool.product_fields import resolve_declared_price, set_product_fields


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    ad = tmp_path / "ad"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    init_master_db()
    upsert_product({"sku": "S1", "fnsku": "X00OLD0001", "asin": "B000000001", "brand": "OLD", "english_name": "Old Mat"})
    set_product_fields("S1", {"listing_price": 15.99, "listing_price_basis": "listing_price"}, "lingxing")


def _client():
    from main import app

    return TestClient(app)


REMOTE = {"fnsku": "X00LX", "asin": "B0LX", "brand": "LXBRAND", "english_name": "LX Name", "local_name": "领星名",
          "listing_price": 19.99, "listing_price_currency": "USD"}


@pytest.mark.parametrize("field,value", [("fnsku", "X00MAN"), ("brand", "MYBRAND"), ("english_name", "My Name"),
                                         ("local_name", "我的名"), ("asin", "B0MAN")])
def test_web_edit_any_field_marked_manual_and_locked(field, value):
    r = _client().patch("/api/master-data/sku/S1", json={"field": field, "value": value})
    assert r.status_code == 200, r.text
    rec = get_product_by_sku("S1")
    assert rec[field] == value and rec["field_sources"][field] == "manual"
    res = apply_lingxing_field_patch("S1", [field], REMOTE)
    assert field in res["skipped_locked"]
    assert get_product_by_sku("S1")[field] == value


@pytest.mark.parametrize("field,value", [("chinese_name", "宠物垫"), ("hs_code", "6307909000"), ("declare_unit", "个"),
                                         ("category", "fishing"), ("carton_gross_weight", 12.5), ("unit_weight", 0.8)])
def test_web_edit_local_only_fields_saved_manual(field, value):
    patch_product_field("S1", field, value)
    rec = get_product_by_sku("S1")
    assert rec[field] == value and rec["field_sources"][field] == "manual"


def test_web_edit_price_not_locked_listing_price_first():
    patch_product_field("S1", "listing_price", 9.99)
    patch_product_field("S1", "declared_value", 12.0)
    apply_lingxing_field_patch("S1", ["listing_price"], REMOTE)
    rec = get_product_by_sku("S1")
    assert rec["listing_price"] == 19.99
    assert resolve_declared_price(rec) == (19.99, "lingxing")


def _export_bytes():
    r = _client().get("/api/master-data/export.xlsx")
    assert r.status_code == 200
    return r.content


def _edit(content, changes):
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[head.index("sku")].value in changes:
            for f, v in changes[row[head.index("sku")].value].items():
                row[head.index(f)].value = v
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_export_contains_products_and_fields():
    wb = load_workbook(io.BytesIO(_export_bytes()))
    ws = wb.active
    head = [c.value for c in ws[1]]
    for f in ("sku", "fnsku", "chinese_name", "hs_code", "declare_unit", "category", "listing_price", "declared_value"):
        assert f in head
    assert ws.cell(2, head.index("sku") + 1).value == "S1"


def test_import_preview_then_apply_marks_manual():
    data = _edit(_export_bytes(), {"S1": {"fnsku": "X00BULK", "chinese_name": "批量名", "hs_code": "6307909000"}})
    c = _client()
    p = c.post("/api/master-data/import/edits/preview", files={"file": ("p.xlsx", data)})
    assert p.status_code == 200, p.text
    body = p.json()
    changed = {(d["sku"], d["field"]) for d in body["changes"]}
    assert changed == {("S1", "fnsku"), ("S1", "chinese_name"), ("S1", "hs_code")}
    assert get_product_by_sku("S1")["fnsku"] == "X00OLD0001"  # 预览不写库
    a = c.post("/api/master-data/import/edits/apply", files={"file": ("p.xlsx", data)})
    assert a.status_code == 200 and a.json()["written"] == 3
    rec = get_product_by_sku("S1")
    assert rec["fnsku"] == "X00BULK" and rec["field_sources"]["fnsku"] == "manual"
    assert apply_lingxing_field_patch("S1", ["fnsku"], REMOTE)["skipped_locked"] == ["fnsku"]


def test_import_unchanged_file_no_changes_and_unknown_sku_reported():
    data = _edit(_export_bytes(), {})
    p = _client().post("/api/master-data/import/edits/preview", files={"file": ("p.xlsx", data)}).json()
    assert p["changes"] == []
    wb = load_workbook(io.BytesIO(data))
    wb.active.append(["NEWSKU"])
    out = io.BytesIO(); wb.save(out)
    p = _client().post("/api/master-data/import/edits/preview", files={"file": ("p.xlsx", out.getvalue())}).json()
    assert any(e["sku"] == "NEWSKU" for e in p["errors"])


def test_import_bad_number_reported_not_zero():
    data = _edit(_export_bytes(), {"S1": {"carton_gross_weight": "abc"}})
    p = _client().post("/api/master-data/import/edits/preview", files={"file": ("p.xlsx", data)}).json()
    assert any(e["field"] == "carton_gross_weight" for e in p["errors"])
    assert not any(d["field"] == "carton_gross_weight" for d in p["changes"])


def test_import_price_edit_keeps_lingxing_first():
    data = _edit(_export_bytes(), {"S1": {"declared_value": 30}})
    _client().post("/api/master-data/import/edits/apply", files={"file": ("p.xlsx", data)})
    rec = get_product_by_sku("S1")
    assert rec["declared_value"] == 30 and resolve_declared_price(rec) == (15.99, "lingxing")
