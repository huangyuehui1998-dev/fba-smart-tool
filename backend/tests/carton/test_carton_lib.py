"""PR-5 箱规库验收。只用合成编号（TEST_* / FBATEST*），不读真实批次。"""
from __future__ import annotations

import hashlib
import io
import sqlite3
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "scripts"))

import pytest

from fba_tool.carton_math import (
    bill_warnings,
    box_status,
    classify_measurements,
    decompose,
    dim_deltas,
    format_volume_bias_pct,
    flag_suspected_mixups,
    inches_to_cm,
    lb_to_kg,
    light_heavy_check,
    net_from_tares,
    published_batch13_identity,
    roundtrip_cm_delta,
    side_level,
    summary_lines,
    suggest_spec_update,
    volumetric_kg,
)


def _init():
    from fba_tool.master_db import init_master_db

    init_master_db()


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    ad = tmp_path / "app_data"
    ad.mkdir()
    monkeypatch.setenv("FBA_APP_DATA_DIR", str(ad))
    monkeypatch.delenv("FBA_SKIP_MIGRATION", raising=False)
    monkeypatch.delenv("FBA_MIGRATION_FAIL_STEP", raising=False)
    return ad


def _product_digest(db_file: Path) -> str:
    conn = sqlite3.connect(db_file)
    try:
        rows = conn.execute("SELECT * FROM product_master ORDER BY sku").fetchall()
    finally:
        conn.close()
    return hashlib.sha256(repr(rows).encode()).hexdigest()


def _user_version(db_file: Path) -> int:
    conn = sqlite3.connect(db_file)
    try:
        return int(conn.execute("PRAGMA user_version").fetchone()[0])
    finally:
        conn.close()


def _xlsx(headers, rows) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _spec(**overrides):
    payload = {
        "sku": "TEST_SKU",
        "spec_name": "装量20",
        "length_cm": 60,
        "width_cm": 40,
        "height_cm": 40,
        "qty_per_carton": 20,
        "gross_kg": 12.5,
        "source": "manual",
    }
    payload.update(overrides)
    return payload


def test_cd01_inches_and_pounds():
    assert inches_to_cm("17.72") == Decimal("45.01")
    assert inches_to_cm("16.14") == Decimal("41.00")
    assert lb_to_kg("16.53") == Decimal("7.498")


def test_cd02_roundtrip_within_half_millimetre_is_zero():
    assert roundtrip_cm_delta(45) == Decimal("0")


def test_cd03_sorted_dim_deltas():
    declared = [Decimal("45.01"), Decimal("45.01"), Decimal("41.00")]
    measured = [Decimal("45.0"), Decimal("47.5"), Decimal("46.5")]
    assert dim_deltas(declared, measured) == (Decimal("2.49"), Decimal("1.49"), Decimal("4.00"))


def test_cd04_volumetric_weight():
    assert volumetric_kg(47, 46.5, 44) == Decimal("16.03")


def test_cd05_decomposition_identity_and_published_batch():
    rows = [{
        "declared_dims": [40, 40, 40],
        "measured_dims": [50, 40, 40],
        "declared_kg": 10,
        "measured_kg": 10,
    }]
    tickets = [{"bill_kg": 14, "rows": [{"weight_kg": 10, "dims": [50, 40, 40]}]}]
    parts = decompose(rows, tickets)
    assert parts["weight_contrib"] + parts["size_contrib"] == parts["box_total"]
    assert parts["weight_contrib"] + parts["size_contrib"] + parts["carry"] == parts["total"]
    assert "actual_delta" not in (
        parts["weight_contrib"],
        parts["size_contrib"],
        parts["carry"],
    )
    published = published_batch13_identity()
    assert published["weight_contrib"] + published["size_contrib"] + published["carry"] == Decimal("320.01")
    assert published["total"] == Decimal("320.01")
    assert published["reconcile_total"] == Decimal("323.85")
    lines = summary_lines(published, 154, ["FBATESTU000017"])
    assert lines["headline"] == "计费重总差 +320.01kg = 实重贡献 0.00 + 尺寸贡献 +318.71 + 进位 +1.30（全票口径）"
    assert lines["reference"] == "实重差 −28.37kg（仅参考，去重 154 箱）"
    assert "贴错" not in lines["headline"]


def test_cd06_side_levels():
    assert side_level("0.8", 45) == "正常"
    assert side_level("2.0", 45) == "提示"
    assert side_level("3.0", 45) == "异常"


def test_cd07_suspected_mixup_is_inference_only():
    rows = [{"id": "A", "sku": "TEST_A", "weight_kg": Decimal("7.06")} for _ in range(3)]
    rows += [{"id": "B", "sku": "TEST_B", "weight_kg": Decimal("9.73")} for _ in range(3)]
    rows.append({"id": "odd", "sku": "TEST_B", "weight_kg": Decimal("7.19")})
    flagged = flag_suspected_mixups(rows)
    assert [item["id"] for item in flagged] == ["odd"]
    assert flagged[0]["label"] == "疑似装错 SKU（推断，待货代/工厂核实）"
    assert "贴错" not in flagged[0]["label"]


def test_cd08_light_heavy_counts_every_scan_row():
    declared = [Decimal("7.5")] * 100 + [Decimal("9.5")] * 60
    measured = [{"weight_kg": Decimal("7.2"), "match_status": "matched"} for _ in range(99)]
    measured.append({"weight_kg": Decimal("7.2"), "match_status": "duplicate"})
    measured.append({"weight_kg": Decimal("7.1"), "match_status": "invalid_id"})
    measured += [{"weight_kg": Decimal("9.6"), "match_status": "matched"} for _ in range(59)]
    result = light_heavy_check(declared, measured)
    assert result["threshold_kg"] == Decimal("8.500")
    assert result["message"] == "轻箱多 1、重箱少 1"
    assert result["scope"] == "全部过机行"


def test_cd09_batch15_style_carry():
    scans = [{"weight_kg": 20, "dims": [61, 46.5, 37.5]} for _ in range(5)]
    parts = decompose([], [{"bill_kg": 101, "rows": scans}])
    line = sum(max(Decimal("20.000"), volumetric_kg(61, 46.5, 37.5)) for _ in range(5))
    assert line == Decimal("100.00")
    assert parts["carry"] == Decimal("1.00")


def test_cd10_volume_bias_prompt():
    big = [{"declared_dims": [45, 45, 41], "measured_dims": [47.2, 46.6, 43.9]} for _ in range(4)]
    assert suggest_spec_update(big) is True
    assert format_volume_bias_pct(Decimal("0.1495")) == "+14.9%"
    small = [{"declared_dims": [100, 100, 100], "measured_dims": [101, 101, 101]} for _ in range(4)]
    assert suggest_spec_update(small) is False


def test_cl08_net_weight_anchors_stay():
    assert net_from_tares(264, [Decimal("0.3")] * 32) == Decimal("254.4")
    assert net_from_tares(Decimal("115.2"), [Decimal("0.2")] * 12) == Decimal("112.8")
    assert net_from_tares(Decimal("122.4"), [Decimal("0.2")] * 12) == Decimal("120.0")


def test_mt_classify_invalid_unmatched_and_duplicate():
    declared = ["FBATESTU000001", "FBATESTU000017"]
    measured = ["1", "2", "3", "4", "5", "FBATESTU003937", "FBATESTU000017", "FBATESTU000017"]
    statuses = classify_measurements(declared, measured)
    assert statuses[:5] == ["invalid_id"] * 5
    assert statuses[5] == "unmatched"
    assert statuses[6:] == ["duplicate", "duplicate"]


def test_cl01_create_manual_version(isolated):
    _init()
    from fba_tool.carton_store import create_spec

    row = create_spec(_spec())
    assert row["version"] == 1
    assert row["source"] == "manual"
    assert row["length_cm"] == 60
    assert row["gross_kg"] == 12.5
    assert row["is_default"] is True


def test_cl02_edit_creates_history(isolated):
    _init()
    from fba_tool.carton_store import active_specs, create_spec, history_for, replace_spec

    created = create_spec(_spec())
    replace_spec(created["id"], {"length_cm": 61})
    current = [row for row in active_specs("TEST_SKU") if row["status"] == "active"]
    assert len(current) == 1
    assert current[0]["length_cm"] == 61
    assert current[0]["version"] == 2
    assert len(history_for("TEST_SKU")) == 2


def test_cl03_empty_dims_fill_and_filled_dims_stay(isolated):
    _init()
    from fba_tool.carton_store import create_spec, replace_spec
    from fba_tool.master_data_apply import apply_master_to_input_rows

    created = create_spec(_spec())
    replace_spec(created["id"], {"length_cm": 61, "width_cm": 40, "height_cm": 40})
    filled, _sources = apply_master_to_input_rows([{"SKU简称": "TEST_SKU", "单箱装量": 20}])
    assert filled[0]["箱长(cm)"] == 61
    assert filled[0]["箱宽(cm)"] == 40
    assert filled[0]["箱高(cm)"] == 40
    kept, _sources = apply_master_to_input_rows([{
        "SKU简称": "TEST_SKU",
        "单箱装量": 20,
        "箱长(cm)": 50,
        "箱宽(cm)": 30,
        "箱高(cm)": 20,
    }])
    assert kept[0]["箱长(cm)"] == 50
    assert kept[0]["箱宽(cm)"] == 30
    assert kept[0]["箱高(cm)"] == 20


def test_cl04_qty_selects_spec(isolated):
    _init()
    from fba_tool.carton_store import create_spec
    from fba_tool.master_data_apply import apply_master_to_input_rows

    create_spec(_spec(length_cm=60, qty_per_carton=20, spec_name="装量20"))
    create_spec(_spec(length_cm=50, width_cm=30, height_cm=20, qty_per_carton=10, spec_name="装量10", gross_kg=8))
    rows, _sources = apply_master_to_input_rows([
        {"SKU简称": "TEST_SKU", "单箱装量": 10},
        {"SKU简称": "TEST_SKU", "单箱装量": 20},
    ])
    assert rows[0]["箱长(cm)"] == 50
    assert rows[1]["箱长(cm)"] == 60
    from fba_tool.carton_store import mixed_spec_notices

    notices = mixed_spec_notices([
        {"SKU简称": "TEST_SKU", "单箱装量": 10},
        {"SKU简称": "TEST_SKU", "单箱装量": 20},
    ])
    assert notices == ["本批 SKU TEST_SKU 使用 2 套箱规"]


def test_cl05_invalid_spec_is_422(isolated):
    _init()
    from fastapi.testclient import TestClient

    from main import app

    client = TestClient(app)
    cases = [
        {"length_cm": 0},
        {"length_cm": -1},
        {"length_cm": "abc"},
        {"length_cm": 201},
        {"gross_kg": 71},
    ]
    for extra in cases:
        response = client.post("/api/carton-specs", json=_spec(**extra))
        assert response.status_code == 422, response.text
        body = response.json()["detail"]
        assert body["fields"]


def test_cl06_archive_keeps_history(isolated):
    _init()
    from fba_tool.carton_store import archive_spec, create_spec, history_for, list_specs

    created = create_spec(_spec())
    archived = archive_spec(created["id"])
    assert archived["status"] == "archived"
    assert list_specs("TEST_SKU") == []
    assert history_for("TEST_SKU")


def test_cl07_migration_copies_only_skus_with_carton(isolated, monkeypatch):
    from fba_tool.master_db import init_master_db, upsert_product

    isolated_env = isolated
    monkeypatch.setenv("FBA_SKIP_MIGRATION", "1")
    init_master_db()
    upsert_product({"sku": "TEST_WITH", "carton_length": 60, "carton_width": 40, "carton_height": 40, "qty_per_box": 20})
    upsert_product({"sku": "TEST_EMPTY", "chinese_name": "无箱规"})
    monkeypatch.delenv("FBA_SKIP_MIGRATION")
    init_master_db()
    conn = sqlite3.connect(isolated_env / "fba_master.db")
    try:
        count = conn.execute("SELECT COUNT(*) FROM carton_spec WHERE source='migrated_v3' AND version=1").fetchone()[0]
    finally:
        conn.close()
    assert count == 1


def test_m1_product_rows_unchanged(isolated, monkeypatch):
    from fba_tool.master_db import init_master_db, upsert_product

    monkeypatch.setenv("FBA_SKIP_MIGRATION", "1")
    init_master_db()
    upsert_product({"sku": "TEST_WITH", "carton_length": 60, "carton_width": 40, "carton_height": 40, "qty_per_box": 12, "chinese_name": "合成品"})
    before = _product_digest(isolated / "fba_master.db")
    assert _user_version(isolated / "fba_master.db") == 3
    monkeypatch.delenv("FBA_SKIP_MIGRATION")
    init_master_db()
    assert _user_version(isolated / "fba_master.db") == 4
    assert _product_digest(isolated / "fba_master.db") == before
    conn = sqlite3.connect(isolated / "fba_master.db")
    try:
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    finally:
        conn.close()
    assert {"carton_spec", "carton_measurement", "item_measurement", "forwarder_bill"} <= tables


def test_m2_migration_backup_integrity(isolated):
    from fba_tool.db_backup import integrity_ok
    from fba_tool.master_db import init_master_db

    init_master_db()
    backups = list((isolated / "backups").glob("fba_master.v3_to_v4.*.db"))
    assert len(backups) == 1
    assert integrity_ok(backups[0])


def test_m3_failed_step_rolls_back(isolated, monkeypatch):
    from fba_tool.db_migrations import migration_error
    from fba_tool.master_db import init_master_db, upsert_product

    monkeypatch.setenv("FBA_SKIP_MIGRATION", "1")
    init_master_db()
    upsert_product({"sku": "TEST_KEEP", "chinese_name": "保持"})
    before = _product_digest(isolated / "fba_master.db")
    monkeypatch.delenv("FBA_SKIP_MIGRATION")
    monkeypatch.setenv("FBA_MIGRATION_FAIL_STEP", "2")
    init_master_db()
    assert _user_version(isolated / "fba_master.db") == 3
    assert _product_digest(isolated / "fba_master.db") == before
    assert migration_error(str(isolated / "fba_master.db")) == "升级失败，数据未改动"
    conn = sqlite3.connect(isolated / "fba_master.db")
    try:
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    finally:
        conn.close()
    assert "carton_spec" not in tables


def test_m4_second_start_does_not_migrate_again(isolated):
    from fba_tool.master_db import init_master_db

    init_master_db()
    init_master_db()
    init_master_db()
    backups = list((isolated / "backups").glob("fba_master.v*_to_v*.db"))
    assert len(backups) == 1
    dailies = list((isolated / "backups").glob("fba_master.daily.*.db"))
    assert len(dailies) == 1


def test_a1_legacy_import_copies(tmp_path, monkeypatch):
    from fba_tool.app_paths import import_legacy_data_dir, legacy_import_status

    old = tmp_path / "old"
    new = tmp_path / "new"
    local = tmp_path / "local"
    old.mkdir()
    new.mkdir()
    (old / "marker.txt").write_text("synthetic", encoding="utf-8")
    monkeypatch.setenv("FBA_APP_DATA_DIR", str(new))
    monkeypatch.setenv("LOCALAPPDATA", str(local))
    pointer = local / "FBA_Smart_Tool" / "last_data_dir.txt"
    pointer.parent.mkdir(parents=True)
    pointer.write_text(str(old), encoding="utf-8")
    status = legacy_import_status()
    assert status["needs_import"] is True
    assert status["message"] == f"检测到旧数据目录 {old}，是否导入"
    copied = import_legacy_data_dir()
    assert copied["copied"] == 1
    assert (new / "marker.txt").read_text(encoding="utf-8") == "synthetic"
    assert (old / "marker.txt").is_file()


def test_a2_second_start_keeps_user_file(isolated):
    from fba_tool.master_db import init_master_db

    note = isolated / "config" / "user_settings.json"
    note.parent.mkdir(parents=True)
    note.write_text('{"theme":"test"}', encoding="utf-8")
    digest = hashlib.sha256(note.read_bytes()).hexdigest()
    init_master_db()
    init_master_db()
    assert hashlib.sha256(note.read_bytes()).hexdigest() == digest


def test_a3_health_fields(isolated):
    _init()
    from fastapi.testclient import TestClient

    from main import app

    response = TestClient(app).get("/api/health")
    assert response.status_code == 200
    body = response.json()
    assert body["data_dir"] == str(isolated)
    assert body["schema_version"] == 4
    assert body["carton_spec_count"] == 0


def test_a4_portable_flag_skips_localappdata(tmp_path, monkeypatch):
    from fba_tool import app_paths

    exe = tmp_path / "exe"
    local = tmp_path / "local"
    exe.mkdir()
    monkeypatch.delenv("FBA_APP_DATA_DIR", raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(local))
    monkeypatch.setattr(app_paths, "project_root", lambda: exe)
    (exe / "portable.flag").write_text("1", encoding="utf-8")
    assert app_paths.app_data_dir() == exe / "app_data"
    assert not (local / "FBA_Smart_Tool").exists()
    (exe / "portable.flag").unlink()
    assert app_paths.app_data_dir() == local / "FBA_Smart_Tool" / "app_data"


def test_b1_export_import_roundtrip(isolated):
    _init()
    from fba_tool.carton_import import parse_carton_spec_rows
    from fba_tool.carton_store import apply_import, create_spec, export_specs_xlsx, list_specs

    create_spec(_spec(tare_kg=0.3, note="合成"))
    blob = export_specs_xlsx()
    conn = sqlite3.connect(isolated / "fba_master.db")
    conn.execute("DELETE FROM carton_spec")
    conn.execute("DELETE FROM carton_spec_history")
    conn.commit()
    conn.close()
    records, errors = parse_carton_spec_rows(blob, "箱规.xlsx")
    assert errors == []
    apply_import(records, "箱规.xlsx")
    current = list_specs("TEST_SKU")
    assert len(current) == 1
    assert current[0]["length_cm"] == 60
    assert current[0]["gross_kg"] == 12.5
    assert current[0]["tare_kg"] == 0.3
    assert current[0]["note"] == "合成"


def test_b2_second_import_is_empty(isolated):
    _init()
    from fba_tool.carton_import import parse_carton_spec_rows
    from fba_tool.carton_store import apply_import, preview_import

    blob = _xlsx(
        ["SKU", "箱规名称", "箱长(cm)", "箱宽(cm)", "箱高(cm)", "单箱数量", "单箱重(kg)", "空箱重", "备注"],
        [["TEST_SKU", "装量20", 60, 40, 40, 20, 12.5, 0.3, "合成"]],
    )
    records, errors = parse_carton_spec_rows(blob, "箱规.xlsx")
    assert errors == []
    apply_import(records, "箱规.xlsx")
    again = preview_import(records)
    assert again["added_count"] == 0
    assert again["changed_count"] == 0
    assert again["unchanged"] == 1


def test_b3_different_spec_previews_old_and_new(isolated):
    _init()
    from fba_tool.carton_import import parse_carton_spec_rows
    from fba_tool.carton_store import apply_import, create_spec, history_for, preview_import

    create_spec(_spec(length_cm=60))
    blob = _xlsx(
        ["SKU", "箱规名称", "箱长(cm)", "箱宽(cm)", "箱高(cm)", "单箱数量", "单箱重(kg)"],
        [["TEST_SKU", "装量10", 50, 30, 20, 10, 8]],
    )
    records, errors = parse_carton_spec_rows(blob, "箱规.xlsx")
    assert errors == []
    preview = preview_import(records)
    assert preview["added_count"] == 1
    assert preview["added"][0]["old"]["length_cm"] == 60
    assert preview["added"][0]["new"]["length_cm"] == 50
    apply_import(records, "箱规.xlsx")
    blob_history = " ".join(str(row.get("before_json") or "") for row in history_for("TEST_SKU"))
    assert "60" in blob_history


def test_b4_zip_restore_counts(isolated, tmp_path):
    _init()
    from fba_tool.carton_store import count_products_and_specs, create_spec, restore_app_data, zip_app_data
    from fba_tool.master_db import upsert_product

    upsert_product({"sku": "TEST_SKU", "chinese_name": "合成"})
    create_spec(_spec())
    before = count_products_and_specs(isolated / "fba_master.db")
    archive = tmp_path / "backup.zip"
    zip_app_data(archive)
    dest = tmp_path / "restored"
    restore_app_data(archive, dest)
    assert count_products_and_specs(dest / "fba_master.db") == before


def test_p1_p2_release_guard(tmp_path):
    from release_guard import find_carton_data, find_user_data, refuse_user_data

    dist = tmp_path / "dist"
    (dist / "app_data").mkdir(parents=True)
    (dist / "app_data" / "fba_master.db").write_bytes(b"not-a-real-db")
    hits = find_user_data(dist)
    assert any(hit.endswith("app_data") or hit.endswith(".db") for hit in hits)
    with pytest.raises(SystemExit) as exc:
        refuse_user_data(dist)
    assert "发现发布目录里已有用户数据" in str(exc.value)
    config = tmp_path / "config"
    config.mkdir()
    (config / "carton_spec.json").write_text('{"qty_per_carton": 1, "length_cm": 1}', encoding="utf-8")
    assert find_carton_data(config)
    assert find_user_data(tmp_path / "missing") == []


def test_p3_dev_db_is_app_data(isolated):
    from fba_tool.app_paths import app_data_dir, master_db_path

    assert master_db_path() == app_data_dir() / "fba_master.db"
    assert master_db_path().parent == isolated


def test_r1_r2_backup_still_has_products(isolated, monkeypatch):
    from fba_tool.master_db import init_master_db, upsert_product

    monkeypatch.setenv("FBA_SKIP_MIGRATION", "1")
    init_master_db()
    upsert_product({"sku": "TEST_KEEP", "chinese_name": "合成"})
    before = _product_digest(isolated / "fba_master.db")
    monkeypatch.delenv("FBA_SKIP_MIGRATION")
    init_master_db()
    backup = next((isolated / "backups").glob("fba_master.v3_to_v4.*.db"))
    assert _user_version(backup) == 3
    assert _product_digest(backup) == before
    conn = sqlite3.connect(isolated / "fba_master.db")
    try:
        names = [row[0] for row in conn.execute("SELECT sku FROM product_master ORDER BY sku")]
    finally:
        conn.close()
    assert names == ["TEST_KEEP"]


def test_r3_newer_db_is_readonly(isolated):
    _init()
    conn = sqlite3.connect(isolated / "fba_master.db")
    conn.execute("PRAGMA user_version = 5")
    conn.commit()
    conn.close()
    from fastapi.testclient import TestClient

    from fba_tool.db_migrations import is_readonly
    from main import app

    assert is_readonly(isolated / "fba_master.db") is True
    client = TestClient(app)
    response = client.post("/api/carton-specs", json=_spec())
    assert response.status_code == 409
    assert response.json()["code"] == "DB_NEWER_THAN_APP"
    health = client.get("/api/health")
    assert health.status_code == 200
    assert health.json()["readonly"] is True


def test_mt02_manual_assign(isolated):
    _init()
    from fba_tool.carton_store import assign_measurement, compare_carton, import_measurements, record_box_snapshot

    record_box_snapshot("FBATESTSHIP", "sta_export", [{
        "fba_box_id": "FBATESTU000001",
        "length_cm": 61,
        "width_cm": 46.5,
        "height_cm": 37.5,
        "gross_kg": 20,
        "sku_items": [{"sku": "TEST_SKU"}],
    }])
    imported = import_measurements(
        [{"kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": str(n), "weight_kg": 20,
          "length_cm": 61, "width_cm": 46.5, "height_cm": 37.5, "sku": "TEST_SKU", "source_row": n}
         for n in range(1, 6)],
        measured_by="forwarder_jiufang",
        source_file="batch15.xlsx",
    )
    report = compare_carton("FBATESTSHIP")
    assert report["rows"] == []
    assert sum(1 for row in report["quality"] if row["match_status"] == "invalid_id") == 5
    conn = sqlite3.connect(isolated / "fba_master.db")
    measure_id = conn.execute("SELECT id FROM carton_measurement WHERE source_row=1").fetchone()[0]
    conn.close()
    assign_measurement(measure_id, "FBATESTU000001")
    after = compare_carton("FBATESTSHIP")
    assert len(after["rows"]) == 1
    assert after["rows"][0]["箱号"] == "FBATESTU000001"
    assert imported["inserted"] == 5


def test_mt03_unmatched_does_not_raise(isolated):
    _init()
    from fba_tool.carton_store import compare_carton, import_measurements, record_box_snapshot

    record_box_snapshot("FBATESTSHIP", "sta_export", [{
        "fba_box_id": "FBATESTU000001", "length_cm": 40, "width_cm": 40, "height_cm": 40, "gross_kg": 10,
    }])
    import_measurements(
        [{"kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "FBATESTU003937", "weight_kg": 10,
          "length_cm": 40, "width_cm": 40, "height_cm": 40, "source_row": 1}],
        measured_by="forwarder_jiufang",
        source_file="scan.xlsx",
    )
    report = compare_carton("FBATESTSHIP")
    assert report["rows"] == []
    assert report["quality"][0]["match_status"] == "unmatched"


def test_mt04_second_measurement_import_inserts_nothing(isolated):
    _init()
    from fba_tool.measurement_adapters import detect_adapter
    from fba_tool.carton_store import import_measurements

    blob = _xlsx(["外箱号", "实重", "长", "宽", "高"], [["FBATESTU000001", 10, 40, 40, 40]])
    adapter = detect_adapter("jiufang.xlsx", blob)
    rows = adapter.parse(blob)
    first = import_measurements(rows, measured_by="forwarder_jiufang", source_file="jiufang.xlsx")
    second = import_measurements(rows, measured_by="forwarder_jiufang", source_file="jiufang.xlsx")
    assert first["inserted"] == 1
    assert second["inserted"] == 0
    assert second["skipped"] == 1


def test_mt05_sources_stay_separate(isolated):
    _init()
    from fba_tool.carton_store import compare_carton, import_measurements, record_box_snapshot

    record_box_snapshot("FBATESTSHIP", "sta_export", [{
        "fba_box_id": "FBATESTU000001", "length_cm": 40, "width_cm": 40, "height_cm": 40, "gross_kg": 10,
        "sku_items": [{"sku": "TEST_SKU"}],
    }])
    row = {"kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "FBATESTU000001", "weight_kg": 10,
           "length_cm": 40, "width_cm": 40, "height_cm": 40, "sku": "TEST_SKU", "source_row": 1}
    import_measurements([dict(row, weight_kg=10)], measured_by="forwarder_jiufang", source_file="jiufang.xlsx")
    import_measurements([dict(row, weight_kg=11)], measured_by="forwarder_send", source_file="send.xlsx")
    jiufang = compare_carton("FBATESTSHIP")
    send = compare_carton("FBATESTSHIP", source="forwarder_send")
    assert jiufang["rows"][0]["实测 L/W/H/kg"].endswith("10.000")
    assert send["rows"][0]["实测 L/W/H/kg"].endswith("11.000")
    assert jiufang["dim_source_label"] == "forwarder"


def test_mt06_ticket_only(isolated):
    _init()
    from fba_tool.carton_store import compare_carton, import_measurements, record_box_snapshot

    record_box_snapshot("FBATESTSHIP", "sta_export", [{
        "fba_box_id": "FBATESTU000001", "length_cm": 40, "width_cm": 40, "height_cm": 40, "gross_kg": 10,
    }])
    import_measurements(
        [{"kind": "forwarder_bill", "fba_id": "FBATESTSHIP", "bill_chargeable_kg": 12.5, "forwarder": "jiufang"}],
        measured_by="forwarder_jiufang",
        source_file="bill.xlsx",
    )
    report = compare_carton("FBATESTSHIP")
    assert report["mode"] == "ticket_only"
    assert report["message"] == "只有整票数据"
    assert report["rows"] == []
    assert report["summary"]["ticket_delta"] == 2.5


def test_ui_c1_deduped_main_table(isolated):
    _init()
    from fba_tool.carton_math import summary_lines as lines_fn
    from fba_tool.carton_store import COMPARE_COLUMNS, compare_carton, import_measurements, record_box_snapshot

    declared = []
    measured = []
    for index in range(1, 155):
        box_id = f"FBATESTU1{index:05d}"
        declared.append({
            "fba_box_id": box_id,
            "length_cm": 45.01,
            "width_cm": 45.01,
            "height_cm": 41,
            "gross_kg": 7.5,
            "sku_items": [{"sku": "TEST_SKU"}],
        })
        dims = (57.5, 54.5, 44) if index == 31 else (45.01, 45.01, 41)
        measured.append({
            "kind": "box",
            "fba_id": "FBATESTSHIP",
            "fba_box_id": box_id,
            "sku": "TEST_SKU",
            "length_cm": dims[0],
            "width_cm": dims[1],
            "height_cm": dims[2],
            "weight_kg": 7.5,
            "source_row": index,
        })
    declared.append({
        "fba_box_id": "FBATESTU000017",
        "length_cm": 45.01, "width_cm": 45.01, "height_cm": 41, "gross_kg": 9.5,
        "sku_items": [{"sku": "TEST_SKU"}],
    })
    for extra in range(1, 6):
        declared.append({
            "fba_box_id": f"FBATESTU2{extra:05d}",
            "length_cm": 45.01, "width_cm": 45.01, "height_cm": 41, "gross_kg": 7.5,
        })
    measured.append({
        "kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "FBATESTU000017", "sku": "TEST_B",
        "length_cm": 45, "width_cm": 45, "height_cm": 41, "weight_kg": 9.9, "source_row": 201,
    })
    measured.append({
        "kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "FBATESTU000017", "sku": "TEST_B",
        "length_cm": 45, "width_cm": 45, "height_cm": 41, "weight_kg": 7.19, "source_row": 202,
    })
    for n in range(3):
        measured.append({
            "kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "1", "weight_kg": 1, "source_row": 300 + n,
            "length_cm": 1, "width_cm": 1, "height_cm": 1,
        })
    measured.append({
        "kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "FBATESTU003937", "weight_kg": 8,
        "length_cm": 40, "width_cm": 40, "height_cm": 40, "source_row": 400,
    })
    record_box_snapshot("FBATESTSHIP", "sta_export", declared)
    import_measurements(measured, measured_by="forwarder_jiufang", source_file="scan.xlsx")
    report = compare_carton("FBATESTSHIP")
    assert report["columns"] == COMPARE_COLUMNS
    assert len(report["rows"]) == 154
    assert 156 - 2 == 154
    assert all(row["箱号"] != "FBATESTU000017" for row in report["rows"])
    quality = report["quality"]
    assert sum(1 for row in quality if row["match_status"] == "invalid_id") == 3
    assert sum(1 for row in quality if row["match_status"] == "unmatched") == 1
    assert sum(1 for row in quality if row["match_status"] == "duplicate") == 2
    assert sum(1 for row in quality if row["match_status"] == "missing") == 5
    assert report["summary"]["dup_note"] == "1 个重复箱号未计入汇总"
    abnormal = next(row for row in report["rows"] if row["箱号"] == "FBATESTU100031")
    assert abnormal["状态"] == "异常"
    assert abnormal["status_class"] == "text-red-600"
    assert box_status([45.01, 45.01, 41], [57.5, 54.5, 44], 7.5, 7.5) == "异常"
    published = published_batch13_identity()
    text = lines_fn(published, 154, ["FBATESTU000017"])
    assert "去重 154 箱" in text["reference"]


def test_ui_c4_measured_average_is_not_default(isolated):
    _init()
    from fba_tool.carton_store import create_measured_average, create_spec

    create_spec(_spec())
    created = create_measured_average("TEST_SKU", [47, 46, 44], 13, 20, "实测均值")
    assert created["source"] == "from_measure_avg"
    assert created["is_default"] is False


def test_ui_c5_export_columns_match_api(isolated):
    _init()
    from openpyxl import load_workbook

    from fba_tool.carton_store import compare_carton, compare_export_xlsx, import_measurements, record_box_snapshot

    record_box_snapshot("FBATESTSHIP", "sta_export", [{
        "fba_box_id": "FBATESTU000001", "length_cm": 40, "width_cm": 40, "height_cm": 40, "gross_kg": 10,
        "sku_items": [{"sku": "TEST_SKU"}],
    }])
    import_measurements([{
        "kind": "box", "fba_id": "FBATESTSHIP", "fba_box_id": "FBATESTU000001", "sku": "TEST_SKU",
        "length_cm": 41, "width_cm": 40, "height_cm": 40, "weight_kg": 10.2, "source_row": 1,
    }], measured_by="forwarder_jiufang", source_file="scan.xlsx")
    report = compare_carton("FBATESTSHIP")
    wb = load_workbook(io.BytesIO(compare_export_xlsx("FBATESTSHIP")))
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    assert headers == report["columns"]
    values = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
    assert values == [report["rows"][0][column] for column in report["columns"]]


def test_it_fee_preview_and_policy(isolated):
    _init()
    from fastapi.testclient import TestClient

    from fba_tool.carton_policy import item_policy
    from fba_tool.carton_store import compare_items, import_measurements
    from fba_tool.master_db import upsert_product
    from fba_tool.measurement_adapters import ParseError, detect_adapter
    from main import app

    policy = item_policy()
    assert policy["source_type"] == "amazon_forum"
    assert policy["effective_date"] is None
    assert policy["verified_at"] == "2026-10-20"
    assert policy["reimbursement_window_days"] == 90
    upsert_product({
        "sku": "TEST_ITEM",
        "unit_length": 20,
        "unit_width": 15,
        "unit_height": 5,
        "unit_weight": 0.4,
        "extra_data": '{"size_tier": "Large standard"}',
    })
    header = "sku\tasin\tlongest-side\tmedian-side\tshortest-side\titem-package-weight\tunit-of-dimension\tunit-of-weight\tproduct-size-tier\tamazon-store"
    same = header + "\nTEST_ITEM\tB0TESTITEM\t20\t15\t5\t0.40\tcentimeters\tkilograms\tLarge standard\tUS"
    higher = header + "\nTEST_ITEM\tB0TESTITEM\t20\t15\t5\t0.40\tcentimeters\tkilograms\tSmall oversize\tUS"
    critical = header + "\nTEST_EDGE\tB0TESTEDGE\t20\t15\t5\t0.445\tcentimeters\tkilograms\tLarge standard\tUS"
    both = header + "\nTEST_ITEM\tB0TESTITEM\t20\t15\t5\t0.40\tcentimeters\tkilograms\tLarge standard\tUS\nTEST_ITEM\tB0TESTITEM\t20\t15\t5\t0.42\tcentimeters\tkilograms\tLarge standard\tCA"
    missing = "sku\tlongest-side\nTEST_ITEM\t20"
    with pytest.raises(ParseError) as exc:
        detect_adapter("fees.txt", missing.encode()).parse(missing.encode())
    assert "缺少列：item-package-weight" in str(exc.value)
    client = TestClient(app)
    rejected = client.post(
        "/api/measurements/import/apply",
        files={"file": ("fees.txt", missing.encode(), "text/plain")},
    )
    assert rejected.status_code == 422
    assert "item-package-weight" in rejected.text
    same_rows = detect_adapter("fees.txt", same.encode()).parse(same.encode())
    import_measurements(same_rows, measured_by="fee_preview", source_file="same.txt")
    matched = compare_items(marketplace="US")
    assert matched["rows"][0]["status"] == "一致"
    conn = sqlite3.connect(isolated / "fba_master.db")
    conn.execute("DELETE FROM item_measurement")
    conn.commit()
    conn.close()
    higher_rows = detect_adapter("fees.txt", higher.encode()).parse(higher.encode())
    import_measurements(higher_rows, measured_by="fee_preview", source_file="higher.txt")
    mismatch = compare_items()
    assert mismatch["rows"][0]["status"] == "分段不一致"
    assert "申请重新测量" in mismatch["rows"][0]["guide"]
    assert "90 天" in mismatch["rows"][0]["reimbursement"]
    conn = sqlite3.connect(isolated / "fba_master.db")
    conn.execute("DELETE FROM item_measurement")
    conn.commit()
    conn.close()
    edge_rows = detect_adapter("fees.txt", critical.encode()).parse(critical.encode())
    import_measurements(edge_rows, measured_by="fee_preview", source_file="edge.txt")
    edge = compare_items()
    assert edge["rows"][0]["status"] == "临界"
    assert edge["rows"][0]["margin_text"].startswith("临界，距下一档 ")
    conn = sqlite3.connect(isolated / "fba_master.db")
    conn.execute("DELETE FROM item_measurement")
    conn.commit()
    conn.close()
    site_rows = detect_adapter("fees.txt", both.encode()).parse(both.encode())
    import_measurements(site_rows, measured_by="fee_preview", source_file="sites.txt")
    us = compare_items(marketplace="US")
    ca = compare_items(marketplace="CA")
    assert len(us["rows"]) == 1 and us["rows"][0]["marketplace"] == "US"
    assert len(ca["rows"]) == 1 and ca["rows"][0]["marketplace"] == "CA"


def test_fb01_fb02_bill_carry(isolated):
    _init()
    from fba_tool.carton_store import import_measurements
    from fba_tool.measurement_adapters import detect_adapter

    tails = [509, 332.5, 604, 584, 517]
    sums = [Decimal("508.70"), Decimal("332.48"), Decimal("603.62"), Decimal("583.76"), Decimal("516.64")]
    rows = []
    for index, (tail, line) in enumerate(zip(tails, sums), start=1):
        fba_id = f"FBATESTSHIP{index}"
        rows.append(["FBATESTU000001", fba_id, line, tail, "jiufang"])
    blob = _xlsx(["外箱号", "货件编号", "单箱计费重", "入库计费重", "货代"], rows)
    parsed = detect_adapter("货代表.xlsx", blob).parse(blob)
    bills = [row for row in parsed if row["kind"] == "forwarder_bill"]
    assert [row["bill_chargeable_kg"] for row in bills] == tails
    imported = import_measurements(parsed, measured_by="forwarder_jiufang", source_file="货代表.xlsx")
    assert imported["inserted"] >= 5
    conn = sqlite3.connect(isolated / "fba_master.db")
    try:
        stored = [row[0] for row in conn.execute("SELECT bill_chargeable_kg FROM forwarder_bill ORDER BY fba_id")]
    finally:
        conn.close()
    assert stored == tails
    gaps = []
    for tail, line in zip(tails, sums):
        parts = decompose([], [{"bill_kg": tail, "rows": [{"chargeable_kg": line, "weight_kg": line}]}])
        gaps.append(parts["carry"])
    assert gaps == [Decimal("0.30"), Decimal("0.02"), Decimal("0.38"), Decimal("0.24"), Decimal("0.36")]
    assert sum(gaps) == Decimal("1.30")
    notes = bill_warnings([{"bill_kg": 20, "rows": [{"weight_kg": 10, "chargeable_kg": 10}]}])
    assert notes == ["货代账单计费重偏高，请核对"]


def test_delete_measurement_batch(isolated):
    _init()
    from fba_tool.carton_store import delete_import_batch, import_measurements

    imported = import_measurements(
        [{"kind": "box", "fba_box_id": "FBATESTU000001", "weight_kg": 1, "length_cm": 10, "width_cm": 10, "height_cm": 10, "source_row": 1}],
        measured_by="manual",
        source_file="manual.xlsx",
    )
    deleted = delete_import_batch(imported["import_batch_id"])
    assert deleted["carton_measurement"] == 1
