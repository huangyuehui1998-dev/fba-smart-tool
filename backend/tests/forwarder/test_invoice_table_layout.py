"""发票表格模式：表头扫描、清洗、计算字段。"""
from io import BytesIO

import pandas as pd
from openpyxl import Workbook

from fba_tool.invoice_automation import (
    _address_block_cells,
    _header_option_cells,
    detect_template_cells,
    load_invoice_config,
)
from fba_tool.invoice_table import (
    build_invoice_line_items,
    detect_invoice_table_layout,
    header_matches_alias,
    sanitize_header_text,
)
from fba_tool.shipment_parser import CsvShipment


def _workbook_with_row9_table_header() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A2"] = "Consignee"
    ws["A3"] = "Country"
    ws["A4"] = "City"
    ws["A5"] = "是否包税*"
    ws["A6"] = "交货方式*"
    ws["A7"] = "是否买保险*"
    headers = [
        "Shipment ID",
        "Reference ID",
        "SKU",
        "海关编码\nHS CODE",
        "材质/用途\nMaterial&Use",
        "是否带电",
        "是否带磁",
        "申报单价/USD",
        "总申报金额/USD",
        "总数量",
        "实际单件重/kg",
    ]
    for col, title in enumerate(headers, start=1):
        ws.cell(9, col, title)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_sanitize_header_strips_newlines_and_spaces():
    raw = "海关编码 \nHS CODE"
    assert sanitize_header_text(raw) == "海关编码hscode"
    assert header_matches_alias(raw, "HS CODE")
    assert header_matches_alias("材质/用途\nMaterial&Use", "材质/用途Material&Use")


def test_detect_header_at_row_9():
    from fba_tool.xlsx_io import load_workbook_fast

    cfg = load_invoice_config()
    wb = load_workbook_fast(_workbook_with_row9_table_header())
    layout = detect_invoice_table_layout(wb, cfg=cfg)
    assert layout is not None
    assert layout.header_row == 9
    assert layout.first_data_row == 10
    assert "sku" in layout.columns
    assert "shipment_id" in layout.columns
    assert "material" in layout.columns
    assert "total_declared_value_usd" in layout.columns
    assert "unit_weight_kg" in layout.columns


def test_computed_fields_in_line_items():
    ship = CsvShipment(
        shipment_id="FBA123456789U123456",
        warehouse_code="FWA4",
        box_count=2,
        box_items=pd.DataFrame(
            {"箱号": [1, 2], "SKU": ["ABC-2PCS", "ABC-2PCS"], "数量": [10, 10]}
        ),
        box_specs=pd.DataFrame(
            {"箱号": [1, 2], "长(cm)": [30, 30], "宽(cm)": [20, 20], "高(cm)": [10, 10], "实重(kg)": [5.0, 5.0]}
        ),
    )
    csv_bytes = "SKU,Product\nABC-2PCS,Test Product\n".encode("utf-8")
    customs = {"ABC-2PCS": {"单个货值": 5.0}}
    items = build_invoice_line_items(ship, csv_bytes, catalog=None, customs_by_sku=customs)
    assert len(items) == 1
    item = items[0]
    assert item.values["total_qty"] == 20
    assert item.values["battery"] == "否"
    assert item.values["magnet"] == "否"
    assert item.values["total_declared_value_usd"] == 100.0
    assert item.values["unit_weight_kg"] == 0.5


def test_address_cells_scanned_above_table_header():
    from fba_tool.xlsx_io import load_workbook_fast

    cfg = load_invoice_config()
    wb = load_workbook_fast(_workbook_with_row9_table_header())
    cell_map = detect_template_cells(wb, cfg, max_row=8)
    assert "Consignee" in cell_map or "地址" in cell_map

    cells = _address_block_cells(
        wb,
        cfg,
        pdf_bytes=b"",
        api_key=None,
        header_row=9,
    )
    assert isinstance(cells, list)


def test_header_option_defaults_above_table():
    from fba_tool.shipment_parser import CsvShipment
    from fba_tool.xlsx_io import load_workbook_fast

    cfg = load_invoice_config()
    wb = load_workbook_fast(_workbook_with_row9_table_header())
    ship = CsvShipment(shipment_id="FBA123", warehouse_code="FWA4", box_count=2)
    cells = _header_option_cells(wb, cfg, ship=ship, header_row=9)
    by_key = {c["field_key"]: c for c in cells}
    assert by_key["是否包税"]["value"] == "是"
    assert by_key["交货方式"]["value"] == "海运"
    assert "精准提取" in by_key["是否包税"]["source"]


def _label_mode_template() -> bytes:
    """无表格特征的模板 → detect_invoice_table_layout 返回 None，走 label 分支。"""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "商业发票"
    ws["A2"] = "FBA单号"
    ws["A3"] = "仓库代码"
    ws["A4"] = "计费吨"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _minimal_shipment_csv() -> bytes:
    return (
        "货件编号,FBA123ABC\n"
        "货件名称,Shipment to SCK4\n"
        "箱子数量,1\n"
        "箱号,FBA123ABCU001\n"
        "SKU,箱子 1\n"
        "SKU-A,10\n"
        "包装箱长度,40\n"
        "包装箱宽度,30\n"
        "包装箱高度,30\n"
        "包装箱重量,12\n"
    ).encode("utf-8-sig")


def test_label_mode_draft_populates_revenue_ton():
    """label 模式曾因引用未定义的 revenue_ton 直接 NameError，非表格式货代模板全部填充失败。"""
    from fba_tool.invoice_automation import build_invoice_draft
    from fba_tool.xlsx_io import load_workbook_fast

    tpl = _label_mode_template()
    assert detect_invoice_table_layout(load_workbook_fast(tpl), cfg=load_invoice_config()) is None

    draft, _ = build_invoice_draft(
        template_bytes=tpl,
        csv_bytes=_minimal_shipment_csv(),
        pdf_bytes=b"",
    )
    assert draft.is_table_mode is False
    assert draft.shipment_id == "FBA123ABC"

    values = {f.key: f.value for f in draft.fields}
    assert "计费吨" in values
    # 单箱 40×30×30 / 6000 = 6kg 抛重 < 12kg 实重 → 计费吨取实重
    assert values["计费吨"] == 12.0


def test_detect_qiluo_template_layout():
    from pathlib import Path

    from fba_tool.xlsx_io import load_workbook_fast

    tpl = Path(__file__).resolve().parents[2] / "templates" / "forwarders" / "qiluo.xlsx"
    assert tpl.is_file(), f"missing {tpl}; run scripts/generate_qiluo_template.py"

    cfg = load_invoice_config()
    wb = load_workbook_fast(tpl.read_bytes())
    layout = detect_invoice_table_layout(wb, cfg=cfg)
    assert layout is not None
    assert layout.header_row == 1
    assert layout.first_data_row == 2

    expected = {
        "box_range",
        "fba_id",
        "po_number",
        "chinese_name",
        "english_name",
        "qty_per_box",
        "declared_price",
        "material",
        "purpose",
        "brand",
        "model",
        "hs_code",
        "unit_weight_kg",
        "length",
        "width",
        "height",
    }
    assert expected.issubset(set(layout.columns.keys()))
