"""发票表格脏数据清理测试。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from fba_tool.invoice_automation import export_invoice_xlsx, load_invoice_config
from fba_tool.invoice_automation import InvoiceDraft
from fba_tool.invoice_table import (
    clear_cells_for_table_zone,
    detect_invoice_table_layout,
    table_cells_from_line_items,
    InvoiceLineItem,
)
from fba_tool.xlsx_io import load_workbook_fast


def test_qiluo_template_has_header_only():
    tpl = Path(__file__).resolve().parents[2] / "templates" / "forwarders" / "qiluo.xlsx"
    wb = load_workbook_fast(tpl.read_bytes())
    ws = wb.active
    assert ws.max_row == 1
    assert str(ws.cell(2, 1).value or "").strip() == ""


def test_clear_cells_cover_data_zone():
    tpl = Path(__file__).resolve().parents[2] / "templates" / "forwarders" / "qiluo.xlsx"
    wb = load_workbook_fast(tpl.read_bytes())
    layout = detect_invoice_table_layout(wb, cfg=load_invoice_config(), forwarder_id="qiluo")
    assert layout is not None
    clears = clear_cells_for_table_zone(layout, max_data_rows=5)
    assert len(clears) == len(layout.columns) * 5
    assert all(c["value"] == "" and c.get("clear") for c in clears)


def test_table_cells_only_data_rows():
    tpl = Path(__file__).resolve().parents[2] / "templates" / "forwarders" / "qiluo.xlsx"
    wb = load_workbook_fast(tpl.read_bytes())
    layout = detect_invoice_table_layout(wb, cfg=load_invoice_config(), forwarder_id="qiluo")
    item = InvoiceLineItem(
        values={"box_range": "1-3", "fba_id": "FBA1", "sku": "A"},
        sources={"box_range": "green", "fba_id": "green", "sku": "green"},
    )
    cells = table_cells_from_line_items(layout, [item], {"green": "精准提取"})
    assert cells
    assert not any(c.get("clear") for c in cells)
    assert any(c.get("field_key") == "box_range" and c.get("value") == "1-3" for c in cells)


def test_export_clears_dirty_template_rows():
    tpl = Path(__file__).resolve().parents[2] / "templates" / "forwarders" / "qiluo.xlsx"
    raw = tpl.read_bytes()
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    ws.cell(2, 1, "脏数据")
    ws.cell(3, 2, "旧FBA")
    dirty = BytesIO()
    wb.save(dirty)

    draft = InvoiceDraft(
        shipment_id="FBA1",
        warehouse_code="ONT8",
        is_table_mode=True,
        table_cells=[{"cell": "A2", "value": "1-5", "field_key": "box_range"}],
    )
    out = export_invoice_xlsx(dirty.getvalue(), draft, forwarder_id="qiluo")
    result = load_workbook(BytesIO(out))
    ws2 = result.active
    assert ws2.cell(2, 1).value == "1-5"
    assert ws2.cell(3, 2).value in (None, "")
