"""Generic export test: variable SKU count and box count."""
import pandas as pd

from fba_tool.boxing import SkuRow, plans_to_editor_df
from fba_tool.excel_export import build_box_matrix, export_amazon_template, export_internal_archive
from fba_tool.sta_template_export import pick_sta_packing_sheet
from fba_tool.template_layout import detect_internal_layout, detect_sta_layout
import openpyxl
from io import BytesIO

# 12 SKUs, 25 boxes
sku_rows = []
editor_rows = []
for s in range(1, 13):
    sku = f"TEST-SKU-{s:03d}-FBA"
    sku_rows.append(
        SkuRow(sku=sku, total_qty=50, carton_qty=5, length_cm=50, width_cm=30, height_cm=40, unit_weight_kg=1.0, carton_tare_kg=0.5)
    )
    for b in range(1, 26):
        if (s + b) % 3 == 0:  # sparse fill
            editor_rows.append({"箱号": b, "来源": "整箱", "SKU简称": sku, "数量": 5 if b % 2 else 2})

editor_df = pd.DataFrame(editor_rows)
plans_df = plans_to_editor_df([])  # unused
internal = export_internal_archive(editor_df, sku_rows)
amazon = export_amazon_template(editor_df, sku_rows)

wb1 = openpyxl.load_workbook(BytesIO(internal))
ws1 = wb1.active
layout1 = detect_internal_layout(ws1)
print("INTERNAL skus filled", sum(1 for r in range(layout1.sku_start_row, layout1.footer_start_row) if str(ws1.cell(r,1).value or "").startswith("TEST")))
print("INTERNAL footer row", layout1.footer_start_row, "max box hdr", layout1.template_max_boxes)

wb2 = openpyxl.load_workbook(BytesIO(amazon))
ws2 = pick_sta_packing_sheet(wb2)
print("AMAZON L3", ws2.cell(3, 12).value)
amazon_skus = sum(
    1
    for r in range(6, ws2.max_row + 1)
    if str(ws2.cell(r, 1).value or "").startswith("TEST")
)
print("AMAZON skus filled", amazon_skus)
mat = build_box_matrix(editor_df)
print("MATRIX shape", mat.shape)
assert mat.shape[1] == 26
assert ws2.cell(3, 12).value == 25
assert amazon_skus == 12
print("ALL OK")
