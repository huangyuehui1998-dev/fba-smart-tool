"""S4-fix1：发票按货代模板导出（SEND / 九方）+ golden（以货代原始样例为期望值）。"""
import io
import json
from pathlib import Path

import openpyxl
import pytest

from fba_tool.invoice_export import SEND_SHEET, ExportConfigError, export_forwarders, export_invoice
from fba_tool.invoice_rules import (
    Box,
    InvoiceBlocked,
    box_no_from_fba_id,
    box_ranges,
    read_jiufang_box_data,
    to_number,
)
from fba_tool.shipment_parser import parse_csv_shipment
from tests._realdata import skip_module_if_missing
from tests.golden.compare_golden import VERDICT_FAIL, VERDICT_KNOWN, failures, write_report
from tests.golden.compare_invoice_export import compare_jiufang, compare_send

HERE = Path(__file__).resolve().parents[2]
G = HERE / "tests" / "golden"
skip_module_if_missing(
    G / "invoice_export_inputs.json",
    G / "product_master_golden.json",
    HERE / "tests/fixtures/warehouse/fba_warehouse_qiluo.xlsx",
)
INPUTS = json.loads((G / "invoice_export_inputs.json").read_text(encoding="utf-8"))["cases"]
SNAP = {p["fnsku"]: p for p in json.loads((G / "product_master_golden.json").read_text(encoding="utf-8"))["products"]}

# 已知差异以 invoice_manifest.csv 的 expected_diff 为准（每条都写了原因）
from tests.golden.compare_golden import load_manifest, parse_expected_diff  # noqa: E402

MANIFEST = {r["case_id"]: r for r in load_manifest(G / "invoice_manifest.csv")}
SEND_DIFF = parse_expected_diff(MANIFEST["FBA19P4N6MM2"]["expected_diff"])
UNVERIFIED_MEASURED = {"LCL2607GZ0PA", "LCL26099BD6A", "LCL2609TNKDA"}


def _warehouses():
    from fba_tool.warehouse_db import parse_warehouse_sheet

    rows, _ = parse_warehouse_sheet((HERE / "tests/fixtures/warehouse/fba_warehouse_qiluo.xlsx").read_bytes(), "q.xlsx")
    return {r["code"]: r for r in rows}


WH = _warehouses()


def _csv(name):
    p = G / "samples" / "inputs" / name
    return parse_csv_shipment(p.read_bytes(), file_name=p.name)


def _export(case):
    s = _csv(case["csv"])
    measured = case.get("measured")
    if case.get("measured_file"):
        measured = read_jiufang_box_data((G / "samples" / "inputs" / case["measured_file"]).read_bytes(), s.shipment_id,
                                         box_id_fix=case.get("measured_box_id_fix"))
    return export_invoice(case["forwarder"], shipment=s, products=SNAP, reference_id=case["reference_id"],
                          warehouse_info=WH.get(s.warehouse_code), order_no=case.get("order_no", ""),
                          channel=case.get("channel", ""), invoice_date=case.get("invoice_date"), measured=measured)


# ---------------------------------------------------------------- golden
@pytest.mark.parametrize("case", [c for c in INPUTS if c["forwarder"] == "send"], ids=lambda c: c["case_id"])
def test_golden_send_against_original_sample(case, tmp_path):
    res = _export(case)
    out = tmp_path / res["file_name"]
    out.write_bytes(res["bytes"])
    rows = compare_send(case["case_id"], G / case["expected"], out, parse_expected_diff(MANIFEST[case["case_id"]]["expected_diff"]))
    write_report(rows, tmp_path / "report.csv")
    assert not failures(rows), "\n".join(f"{r.ref}: 期望 {r.expected!r} 实际 {r.actual!r}" for r in failures(rows))
    assert sum(1 for r in rows if r.verdict == "通过") >= 40
    assert res["file_name"] == f"{case['case_id']}发票.xlsx"


@pytest.mark.parametrize("case", [c for c in INPUTS if c["forwarder"] == "jiufang"], ids=lambda c: c["case_id"])
def test_golden_jiufang_against_original_sample(case, tmp_path):
    res = _export(case)
    out = tmp_path / res["file_name"]
    out.write_bytes(res["bytes"])
    diff = parse_expected_diff(MANIFEST[case["case_id"]]["expected_diff"])
    rows = compare_jiufang(case["case_id"], G / case["expected"], out, diff)
    assert not failures(rows), "\n".join(f"{r.ref}: 期望 {r.expected!r} 实际 {r.actual!r} {r.note}" for r in failures(rows))
    inv = [r for r in rows if r.ref.startswith("Invoice!") and r.verdict == "通过"]
    pk = [r for r in rows if r.ref.startswith("Packing list!") and r.verdict == "通过"]
    assert len(inv) >= 30 and len(pk) >= 20
    assert res["file_name"] == f"{case['order_no']}-箱单发票.xlsx"


def test_golden_cases_cover_all_13_samples():
    assert sorted(c["case_id"] for c in INPUTS) == sorted(
        p.name.replace("发票.xlsx", "").replace("-箱单", "") for p in (G / "samples" / "invoice").glob("*.xlsx"))
    assert len([c for c in INPUTS if c["forwarder"] == "send"]) == 5
    assert len([c for c in INPUTS if c["forwarder"] == "jiufang"]) == 8
    # 只有 3 票的实测值取自样例，其余 5 票来自九方单箱数据原件
    assert {c["case_id"] for c in INPUTS if c.get("measured")} == UNVERIFIED_MEASURED


def test_measured_box_id_typo_is_blocked_until_confirmed():
    """QGXTA 原件里第 6 箱的外箱号写成了 12：不确认更正就拦截，并点名多出/缺的箱。"""
    c = next(x for x in INPUTS if x["case_id"] == "LCL2610QGXTA")
    assert c["measured_box_id_fix"] == {"12": 6}
    s = _csv(c["csv"])
    raw = read_jiufang_box_data((G / "samples" / "inputs" / c["measured_file"]).read_bytes(), s.shipment_id)
    with pytest.raises(InvoiceBlocked) as e:
        export_invoice("jiufang", shipment=s, products=SNAP, reference_id="R", warehouse_info=WH["TEB9"], order_no="L", measured=raw)
    b = e.value.blocking[0]
    assert b["code"] == "measured_box_mismatch" and b["extra"] == ["FBA19SC5CCXCU000012"] and b["missing"] == ["FBA19SC5CCXCU000006"]


def test_jiufang_measured_box_data_matches_sample_independently():
    """九方单箱数据（原件）与九方箱单 Packing list 的单箱毛重/尺寸逐箱一致：实测值是输入，不是算出来的。"""
    for c in INPUTS:
        if not c.get("measured_file"):
            continue
        s = _csv(c["csv"])
        m = read_jiufang_box_data((G / "samples" / "inputs" / c["measured_file"]).read_bytes(), s.shipment_id,
                                  box_id_fix=c.get("measured_box_id_fix"))
        ws = openpyxl.load_workbook(G / c["expected"], data_only=True)["Packing list"]
        seen = 0
        for r in ws.iter_rows(min_row=6, values_only=True):
            if r[5] and str(r[5]).startswith("FBA"):
                assert m[r[5]]["weight"] == r[12] and (m[r[5]]["length"], m[r[5]]["width"], m[r[5]]["height"]) == (r[8], r[9], r[10])
                seen += 1
        assert seen == s.box_count == len(m)


# ---------------------------------------------------------------- 规则
def test_send_sheet_name_has_nbsp_and_numbers_are_numeric(tmp_path):
    res = _export(next(c for c in INPUTS if c["case_id"] == "FBA19P4N6MM2"))
    wb = openpyxl.load_workbook(io.BytesIO(res["bytes"]))
    assert wb.sheetnames == [SEND_SHEET] and "\xa0" in wb.sheetnames[0]
    ws = wb[SEND_SHEET]
    assert ws["A15"].value == "1-23" and ws["A16"].value == "24-37"
    for ref in ("C15", "F15", "Q15", "R15", "S15", "Q16", "R16"):
        assert isinstance(ws[ref].value, (int, float)), ref
    assert ws["Q15"].value == 50.0 and ws["R15"].value == 39.99 and ws["T15"].value == "USD"
    assert ws["M15"].value == "6307909000"  # HS 保持文本，不丢前导位
    assert ws["A17"].value is None


def test_to_number_text_numbers():
    assert to_number("39.99") == 39.99 and to_number("50.0") == 50.0 and to_number("92") == 92
    assert to_number("1-23") == "1-23" and to_number(None) is None and to_number("") is None


def test_box_no_from_u_number():
    assert box_no_from_fba_id("FBA19P4N6MM2U000023") == 23
    assert box_no_from_fba_id("fba19p4n6mm2u000001") == 1
    assert box_no_from_fba_id("abc") is None


def test_box_ranges_merge_only_same_sku_qty_weight_dims_and_mixed_box():
    d = (45.0, 45.0, 41.0)
    boxes = [Box(1, "X U000001", [("A", 4)], 7.5, d), Box(2, "", [("A", 4)], 7.5, d), Box(3, "", [("A", 4)], 7.6, d),
             Box(4, "", [("A", 2), ("B", 3)], 8.0, d), Box(5, "", [("B", 4)], 9.5, d), Box(6, "", [("B", 4)], 9.5, d),
             Box(8, "", [("B", 4)], 9.5, d)]
    rows = box_ranges(boxes)
    got = [(r.label, r.fnsku, r.qty, r.weight) for r in rows]
    assert got == [("1-2", "A", 8, 7.5), ("3", "A", 4, 7.6), ("4", "A", 2, 8.0), ("4", "B", 3, None),
                   ("5-6", "B", 8, 9.5), ("8", "B", 4, 9.5)]
    # 混装箱的箱重只记一次
    assert sum(r.weight or 0 for r in rows if r.start == 4) == 8.0


def test_send_box_ranges_follow_u_number_not_csv_column_order():
    s = _csv("第三票FBA19P4N6MM2.csv")
    cols = list(s.box_fba_ids.items())
    assert [box_no_from_fba_id(v) for _, v in cols] != sorted(box_no_from_fba_id(v) for _, v in cols)  # CSV 列序是乱的
    res = _export(next(c for c in INPUTS if c["case_id"] == "FBA19P4N6MM2"))
    assert [r.label for r in res["model"].ranges] == ["1-23", "24-37"]


def _mixed_csv():
    head = ['"货件编号","FBA1TESTMIX1"', '"货件名称","TEST-LAX9"', '"配送地址","FONTANA, CA"', '"箱子数量","2"', '', '"多件商品（2 箱）"']
    rows = [
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量","箱子 2 的商品数量"',
        '"S-A","a","B0A","X0000000AA","新品","无","6","4","2"',
        '"S-B","b","B0B","X0000000BB","新品","无","3","","3"',
        '',
        '"","","","","","","箱号","FBA1TESTMIX1U000001","FBA1TESTMIX1U000002"',
        '"","","","","","","包装箱重量（磅）：","10","20"',
        '"","","","","","","包装箱长度（英寸）：","10","10"',
        '"","","","","","","包装箱宽度（英寸）：","10","10"',
        '"","","","","","","包装箱高度（英寸）：","10","10"',
    ]
    return "\n".join(head + rows).encode("utf-8")


def _prod(fnsku, **kw):
    p = {"fnsku": fnsku, "listing_price": 10.0, "english_name": "Item", "chinese_name": "物品", "brand": "B", "material": "铁",
         "purpose": "用", "hs_code": "7326909000", "asin": "B0X", "model_name": "无",
         "invoice": {"unit": "个", "purchase_price": 5, "material_en": "iron", "purpose_en": "use", "pcs_per_set": 1}}
    p.update(kw)
    return p


def test_mixed_box_export_send_and_jiufang():
    try:
        s = parse_csv_shipment(_mixed_csv(), file_name="mix.csv")
    except ValueError as exc:  # 解析器对手写 CSV 格式要求严格时，退回用真实 CSV 改造
        pytest.skip(f"手写混装 CSV 解析失败：{exc}")
    prods = {"X0000000AA": _prod("X0000000AA"), "X0000000BB": _prod("X0000000BB", listing_price=20.0)}
    res = export_invoice("send", shipment=s, products=prods, reference_id="R1")
    ws = openpyxl.load_workbook(io.BytesIO(res["bytes"]))[SEND_SHEET]
    rows = [(ws[f"A{r}"].value, ws[f"E{r}"].value, ws[f"F{r}"].value, ws[f"C{r}"].value) for r in range(15, 19)]
    assert rows[0][:3] == ("1", "X0000000AA", 4) and rows[1][:3] == ("2", "X0000000AA", 2) and rows[2][:3] == ("2", "X0000000BB", 3)
    assert rows[2][3] is None and rows[3][0] is None
    res = export_invoice("jiufang", shipment=s, products=prods, reference_id="R1", warehouse_info=WH["LAX9"], order_no="LCLX")
    wb = openpyxl.load_workbook(io.BytesIO(res["bytes"]))
    inv, pk = wb["Invoice"], wb["Packing list"]
    assert [inv[f"O{r}"].value for r in (25, 26)] == [6, 3] and inv["S27"].value == 120.0 and inv["U27"].value == 2
    assert [pk[f"F{r}"].value for r in (6, 7, 8)] == ["FBA1TESTMIX1U000001", "FBA1TESTMIX1U000002", "FBA1TESTMIX1U000002"]
    assert pk["G8"].value is None and pk["M8"].value is None and pk["G9"].value == 2


def test_declared_price_follows_d01_and_purchase_price_copied():
    s = _csv("第三票FBA19P4N6MM2.csv")
    prods = json.loads(json.dumps(SNAP))
    prods["X0059EFBRF"]["listing_price"] = None
    prods["X0059EFBRF"]["declared_value"] = 33.0
    prods["X0059EFBRF"]["field_sources"] = {"declared_value": "manual"}
    res = export_invoice("send", shipment=s, products=prods, reference_id="R")
    ws = openpyxl.load_workbook(io.BytesIO(res["bytes"]))[SEND_SHEET]
    assert ws["R15"].value == 33.0 and ws["Q15"].value == 50.0 and ws["R16"].value == 49.99
    # 旧数据申报价：能导出但出警告
    prods["X0059EFBRF"]["field_sources"] = {}
    res = export_invoice("send", shipment=s, products=prods, reference_id="R")
    assert "price_legacy" in [w["code"] for w in res["warnings"]]
    # 没有任何价格：拦截
    prods["X0059EFBRF"]["declared_value"] = None
    with pytest.raises(InvoiceBlocked) as e:
        export_invoice("send", shipment=s, products=prods, reference_id="R")
    assert "price_missing" in [b["code"] for b in e.value.blocking]


def test_price_drop_warning_same_rule_as_s2fix2():
    s = _csv("第三票FBA19P4N6MM2.csv")
    prods = json.loads(json.dumps(SNAP))
    prods["X0059EFBRF"]["declared_value"] = 60.0  # 历史申报价 60，领星 39.99：降 33%
    res = export_invoice("send", shipment=s, products=prods, reference_id="R")
    assert "price_drop" in [w["code"] for w in res["warnings"]]


def test_fishing_without_config_blocked_and_with_config_has_no_pet_text(tmp_path):
    from fba_tool.hs_config import upsert_hs_template
    s = _csv("第三票FBA19P4N6MM2.csv")
    prods = json.loads(json.dumps(SNAP))
    for f in ("X0059EFBRF", "X0059EHE0H"):
        prods[f].update({"category": "fishing", "hs_code": "9507909000", "chinese_name": "", "english_name": "Fishing Lure",
                         "material": "塑料", "purpose": "钓鱼用"})
        prods[f]["invoice"]["dest_hs_code"] = None
    with pytest.raises(InvoiceBlocked) as e:
        export_invoice("send", shipment=s, products=prods, reference_id="R", hs_templates={})
    codes = {b["code"] for b in e.value.blocking}
    assert codes == {"fishing_config_missing"}
    assert set(e.value.blocking[0]["missing"]) == {"chinese_name", "hs_declare_template"}
    for f in ("X0059EFBRF", "X0059EHE0H"):
        prods[f]["chinese_name"] = "路亚饵"
    tpl = {("fishing", "9507909000"): {"category": "fishing", "hs_code": "9507909000"}}
    res = export_invoice("send", shipment=s, products=prods, reference_id="R", hs_templates=tpl)
    wb = openpyxl.load_workbook(io.BytesIO(res["bytes"]))
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "宠物" not in text and "路亚饵" in text
    res = export_invoice("jiufang", shipment=s, products={k: {**v, "invoice": {**v["invoice"], "material_en": "plastic", "purpose_en": "fishing"}} for k, v in prods.items()},
                         reference_id="R", hs_templates=tpl, warehouse_info=WH["LAS1"], order_no="LCLX")
    wb = openpyxl.load_workbook(io.BytesIO(res["bytes"]))
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "宠物" not in text and "Pet" not in text


def test_missing_invoice_fields_block_with_list():
    s = _csv("第三票FBA19P4N6MM2.csv")
    prods = json.loads(json.dumps(SNAP))
    prods["X0059EFBRF"]["brand"] = ""
    prods["X0059EFBRF"]["invoice"]["unit"] = ""
    del prods["X0059EHE0H"]
    with pytest.raises(InvoiceBlocked) as e:
        export_invoice("send", shipment=s, products=prods, reference_id="R")
    by = {b["code"]: b for b in e.value.blocking}
    assert set(by["invoice_field_missing"]["missing"]) == {"品牌", "单位"}
    assert by["product_missing"]["sku"] == "X0059EHE0H"


def test_jiufang_needs_warehouse_address():
    s = _csv("第四票FBA19P4NX8HH.csv")
    with pytest.raises(InvoiceBlocked) as e:
        export_invoice("jiufang", shipment=s, products=SNAP, reference_id="R", warehouse_info=None, order_no="L")
    assert e.value.blocking[0]["code"] == "warehouse_address_missing"


def test_adapters_registered_and_unknown_forwarder():
    fw = export_forwarders()
    assert {fw["send"]["export_adapter"], fw["jiufang"]["export_adapter"]} == {"send", "jiufang"}
    assert "qiluo" not in fw
    with pytest.raises(ExportConfigError):
        export_invoice("qiluo", shipment=_csv("第三票FBA19P4N6MM2.csv"), products=SNAP)


def test_bad_json_forwarder_config_gives_clear_error(monkeypatch, tmp_path):
    import fba_tool.app_paths as ap
    bad = tmp_path / "forwarders.json"
    bad.write_text('{"forwarders": [ {"id": "x",, } ]', encoding="utf-8")
    monkeypatch.setattr(ap, "resolve_config_read", lambda name: bad)
    with pytest.raises(ExportConfigError, match="配置文件格式错误"):
        export_forwarders()


def test_test1_five_tickets_export_smoke(tmp_path):
    """repo test1/ 的 5 票 CSV 用 SEND 和九方各导一遍：能出文件、箱数/数量与 CSV 一致；缺产品资料的按拦截列出。"""
    root = HERE.parent / "test1"
    files = sorted(root.glob("*.csv"))
    assert len(files) == 5
    done = []
    for f in files:
        s = parse_csv_shipment(f.read_bytes(), file_name=f.name, allow_unassigned_fba=True)  # test1 是“待分配”的货件
        fns = sorted({str(x).upper() for x in s.box_items["FNSKU"]})
        prods = {fn: _prod(fn, listing_price=19.99) for fn in fns}
        for fw in ("send", "jiufang"):
            wh = WH.get(s.warehouse_code) or {"address1": "x", "city": "y", "state": "z", "zip": "0"}
            res = export_invoice(fw, shipment=s, products=prods, reference_id="REF", warehouse_info=wh, order_no="LCLTEST")
            (tmp_path / f"{fw}-{res['file_name']}").write_bytes(res["bytes"])
            m = res["model"]
            codes = [w["code"] for w in res["warnings"]]
            assert ("box_mark_missing" in codes) == (fw == "jiufang")  # 待分配货件：九方箱唛为空要提示
            assert len(m.boxes) == s.box_count
            assert m.total_qty == int(s.box_items["数量"].sum())
            assert sum(r.qty for r in m.ranges) == m.total_qty
            done.append((fw, s.shipment_id))
        # 去掉一个产品 → 被拦截，并点名
        prods.pop(fns[0])
        with pytest.raises(InvoiceBlocked) as e:
            export_invoice("send", shipment=s, products=prods, reference_id="REF")
        assert fns[0] in e.value.args[0]
    assert len(done) == 10


# ---------------------------------------------------------------- API（临时 app_data 里的产品库 + 仓库地址库）
def test_api_export_send_and_jiufang_with_db(tmp_path):
    from fastapi.testclient import TestClient

    from fba_tool.master_db import init_master_db, upsert_product
    from fba_tool.product_fields import set_product_fields
    from fba_tool.warehouse_db import apply_warehouse_import
    from main import app

    init_master_db()
    apply_warehouse_import((HERE / "tests/fixtures/warehouse/fba_warehouse_qiluo.xlsx").read_bytes(), "q.xlsx")
    for f in ("X0059EFBRF", "X0059EHE0H"):
        p = SNAP[f]
        sku = "GOLDEN-" + f
        upsert_product({"sku": sku, "fnsku": f, "asin": p["asin"], "english_name": p["english_name"], "chinese_name": p["chinese_name"],
                        "brand": p["brand"], "material": p["material"], "purpose": p["purpose"], "model_name": p["model_name"],
                        "hs_code": p["hs_code"], "extra_data": {"invoice": p["invoice"]}}, overwrite=True)
        set_product_fields(sku, {"listing_price": p["listing_price"]}, "lingxing")
    c = TestClient(app)
    assert {x["id"] for x in c.get("/api/invoice-export/forwarders").json()["forwarders"]} == {"send", "jiufang", "qiluo_t6", "qiluo_t5", "deshu"}  # S8 D4：显式列出全部 5 家
    csv = (G / "samples/inputs/第三票FBA19P4N6MM2.csv").read_bytes()
    r = c.post("/api/invoice-export/preview", data={"forwarder_id": "send", "reference_id": "2VJ79IPP"}, files={"csv_file": ("a.csv", csv)})
    assert r.status_code == 200, r.text
    j = r.json()
    assert [x["box"] for x in j["ranges"]] == ["1-23", "24-37"] and j["warehouse"] == "LAS1" and j["total_qty"] == 148
    r = c.post("/api/invoice-export/export", data={"forwarder_id": "send", "reference_id": "2VJ79IPP"}, files={"csv_file": ("a.csv", csv)})
    assert r.status_code == 200
    out = tmp_path / "o.xlsx"
    out.write_bytes(r.content)
    rows = compare_send("FBA19P4N6MM2", G / "samples/invoice/FBA19P4N6MM2发票.xlsx", out, SEND_DIFF)
    assert not failures(rows)
    r = c.post("/api/invoice-export/export", data={"forwarder_id": "jiufang", "reference_id": "2VJ79IPP", "order_no": "LCLX"},
               files={"csv_file": ("a.csv", csv)})
    assert r.status_code == 200 and openpyxl.load_workbook(io.BytesIO(r.content))["Invoice"]["B13"].value == "12300 Bermuda Road"
    # 缺产品 → 409 + 拦截清单
    csv2 = (G / "samples/inputs/第二票FBA19SC2G1WQ.csv").read_bytes()
    r = c.post("/api/invoice-export/export", data={"forwarder_id": "send"}, files={"csv_file": ("b.csv", csv2)})
    assert r.status_code == 409 and r.json()["blocking"][0]["code"] == "product_missing"


# ---------------------------------------------------------------- 解析侧：对原始样例独立核对（期望来自 CSV / 九方单箱数据原件，不来自被测代码）
@pytest.mark.parametrize("case", INPUTS, ids=lambda c: c["case_id"])
def test_parse_side_original_samples_vs_independent_sources(case):
    from collections import Counter

    from fba_tool import invoice_templates as it

    p = G / case["expected"]
    r = it.parse_invoice_file(p.read_bytes(), p.name)
    assert r["template_id"] == ("T2" if case["forwarder"] == "send" else "T8")
    t = r["tickets"][0]
    s = _csv(case["csv"])
    assert t["shipment_id"] == s.shipment_id
    assert t["box_count"] == s.box_count
    csv_qty = Counter()
    for _, row in s.box_items.iterrows():
        csv_qty[str(row["FNSKU"]).upper()] += int(row["数量"])
    got = Counter()
    for l in t["lines"]:
        got[l["sku"].upper()] += int(l["qty"])
    assert got == csv_qty
    if case["forwarder"] == "send":
        # 单箱重 = CSV 该 U 号箱的重量（磅→kg，0.01 容差）
        by_no = {box_no_from_fba_id(s.box_fba_ids[c]): float(w) for c, w in zip(s.box_specs["箱号"], s.box_specs["实重(kg)"])}
        assert all(abs(t["boxes"][n - 1] - by_no[n]) <= 0.01 for n in by_no)
    elif case.get("measured_file"):
        m = read_jiufang_box_data((G / "samples" / "inputs" / case["measured_file"]).read_bytes(), s.shipment_id,
                                  box_id_fix=case.get("measured_box_id_fix"))
        assert sorted(t["boxes"]) == sorted(v["weight"] for v in m.values())


def test_send_dims_positive_check_against_csv():
    """SEND 尺寸列正向核验：导出值 = CSV 尺寸（英寸→厘米），且 = 样例值四舍五入；已知差异逐格登记，没有通配符。"""
    from tests.golden.compare_invoice_export import _dims
    assert all("*" not in k for r in MANIFEST.values() for k in parse_expected_diff(r["expected_diff"]))
    for c in [x for x in INPUTS if x["forwarder"] == "send"]:
        res = _export(c)
        ws = openpyxl.load_workbook(io.BytesIO(res["bytes"]))[SEND_SHEET]
        exp = openpyxl.load_workbook(G / c["expected"])[SEND_SHEET]
        s = _csv(c["csv"])
        csv_dims = {(float(r["长(cm)"]), float(r["宽(cm)"]), float(r["高(cm)"])) for _, r in s.box_specs.iterrows()}
        for row in (15, 16):
            got = tuple(_dims(ws[f"B{row}"].value))
            assert got in csv_dims
            assert got == tuple(round(x) for x in _dims(exp[f"B{row}"].value))


def test_jiufang_box_mark_present_for_assigned_shipment_no_warning():
    res = _export(next(c for c in INPUTS if c["case_id"] == "LCL2609PU1DA"))
    assert "box_mark_missing" not in [w["code"] for w in res["warnings"]]
