"""golden 回归比对工具（方案 v7 §5.1）。

只做“期望文件 vs 实际文件”的单元格比对，不依赖导出代码；S3/S4 的导出完成后，
用 run_case() 传入实际生成的文件即可接入。规则：
- 数值转 Decimal 比较：金额容差 0，重量容差 0.01kg；
- 文本 strip 后全等；
- 期望值是文本数字（'39.99'、'50.0'）时先转 Decimal；
- “N箱”这类文本件数先清洗成整数，结论记为“通过（期望值已清洗）”；
- 工作表名先 NFKC 规范化，并把 \\xa0 换成普通空格；
- manifest 里列入 expected_diff 的单元格，结论记为“已知差异”，不算失败；
- 警告列表（expected_warnings）单独比对。
"""
from __future__ import annotations

import csv
import datetime as _dt
import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

GOLDEN_DIR = Path(__file__).resolve().parent
MANIFEST_COLUMNS = [
    "case_id",
    "kind",
    "input_paths",
    "expected_path",
    "level",
    "expected_diff",
    "expected_warnings",
]

VERDICT_PASS = "通过"
VERDICT_PASS_CLEANED = "通过（期望值已清洗）"
VERDICT_KNOWN = "已知差异"
VERDICT_FAIL = "失败"

MONEY_TOLERANCE = Decimal("0")
WEIGHT_TOLERANCE = Decimal("0.01")

_BOX_TEXT = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*箱\s*$")
_NUMBER_TEXT = re.compile(r"^\s*-?\d+(?:\.\d+)?\s*$")


def normalize_sheet_name(name: str) -> str:
    return unicodedata.normalize("NFKC", str(name)).replace("\xa0", " ").strip()


def find_sheet_name(names: Iterable[str], target: str) -> Optional[str]:
    want = normalize_sheet_name(target)
    for n in names:
        if normalize_sheet_name(n) == want:
            return n
    return None


def to_decimal(value: Any) -> Tuple[Optional[Decimal], bool]:
    """返回 (Decimal 或 None, 是否经过“N箱”清洗)。非数字返回 (None, False)。"""
    if value is None or isinstance(value, bool):
        return None, False
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)), False
    s = str(value)
    m = _BOX_TEXT.match(s)
    if m:
        return Decimal(m.group(1)), True
    if _NUMBER_TEXT.match(s):
        try:
            return Decimal(s.strip()), False
        except InvalidOperation:
            return None, False
    return None, False


def xls_serial_to_iso(serial: float) -> str:
    base = _dt.datetime(1899, 12, 30)
    return (base + _dt.timedelta(days=float(serial))).date().isoformat()


def compare_value(expected: Any, actual: Any, *, kind: str = "text") -> str:
    """kind: money | weight | count | text | date。返回结论字符串。"""
    if kind == "date":
        e = xls_serial_to_iso(expected) if isinstance(expected, (int, float)) else (str(expected or "").strip())
        a = xls_serial_to_iso(actual) if isinstance(actual, (int, float)) else (str(actual or "").strip())
        return VERDICT_PASS if e == a else VERDICT_FAIL
    if kind in ("money", "weight", "count"):
        e, cleaned = to_decimal(expected)
        a, _ = to_decimal(actual)
        if e is None and a is None:
            return VERDICT_PASS if _blank(expected) and _blank(actual) else VERDICT_FAIL
        if e is None or a is None:
            return VERDICT_FAIL
        tol = WEIGHT_TOLERANCE if kind == "weight" else MONEY_TOLERANCE
        if abs(e - a) <= tol:
            return VERDICT_PASS_CLEANED if cleaned else VERDICT_PASS
        return VERDICT_FAIL
    e_num, _ = to_decimal(expected)
    a_num, _ = to_decimal(actual)
    if e_num is not None and a_num is not None:
        return VERDICT_PASS if e_num == a_num else VERDICT_FAIL
    return VERDICT_PASS if _text(expected) == _text(actual) else VERDICT_FAIL


def _blank(v: Any) -> bool:
    return v is None or str(v).strip() == ""


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- 读单元格
def _col_row(ref: str) -> Tuple[int, int]:
    m = re.match(r"^([A-Z]+)(\d+)$", ref.strip().upper())
    if not m:
        raise ValueError(f"非法单元格坐标：{ref}")
    col = 0
    for ch in m.group(1):
        col = col * 26 + (ord(ch) - 64)
    return int(m.group(2)) - 1, col - 1


def read_cells(path: Path, refs: Sequence[str]) -> Dict[str, Any]:
    """refs 形如 '报关单!E12'。支持 .xls（xlrd）与 .xlsx（openpyxl）。"""
    path = Path(path)
    out: Dict[str, Any] = {}
    if path.suffix.lower() == ".xls":
        import xlrd

        wb = xlrd.open_workbook(str(path))
        names = wb.sheet_names()
        for ref in refs:
            sheet, cell = ref.split("!", 1)
            real = find_sheet_name(names, sheet)
            if real is None:
                out[ref] = None
                continue
            ws = wb.sheet_by_name(real)
            r, c = _col_row(cell)
            out[ref] = ws.cell_value(r, c) if r < ws.nrows and c < ws.ncols else None
    else:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), data_only=True)
        for ref in refs:
            sheet, cell = ref.split("!", 1)
            real = find_sheet_name(wb.sheetnames, sheet)
            out[ref] = None if real is None else wb[real][cell.upper()].value
    return out


# ---------------------------------------------------------------- manifest
def load_manifest(path: Path) -> List[Dict[str, str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != MANIFEST_COLUMNS:
            raise ValueError(f"manifest 列不对：{reader.fieldnames}，应为 {MANIFEST_COLUMNS}")
        return list(reader)


def parse_expected_diff(text: str) -> Dict[str, str]:
    """'报关单!G12=原因 ; 合同!A33=原因' → {ref: 原因}。"""
    out: Dict[str, str] = {}
    for part in (text or "").split(";"):
        part = part.strip()
        if not part or "=" not in part:
            continue
        ref, reason = part.split("=", 1)
        out[ref.strip()] = reason.strip()
    return out


# ---------------------------------------------------------------- 比对
@dataclass
class DiffRow:
    case_id: str
    ref: str
    expected: Any
    actual: Any
    verdict: str
    note: str = ""


# 报关单 A 级字段（坐标来自规则清单 §4.5 证据坐标表）
CUSTOMS_FIELDS_A: Dict[str, str] = {
    "报关单!E12": "count",
    "报关单!F12": "weight",
    "报关单!G12": "weight",
    "报关单!B20": "text",
    "报关单!D20": "text",
    "报关单!G20": "count",
    "报关单!H20": "text",
    "报关单!I20": "money",
    "报关单!J20": "money",
    "报关单!H7": "text",
    "报关单!B16": "text",
    "报关单!K12": "text",
    "申报要素!I10": "text",
    "合同!A33": "text",
    "委托书!E12": "text",
}
# B 级只比结构字段（方案 §5.1）
CUSTOMS_FIELDS_B: Dict[str, str] = {
    k: v
    for k, v in CUSTOMS_FIELDS_A.items()
    if k in {"报关单!E12", "报关单!B20", "报关单!D20", "报关单!H20", "报关单!H7", "报关单!B16", "申报要素!I10"}
}


def compare_files(
    case_id: str,
    expected_path: Path,
    actual_path: Path,
    fields: Dict[str, str],
    expected_diff: Dict[str, str],
) -> List[DiffRow]:
    refs = list(fields)
    exp = read_cells(expected_path, refs)
    act = read_cells(actual_path, refs)
    rows: List[DiffRow] = []
    for ref in refs:
        verdict = compare_value(exp[ref], act[ref], kind=fields[ref])
        note = ""
        if verdict == VERDICT_FAIL and ref in expected_diff:
            verdict, note = VERDICT_KNOWN, expected_diff[ref]
        rows.append(DiffRow(case_id, ref, exp[ref], act[ref], verdict, note))
    return rows


def compare_warnings(case_id: str, expected: Sequence[str], actual: Sequence[str]) -> List[DiffRow]:
    e, a = sorted(set(map(str, expected))), sorted(set(map(str, actual)))
    verdict = VERDICT_PASS if e == a else VERDICT_FAIL
    return [DiffRow(case_id, "WARNINGS", e, a, verdict)]


def write_report(rows: Sequence[DiffRow], path: Path) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["case_id", "单元格", "期望值", "实际值", "结论", "说明"])
        for r in rows:
            w.writerow([r.case_id, r.ref, r.expected, r.actual, r.verdict, r.note])


def failures(rows: Sequence[DiffRow]) -> List[DiffRow]:
    return [r for r in rows if r.verdict == VERDICT_FAIL]
