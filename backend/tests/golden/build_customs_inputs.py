"""从发票样例（A 级）和报关样例本身（B 级）生成报关 golden 输入 customs_inputs.json（S3）。

只读 samples/ 下的副本，不碰用户测试目录。重新生成：
    cd backend && python -m tests.golden.build_customs_inputs
A 级：输入 = 发票（数量、申报额、单箱重/票总重、FBA、仓库），产品库字段（品名/HS/单位/品牌/品类）取发票或产品库快照；
      期望 = 历史报关单。数值严格比对，口径差异写在 manifest 的 expected_diff。
B 级：没有发票，输入取报关样例自身的结构字段（品名、HS、单位、件数、数量、单价、FBA/仓库），只比结构字段（方案 §5.1）。
"""
from __future__ import annotations

import csv
import json
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from tests.golden.compare_golden import GOLDEN_DIR, find_sheet_name, load_manifest, read_cells

S = GOLDEN_DIR / "samples"
OUT = GOLDEN_DIR / "customs_inputs.json"
# 产品库快照里的报关字段（发票里的货代 HS 可能不同，报关以产品库为准）
PRODUCT = {
    "宠物垫": {"hs_code": "6307909000", "unit": "包"},
    "车载狗栏": {"hs_code": "7326909000", "unit": "个"},
    "宠物围栏": {"hs_code": "6307909000", "unit": "个"},
}
CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "1": 1, "2": 2, "3": 3, "4": 4, "5": 5}


def _d(v) -> str:
    return str(Decimal(str(v)).normalize()) if v not in (None, "") else ""


def _line(cn, qty, price, brand="JoyBlooms", sku=""):
    p = PRODUCT.get(cn, {})
    return {"sku": sku or cn, "chinese_name": cn, "hs_code": p.get("hs_code", ""), "unit": p.get("unit", ""),
            "brand": brand, "category": "pet", "qty": float(qty), "unit_price": float(price)}


def _csv_warehouse(fba: str) -> str:
    for p in (S / "inputs").glob(f"*{fba}.csv"):
        with open(p, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if row and row[0] == "货件名称":
                    return row[1].rsplit("-", 1)[-1].strip()
    return ""


def send_ticket(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb[find_sheet_name(wb.sheetnames, "Invoice Template")]
    fba = ""
    for row in ws.iter_rows(min_row=1, max_row=14):
        for c in row:
            if isinstance(c.value, str) and re.fullmatch(r"FBA[0-9A-Z]{9}", c.value.strip()):
                fba = c.value.strip()
    boxes, lines = [], []
    for r in range(15, ws.max_row + 1):
        rng = ws[f"A{r}"].value
        if not rng:
            continue
        a, b = (int(x) for x in str(rng).split("-"))
        boxes += [float(ws[f"C{r}"].value)] * (b - a + 1)
        lines.append(_line(str(ws[f"H{r}"].value).strip(), ws[f"F{r}"].value, ws[f"R{r}"].value, str(ws[f"I{r}"].value), str(ws[f"E{r}"].value)))
    return {"shipment_id": fba, "warehouse": _csv_warehouse(fba), "boxes": boxes, "lines": lines,
            "provenance": f"{path.name}!Invoice Template A15:R{ws.max_row}（箱号区间×单箱重、F 数量、R 售价）；仓库取 inputs/*{fba}.csv 货件名称"}


def jiufang_ticket(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    inv, pl = wb["Invoice"], wb["Packing list"]
    head_row = next(r for r in range(1, 40) if str(inv.cell(r, 1).value or "").startswith("*Shipment"))
    heads = [str(c.value or "") for c in inv[head_row]]
    col = lambda kw: next(i for i, h in enumerate(heads) if kw in h)  # noqa: E731
    c_sku, c_cn, c_brand, c_qty = col("SKU"), col("中文品名"), col("Brand"), col("Unit\n个数")
    c_price = next(i for i, h in enumerate(heads) if "Price" in h or "单价" in h)
    lines, fba = [], ""
    for r in range(head_row + 1, inv.max_row + 1):
        v = [c.value for c in inv[r]]
        if str(v[0] or "").startswith("*Total") or not v[0]:
            break
        fba = str(v[0]).strip()
        lines.append(_line(str(v[c_cn]).strip(), v[c_qty], v[c_price], str(v[c_brand]), str(v[c_sku])))
    wh = ""
    for r in range(1, head_row):
        if str(inv.cell(r, 1).value or "").startswith("FBA CODE"):
            wh = next(str(c.value).strip() for c in inv[r][1:] if c.value)
    ph = next(r for r in range(1, 20) if str(pl.cell(r, 1).value or "").startswith("*Shipment"))
    pheads = [str(c.value or "") for c in pl[ph]]
    c_ctn = next(i for i, h in enumerate(pheads) if "Carton" in h)
    c_gw = next(i for i, h in enumerate(pheads) if "G.W" in h)
    boxes = []
    for r in range(ph + 1, pl.max_row + 1):
        v = [c.value for c in pl[r]]
        if str(v[0] or "").startswith("*Total") or not v[0]:
            break
        boxes += [float(v[c_gw])] * int(v[c_ctn])
    return {"shipment_id": fba, "warehouse": wh, "boxes": boxes, "lines": lines,
            "provenance": f"{path.name}!Invoice（数量/单价/FBA CODE）+ Packing list（箱数×单箱毛重 G.W）"}


def _wh_from_name(path: Path) -> str:
    """第5/6批：票的仓库以文件名为准（第6批第5票发票 E3 误抄成 ACY2，文件名、货件 CSV、运费表、报关单都是 MGE3，见 D-2026-10-16-09 第4点）。"""
    m = re.search(r"海运([A-Z0-9]{4})-", path.name)
    return m.group(1) if m else ""


def batch5_ticket(path: Path) -> dict:
    ws = load_workbook(path, data_only=True)["Sheet1"]
    return {"shipment_id": str(ws["B3"].value).strip(), "warehouse": _wh_from_name(path) or str(ws["B7"].value).strip(),
            "box_count": int(ws["J10"].value), "gross_kg": float(ws["N10"].value),
            "lines": [_line(str(ws["B10"].value).strip(), ws["H10"].value, ws["I10"].value, str(ws["G10"].value))],
            "provenance": f"{path.name}!Sheet1 B3/仓库取文件名/H10 个数/I10 单个货值/J10 件数/N10 总实重"}


def batch6_ticket(path: Path) -> dict:
    ws = load_workbook(path, data_only=True)["发票"]
    boxes = int(ws["B6"].value)
    per_box = int(ws["F14"].value)
    unit_w = Decimal(str(ws["M14"].value))
    return {"shipment_id": str(ws["B3"].value).strip(), "warehouse": _wh_from_name(path) or str(ws["E3"].value).strip(),
            "boxes": [float(unit_w * per_box)] * boxes,
            "lines": [_line(str(ws["D14"].value).strip(), boxes * per_box, ws["G14"].value, str(ws["J14"].value))],
            "provenance": f"{path.name}!发票 B3/E3/B6 预报件数×F14 单箱内数量、G14 申报单价、M14 单件实重×单箱数量=单箱重；HS 取产品库（发票 L14 为货代 HS）；仓库取文件名"}


def _ticket_order(p: Path) -> int:
    m = re.search(r"第(.)票", p.name)
    return CN_NUM.get(m.group(1), 9) if m else 9


A_SOURCES = {
    "JUS11DD25W": ("jiufang", ["invoice/LCL2607GZ0PA-箱单发票.xlsx"]),
    "JUS13DD21": ("send", ["invoice/FBA19P4NX8HH发票.xlsx"]),
    "JUS13DD32CLT": ("send", ["invoice/FBA19P4P4CL2发票.xlsx"]),
    "JUS13DD32LAX": ("send", ["invoice/FBA19P4P4MBQ发票.xlsx"]),
    "JUS13DD37": ("send", ["invoice/FBA19P4N6MM2发票.xlsx"]),
    "JUS13DD38": ("send", ["invoice/FBA19P4P3PZC发票.xlsx"]),
    "JUS14DD12": ("jiufang", ["invoice/LCL2609TNKDA-箱单发票.xlsx"]),
    "JUS14DD28": ("jiufang", ["invoice/LCL26099BD6A-箱单发票.xlsx"]),
    "JUS15DD5SBD": ("jiufang", ["invoice/LCL2610S97BA-箱单发票.xlsx"]),
    "JUS15DD6": ("jiufang", ["invoice/LCL2610DHTZA-箱单发票.xlsx"]),
    "JUS15DD6TEB": ("jiufang", ["invoice/LCL2610QGXTA-箱单发票.xlsx"]),
    "JUS15DD8": ("jiufang", ["invoice/LCL2610UKEKA-箱单发票.xlsx"]),
    "JUS5HB171": ("batch5", sorted((p.relative_to(S).as_posix() for p in (S / "invoice_legacy").glob("*第五批*")), key=lambda x: _ticket_order(Path(x)))),
    "JUS6HB30": ("batch6", sorted((p.relative_to(S).as_posix() for p in (S / "invoice_legacy").glob("*第6批*")), key=lambda x: _ticket_order(Path(x)))),
}
READERS = {"send": send_ticket, "jiufang": jiufang_ticket, "batch5": batch5_ticket, "batch6": batch6_ticket}


def b_case(expected: Path) -> list:
    c = read_cells(expected, ["报关单!B16", "报关单!E12", "报关单!B20", "报关单!D20", "报关单!G20", "报关单!H20", "报关单!I20", "报关单!F12"])
    rem = str(c["报关单!B16"] or "")
    whs = re.search(r"仓库编码：([^。]*)", rem)
    fbas = re.search(r"FBA编码：([^。]*)", rem)
    whs = [w.strip() for w in re.split(r"[，,]", whs.group(1))] if whs else []
    fbas = [f.strip() for f in re.split(r"[，,]", fbas.group(1))] if fbas else []
    pieces = int(Decimal(re.sub(r"[^\d.]", "", str(c["报关单!E12"])) or "0"))
    hs = str(c["报关单!B20"]).split(".")[0]
    line = {"sku": str(c["报关单!D20"]), "chinese_name": str(c["报关单!D20"]), "hs_code": hs, "unit": str(c["报关单!H20"]),
            "brand": "JoyBlooms", "category": "pet", "qty": float(c["报关单!G20"]), "unit_price": float(c["报关单!I20"] or 0)}
    n = max(len(fbas), 1)
    tickets = []
    for i in range(n):
        tickets.append({"shipment_id": fbas[i] if i < len(fbas) else "", "warehouse": whs[i] if i < len(whs) else "",
                        "box_count": pieces if i == 0 else 0, "gross_kg": float(c["报关单!F12"] or 0) if i == 0 else 0,
                        "lines": [line] if i == 0 else []})
    for t in tickets:
        t["provenance"] = f"{expected.name}（B 级：结构字段取样例本身，无发票）"
    return tickets


def build() -> dict:
    out = {}
    for m in load_manifest(GOLDEN_DIR / "customs_manifest.csv"):
        cid = m["case_id"]
        order = cid
        if m["level"] == "A":
            kind, paths = A_SOURCES[cid]
            tickets = [READERS[kind](S / p) for p in paths]
            out[cid] = {"order_no": order, "level": "A", "inputs": ["samples/" + p for p in paths], "tickets": tickets}
        else:
            out[cid] = {"order_no": order, "level": "B", "inputs": [m["expected_path"]], "tickets": b_case(GOLDEN_DIR / m["expected_path"])}
    return out


if __name__ == "__main__":
    data = build()
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"写入 {OUT}：{len(data)} 份")
