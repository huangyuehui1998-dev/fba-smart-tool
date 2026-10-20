"""发票导出 golden：以货代原始样例文件为期望值，逐格比对导出结果（S4-fix1）。

- SEND：表头 C4/C5/C6 + 第 15 行起每个区间行 A–T 全列；多出/缺少的行都算失败；
- 九方 Invoice：表头（渠道、单号、日期、发货人、派送地址、FBA CODE…）+ 每个 SKU 行 A–V + 合计行；
- 九方 Packing list：表头 + 每箱一行 A–N + 合计行；
- 九方系统的 Invoice 行顺序无法从输入推出（不是 CSV 顺序也不是字母序），所以 Invoice 数据行按 SKU 对齐后再逐格比；
  Packing list 按箱号对齐。行数不同直接算失败。
图片（T 列 Picture、SEND 产品图）不比。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Tuple

import openpyxl

from tests.golden.compare_golden import (
    VERDICT_FAIL,
    VERDICT_KNOWN,
    DiffRow,
    compare_value,
    find_sheet_name,
)

SEND_KINDS = {"A": "text", "B": "dims", "C": "weight", "D": "text", "E": "text", "F": "count", "G": "text", "H": "text",
              "I": "text", "J": "text", "K": "text", "L": "text", "M": "text", "N": "text", "O": "text", "P": "count",
              "Q": "money", "R": "money", "S": "money", "T": "text"}
JF_INV_KINDS = {c: "text" for c in "ABCDEFGHIJKLMNPUV"} | {"O": "count", "Q": "count", "R": "money", "S": "money"}
JF_INV_HEAD = ["B2", "P2", "P3", "P4", "P5", "V4", "B4", "B5", "B7", "B8", "B12", "B13", "B15", "B16", "B17", "B18", "B19", "B22", "P20"]
JF_INV_TOTAL = {"M": "text", "O": "count", "Q": "text", "R": "text", "S": "money", "T": "text", "U": "count"}
JF_PK_KINDS = {c: "text" for c in "ABCDEFN"} | {"G": "count", "H": "count", "I": "weight", "J": "weight", "K": "weight", "L": "weight", "M": "weight"}
JF_PK_TOTAL = {"F": "text", "G": "count", "L": "weight", "M": "weight"}


def _ws(path: Path, name: str):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    real = find_sheet_name(wb.sheetnames, name)
    if real is None:
        raise AssertionError(f"{path.name} 没有工作表 {name!r}（有 {wb.sheetnames}）")
    return wb[real]


def _rows(ws, first: int, key_col: str, stop_col: str = "", stop_text: str = "") -> Dict[int, Any]:
    out = {}
    r = first
    while r <= ws.max_row:
        if stop_col and str(ws[f"{stop_col}{r}"].value or "").startswith(stop_text):
            break
        if ws[f"{key_col}{r}"].value in (None, ""):
            break
        out[r] = ws[f"{key_col}{r}"].value
        r += 1
    return out


def _total_row(ws, col: str, text: str, start: int) -> int:
    for r in range(start, ws.max_row + 1):
        if str(ws[f"{col}{r}"].value or "").startswith(text):
            return r
    return -1


def _dims(v):
    try:
        return [float(x) for x in str(v or "").split("*")]
    except ValueError:
        return None


def _cmp(case, sheet, exp_ws, act_ws, er, ar, kinds, diff, out, tag=""):
    """已知差异只认逐格登记（ref 精确匹配），不支持通配符。
    尺寸列（kind=dims）另做正向核验：只有“每一边四舍五入到整厘米后与样例一致”且逐格登记过，才算已知差异，否则失败。"""
    for col, kind in kinds.items():
        ref = f"{sheet}!{col}{er}"
        e, a = exp_ws[f"{col}{er}"].value, act_ws[f"{col}{ar}"].value
        note = f"实际行 {ar}{tag}" if ar != er else tag
        if kind == "dims":
            de, da = _dims(e), _dims(a)
            if de is not None and de == da:
                v = "通过"
            elif de and da and len(de) == len(da) == 3 and all(round(x) == round(y) for x, y in zip(de, da)) and ref in diff:
                v, note = VERDICT_KNOWN, f"{diff[ref]}（正向核验：样例 {e} 四舍五入 = {'*'.join(str(round(x)) for x in de)} = 导出 {a}）"
            else:
                v = VERDICT_FAIL
            out.append(DiffRow(case, ref, e, a, v, note))
            continue
        v = compare_value(e, a, kind=kind)
        if v == VERDICT_FAIL and ref in diff:
            v, note = VERDICT_KNOWN, diff[ref]
        out.append(DiffRow(case, ref, e, a, v, note))


def compare_send(case: str, expected: Path, actual: Path, diff: Dict[str, str]) -> List[DiffRow]:
    sheet = "Invoice\xa0Template"
    e, a = _ws(expected, sheet), _ws(actual, sheet)
    out: List[DiffRow] = []
    out.append(DiffRow(case, "SHEETNAME", "Invoice\\xa0Template", repr(a.title), "通过" if a.title == sheet else VERDICT_FAIL))
    for ref in ("C4", "C5", "C6"):
        out.append(DiffRow(case, f"{sheet}!{ref}", e[ref].value, a[ref].value, compare_value(e[ref].value, a[ref].value)))
    er, ar = _rows(e, 15, "E"), _rows(a, 15, "E")
    out.append(DiffRow(case, "ROWCOUNT", len(er), len(ar), "通过" if len(er) == len(ar) else VERDICT_FAIL))
    for (r1, _), (r2, _) in zip(er.items(), ar.items()):
        _cmp(case, sheet, e, a, r1, r2, SEND_KINDS, diff, out)
    return out


def compare_jiufang(case: str, expected: Path, actual: Path, diff: Dict[str, str]) -> List[DiffRow]:
    out: List[DiffRow] = []
    ei, ai = _ws(expected, "Invoice"), _ws(actual, "Invoice")
    for ref in JF_INV_HEAD:
        v = compare_value(ei[ref].value, ai[ref].value)
        note = ""
        if v == VERDICT_FAIL and f"Invoice!{ref}" in diff:
            v, note = VERDICT_KNOWN, diff[f"Invoice!{ref}"]
        out.append(DiffRow(case, f"Invoice!{ref}", ei[ref].value, ai[ref].value, v, note))
    er = _rows(ei, 25, "C", "M", "*Total")
    ar = _rows(ai, 25, "C", "M", "*Total")
    out.append(DiffRow(case, "Invoice!ROWCOUNT", len(er), len(ar), "通过" if len(er) == len(ar) else VERDICT_FAIL))
    # T 列图片占位符只在第一行：按位置比，不按 SKU 对齐
    for r1, r2 in zip(er, ar):
        _cmp(case, "Invoice", ei, ai, r1, r2, {"T": "text"}, diff, out, tag="（按位置）")
    a_by_sku = {v: r for r, v in ar.items()}
    for r1, sku in er.items():
        if sku not in a_by_sku:
            out.append(DiffRow(case, f"Invoice!C{r1}", sku, None, VERDICT_FAIL, "导出里没有这个 SKU"))
            continue
        _cmp(case, "Invoice", ei, ai, r1, a_by_sku[sku], JF_INV_KINDS, diff, out, tag="（按 SKU 对齐）")
    et, at = _total_row(ei, "M", "*Total", 25), _total_row(ai, "M", "*Total", 25)
    _cmp(case, "Invoice", ei, ai, et, at, JF_INV_TOTAL, diff, out, tag="（合计行）")

    ep, ap = _ws(expected, "Packing list"), _ws(actual, "Packing list")
    for ref in ("C2", "C3"):
        out.append(DiffRow(case, f"Packing list!{ref}", ep[ref].value, ap[ref].value, compare_value(ep[ref].value, ap[ref].value)))
    er = _rows(ep, 6, "F", "F", "*Total")
    ar = _rows(ap, 6, "F", "F", "*Total")
    out.append(DiffRow(case, "Packing list!ROWCOUNT", len(er), len(ar), "通过" if len(er) == len(ar) else VERDICT_FAIL))
    a_by_box = {v: r for r, v in ar.items()}
    for r1, box in er.items():
        if box not in a_by_box:
            out.append(DiffRow(case, f"Packing list!F{r1}", box, None, VERDICT_FAIL, "导出里没有这个箱"))
            continue
        _cmp(case, "Packing list", ep, ap, r1, a_by_box[box], JF_PK_KINDS, diff, out, tag="（按箱号对齐）")
    et, at = _total_row(ep, "F", "*Total", 6), _total_row(ap, "F", "*Total", 6)
    _cmp(case, "Packing list", ep, ap, et, at, JF_PK_TOTAL, diff, out, tag="（合计行）")
    return out
