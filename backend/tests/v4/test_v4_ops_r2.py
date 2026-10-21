"""运营 r2：Step 4.5 与货代适配器逐格一致，手填售价默认不覆盖，回滚保留 None，箱规保留 44.96。"""
from __future__ import annotations

import base64
import io
import json
from datetime import datetime

import openpyxl
import pytest
from openpyxl.cell.cell import MergedCell

FORWARDERS = ("send", "jiufang", "qiluo", "qiluo_t5", "qiluo_t6", "deshu")
ADAPTERS = ("send", "jiufang", "qiluo_t5", "qiluo_t6", "deshu")


def _csv() -> bytes:
    return "\n".join([
        '"货件编号","FBASTEP4501"',
        '"货件名称","STEP45-TST1"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","1"',
        "",
        '"单件商品（1 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        '"STEP45-SKU","rod","B0STEP4501","X00TEST45A","新品","无","4","4"',
        "",
        '"","","","","","","箱号","FBASTEP4501U000001"',
        '"","","","","","","包装箱重量（磅）：","10"',
        '"","","","","","","包装箱长度（英寸）：","17.70"',
        '"","","","","","","包装箱宽度（英寸）：","17.70"',
        '"","","","","","","包装箱高度（英寸）：","16.10"',
    ]).encode("utf-8")


def _seed():
    from datetime import timezone

    from fba_tool.master_db import _connect, init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields

    init_master_db()
    now = datetime.now(timezone.utc).isoformat()
    with _connect() as conn:
        conn.execute("DELETE FROM fba_warehouse WHERE code = ?", ("TST1",))
        conn.execute(
            """
            INSERT INTO fba_warehouse
                (code, name, company, address1, city, state, zip, country, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("TST1", "TST1", "Amazon", "1 Test Road", "STEP45VILLE", "ZZ", "00000", "US", now),
        )
        conn.commit()
    upsert_product(
        {
            "sku": "STEP45-SKU",
            "fnsku": "X00TEST45A",
            "asin": "B0STEP4501",
            "english_name": "Fishing rod",
            "chinese_name": "鱼竿",
            "brand": "Acme",
            "material": "棉",
            "purpose": "宠物",
            "hs_code": "9507100090",
            "model_name": "型号甲",
            "extra_data": json.dumps(
                {"invoice": {"unit": "个", "purchase_price": 5, "material_en": "cotton", "purpose_en": "pet", "pcs_per_set": 1}},
                ensure_ascii=False,
            ),
        }
    )
    set_product_fields(
        "STEP45-SKU",
        {"listing_price": 12.5, "listing_price_currency": "USD"},
        SOURCE_MANUAL,
    )


def _values(data: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    out = {}
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    value = cell.value
                    if isinstance(value, datetime):
                        value = value.isoformat()
                    out[(ws.title, cell.coordinate)] = value
        return out
    finally:
        wb.close()


def test_to_cm_keeps_carton_hundredths():
    from fba_tool.volume_weight import to_cm

    assert to_cm(44.96, "cm") == 44.96
    assert to_cm(17.70, "in") == 44.96
    assert to_cm(17.70, "in") != 45.0


def test_manual_listing_price_defaults_unchecked(monkeypatch, tmp_path):
    from fba_tool.lingxing_preview import build_listing_preview, serialize_preview
    from fba_tool.master_db import init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields

    ad = tmp_path / "app_data"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    init_master_db()
    upsert_product({"sku": "PRICE-HAND", "fnsku": "X00TESTPH1", "asin": "B0PRICEH01"})
    set_product_fields("PRICE-HAND", {"listing_price": 52, "listing_price_currency": "USD"}, SOURCE_MANUAL)
    run_id = 545701
    build_listing_preview(
        run_id,
        [{
            "sku": "PRICE-HAND",
            "asin": "B0PRICEH01",
            "fnsku": "X00TESTPH1",
            "listing_price": 49.99,
            "listing_price_currency": "USD",
            "listing_price_basis": "listing_price",
        }],
    )
    payload = serialize_preview(run_id)
    rows = [row for group in payload["groups"].values() for row in group if row["sku"] == "PRICE-HAND"]
    assert rows
    price = next(field for field in rows[0]["fields"] if field["field"] == "listing_price")
    assert price["default_checked"] is False
    assert price["manual_diff"] is True
    assert "52" in price["difference"] and "49.99" in price["difference"]


def test_rollback_restores_null_not_empty_string():
    from fba_tool.lingxing_rollback import record_import_changes, rollback_lingxing_run
    from fba_tool.master_db import delete_product_by_sku, get_product_by_sku, upsert_product
    from fba_tool.master_db_lingxing import create_lingxing_sync_run, finish_lingxing_sync_run

    sku = "RB-NONE-STEP45"
    delete_product_by_sku(sku)
    upsert_product({"sku": sku})
    before = get_product_by_sku(sku)
    assert before["asin"] is None
    assert before["chinese_name"] is None
    run_id = create_lingxing_sync_run()
    finish_lingxing_sync_run(run_id, "confirmed", [], {})
    changed = dict(before)
    changed["asin"] = "B0RBNONE01"
    changed["chinese_name"] = "临时品名"
    upsert_product(changed, overwrite=True)
    record_import_changes(run_id, sku, before, get_product_by_sku(sku))
    try:
        rollback_lingxing_run(run_id)
        restored = get_product_by_sku(sku)
        assert restored["asin"] is None
        assert restored["chinese_name"] is None
        assert restored["asin"] != ""
        assert restored["chinese_name"] != ""
    finally:
        delete_product_by_sku(sku)


@pytest.mark.parametrize("forwarder_id", FORWARDERS)
def test_step45_matches_adapter_cells(forwarder_id):
    from fastapi.testclient import TestClient

    from fba_tool.invoice_export import export_invoice
    from fba_tool.invoice_step45 import render_adapter_invoice
    from fba_tool.master_db import delete_product_by_sku
    from fba_tool.shipment_parser import parse_csv_shipment
    from main import app

    _seed()
    csv_bytes = _csv()
    ship = parse_csv_shipment(csv_bytes, file_name="step45.csv", allow_unassigned_fba=True)
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": ship.shipment_id,
                "forwarder_id": forwarder_id,
                "csv_base64": base64.b64encode(csv_bytes).decode(),
                "po_number": "PO-45",
            },
        )
        assert generated.status_code == 200, generated.text
        body = generated.json()
        cells = {item["cell"]: item.get("value") for item in body["cells"]}
        if forwarder_id == "send":
            assert cells.keys().isdisjoint({"B4", "B5", "B6", "E5"})
        if forwarder_id == "qiluo_t5":
            assert "D6" not in cells and "H3" not in cells
        if forwarder_id == "qiluo":
            assert cells.get("K2") == "型号甲"
            assert cells.get("H2") == "棉"
            assert cells.get("I2") == "宠物"
            assert "/" not in str(cells.get("H2") or "")
        exported = client.post(
            "/api/export_invoice",
            json={
                "template_hash": body["template_hash"],
                "forwarder_id": forwarder_id,
                "shipment_id": body["shipment_id"],
                "warehouse_code": body["warehouse_code"],
                "template_sheet": body["template_sheet"],
                "fields": body["cells"],
            },
        )
        assert exported.status_code == 200, exported.text
        from fba_tool.forwarder_registry import get_template_bytes
        from fba_tool.univer_snapshot import workbook_to_snapshot
        from openpyxl.utils.cell import coordinate_to_tuple

        snap = workbook_to_snapshot(get_template_bytes(forwarder_id))["snapshot"]
        sheet = snap["sheets"][snap["sheetOrder"][0]]
        for coord, value in cells.items():
            if value in (None, ""):
                continue
            row, col = coordinate_to_tuple(str(coord))
            sheet["cellData"].setdefault(str(row - 1), {})[str(col - 1)] = {"v": value}
        with_snapshot = client.post(
            "/api/export_invoice",
            json={
                "template_hash": body["template_hash"],
                "forwarder_id": forwarder_id,
                "shipment_id": body["shipment_id"],
                "warehouse_code": body["warehouse_code"],
                "template_sheet": body["template_sheet"],
                "fields": body["cells"],
                "snapshot": snap,
            },
        )
        assert with_snapshot.status_code == 200, with_snapshot.text
        if forwarder_id in ADAPTERS:
            from fba_tool.master_db import lookup_by_fnsku
            from fba_tool.warehouse_db import get_warehouse
            from routers.invoice_export import resolve_export_warehouse

            product = lookup_by_fnsku("X00TEST45A")
            code, _source = resolve_export_warehouse(forwarder_id, "", ship.ship_to, ship.warehouse_code)
            direct = export_invoice(
                forwarder_id,
                shipment=ship,
                products={"X00TEST45A": product},
                reference_id="",
                warehouse=code,
                warehouse_info=get_warehouse(code) if code else None,
                invoice_date=None,
            )
            same_path = render_adapter_invoice(forwarder_id, csv_bytes, None)
            assert _values(exported.content) == _values(direct["bytes"])
            assert _values(exported.content) == _values(same_path["bytes"])
            assert _values(with_snapshot.content) == _values(direct["bytes"])
        else:
            for payload in (exported.content, with_snapshot.content):
                wb = openpyxl.load_workbook(io.BytesIO(payload))
                assert wb.active["K2"].value == "型号甲"
                assert wb.active["H2"].value == "棉"
                assert wb.active["I2"].value == "宠物"
                wb.close()
    finally:
        delete_product_by_sku("STEP45-SKU")
