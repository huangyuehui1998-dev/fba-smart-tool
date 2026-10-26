"""待分配刷箱 / FIST：send、jiufang、qiluo 的发票 200，以及空价 / 缺价 409。

合成 CSV 在 tests/fixtures/synthetic/rebox/。不是 realdata，也不把这次覆盖写成刷箱交付。
命名 FBA 的批量扫描仍跳过刷箱、FIST、golden 和文件名里的待分配。
"""
from __future__ import annotations

import base64
import io
import json
import os
from pathlib import Path
from urllib.parse import unquote

import pytest

FORWARDERS = ("send", "jiufang", "qiluo")
WAREHOUSE = "ONL1"
FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic" / "rebox" / "unassigned_send_min.csv"
SKUS = (
    ("FIST48-A", "X00TEST48A", "B0SYNTH48A"),
    ("REBOX80-B", "X00TEST80B", "B0SYNTH80B"),
)
_PRICE_KEYS = (
    "listing_price",
    "declared_value",
    "regular_price",
    "price",
    "landed_price",
    "last_nonpromo_price",
)


def _raw() -> bytes:
    data = FIXTURE.read_bytes()
    assert FIXTURE.name.casefold().find("listing") < 0
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    assert "待分配" in text
    assert "FIST48-A" in text and "REBOX80-B" in text
    return data


def _shipment(raw: bytes):
    from fba_tool.shipment_parser import parse_csv_shipment

    ship = parse_csv_shipment(raw, file_name=FIXTURE.name, allow_unassigned_fba=True)
    assert ship.shipment_id == "待分配"
    totals = {str(sku): int(qty) for sku, qty in ship.box_items.groupby("SKU")["数量"].sum().items()}
    assert totals == {"FIST48-A": 48, "REBOX80-B": 80}
    lengths = {round(float(value), 2) for value in ship.box_specs["长(cm)"].tolist()}
    assert len(lengths) > 1
    assert not any(str(value or "").strip() for value in ship.box_fba_ids.values())
    return ship


def _prepare(monkeypatch, tmp_path, *, listing_price: float | None):
    import json as _json
    from datetime import datetime, timezone

    from fba_tool.master_db import _connect, init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields

    monkeypatch.setenv("FBA_APP_DATA_DIR", str(tmp_path / "app"))
    monkeypatch.setenv("FBA_DISABLE_SCHEDULER", "1")
    init_master_db()
    now = datetime.now(timezone.utc).isoformat()
    with _connect() as conn:
        conn.execute("DELETE FROM fba_warehouse WHERE code = ?", (WAREHOUSE,))
        conn.execute(
            """
            INSERT INTO fba_warehouse
                (code, name, company, address1, city, state, zip, country, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (WAREHOUSE, WAREHOUSE, "Amazon", "1 Only Road", "ONLYVILLE", "ZZ", "00001", "US", now),
        )
        conn.commit()
    for sku, fnsku, asin in SKUS:
        upsert_product(
            {
                "sku": sku,
                "fnsku": fnsku,
                "asin": asin,
                "english_name": "Synthetic rod",
                "chinese_name": "合成鱼竿",
                "brand": "Acme",
                "material": "棉",
                "purpose": "宠物",
                "hs_code": "9507100090",
                "model_name": "型号甲",
                "extra_data": _json.dumps(
                    {
                        "invoice": {
                            "unit": "个",
                            "purchase_price": 5,
                            "material_en": "cotton",
                            "purpose_en": "pet",
                            "pcs_per_set": 1,
                        }
                    },
                    ensure_ascii=False,
                ),
            }
        )
        if listing_price is not None:
            set_product_fields(
                sku,
                {"listing_price": listing_price, "listing_price_currency": "USD"},
                SOURCE_MANUAL,
            )


def _cleanup():
    from fba_tool.master_db import delete_product_by_sku

    for sku, _fnsku, _asin in SKUS:
        delete_product_by_sku(sku)


def _client():
    from fastapi.testclient import TestClient

    from main import create_app

    return TestClient(create_app())


def _generate(client, raw: bytes, forwarder_id: str):
    return client.post(
        "/api/invoice/generate",
        json={
            "shipment_id": "待分配",
            "forwarder_id": forwarder_id,
            "warehouse": WAREHOUSE,
            "csv_base64": base64.b64encode(raw).decode(),
        },
    )


def _export_generated(client, forwarder_id: str, body: dict):
    return client.post(
        "/api/export_invoice",
        json={
            "template_hash": body["template_hash"],
            "forwarder_id": forwarder_id,
            "shipment_id": body["shipment_id"],
            "warehouse_code": body.get("warehouse_code") or WAREHOUSE,
            "template_sheet": body["template_sheet"],
            "fields": body["cells"],
        },
    )


def _adapter_export(client, raw: bytes, forwarder_id: str):
    data = {"forwarder_id": forwarder_id, "reference_id": "SYNREF", "warehouse": WAREHOUSE}
    if forwarder_id == "jiufang":
        data["order_no"] = "SYNORDER"
    return client.post(
        "/api/invoice-export/export",
        data=data,
        files={"csv_file": (FIXTURE.name, raw, "text/csv")},
    )


def _assert_export_bytes(content: bytes, where: str) -> None:
    assert content, where
    assert content != b"hello", where
    assert not content.startswith(b"hello"), where
    assert content[:2] == b"PK", where


def _blocking(response) -> list:
    payload = response.json()
    if isinstance(payload.get("blocking"), list):
        return payload["blocking"]
    detail = payload.get("detail")
    if isinstance(detail, dict) and isinstance(detail.get("blocking"), list):
        return detail["blocking"]
    return []


def _assert_price_missing(response, where: str) -> None:
    if response.status_code == 200:
        raise AssertionError(f"{where} returned 200; missing price must not export as 0.00. {response.text[:400]}")
    assert response.status_code == 409, f"{where} {response.status_code} {response.text[:500]}"
    blocking = _blocking(response)
    assert any(item.get("code") == "price_missing" for item in blocking), blocking
    text = json.dumps(blocking, ensure_ascii=False)
    assert "不能导出" in text
    for sku, fnsku, _asin in SKUS:
        assert sku in text, text
        assert fnsku in text, text
    assert "0.00" not in text


def _warning_codes(response) -> list:
    raw = response.headers.get("x-invoice-warnings") or ""
    if not raw:
        return []
    return json.loads(unquote(raw))


def _qty_column(data: bytes, header_token: str) -> list:
    import openpyxl

    from fba_tool.invoice_templates import norm

    book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    try:
        for sheet in book.worksheets:
            rows = list(sheet.iter_rows(
                max_row=min(sheet.max_row or 1, 80),
                max_col=min(sheet.max_column or 1, 30),
                values_only=True,
            ))
            header_at = None
            column = None
            for index, row in enumerate(rows[:40]):
                for col, value in enumerate(row):
                    if isinstance(value, str) and header_token in norm(value):
                        header_at, column = index, col
                        break
                if header_at is not None:
                    break
            if header_at is None:
                continue
            found = []
            blank = 0
            for row in rows[header_at + 1:]:
                if not any(cell not in (None, "") for cell in row):
                    blank += 1
                    if blank >= 3:
                        break
                    continue
                blank = 0
                if column < len(row) and row[column] not in (None, ""):
                    found.append(row[column])
            if found:
                return found
    finally:
        book.close()
    return []


def _whole_numbers(values: list) -> list:
    out = []
    for value in values:
        if isinstance(value, bool):
            out.append(value)
            continue
        if isinstance(value, (int, float)) and float(value).is_integer():
            out.append(int(value))
            continue
        out.append(value)
    return out


def _blank_prices(skus: tuple[tuple[str, str, str], ...] = SKUS) -> None:
    """listing_price 和 declared_value 两列都在，值是空的。"""
    from fba_tool.master_db import _connect, get_product_by_sku

    for sku, _fnsku, _asin in skus:
        with _connect() as conn:
            conn.execute(
                "UPDATE product_master SET listing_price = '', declared_value = '' WHERE sku = ?",
                (sku,),
            )
            conn.commit()
        stored = get_product_by_sku(sku) or {}
        assert "listing_price" in stored and "declared_value" in stored
        assert stored["listing_price"] in (None, "", 0, 0.0)
        assert stored["declared_value"] in (None, "", 0, 0.0)


def _drop_price_fields(product):
    if not product:
        return product
    out = dict(product)
    for key in _PRICE_KEYS:
        out.pop(key, None)
    sources = dict(out.get("field_sources") or {})
    for key in _PRICE_KEYS:
        sources.pop(key, None)
    out["field_sources"] = sources
    return out


def test_unassigned_rebox_send_jiufang_qiluo_generate_export_200(monkeypatch, tmp_path):
    """待分配多箱：三家生成和导出都是 200。琪珞写出 SKU 合计，不写一串单箱 4。

    九方先种唯一仓库，避免 warehouse_undecided 的 422 被当成这条用例失败。
    待分配箱唛的 box_mark_missing 只是警告。
    """
    raw = _raw()
    _shipment(raw)
    _prepare(monkeypatch, tmp_path, listing_price=10)
    client = _client()
    try:
        for forwarder_id in FORWARDERS:
            generated = _generate(client, raw, forwarder_id)
            assert generated.status_code == 200, f"{forwarder_id} generate {generated.status_code} {generated.text[:500]}"
            body = generated.json()
            assert body["shipment_id"] == "待分配"
            if forwarder_id == "jiufang":
                warning_text = json.dumps(body.get("warnings") or [], ensure_ascii=False)
                assert "warehouse_undecided" not in warning_text
            exported = _export_generated(client, forwarder_id, body)
            assert exported.status_code == 200, f"{forwarder_id} export {exported.status_code} {exported.text[:400]}"
            _assert_export_bytes(exported.content, f"{forwarder_id} /api/export_invoice")
            adapter = _adapter_export(client, raw, forwarder_id)
            if forwarder_id == "qiluo":
                numbers = _whole_numbers(_qty_column(exported.content, "单箱内数量"))
                assert numbers == [48, 80], numbers
                assert numbers != [4] * 32
                assert 4 not in numbers
                # 货代 id qiluo 走琪珞模板 ETL，合计在 /api/export_invoice。
                # 它没有 export_adapter；若 /api/invoice-export/export 仍返回 200，也必须是同一组合计。
                if adapter.status_code == 200:
                    _assert_export_bytes(adapter.content, "qiluo /api/invoice-export/export")
                    adapter_numbers = _whole_numbers(_qty_column(adapter.content, "单箱内数量"))
                    assert adapter_numbers == [48, 80], adapter_numbers
            else:
                assert adapter.status_code == 200, f"{forwarder_id} invoice-export {adapter.status_code} {adapter.text[:400]}"
                _assert_export_bytes(adapter.content, f"{forwarder_id} /api/invoice-export/export")
                if forwarder_id == "jiufang":
                    codes = _warning_codes(adapter)
                    assert "warehouse_undecided" not in codes
                    assert "box_mark_missing" in codes
    finally:
        _cleanup()


def test_unassigned_rebox_empty_price_is_409(monkeypatch, tmp_path):
    """listing_price 和 declared_value 都是空值。三家生成 409，blocking 含 price_missing。"""
    raw = _raw()
    _prepare(monkeypatch, tmp_path, listing_price=None)
    _blank_prices()
    client = _client()
    try:
        for forwarder_id in FORWARDERS:
            generated = _generate(client, raw, forwarder_id)
            _assert_price_missing(generated, f"{forwarder_id} generate empty price")
            if forwarder_id == "qiluo":
                continue
            exported = _adapter_export(client, raw, forwarder_id)
            _assert_price_missing(exported, f"{forwarder_id} invoice-export empty price")
    finally:
        _cleanup()


def test_unassigned_rebox_missing_price_field_is_409(monkeypatch, tmp_path):
    """产品记录上没有申报价 / 售价字段时必须 409，不能生成 200 再填 0.00。"""
    import fba_tool.invoice_etl as invoice_etl
    import fba_tool.master_db as master_db

    raw = _raw()
    _prepare(monkeypatch, tmp_path, listing_price=None)
    real_get = master_db.get_product_by_sku
    real_lookup = master_db.lookup_by_fnsku

    def stripped_get(sku):
        return _drop_price_fields(real_get(sku))

    def stripped_lookup(fnsku):
        return _drop_price_fields(real_lookup(fnsku))

    sample = stripped_get(SKUS[0][0])
    assert sample
    for key in ("listing_price", "declared_value"):
        assert key not in sample
    monkeypatch.setattr(master_db, "get_product_by_sku", stripped_get)
    monkeypatch.setattr(master_db, "lookup_by_fnsku", stripped_lookup)
    monkeypatch.setattr(invoice_etl, "get_product_by_sku", stripped_get)
    client = _client()
    try:
        for forwarder_id in FORWARDERS:
            generated = _generate(client, raw, forwarder_id)
            _assert_price_missing(generated, f"{forwarder_id} generate missing price field")
            if forwarder_id == "qiluo":
                continue
            exported = _adapter_export(client, raw, forwarder_id)
            _assert_price_missing(exported, f"{forwarder_id} invoice-export missing price field")
    finally:
        _cleanup()


def test_named_fba_scan_skips_rebox_dirs_when_realdata_dir_set():
    """真实目录在时，命名扫描结果里不能出现刷箱 / FIST / golden / 待分配。目录没设就跳过。"""
    raw_root = os.environ.get("FBA_REALDATA_DIR", "").strip()
    if not raw_root:
        pytest.skip("FBA_REALDATA_DIR is not set")
    root = Path(raw_root)
    if not root.is_dir():
        pytest.skip("FBA_REALDATA_DIR is not a directory")

    from tests.v4.rebox_scan import batch_csvs, invoice_files

    chosen = [path for group in batch_csvs(root).values() for path in group]
    chosen.extend(invoice_files(root))
    leaked = [str(path) for path in chosen if _named_scan_should_skip(path)]
    assert leaked == []


def _named_scan_should_skip(path: Path) -> bool:
    """和 rebox_scan 同一条边界，单独写在这里，避免扫描函数自己放行自己。"""
    if "待分配" in path.name:
        return True
    folded_parts = [part.strip().casefold() for part in path.parts]
    for part, folded in zip(path.parts, folded_parts):
        text = part.strip()
        if "刷箱" in text:
            return True
        stem, dot, _ext = folded.rpartition(".")
        if folded == "fist" or (dot and stem == "fist"):
            return True
        if folded.startswith("fist") and len(folded) > 4 and not folded[4].isalnum():
            return True
        if folded == "golden" or ("golden" in folded and "sample" in folded):
            return True
    for left, right in zip(folded_parts, folded_parts[1:]):
        if "golden" in left and "sample" in right:
            return True
    return False
