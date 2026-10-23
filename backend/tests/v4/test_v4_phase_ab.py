"""V4 Phase A/B：快照、回写、任务、报关币种、领星表格、箱规与打包约束。"""
from __future__ import annotations

import base64
import io
import os
import sys
import time
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

os.environ.setdefault("FBA_DISABLE_SCHEDULER", "1")

BACKEND = Path(__file__).resolve().parents[2]
ROOT = BACKEND.parent
TEMPLATES = BACKEND / "templates" / "forwarders"
SYNTHETIC_LINGXING = BACKEND / "tests" / "fixtures" / "synthetic" / "lingxing_us_columns_82.xlsx"
FORWARDERS = {
    "send": "send.xlsx",
    "qiluo": "qiluo.xlsx",
    "jiufang": "jiufang.xlsx",
    "qiluo_t5": "qiluo_t5.xlsx",
    "qiluo_t6": "qiluo_t6.xlsx",
    "deshu": "deshu_t3.xlsx",
}


def _bytes(name: str) -> bytes:
    return (TEMPLATES / name).read_bytes()


def _snapshot(name: str, clamp: bool = True):
    from fba_tool.univer_snapshot import workbook_to_snapshot

    return workbook_to_snapshot(_bytes(name), clamp=clamp)


def _rules(snapshot: dict) -> list:
    resources = snapshot.get("resources") or []
    import json

    for item in resources:
        if item.get("name") == "SHEET_DATA_VALIDATION_PLUGIN":
            return json.loads(item["data"])
    return {}


def _semantic(data: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    try:
        sheets = {}
        for ws in wb.worksheets:
            values = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        values[cell.coordinate] = cell.value
            dvs = []
            for dv in ws.data_validations.dataValidation:
                dvs.append((str(dv.sqref), str(dv.type), str(dv.formula1), str(dv.formula2), str(dv.operator)))
            cfs = []
            for sqref, rules in ws.conditional_formatting._cf_rules.items():
                cfs.append((str(sqref), len(rules)))
            sheets[ws.title] = {
                "state": ws.sheet_state,
                "values": values,
                "merges": sorted(str(item) for item in ws.merged_cells.ranges),
                "dv": sorted(dvs),
                "images": len(getattr(ws, "_images", []) or []),
                "cf": sorted(cfs),
            }
        return sheets
    finally:
        wb.close()


def _diff(left: dict, right: dict) -> list:
    diffs = []
    for name in sorted(set(left) | set(right)):
        a = left.get(name)
        b = right.get(name)
        if a is None or b is None:
            diffs.append(f"sheet {name}")
            continue
        for key in ("state", "merges", "dv", "images", "cf"):
            if a[key] != b[key]:
                diffs.append(f"{name}.{key}")
        for coord in sorted(set(a["values"]) | set(b["values"])):
            if a["values"].get(coord) != b["values"].get(coord):
                diffs.append(f"{name}!{coord}")
    return diffs


def test_v4_t5_send_clamp_and_noclamp_stay_small():
    clamped = _snapshot("send.xlsx", True)["snapshot"]
    plain = _snapshot("send.xlsx", False)["snapshot"]
    for sheet in clamped["sheets"].values():
        assert sheet["rowCount"] < 100000
        assert sheet["rowCount"] != 1048576
    rules = _rules(clamped)
    million = []
    for sheet_rules in rules.values():
        for rule in sheet_rules:
            for item in rule["ranges"]:
                if item["endRow"] >= 100000:
                    million.append(item)
    assert million == []
    noclamp_rows = [sheet["rowCount"] for sheet in plain["sheets"].values()]
    assert max(noclamp_rows) < 1000
    huge = [
        item["endRow"]
        for sheet_rules in _rules(plain).values()
        for rule in sheet_rules
        for item in rule["ranges"]
        if item["endRow"] >= 1048570
    ]
    assert huge


def test_v4_t5b_hidden_sheets_and_formula1():
    # qiluo_t6 的四张表在模板里都是 visible。隐藏标记必须跟 Excel 一致，不能为了断言改成 hidden。
    for filename, expect_hidden in (("deshu_t3.xlsx", True), ("qiluo_t6.xlsx", False)):
        original = openpyxl.load_workbook(TEMPLATES / filename)
        snap = _snapshot(filename)["snapshot"]
        hidden = [sheet["name"] for sheet in snap["sheets"].values() if sheet.get("hidden") == 1]
        excel_hidden = [ws.title for ws in original.worksheets if ws.sheet_state != "visible"]
        assert hidden == excel_hidden
        assert bool(hidden) is expect_hidden
        expected = {}
        for ws in original.worksheets:
            for dv in ws.data_validations.dataValidation:
                if dv.formula1 not in (None, ""):
                    expected.setdefault(ws.title, []).append(str(dv.formula1))
        got = {}
        rules = _rules(snap)
        by_id = {sheet["id"]: sheet["name"] for sheet in snap["sheets"].values()}
        for sheet_id, sheet_rules in rules.items():
            got[by_id[sheet_id]] = [rule["formula1"] for rule in sheet_rules if rule.get("formula1")]
        for name, formulas in expected.items():
            assert formulas == got.get(name)
        original.close()


@pytest.mark.parametrize("forwarder", list(FORWARDERS))
def test_v4_roundtrip_and_single_cell_edit(forwarder):
    from fba_tool.univer_apply import apply_snapshot_values

    raw = _bytes(FORWARDERS[forwarder])
    payload = _snapshot(FORWARDERS[forwarder])
    same = apply_snapshot_values(raw, payload["snapshot"])
    assert _diff(_semantic(raw), _semantic(same)) == []

    snapshot = payload["snapshot"]
    edited = None
    for sheet in snapshot["sheets"].values():
        for row, cols in (sheet.get("cellData") or {}).items():
            for col, cell in cols.items():
                if not isinstance(cell, dict) or cell.get("f"):
                    continue
                if cell.get("v") in (None, ""):
                    continue
                cell["v"] = "V4-EDIT"
                cell["t"] = 1
                edited = (sheet["name"], int(row) + 1, int(col) + 1)
                break
            if edited:
                break
        if edited:
            break
    assert edited
    changed = apply_snapshot_values(raw, snapshot)
    diffs = _diff(_semantic(raw), _semantic(changed))
    from openpyxl.utils import get_column_letter

    coord = f"{edited[0]}!{get_column_letter(edited[2])}{edited[1]}"
    assert diffs == [coord]


def test_v4_t4_send_slave_write_is_rejected():
    """B4 是「FBA Shipment ID」标签合并的从属格。显式写入返回 400，不落到 C4。"""
    from fba_tool.univer_apply import ApplyRejected, apply_snapshot_values

    raw = _bytes("send.xlsx")
    payload = _snapshot("send.xlsx")["snapshot"]
    sheet = next(iter(payload["sheets"].values()))
    sheet.setdefault("cellData", {}).setdefault("3", {})["1"] = {"v": "SLAVE", "t": 1}
    with pytest.raises(ApplyRejected) as caught:
        apply_snapshot_values(raw, payload)
    assert caught.value.code == "MERGED_SLAVE"
    assert "B4" in caught.value.cells
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    try:
        assert wb.active["C4"].value in (None, "")
        assert "FBA Shipment ID" in str(wb.active["A4"].value)
    finally:
        wb.close()


def test_v4_t8_formula_plain_value_rejected_and_formula_kept():
    from fba_tool.univer_apply import ApplyRejected, apply_snapshot_values

    raw = _bytes("qiluo_t6.xlsx")
    payload = _snapshot("qiluo_t6.xlsx")["snapshot"]
    target = None
    for sheet in payload["sheets"].values():
        for row, cols in (sheet.get("cellData") or {}).items():
            for col, cell in cols.items():
                if isinstance(cell, dict) and cell.get("f"):
                    target = (sheet, row, col, cell["f"])
                    break
            if target:
                break
        if target:
            break
    assert target
    sheet, row, col, formula = target
    kept = apply_snapshot_values(raw, payload)
    assert _diff(_semantic(raw), _semantic(kept)) == []
    cell = sheet["cellData"][row][col]
    cell.pop("f", None)
    cell["v"] = "123"
    cell["t"] = 2
    with pytest.raises(ApplyRejected) as caught:
        apply_snapshot_values(raw, payload)
    assert caught.value.code == "FORMULA_OVERWRITE"
    assert caught.value.cells
    cell["f"] = formula + "+1"
    cell.pop("v", None)
    rewritten = apply_snapshot_values(raw, payload)
    assert _diff(_semantic(raw), _semantic(rewritten)) == []


def test_v4_t7_large_sheet_job_progress_and_cancel(monkeypatch):
    from fastapi.testclient import TestClient

    from fba_tool.univer_snapshot import SnapshotCancelled
    from main import app

    state = {"calls": 0}

    def slow(data, clamp=True, progress=None, cancel_event=None):
        state["calls"] += 1
        for step in range(20):
            if cancel_event is not None and cancel_event.is_set():
                raise SnapshotCancelled()
            if progress:
                progress(step * 5)
            time.sleep(0.05)
        return {"snapshot": {"id": "wb", "sheets": {}}, "meta": {"used_rows": 6001}}

    monkeypatch.setattr("fba_tool.univer_snapshot.estimate_used_rows", lambda data: 6001)
    monkeypatch.setattr("fba_tool.univer_snapshot.workbook_to_snapshot", slow)
    client = TestClient(app)
    started = time.perf_counter()
    response = client.post("/api/univer/snapshot", data={"forwarder_id": "send"})
    assert time.perf_counter() - started < 2
    assert response.status_code == 202
    body = response.json()
    assert body["status"] in {"queued", "running", "cancelled"}
    job_id = body["job_id"]
    seen = client.get(f"/api/jobs/{job_id}")
    assert seen.status_code == 200
    cancelled = client.post(f"/api/jobs/{job_id}/cancel")
    assert cancelled.status_code == 200
    deadline = time.time() + 3
    status = cancelled.json()["status"]
    while status in {"queued", "running"} and time.time() < deadline:
        status = client.get(f"/api/jobs/{job_id}").json()["status"]
        time.sleep(0.05)
    assert status == "cancelled"


def test_v4_t7_sync_template_returns_without_blocking_past_two_seconds():
    from fastapi.testclient import TestClient

    from main import app

    client = TestClient(app)
    started = time.perf_counter()
    response = client.post("/api/univer/snapshot", data={"forwarder_id": "send"})
    elapsed = time.perf_counter() - started
    assert response.status_code == 200
    assert elapsed < 2
    assert response.json()["snapshot"]["sheets"]


def test_v4_c1_export_send_keeps_model_name_wu_and_records_fx():
    from fba_tool.invoice_export import export_invoice
    from fba_tool.shipment_parser import parse_csv_shipment

    csv_text = "\n".join([
        '"货件编号","FBATESTMIX01"',
        '"货件名称","TEST-LAX9"',
        '"配送地址","FONTANA, CA"',
        '"箱子数量","1"',
        "",
        '"多件商品（1 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        '"S-A","a","B0A","X00TEST0AA","新品","无","4","4"',
        "",
        '"","","","","","","箱号","FBATESTMIX01U000001"',
        '"","","","","","","包装箱重量（磅）：","10"',
        '"","","","","","","包装箱长度（英寸）：","10"',
        '"","","","","","","包装箱宽度（英寸）：","10"',
        '"","","","","","","包装箱高度（英寸）：","10"',
    ]).encode("utf-8")
    shipment = parse_csv_shipment(csv_text, file_name="mix.csv")
    product = {
        "fnsku": "X00TEST0AA",
        "listing_price": 10.0,
        "english_name": "Item",
        "chinese_name": "物品",
        "brand": "B",
        "material": "铁",
        "purpose": "用",
        "hs_code": "7326909000",
        "asin": "B0X",
        "model_name": "无",
        "invoice": {"unit": "个", "purchase_price": 5, "material_en": "iron", "purpose_en": "use", "pcs_per_set": 1},
    }
    result = export_invoice("send", shipment=shipment, products={"X00TEST0AA": product}, reference_id="R1", invoice_date="2026-10-21")
    assert result["bytes"][:2] == b"PK"
    assert result["model"].meta["fx"]["fx_pair"] == "USD/CNY"
    blank = dict(product, model_name="")
    from fba_tool.invoice_rules import InvoiceBlocked

    with pytest.raises(InvoiceBlocked):
        export_invoice("send", shipment=shipment, products={"X00TEST0AA": blank}, reference_id="R1")


def test_v4_c1_customs_currency_and_model_blank():
    from fba_tool.export_compliance import review_products
    from fba_tool.fx_rates import upsert_rate

    product = {
        "chinese_name": "垫",
        "material": "布",
        "purpose": "宠物",
        "brand": "Joy",
        "model_name": "无",
        "hs_code": "630790",
        "product_cost_currency": "CNY",
        "listing_price_currency": "USD",
    }
    blockers, trace = review_products({"SKU": product}, "2026-10-21")
    assert blockers == []
    blank = dict(product, model_name="")
    blockers, _trace = review_products({"SKU": blank}, "2026-10-21")
    assert any(item["field"] == "model_name" for item in blockers)
    usd_cost = dict(product, product_cost_currency="USD")
    blockers, _trace = review_products({"SKU": usd_cost}, "2026-10-21")
    assert any(item["code"] == "cost_currency_mismatch" for item in blockers)
    cad = dict(product, listing_price_currency="CAD")
    blockers, _trace = review_products({"SKU": cad}, "2026-10-21")
    assert any(item["code"] == "retail_currency_mismatch" for item in blockers)
    upsert_rate({
        "currency_pair": "USD/CNY",
        "rate": 7.1,
        "rate_type": "customs_monthly",
        "applicable_month": "2026-10",
        "publish_date": "2026-10-01",
        "source": "测试海关月汇率",
    })
    _blockers, trace = review_products({"SKU": product}, "2026-10-21")
    assert trace["fx_source"] == "测试海关月汇率"
    assert trace["fx_publish_date"] == "2026-10-01"


def test_v4_a6_frozen_exe_data_dir_beside_exe(tmp_path, monkeypatch):
    from fba_tool import app_paths

    exe = tmp_path / "exe"
    local = tmp_path / "local"
    exe.mkdir()
    monkeypatch.delenv("FBA_APP_DATA_DIR", raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(local))
    monkeypatch.setattr(app_paths, "project_root", lambda: exe)
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    assert app_paths.app_data_dir() == exe / "app_data"
    assert (exe / "app_data").is_dir()
    assert not (local / "FBA_Smart_Tool").exists()


def test_v4_a6_upx_disabled():
    text = (ROOT / "packaging" / "FBA_Smart_Tool.spec").read_text(encoding="utf-8")
    assert "upx=True" not in text
    assert text.count("upx=False") >= 2


def test_v4_a7_dependency_scan():
    import subprocess

    proc = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_spreadsheet_deps.py")], capture_output=True, text=True)
    assert proc.returncode == 0, proc.stdout + proc.stderr


def _wipe_listing(skus: list[str], run_ids: list[int]) -> None:
    from fba_tool.master_db import _connect

    with _connect() as conn:
        for run_id in run_ids:
            conn.execute("DELETE FROM lingxing_import_change WHERE run_id = ?", (run_id,))
            conn.execute("DELETE FROM lingxing_sync_diff WHERE run_id = ?", (run_id,))
            conn.execute("DELETE FROM lingxing_sync_preview_field WHERE run_id = ?", (run_id,))
            conn.execute("DELETE FROM lingxing_sync_preview_sku WHERE run_id = ?", (run_id,))
            conn.execute("DELETE FROM lingxing_sync_run WHERE id = ?", (run_id,))
        for sku in skus:
            conn.execute("DELETE FROM lingxing_import_change WHERE sku = ?", (sku,))
            conn.execute("DELETE FROM product_field_history WHERE sku = ?", (sku,))
            conn.execute("DELETE FROM product_master WHERE sku = ?", (sku,))
        conn.commit()


def _synthetic_listing_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append([
        "MSKU", "FNSKU", "ASIN", "父ASIN", "状态", "店铺", "国家", "配送方式",
        "币种", "价格", "优惠价", "售价(总价)", "List Price", "划线价",
        "预估FBA费（API）", "平台费率", "SKU", "品名", "标题", "亚马逊品牌", "大类名称",
    ])
    ws.append([
        "SYN-MAT-1", "X00TESTMAT", "B0SYNMAT01", "B0SYNPARENT", "在售", "Shop-US", "美国", "FBA",
        "USD", 37.99, 37.99, 37.99, 39.99, 39.99,
        8.93, "15.0%", "SYN-LOCAL", "本地品名不入库", "English title", "AmazonBrand", "Pet Supplies",
    ])
    ws.append([
        "SYN-MAT-2", "X00TESTMA2", "B0SYNMAT02", "B0SYNPARENT", "在售", "Shop-US", "美国", "FBM",
        "USD", 12, 12, 12, 14, 14,
        3.2, 0.15, "SYN-LOCAL", "另一条本地名", "Other title", "AmazonBrand", "Pet Supplies",
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_v4_d1_listing_preview_confirm_rollback():
    from fastapi.testclient import TestClient

    from fba_tool.lingxing_listing_xlsx import parse_listing_xlsx
    from fba_tool.master_db import get_product_by_sku, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, mark_field_sources, set_product_fields
    from main import app

    data = _synthetic_listing_xlsx()
    filename = "synthetic-us-listing-2026-10-21.xlsx"
    parsed = parse_listing_xlsx(data, filename)
    assert parsed["export_at"] == "2026-10-21"
    assert len(parsed["products"]) == 2
    assert any("SYN-LOCAL" in item and "2 个 MSKU" in item for item in parsed["warnings"])
    sample = next(item for item in parsed["products"] if item["sku"] == "SYN-MAT-1")
    assert sample["listing_price_currency"] == "USD"
    assert sample["listing_price_site"] == "US"
    assert sample["listing_price"] == 37.99
    assert sample["fee_fba_est"] == 8.93
    assert sample["fee_referral_rate"] == pytest.approx(0.15)
    assert sample["local_sku"] == "SYN-LOCAL"
    assert "chinese_name" not in sample
    assert "brand" not in sample
    assert "english_name" not in sample
    skus = [item["sku"] for item in parsed["products"]]
    run_ids: list[int] = []
    try:
        upsert_product({
            "sku": sample["sku"],
            "chinese_name": "手填品名",
            "material": "手填材质",
            "carton_length": 40,
        })
        set_product_fields(sample["sku"], {"listing_price": 1, "listing_price_currency": "USD"}, SOURCE_MANUAL)
        mark_field_sources(sample["sku"], ["chinese_name", "material", "carton_length"], SOURCE_MANUAL)
        client = TestClient(app)
        before = get_product_by_sku(sample["sku"])
        response = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200
        body = response.json()
        run_ids.append(body["run_id"])
        assert body["wrote_product_master"] is False
        assert body["export_at"] == "2026-10-21"
        assert get_product_by_sku(sample["sku"])["chinese_name"] == before["chinese_name"]
        preview = client.get("/api/lingxing/sync/preview", params={"run_id": body["run_id"]}).json()
        items = []
        for rows in preview["groups"].values():
            for row in rows:
                for field in row.get("fields") or []:
                    if row["sku"] == sample["sku"] and field["field"] == "listing_price":
                        assert field["default_checked"] is False
                        assert field["manual_diff"] is True
                        assert "1" in field["difference"] and "37.99" in field["difference"]
                fields = [field["field"] for field in row.get("fields") or [] if field.get("default_checked")]
                if "carton_length" in fields or "chinese_name" in fields or "material" in fields:
                    raise AssertionError(fields)
                if fields:
                    items.append({"sku": row["sku"], "fields": fields})
        confirmed = client.post("/api/lingxing/sync/confirm", json={"run_id": body["run_id"], "items": items})
        assert confirmed.status_code == 200
        saved = get_product_by_sku(sample["sku"])
        assert saved["chinese_name"] == "手填品名"
        assert saved["material"] == "手填材质"
        assert saved["carton_length"] == 40
        assert saved["listing_price"] == 1  # 手填 1 与领星 37.99 不同，默认确认不覆盖
        assert saved["listing_price_currency"] == "USD"
        again = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        run_ids.append(again.json()["run_id"])
        second_items = []
        second_preview = client.get("/api/lingxing/sync/preview", params={"run_id": again.json()["run_id"]}).json()
        for rows in second_preview["groups"].values():
            for row in rows:
                fields = [field["field"] for field in row.get("fields") or [] if field.get("default_checked")]
                if fields:
                    second_items.append({"sku": row["sku"], "fields": fields})
        assert client.post("/api/lingxing/sync/confirm", json={"run_id": again.json()["run_id"], "items": second_items}).status_code == 200
        from fba_tool.master_db import _connect

        with _connect() as conn:
            count = conn.execute("SELECT COUNT(*) FROM product_master WHERE sku = ?", (sample["sku"],)).fetchone()[0]
        assert count == 1
        rolled = client.post(f"/api/lingxing/runs/{body['run_id']}/rollback")
        assert rolled.status_code == 200
        restored = get_product_by_sku(sample["sku"])
        assert restored["listing_price"] == 1
        assert restored["chinese_name"] == "手填品名"
    finally:
        _wipe_listing(skus, run_ids)


def _owner_listing_path() -> Path:
    """本机私有的美国站导出。仓库不带这份文件。

    设置 FBA_LINGXING_OWNER_XLSX 指向本地路径才跑；未设置或文件不存在就跳过。
    """
    override = os.environ.get("FBA_LINGXING_OWNER_XLSX", "").strip()
    if not override:
        pytest.skip("未设置 FBA_LINGXING_OWNER_XLSX，跳过本机领星导出回归")
    path = Path(override)
    if not path.is_file():
        pytest.skip(f"FBA_LINGXING_OWNER_XLSX 指向的文件不存在：{path}")
    return path


def _default_confirm_items(preview: dict) -> list[dict]:
    items = []
    for rows in preview["groups"].values():
        for row in rows:
            fields = [field["field"] for field in row.get("fields") or [] if field.get("default_checked")]
            if fields:
                items.append({"sku": row["sku"], "fields": fields})
    return items


def test_v4_d1_owner_listing_default_confirm():
    """本机私有表按预览默认勾选确认。未提供 FBA_LINGXING_OWNER_XLSX 时跳过。

    文件在时：0 费用不导致 500，手填六要素/箱规/0 价保护还在，失败时不留半写入。
    """
    from fastapi.testclient import TestClient

    from fba_tool.lingxing_listing_xlsx import parse_listing_xlsx
    from fba_tool.master_db import get_product_by_sku, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, mark_field_sources, set_product_fields
    from main import app

    upload = _owner_listing_path()
    parsed = parse_listing_xlsx(upload.read_bytes(), upload.name)
    products = parsed["products"]
    assert parsed["export_at"] == "2026-10-21"
    assert len(products) == 44
    assert parsed["warnings"]
    assert {item.get("listing_price_currency") for item in products} == {"USD"}
    assert {item.get("listing_price_site") for item in products} == {"US"}
    assert all("chinese_name" not in item and "brand" not in item and "english_name" not in item for item in products)
    assert any(item.get("fee_fba_est") == 0 for item in products)
    assert any(isinstance(item.get("fee_fba_est"), (int, float)) and item["fee_fba_est"] > 0 for item in products)
    assert any(item.get("fee_referral_rate") == pytest.approx(0.15) for item in products)

    priced = [item for item in products if isinstance(item.get("listing_price"), (int, float)) and item["listing_price"] > 0]
    unpriced = [item for item in products if item.get("listing_price") in (0, 0.0, None, "")]
    manual = priced[0]
    protected = unpriced[0]
    regular_host = next(item for item in priced if item["sku"] != manual["sku"] and not item.get("regular_price"))
    customs = {
        "chinese_name": "手填品名",
        "material": "手填材质",
        "purpose": "手填用途",
        "brand": "手填品牌",
        "model_name": "手填型号",
        "hs_code": "3926909090",
    }
    carton = {"carton_length": 40, "carton_width": 30, "carton_height": 20, "carton_weight": 5}
    skus = [item["sku"] for item in products]
    run_ids: list[int] = []
    try:
        upsert_product({"sku": manual["sku"], **customs, **carton})
        mark_field_sources(manual["sku"], [*customs, *carton], SOURCE_MANUAL)
        set_product_fields(manual["sku"], {"listing_price": 1, "listing_price_currency": "USD"}, SOURCE_MANUAL)
        upsert_product({"sku": protected["sku"]})
        set_product_fields(protected["sku"], {"listing_price": 45, "listing_price_currency": "USD"}, SOURCE_MANUAL)
        upsert_product({"sku": regular_host["sku"]})
        set_product_fields(regular_host["sku"], {"regular_price": 42}, SOURCE_MANUAL)
        client = TestClient(app)
        response = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": (upload.name, upload.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        run_ids.append(body["run_id"])
        assert body["wrote_product_master"] is False
        assert get_product_by_sku(manual["sku"])["chinese_name"] == "手填品名"
        preview = client.get("/api/lingxing/sync/preview", params={"run_id": body["run_id"]}).json()
        zero_fee_seen = False
        for rows in preview["groups"].values():
            for row in rows:
                for field in row.get("fields") or []:
                    if field["field"] == "fee_fba_est" and field["remote_value"] in {"0", "0.0"}:
                        zero_fee_seen = True
                        assert field["default_checked"] is False
                    if row["sku"] == manual["sku"] and field["field"] == "listing_price":
                        assert field["default_checked"] is False
                        assert field["manual_diff"] is True
                        assert "1" in field["difference"]
                    if field["field"] in {"chinese_name", "material", "purpose", "brand", "model_name", "hs_code", "carton_length"}:
                        raise AssertionError(field["field"])
        assert zero_fee_seen
        confirmed = client.post(
            "/api/lingxing/sync/confirm",
            json={"run_id": body["run_id"], "items": _default_confirm_items(preview)},
        )
        assert confirmed.status_code == 200, confirmed.text
        saved = get_product_by_sku(manual["sku"])
        for key, value in {**customs, **carton}.items():
            assert saved[key] == value
        assert saved["listing_price"] == 1  # 手填售价和领星不同，默认确认不覆盖
        assert get_product_by_sku(protected["sku"])["listing_price"] == 45
        assert get_product_by_sku(regular_host["sku"])["regular_price"] == 42
        for item in _default_confirm_items(preview):
            row = get_product_by_sku(item["sku"])
            assert row is not None
            remote = next(product for product in products if product["sku"] == item["sku"])
            if remote.get("fee_fba_est") in (0, 0.0):
                assert row.get("fee_fba_est") in (None, "")
        rolled = client.post(f"/api/lingxing/runs/{body['run_id']}/rollback")
        assert rolled.status_code == 200, rolled.text
        restored = get_product_by_sku(manual["sku"])
        assert restored["listing_price"] == 1
        assert restored["chinese_name"] == "手填品名"
        assert restored["carton_length"] == 40
        assert get_product_by_sku(protected["sku"])["listing_price"] == 45
    finally:
        _wipe_listing(skus, run_ids)


def test_v4_d1_synthetic_sheet_default_confirm_and_rollback():
    """82 列脱敏表在 CI 里必跑：默认确认 200，0 费用不入库，回滚不影响后来的导入和手改。"""
    from fastapi.testclient import TestClient

    from fba_tool.lingxing_listing_xlsx import parse_listing_xlsx
    from fba_tool.master_db import get_product_by_sku, upsert_product
    from fba_tool.master_db_lingxing import get_lingxing_sync_run
    from fba_tool.product_fields import SOURCE_MANUAL, mark_field_sources, set_product_fields
    from main import app

    raw = SYNTHETIC_LINGXING.read_bytes()
    parsed = parse_listing_xlsx(raw, SYNTHETIC_LINGXING.name)
    products = {item["sku"]: item for item in parsed["products"]}
    assert len(parsed["columns"]) == 82
    assert set(products) == {"TEST_PRICED", "TEST_ZERO_FEE", "TEST_NO_FNSKU", "TEST_UNPRICED"}
    assert products["TEST_ZERO_FEE"]["fee_fba_est"] == 0
    assert products["TEST_NO_FNSKU"].get("fnsku", "") == ""
    assert products["TEST_PRICED"]["listing_price"] == 10
    assert products["TEST_PRICED"]["listing_price_currency"] == "USD"
    customs = {
        "chinese_name": "手填品名",
        "material": "手填材质",
        "purpose": "手填用途",
        "brand": "手填品牌",
        "model_name": "手填型号",
        "hs_code": "3926909090",
    }
    carton = {"carton_length": 40, "carton_width": 30, "carton_height": 20, "carton_weight": 5}
    skus = ["TEST_PRICED", "TEST_ZERO_FEE", "TEST_NO_FNSKU", "TEST_UNPRICED", "TEST_LATER", "KEEP-HAND"]
    run_ids: list[int] = []
    client = TestClient(app)
    try:
        upsert_product({"sku": "TEST_PRICED", **customs, **carton})
        mark_field_sources("TEST_PRICED", [*customs, *carton], SOURCE_MANUAL)
        set_product_fields("TEST_PRICED", {"listing_price": 1, "listing_price_currency": "USD"}, SOURCE_MANUAL)
        upsert_product({"sku": "TEST_UNPRICED"})
        set_product_fields("TEST_UNPRICED", {"listing_price": 45, "listing_price_currency": "USD"}, SOURCE_MANUAL)
        upsert_product({"sku": "KEEP-HAND", "chinese_name": "旁路品名"})
        response = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": (SYNTHETIC_LINGXING.name, raw, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200, response.text
        run1 = response.json()["run_id"]
        run_ids.append(run1)
        assert response.json()["wrote_product_master"] is False
        preview = client.get("/api/lingxing/sync/preview", params={"run_id": run1}).json()
        zero_fee_seen = False
        no_fnsku_checked = False
        for rows in preview["groups"].values():
            for row in rows:
                fields = row.get("fields") or []
                if row["sku"] == "TEST_NO_FNSKU":
                    no_fnsku_checked = any(field.get("default_checked") for field in fields)
                for field in fields:
                    if field["field"] == "fee_fba_est" and field["remote_value"] in {"0", "0.0"}:
                        zero_fee_seen = True
                        assert field["default_checked"] is False
                    if row["sku"] == "TEST_PRICED" and field["field"] == "listing_price":
                        assert field["default_checked"] is False
                        assert field["manual_diff"] is True
                        assert "1" in field["difference"] and "10" in field["difference"]
                    if field["field"] in {"chinese_name", "material", "purpose", "brand", "model_name", "hs_code", "carton_length"}:
                        raise AssertionError(field["field"])
        assert zero_fee_seen
        assert no_fnsku_checked is False
        confirmed = client.post(
            "/api/lingxing/sync/confirm",
            json={"run_id": run1, "items": _default_confirm_items(preview)},
        )
        assert confirmed.status_code == 200, confirmed.text
        saved = get_product_by_sku("TEST_PRICED")
        for key, value in {**customs, **carton}.items():
            assert saved[key] == value
        assert saved["listing_price"] == 1  # 手填 1 与领星 10 不同，默认确认不覆盖
        assert get_product_by_sku("TEST_ZERO_FEE")["listing_price"] == 11
        assert get_product_by_sku("TEST_ZERO_FEE").get("fee_fba_est") in (None, "")
        assert get_product_by_sku("TEST_UNPRICED")["listing_price"] == 45
        assert get_product_by_sku("TEST_NO_FNSKU") is None
        set_product_fields("TEST_PRICED", {"local_name": "手改品名"}, SOURCE_MANUAL)
        later = _tiny_listing([["TEST_LATER", "X00TESTLT1", "B0TESTLT01", "美国", "USD", 33]])
        second = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": ("later.xlsx", later, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert second.status_code == 200, second.text
        run2 = second.json()["run_id"]
        run_ids.append(run2)
        preview2 = client.get("/api/lingxing/sync/preview", params={"run_id": run2}).json()
        confirmed2 = client.post(
            "/api/lingxing/sync/confirm",
            json={"run_id": run2, "items": _default_confirm_items(preview2)},
        )
        assert confirmed2.status_code == 200, confirmed2.text
        rolled = client.post(f"/api/lingxing/runs/{run1}/rollback")
        assert rolled.status_code == 200, rolled.text
        assert run2 in rolled.json()["later_runs"]
        restored = get_product_by_sku("TEST_PRICED")
        assert restored["listing_price"] == 1
        assert restored["local_name"] == "手改品名"
        assert restored["chinese_name"] == "手填品名"
        assert get_product_by_sku("TEST_LATER")["listing_price"] == 33
        assert get_product_by_sku("KEEP-HAND")["chinese_name"] == "旁路品名"
        assert get_lingxing_sync_run(run2)["status"] == "confirmed"
    finally:
        _wipe_listing(skus, run_ids)


def test_v4_d1_synthetic_sheet_confirm_is_atomic(monkeypatch):
    from fastapi.testclient import TestClient

    from fba_tool import product_fields
    from fba_tool.master_db import get_product_by_sku
    from fba_tool.master_db_lingxing import get_lingxing_sync_run
    from main import app

    real = product_fields.set_product_fields
    state = {"n": 0}

    def wrapped(sku, values, source):
        state["n"] += 1
        if state["n"] >= 2:
            raise ValueError("「fee_fba_est」必须大于 0：0.0")
        return real(sku, values, source)

    monkeypatch.setattr(product_fields, "set_product_fields", wrapped)
    raw = SYNTHETIC_LINGXING.read_bytes()
    skus = ["TEST_PRICED", "TEST_ZERO_FEE", "TEST_NO_FNSKU", "TEST_UNPRICED"]
    client = TestClient(app)
    response = client.post(
        "/api/lingxing/import-xlsx",
        files={"file": (SYNTHETIC_LINGXING.name, raw, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    run_id = response.json()["run_id"]
    try:
        preview = client.get("/api/lingxing/sync/preview", params={"run_id": run_id}).json()
        items = _default_confirm_items(preview)
        assert {item["sku"] for item in items} >= {"TEST_PRICED", "TEST_ZERO_FEE"}
        confirmed = client.post("/api/lingxing/sync/confirm", json={"run_id": run_id, "items": items})
        assert confirmed.status_code == 409, confirmed.text
        for sku in skus:
            assert get_product_by_sku(sku) is None
        assert get_lingxing_sync_run(run_id)["status"] != "confirmed"
    finally:
        _wipe_listing(skus, [run_id])


def test_v4_d1_zero_fee_checked_does_not_fail_confirm():
    from fastapi.testclient import TestClient

    from fba_tool.master_db import get_product_by_sku
    from main import app

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(["MSKU", "FNSKU", "ASIN", "国家", "币种", "价格", "预估FBA费（API）", "平台费率"])
    ws.append(["FEE-OK", "X00TESTFEE", "B0TESTFEE1", "美国", "USD", 10, 1.25, "15%"])
    ws.append(["FEE-ZERO", "X00TESTFE0", "B0TESTFEE0", "美国", "USD", 11, 0, "15%"])
    buf = io.BytesIO()
    wb.save(buf)
    payload = buf.getvalue()
    skus = ["FEE-OK", "FEE-ZERO"]
    client = TestClient(app)
    response = client.post(
        "/api/lingxing/import-xlsx",
        files={"file": ("synthetic-fee.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    run_id = response.json()["run_id"]
    try:
        preview = client.get("/api/lingxing/sync/preview", params={"run_id": run_id}).json()
        items = []
        for rows in preview["groups"].values():
            for row in rows:
                fields = [field["field"] for field in row.get("fields") or [] if field.get("default_checked")]
                if row["sku"] == "FEE-ZERO":
                    assert any(
                        field["field"] == "fee_fba_est" and field["default_checked"] is False
                        for field in row.get("fields") or []
                    )
                if row["sku"] in {"FEE-ZERO", "FEE-OK"}:
                    for name in ("listing_price", "fee_fba_est"):
                        if name not in fields:
                            fields.append(name)
                if fields:
                    items.append({"sku": row["sku"], "fields": fields})
        confirmed = client.post("/api/lingxing/sync/confirm", json={"run_id": run_id, "items": items})
        assert confirmed.status_code == 200, confirmed.text
        assert get_product_by_sku("FEE-ZERO")["fee_fba_est"] in (None, "")
        assert get_product_by_sku("FEE-ZERO")["listing_price"] == 11
        assert get_product_by_sku("FEE-OK")["fee_fba_est"] == pytest.approx(1.25)
    finally:
        _wipe_listing(skus, [run_id])


def _tiny_listing(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(["MSKU", "FNSKU", "ASIN", "国家", "币种", "价格"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_v4_d1_rollback_run1_keeps_later_import_and_manual_edits():
    """回滚只撤销这一次写入。之后的导入和手改还在。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import get_product_by_sku, upsert_product
    from fba_tool.master_db_lingxing import get_lingxing_sync_run
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields
    from main import app

    file1 = _tiny_listing([["RUN1-SKU", "X00TESTR01", "B0RUN1SKU1", "美国", "USD", 21]])
    file2 = _tiny_listing([["RUN2-SKU", "X00TESTR02", "B0RUN2SKU2", "美国", "USD", 33]])
    skus = ["RUN1-SKU", "RUN2-SKU", "KEEP-HAND"]
    run_ids: list[int] = []
    client = TestClient(app)
    try:
        upsert_product({"sku": "KEEP-HAND", "chinese_name": "旁路品名"})
        set_product_fields("KEEP-HAND", {"listing_price": 9, "listing_price_currency": "USD"}, SOURCE_MANUAL)
        first = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": ("run1.xlsx", file1, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert first.status_code == 200, first.text
        run1 = first.json()["run_id"]
        run_ids.append(run1)
        preview1 = client.get("/api/lingxing/sync/preview", params={"run_id": run1}).json()
        confirmed1 = client.post(
            "/api/lingxing/sync/confirm",
            json={"run_id": run1, "items": _default_confirm_items(preview1)},
        )
        assert confirmed1.status_code == 200, confirmed1.text
        assert get_product_by_sku("RUN1-SKU")["listing_price"] == 21
        set_product_fields("RUN1-SKU", {"local_name": "手改品名"}, SOURCE_MANUAL)
        second = client.post(
            "/api/lingxing/import-xlsx",
            files={"file": ("run2.xlsx", file2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert second.status_code == 200, second.text
        run2 = second.json()["run_id"]
        run_ids.append(run2)
        preview2 = client.get("/api/lingxing/sync/preview", params={"run_id": run2}).json()
        confirmed2 = client.post(
            "/api/lingxing/sync/confirm",
            json={"run_id": run2, "items": _default_confirm_items(preview2)},
        )
        assert confirmed2.status_code == 200, confirmed2.text
        assert get_product_by_sku("RUN2-SKU")["listing_price"] == 33
        rolled = client.post(f"/api/lingxing/runs/{run1}/rollback")
        assert rolled.status_code == 200, rolled.text
        body = rolled.json()
        assert run2 in body["later_runs"]
        kept = get_product_by_sku("RUN1-SKU")
        assert kept["local_name"] == "手改品名"
        assert kept.get("listing_price") in (None, "")
        assert get_product_by_sku("RUN2-SKU")["listing_price"] == 33
        assert get_product_by_sku("KEEP-HAND")["chinese_name"] == "旁路品名"
        assert get_product_by_sku("KEEP-HAND")["listing_price"] == 9
        assert get_lingxing_sync_run(run2)["status"] == "confirmed"
        again = client.post(f"/api/lingxing/runs/{run1}/rollback")
        assert again.status_code == 409
    finally:
        _wipe_listing(skus, run_ids)


def test_v4_d1_confirm_failure_leaves_no_partial_write(monkeypatch):
    from fastapi.testclient import TestClient

    from fba_tool import product_fields
    from fba_tool.master_db import get_product_by_sku
    from fba_tool.master_db_lingxing import get_lingxing_sync_run
    from main import app

    real = product_fields.set_product_fields
    state = {"n": 0}

    def wrapped(sku, values, source):
        state["n"] += 1
        if state["n"] >= 2:
            raise ValueError("「fee_fba_est」必须大于 0：0.0")
        return real(sku, values, source)

    monkeypatch.setattr(product_fields, "set_product_fields", wrapped)
    data = _synthetic_listing_xlsx()
    skus = ["SYN-MAT-1", "SYN-MAT-2"]
    client = TestClient(app)
    response = client.post(
        "/api/lingxing/import-xlsx",
        files={"file": ("synthetic-us-listing-2026-10-21.xlsx", data, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    run_id = response.json()["run_id"]
    try:
        preview = client.get("/api/lingxing/sync/preview", params={"run_id": run_id}).json()
        confirmed = client.post(
            "/api/lingxing/sync/confirm",
            json={"run_id": run_id, "items": _default_confirm_items(preview)},
        )
        assert confirmed.status_code == 409, confirmed.text
        assert "fee_fba_est" in confirmed.text
        for sku in skus:
            assert get_product_by_sku(sku) is None
        assert get_lingxing_sync_run(run_id)["status"] != "confirmed"
    finally:
        _wipe_listing(skus, [run_id])


def test_v4_d1_currency_anomaly_keeps_other_rows_and_exports_errors():
    from fastapi.testclient import TestClient

    from fba_tool.lingxing_listing_xlsx import parse_listing_xlsx
    from main import app

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["MSKU", "国家", "币种", "价格", "采购成本", "成本币种", "品名"])
    ws.append(["US-1", "美国", "USD", 12.5, 30, "CNY", "本地名"])
    ws.append(["CA-1", "加拿大", "USD", 9, None, None, "加拿大名"])
    ws.append(["US-2", "美国", "USD", 8, 10, "USD", "成本美元"])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_listing_xlsx(buf.getvalue(), "synthetic-listing-anomaly.xlsx")
    skus = {item["sku"]: item for item in parsed["products"]}
    assert skus["US-1"]["listing_price"] == 12.5
    assert skus["US-1"]["listing_price_currency"] == "USD"
    assert skus["US-1"]["product_cost_cny"] == 30
    assert "CA-1" not in skus
    assert "US-2" in skus
    assert "product_cost_cny" not in skus["US-2"]
    assert any(item["sku"] == "CA-1" for item in parsed["anomalies"])
    client = TestClient(app)
    response = client.post(
        "/api/lingxing/import-xlsx",
        files={"file": ("synthetic-listing-anomaly.xlsx", buf.getvalue(), "application/octet-stream")},
    )
    run_id = response.json()["run_id"]
    try:
        csv_text = client.get(f"/api/lingxing/runs/{run_id}/errors.csv").text
        assert "CA-1" in csv_text
        assert "US-2" in csv_text
    finally:
        _wipe_listing(["US-1", "US-2", "CA-1"], [run_id])


def test_v4_d2_readonly_schedule_and_absent_sku_kept(monkeypatch):
    from fba_tool.lingxing_scheduler import DEFAULTS, ReadOnlyClient, ReadOnlyViolation, load_schedule, read_log_tail, run_once, sync_logger
    from fba_tool.master_db import get_product_by_sku, upsert_product

    assert DEFAULTS["enabled"] is False
    assert load_schedule()["enabled"] is False
    client = ReadOnlyClient(object())
    with pytest.raises(ReadOnlyViolation):
        client.action("listing_update", {})
    result = run_once(trigger="scheduled", sleep=lambda _seconds: None)
    assert result["ok"] is False
    assert result["status"] == "failed"
    sync_logger().info("v4 d2 log marker")
    assert any("v4 d2 log marker" in line or "MCP Key" in line for line in read_log_tail(50))
    upsert_product({"sku": "KEEP-ME", "chinese_name": "留下"})
    try:
        assert get_product_by_sku("KEEP-ME")["chinese_name"] == "留下"
        text = (BACKEND / "fba_tool" / "lingxing_sync.py").read_text(encoding="utf-8")
        assert "DELETE FROM product_master" not in text
        assert get_product_by_sku("KEEP-ME")["chinese_name"] == "留下"
    finally:
        from fba_tool.master_db import delete_product_by_sku

        delete_product_by_sku("KEEP-ME")


def test_v4_b2_oversize_does_not_block_existing_api_limit():
    from fastapi.testclient import TestClient

    from main import app

    client = TestClient(app)
    ok = client.post("/api/carton-specs", json={
        "sku": "V4-OVER",
        "length_cm": 70,
        "width_cm": 40,
        "height_cm": 30,
        "qty_per_carton": 4,
        "gross_kg": 24,
        "source": "manual",
    })
    assert ok.status_code == 200
    blocked = client.post("/api/carton-specs", json={
        "sku": "V4-HEAVY",
        "length_cm": 40,
        "width_cm": 40,
        "height_cm": 40,
        "qty_per_carton": 1,
        "gross_kg": 80,
        "source": "manual",
    })
    assert blocked.status_code == 422
    missing = client.post("/api/carton-specs", json={"sku": "", "gross_kg": 1})
    assert missing.status_code == 422


def test_forwarder_dropdown_includes_invoice_templates():
    from fba_tool.forwarder_registry import get_template_bytes, list_forwarders

    ids = {item.id for item in list_forwarders()}
    for forwarder_id in ("qiluo_t5", "qiluo_t6", "deshu", "send"):
        assert forwarder_id in ids
        assert get_template_bytes(forwarder_id)


def _structure(data: bytes) -> dict:
    full = _semantic(data)
    return {
        name: {key: sheet[key] for key in ("state", "merges", "dv", "images", "cf")}
        for name, sheet in full.items()
    }


def _export(client, filename: str, forwarder_id: str, sheet: str, fields: list, snapshot=None):
    raw = _bytes(filename)
    body = {
        "forwarder_id": forwarder_id,
        "template_base64": base64.b64encode(raw).decode(),
        "template_sheet": sheet,
        "shipment_id": "FBAE2E",
        "warehouse_code": "ONT8",
        "fields": fields,
    }
    if snapshot is not None:
        body["snapshot"] = snapshot
    return client.post("/api/export_invoice", json=body), raw


def test_v4_export_invoice_keeps_workbook_and_rejects_conflicts():
    """发出件保留原模板的表、隐藏状态、图片、数据验证；公式格改成纯值、合并从属格都返回 400。"""
    from fastapi.testclient import TestClient

    from main import app

    client = TestClient(app)
    saved, raw = _export(
        client,
        "deshu_t3.xlsx",
        "deshu",
        "发票填写模板",
        [{"cell": "Z40", "value": "TEST_V", "key": "Z40", "label": "Z40"}],
    )
    assert saved.status_code == 200, saved.text[:400]
    assert _structure(saved.content) == _structure(raw)
    book = openpyxl.load_workbook(io.BytesIO(saved.content))
    try:
        assert book["发票填写模板"]["Z40"].value == "TEST_V"
        assert book["渠道"].sheet_state == "hidden"
    finally:
        book.close()

    saved, raw = _export(
        client,
        "qiluo_t6.xlsx",
        "qiluo_t6",
        "发票",
        [{"cell": "E3", "value": "TEST_V", "key": "E3", "label": "E3"}],
    )
    assert saved.status_code == 200, saved.text[:400]
    assert _structure(saved.content) == _structure(raw)
    book = openpyxl.load_workbook(io.BytesIO(saved.content))
    try:
        assert "！" in book.sheetnames
        assert str(book["发票"]["E4"].value).startswith("=VLOOKUP")
        assert book["发票"]["E3"].value == "TEST_V"
    finally:
        book.close()

    saved, raw = _export(client, "send.xlsx", "send", "Sheet1", [])
    assert saved.status_code == 200, saved.text[:400]
    assert _structure(saved.content) == _structure(raw)

    slave, _raw = _export(
        client,
        "send.xlsx",
        "send",
        "Sheet1",
        [{"cell": "B4", "value": "TEST_SLAVE", "key": "B4", "label": "B4"}],
    )
    assert slave.status_code == 400, slave.text[:400]
    assert slave.json()["detail"]["code"] == "MERGED_SLAVE"
    assert "B4" in slave.json()["detail"]["cells"]

    formula, _raw = _export(
        client,
        "qiluo_t6.xlsx",
        "qiluo_t6",
        "发票",
        [{"cell": "E4", "value": "TEST_PURE_123", "key": "E4", "label": "E4"}],
    )
    assert formula.status_code == 400
    assert formula.json()["detail"]["code"] == "FORMULA_OVERWRITE"

    snap = _snapshot("send.xlsx")["snapshot"]
    sheet = next(iter(snap["sheets"].values()))
    sheet.setdefault("cellData", {}).setdefault("3", {})["1"] = {"v": "TEST_SLAVE"}
    moved, _raw = _export(client, "send.xlsx", "send", "Sheet1", [], snap)
    assert moved.status_code == 400, moved.text[:400]
    assert moved.json()["detail"]["code"] == "MERGED_SLAVE"
    assert "B4" in moved.json()["detail"]["cells"]

    snap = _snapshot("qiluo_t6.xlsx")["snapshot"]
    sheet = next(item for item in snap["sheets"].values() if item["name"] == "发票")
    cell = sheet["cellData"]["3"]["4"]
    cell.pop("f", None)
    cell["v"] = "TEST_PURE_123"
    rejected, _raw = _export(client, "qiluo_t6.xlsx", "qiluo_t6", "发票", [], snap)
    assert rejected.status_code == 400
    assert rejected.json()["detail"]["code"] == "FORMULA_OVERWRITE"

    snap = _snapshot("deshu_t3.xlsx")["snapshot"]
    kept, raw = _export(client, "deshu_t3.xlsx", "deshu", "发票填写模板", [], snap)
    assert kept.status_code == 200, kept.text[:400]
    assert _structure(kept.content) == _structure(raw)


def test_v4_a7_flags_injected_workbook_parser(tmp_path):
    import importlib.util

    token = "excel" + "js"
    probe = tmp_path / "injected.ts"
    probe.write_text(f'import Parser from "{token}";\nParser.load(buf);\n', encoding="utf-8")
    spec = importlib.util.spec_from_file_location(
        "check_spreadsheet_deps",
        ROOT / "scripts" / "check_spreadsheet_deps.py",
    )
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    assert module.scan_paths([tmp_path])
    assert module.scan_paths([tmp_path / "missing"]) == []


def test_all_forwarder_slave_cells_reject_merged_slave():
    """每个货代模板的从属格都按原坐标拒绝，不能改写到旁边的格子。"""
    from openpyxl.utils import get_column_letter

    from fba_tool.univer_apply import ApplyRejected, _slave_map, apply_cell_writes, apply_snapshot_values

    scanned = 0
    for path in sorted(TEMPLATES.glob("*.xlsx")):
        raw = path.read_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        try:
            for ws in wb.worksheets:
                slaves = _slave_map(ws)
                if not slaves:
                    continue
                writes = []
                cell_data: dict = {}
                for row, col in slaves:
                    ref = f"{get_column_letter(col)}{row}"
                    writes.append({"cell": ref, "value": "SLAVE_SCAN"})
                    cell_data.setdefault(str(row - 1), {})[str(col - 1)] = {"v": "SLAVE_SCAN"}
                scanned += len(writes)
                with pytest.raises(ApplyRejected) as caught:
                    apply_cell_writes(raw, writes, ws.title)
                assert caught.value.code == "MERGED_SLAVE", path.name
                assert set(caught.value.cells) == {item["cell"] for item in writes}
                snapshot = {
                    "sheetOrder": ["s0"],
                    "sheets": {"s0": {"name": ws.title, "cellData": cell_data}},
                }
                with pytest.raises(ApplyRejected) as snapped:
                    apply_snapshot_values(raw, snapshot)
                assert snapped.value.code == "MERGED_SLAVE"
                assert set(snapped.value.cells) == {item["cell"] for item in writes}
        finally:
            wb.close()
    assert scanned >= 24

    raw = _bytes("qiluo_t6.xlsx")
    with pytest.raises(ApplyRejected) as caught:
        apply_cell_writes(raw, [{"cell": "C7", "value": "HACK"}], "发票")
    assert caught.value.code == "MERGED_SLAVE"
    assert caught.value.cells == ["C7"]
    book = openpyxl.load_workbook(io.BytesIO(raw))
    try:
        assert book["发票"]["D7"].value == "收件人地址1"
        assert book["发票"]["B7"].value == "不需报关"
    finally:
        book.close()


def test_template_base64_formula_or_wrong_forwarder_returns_400():
    from fastapi.testclient import TestClient

    from main import app

    client = TestClient(app)

    def post(forwarder: str, sheet: str, data: bytes, fields: list):
        return client.post(
            "/api/export_invoice",
            json={
                "forwarder_id": forwarder,
                "template_base64": base64.b64encode(data).decode(),
                "template_sheet": sheet,
                "shipment_id": "FBAE2E",
                "warehouse_code": "ONT8",
                "fields": fields,
            },
        )

    mismatch = post("qiluo_t6", "发票", _bytes("send.xlsx"), [{"cell": "E3", "value": "X", "key": "E3", "label": "E3"}])
    assert mismatch.status_code == 400, mismatch.text[:400]
    assert mismatch.json()["detail"]["code"] == "TEMPLATE_MISMATCH"

    other = post("qiluo_t6", "发票", _bytes("qiluo.xlsx"), [])
    assert other.status_code == 400
    assert other.json()["detail"]["code"] == "TEMPLATE_MISMATCH"

    raw = _bytes("qiluo_t6.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    wb["发票"]["E4"] = 123
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    formula = post("qiluo_t6", "发票", buf.getvalue(), [])
    assert formula.status_code == 400, formula.text[:400]
    assert formula.json()["detail"]["code"] == "FORMULA_OVERWRITE"
    assert "E4" in formula.json()["detail"]["cells"]

    wb = openpyxl.load_workbook(io.BytesIO(raw))
    wb["发票"]["E3"] = "FROM_UPLOAD"
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    edited = post("qiluo_t6", "发票", buf.getvalue(), [])
    assert edited.status_code == 200, edited.text[:400]
    book = openpyxl.load_workbook(io.BytesIO(edited.content))
    try:
        assert book["发票"]["E3"].value == "FROM_UPLOAD"
        assert str(book["发票"]["E4"].value).startswith("=VLOOKUP")
    finally:
        book.close()


def test_deshu_currency_dropdown_survives_snapshot_apply():
    from fba_tool.univer_apply import apply_snapshot_values

    raw = _bytes("deshu_t3.xlsx")
    out = apply_snapshot_values(raw, _snapshot("deshu_t3.xlsx")["snapshot"])
    before = _semantic(raw)["发票填写模板"]["dv"]
    after = _semantic(out)["发票填写模板"]["dv"]
    assert before == after
    assert any("USD" in item[2] for item in after)
