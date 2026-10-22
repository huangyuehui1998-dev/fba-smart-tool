"""测试方 r2：琪珞型号/材质分列、SEND 明细区与主格、预览改动导出、回滚 NULL、脱敏第 13 批。"""
from __future__ import annotations

import base64
import io
import json
from datetime import datetime, timezone

import openpyxl
import pytest

FORWARDERS = ("send", "jiufang", "qiluo", "qiluo_t5", "qiluo_t6", "deshu")


def _csv(shipment: str, sku: str, fnsku: str, asin: str) -> bytes:
    return "\n".join([
        f'"货件编号","{shipment}"',
        '"货件名称","B13-DESENS"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","1"',
        "",
        '"单件商品（1 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        f'"{sku}","rod","{asin}","{fnsku}","新品","无","4","4"',
        "",
        f'"","","","","","","箱号","{shipment}U000001"',
        '"","","","","","","包装箱重量（磅）：","10"',
        '"","","","","","","包装箱长度（英寸）：","17.70"',
        '"","","","","","","包装箱宽度（英寸）：","17.70"',
        '"","","","","","","包装箱高度（英寸）：","16.10"',
    ]).encode("utf-8")


def _seed(
    sku: str,
    fnsku: str,
    asin: str,
    model_name: str,
    material: str = "塑料",
    purpose: str = "户外",
    declared_value: float | None = None,
    english_name: str = "Fishing rod",
    listing_price: float | None = 12.5,
):
    from fba_tool.master_db import _connect, init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields

    init_master_db()
    now = datetime.now(timezone.utc).isoformat()
    with _connect() as conn:
        conn.execute("DELETE FROM fba_warehouse WHERE code = ?", ("TST1",))
        conn.execute(
            """
            INSERT INTO fba_warehouse
                (code, name, company, address1, city, state, zip, country, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("TST1", "TST1", "Amazon", "1 Test Road", "STEP45VILLE", "ZZ", "00000", "US", now),
        )
        conn.commit()
    upsert_product(
        {
            "sku": sku,
            "fnsku": fnsku,
            "asin": asin,
            "english_name": english_name,
            "chinese_name": "鱼竿",
            "brand": "Acme",
            "material": material,
            "purpose": purpose,
            "hs_code": "9507100090",
            "model_name": model_name,
            "declared_value": declared_value,
            "extra_data": json.dumps(
                {"invoice": {"unit": "个", "purchase_price": 5, "material_en": "plastic", "purpose_en": "outdoor", "pcs_per_set": 1}},
                ensure_ascii=False,
            ),
        }
    )
    if listing_price is not None:
        set_product_fields(sku, {"listing_price": listing_price, "listing_price_currency": "USD"}, SOURCE_MANUAL)


def _book(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


@pytest.mark.parametrize("forwarder_id", FORWARDERS)
def test_desensitized_batch13_fill_then_export(forwarder_id):
    """脱敏第 13 批形状：Step 4.5 生成再导出，六个模板都要 200，且明细写进表里。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku, fnsku, asin = "B13-SKU", "X00TESTE2E", "B0B13DS001"
    _seed(sku, fnsku, asin, "型号乙")
    raw = _csv("FBA13DS01", sku, fnsku, asin)
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13DS01",
                "forwarder_id": forwarder_id,
                "csv_base64": base64.b64encode(raw).decode(),
                "po_number": "PO-13",
            },
        )
        assert generated.status_code == 200, generated.text
        body = generated.json()
        cells = {item["cell"]: item.get("value") for item in body["cells"]}
        if forwarder_id == "qiluo":
            assert cells.get("K2") == "型号乙"
            assert cells.get("H2") == "塑料"
            assert cells.get("I2") == "户外"
            assert "/" not in str(cells.get("H2") or "")
            assert sku not in cells.values()
        if forwarder_id == "send":
            assert "B4" not in cells and "B5" not in cells and "B6" not in cells and "E5" not in cells
            assert cells.get("C4") == "FBA13DS01"
            assert fnsku in cells.values()
        exported = client.post(
            "/api/export_invoice",
            json={
                "template_hash": body["template_hash"],
                "forwarder_id": forwarder_id,
                "shipment_id": body["shipment_id"],
                "warehouse_code": body["warehouse_code"],
                "template_sheet": body["template_sheet"],
                "fields": body["cells"],
            },
        )
        assert exported.status_code == 200, exported.text[:400]
        assert exported.content[:2] == b"PK"
        wb = _book(exported.content)
        try:
            texts = {
                str(cell.value).strip()
                for ws in wb.worksheets
                for row in ws.iter_rows(max_row=min(ws.max_row or 1, 80), max_col=min(ws.max_column or 1, 24))
                for cell in row
                if cell.value not in (None, "")
            }
            assert "FBA13DS01" in texts
            if forwarder_id == "qiluo":
                assert wb.active["K2"].value == "型号乙"
                assert wb.active["H2"].value == "塑料"
                assert wb.active["I2"].value == "户外"
            if forwarder_id == "send":
                assert wb.active["C4"].value == "FBA13DS01"
                assert "FBA Shipment ID" in str(wb.active["A4"].value)
                assert fnsku in texts
        finally:
            wb.close()
    finally:
        delete_product_by_sku(sku)


def test_qiluo_empty_model_name_does_not_fall_back_to_sku():
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku = "B13-NOMODEL"
    _seed(sku, "X00TESTNON", "B0B13NONE1", "")
    raw = _csv("FBA13NONE", sku, "X00TESTNON", "B0B13NONE1")
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13NONE",
                "forwarder_id": "qiluo",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 200, generated.text
        cells = {item["cell"]: item.get("value") for item in generated.json()["cells"]}
        assert cells.get("K2") in (None, "")
        assert sku not in set(cells.values())
        assert cells.get("H2") == "塑料"
        assert cells.get("I2") == "户外"
        assert "/" not in str(cells.get("H2") or "")
    finally:
        delete_product_by_sku(sku)


def test_export_keeps_preview_edits_and_official_shell():
    from fastapi.testclient import TestClient

    from fba_tool.forwarder_registry import get_template_bytes
    from main import app

    official = get_template_bytes("qiluo")
    wb = _book(official)
    wb.active["K2"] = "手改型号"
    wb.active["H2"] = "手改材质"
    wb.create_sheet("不应出现")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    client = TestClient(app)
    exported = client.post(
        "/api/export_invoice",
        json={
            "forwarder_id": "qiluo",
            "template_base64": base64.b64encode(buf.getvalue()).decode(),
            "shipment_id": "FBA13SHELL",
            "warehouse_code": "TST1",
            "template_sheet": "发票",
            "fields": [],
        },
    )
    assert exported.status_code == 200, exported.text[:400]
    out = _book(exported.content)
    try:
        assert "不应出现" not in out.sheetnames
        assert out.active["K2"].value == "手改型号"
        assert out.active["H2"].value == "手改材质"
        assert out.active["A1"].value == "箱序号*"
        assert str(out.active["K1"].value).startswith("型号")
    finally:
        out.close()


def test_export_field_edit_overrides_rendered_qiluo_model():
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku = "B13-EDIT"
    _seed(sku, "X00TESTEDT", "B0B13EDIT1", "型号乙")
    raw = _csv("FBA13EDIT", sku, "X00TESTEDT", "B0B13EDIT1")
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13EDIT",
                "forwarder_id": "qiluo",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 200, generated.text
        body = generated.json()
        fields = [dict(item) for item in body["cells"]]
        for item in fields:
            if item.get("cell") == "K2":
                item["value"] = "预览改过"
        exported = client.post(
            "/api/export_invoice",
            json={
                "template_hash": body["template_hash"],
                "forwarder_id": "qiluo",
                "shipment_id": body["shipment_id"],
                "warehouse_code": body["warehouse_code"],
                "template_sheet": body["template_sheet"],
                "fields": fields,
            },
        )
        assert exported.status_code == 200, exported.text[:400]
        wb = _book(exported.content)
        try:
            assert wb.active["K2"].value == "预览改过"
            assert wb.active["H2"].value == "塑料"
        finally:
            wb.close()
    finally:
        delete_product_by_sku(sku)


def test_send_note_slave_is_rejected():
    from fastapi.testclient import TestClient

    from fba_tool.forwarder_registry import get_template_bytes
    from main import app

    official = get_template_bytes("send")
    client = TestClient(app)
    exported = client.post(
        "/api/export_invoice",
        json={
            "forwarder_id": "send",
            "template_base64": base64.b64encode(official).decode(),
            "shipment_id": "FBA13NOTE",
            "warehouse_code": "",
            "template_sheet": "Sheet1",
            "fields": [{"cell": "E5", "value": "NOPE", "key": "E5", "label": "E5"}],
        },
    )
    assert exported.status_code == 400, exported.text[:400]
    assert exported.json()["detail"]["code"] == "MERGED_SLAVE"
    assert "E5" in exported.json()["detail"]["cells"]


def test_send_step45_empty_reference_id_asks_for_pdf():
    from fastapi.testclient import TestClient

    from fba_tool.invoice_step45 import REFERENCE_ID_EMPTY_HINT
    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku = "B13-REF"
    delete_product_by_sku(sku)
    _seed(sku, "X00TESTREF", "B0B13REF01", "型号乙")
    raw = _csv("FBA13REF1", sku, "X00TESTREF", "B0B13REF01")
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13REF1",
                "forwarder_id": "send",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 200, generated.text[:400]
        assert REFERENCE_ID_EMPTY_HINT in (generated.json().get("warnings") or [])
    finally:
        delete_product_by_sku(sku)


def test_rollback_restores_sql_null_not_empty_text(monkeypatch, tmp_path):
    from fba_tool.lingxing_rollback import record_import_changes, rollback_lingxing_run
    from fba_tool.master_db import _connect, delete_product_by_sku, get_product_by_sku, init_master_db, upsert_product
    from fba_tool.master_db_lingxing import create_lingxing_sync_run, finish_lingxing_sync_run
    from fba_tool.product_fields import SOURCE_LINGXING, set_product_fields

    ad = tmp_path / "app"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    init_master_db()
    sku = "RB-NULL-B13"
    delete_product_by_sku(sku)
    upsert_product({"sku": sku})
    with _connect() as conn:
        raw = conn.execute(
            "SELECT asin, english_name, brand, local_name, field_source_json FROM product_master WHERE sku = ?",
            (sku,),
        ).fetchone()
    assert list(raw) == [None, None, None, None, None]
    before = get_product_by_sku(sku)
    run_id = create_lingxing_sync_run()
    finish_lingxing_sync_run(run_id, "confirmed", [], {})
    changed = dict(before)
    changed["asin"] = "B0NULLB13"
    changed["english_name"] = "Temp"
    changed["brand"] = "Brand"
    upsert_product(changed, overwrite=True)
    set_product_fields(sku, {"local_name": "本地名"}, SOURCE_LINGXING)
    record_import_changes(run_id, sku, before, get_product_by_sku(sku))
    rollback_lingxing_run(run_id)
    with _connect() as conn:
        restored = conn.execute(
            "SELECT asin, english_name, brand, local_name, field_source_json FROM product_master WHERE sku = ?",
            (sku,),
        ).fetchone()
    assert list(restored) == [None, None, None, None, None]
    delete_product_by_sku(sku)


def _range_boxes(label: str) -> int:
    text = str(label)
    if "-" not in text:
        return 1
    start, end = text.split("-", 1)
    return int(end) - int(start) + 1


def _multi_sku_csv() -> bytes:
    return "\n".join([
        '"货件编号","FBA13DS01"',
        '"货件名称","B13-DESENS"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","4"',
        "",
        '"多件商品（4 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量","箱子 2 的商品数量","箱子 3 的商品数量","箱子 4 的商品数量"',
        '"B13-AAA","rod","B0B13AAA1","X00TESTAAA","新品","无","4","2","2","",""',
        '"B13-TANZI","rod","B0B13TNZ1","X00TESTTNZ","新品","无","4","","","2","2"',
        "",
        '"","","","","","","箱号","FBA13DS01U000001","FBA13DS01U000002","FBA13DS01U000003","FBA13DS01U000004"',
        '"","","","","","","包装箱重量（磅）：","10","10","10","10"',
        '"","","","","","","包装箱长度（英寸）：","17.70","17.70","17.70","17.70"',
        '"","","","","","","包装箱宽度（英寸）：","17.70","17.70","17.70","17.70"',
        '"","","","","","","包装箱高度（英寸）：","16.10","16.10","16.10","16.10"',
    ]).encode("utf-8")


def test_qiluo_multi_sku_rows_and_declared_total():
    """同一型号的两个 SKU 各占一行，申报总额含 TANZI 的 49.99。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    skus = (
        ("B13-AAA", "X00TESTAAA", "B0B13AAA1", 10),
        ("B13-TANZI", "X00TESTTNZ", "B0B13TNZ1", 49.99),
    )
    for sku, fnsku, asin, price in skus:
        delete_product_by_sku(sku)
        _seed(sku, fnsku, asin, "共用型号", listing_price=price)
    raw = _multi_sku_csv()
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13DS01",
                "forwarder_id": "qiluo",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 200, generated.text
        cells = generated.json()["cells"]
        by_field: dict[str, list] = {}
        for item in cells:
            by_field.setdefault(item.get("field_key"), []).append(item.get("value"))
        assert by_field.get("box_range") == ["1-2", "3-4"]
        assert by_field.get("model") == ["共用型号", "共用型号"]
        prices = [float(value) for value in by_field.get("declared_price") or []]
        qty = [float(value) for value in by_field.get("qty_per_box") or []]
        assert prices == [10.0, 49.99]
        assert qty == [2.0, 2.0]
        total = sum(price * per * _range_boxes(label) for price, per, label in zip(prices, qty, by_field["box_range"]))
        assert round(total, 2) == 239.96
        located = {item["cell"]: item.get("value") for item in cells}
        assert located.get("K2") == "共用型号"
        assert located.get("K3") == "共用型号"
        assert located.get("G2") == 10 or located.get("G2") == 10.0
        assert float(located.get("G3")) == 49.99
    finally:
        for sku, *_rest in skus:
            delete_product_by_sku(sku)


@pytest.mark.parametrize("forwarder_id", FORWARDERS)
def test_empty_declared_price_blocks_every_forwarder(forwarder_id):
    """空申报价有报关风险。六家生成都 409，qiluo 不再只警告后继续导出。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku, fnsku, asin = "B13-NOPRICE", "X00TESTNOP", "B0B13NOP01"
    delete_product_by_sku(sku)
    _seed(sku, fnsku, asin, "型号乙", listing_price=None, declared_value=None)
    raw = _csv("FBA13NOPR", sku, fnsku, asin)
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13NOPR",
                "forwarder_id": forwarder_id,
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 409, generated.text[:500]
        detail = generated.json()["detail"]
        blocking = detail["blocking"]
        assert any(item.get("code") == "price_missing" for item in blocking)
        text = json.dumps(blocking, ensure_ascii=False)
        assert "不能导出" in text
        assert sku in text or fnsku in text
    finally:
        delete_product_by_sku(sku)


def test_send_and_jiufang_take_asin_from_csv_when_product_asin_is_empty():
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku, get_product_by_sku
    from main import app

    sku, fnsku, asin = "B13-NOASIN", "X00TESTNAS", "B0CSVASIN1"
    delete_product_by_sku(sku)
    _seed(sku, fnsku, "", "型号乙")
    raw = _csv("FBA13DS01", sku, fnsku, asin)
    client = TestClient(app)
    try:
        for forwarder_id in ("send", "jiufang"):
            generated = client.post(
                "/api/invoice/generate",
                json={
                    "shipment_id": "FBA13DS01",
                    "forwarder_id": forwarder_id,
                    "csv_base64": base64.b64encode(raw).decode(),
                },
            )
            assert generated.status_code == 200, f"{forwarder_id} {generated.text}"
            assert "缺发票必填字段" not in generated.text
        exported = client.post(
            "/api/invoice-export/export",
            data={"forwarder_id": "send", "reference_id": "R"},
            files={"csv_file": ("ship.csv", raw, "text/csv")},
        )
        assert exported.status_code == 200, exported.text[:400]
        wb = _book(exported.content)
        try:
            sheet = next(ws for ws in wb.worksheets if "Invoice" in ws.title)
            assert sheet["N15"].value == asin
        finally:
            wb.close()
        jiufang = client.post(
            "/api/invoice-export/export",
            data={"forwarder_id": "jiufang", "reference_id": "R", "warehouse": "TST1", "order_no": "LCLX"},
            files={"csv_file": ("ship.csv", raw, "text/csv")},
        )
        assert jiufang.status_code == 200, jiufang.text[:400]
        book = _book(jiufang.content)
        try:
            assert book["Invoice"]["U25"].value == asin
        finally:
            book.close()
        stored = get_product_by_sku(sku) or {}
        assert not str(stored.get("asin") or "").strip()
    finally:
        delete_product_by_sku(sku)


def _dim_csv() -> bytes:
    return "\n".join([
        '"货件编号","FBA13DIM1"',
        '"货件名称","B13-DIM"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","1"',
        "",
        '"单件商品（1 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        '"B13-DIM","rod","B0B13DIM1","X00TESTDIM","新品","无","4","4"',
        "",
        '"","","","","","","箱号","FBA13DIM1U000001"',
        '"","","","","","","包装箱重量（磅）：","10"',
        '"","","","","","","包装箱长度（英寸）：","17.72"',
        '"","","","","","","包装箱宽度（英寸）：","17.72"',
        '"","","","","","","包装箱高度（英寸）：","16.14"',
    ]).encode("utf-8")


def test_send_dimension_text_uses_two_decimals():
    import re

    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku = "B13-DIM"
    delete_product_by_sku(sku)
    _seed(sku, "X00TESTDIM", "B0B13DIM1", "型号乙")
    client = TestClient(app)
    try:
        exported = client.post(
            "/api/invoice-export/export",
            data={"forwarder_id": "send"},
            files={"csv_file": ("dim.csv", _dim_csv(), "text/csv")},
        )
        assert exported.status_code == 200, exported.text[:400]
        wb = _book(exported.content)
        try:
            sheet = next(ws for ws in wb.worksheets if "Invoice" in ws.title)
            text = str(sheet["B15"].value)
        finally:
            wb.close()
    finally:
        delete_product_by_sku(sku)
    assert text == "45.01*45.01*41.00"
    assert re.fullmatch(r"\d+\.\d{2}(\*\d+\.\d{2}){2}", text)


def _lines_csv(count: int) -> bytes:
    headers = ["SKU", "商品名称", "ASIN", "FNSKU", "状况", "预处理类型", "商品总数"]
    headers += [f"箱子 {i} 的商品数量" for i in range(1, count + 1)]
    body = [",".join(f'"{cell}"' for cell in headers)]
    for index in range(1, count + 1):
        qty = ["4" if col == index else "" for col in range(1, count + 1)]
        row = [f"B13-L{index}", "rod", f"B0LINE{index:02d}", f"X00TESTL{index}", "新品", "无", "4", *qty]
        body.append(",".join(f'"{cell}"' for cell in row))
    marks = ["", "", "", "", "", "", "箱号"] + [f"FBA13LINEU{index:06d}" for index in range(1, count + 1)]
    weight = ["", "", "", "", "", "", "包装箱重量（磅）："] + ["10"] * count
    length = ["", "", "", "", "", "", "包装箱长度（英寸）："] + ["10"] * count
    width = ["", "", "", "", "", "", "包装箱宽度（英寸）："] + ["10"] * count
    height = ["", "", "", "", "", "", "包装箱高度（英寸）："] + ["10"] * count
    return "\n".join([
        '"货件编号","FBA13LINE"',
        '"货件名称","B13-LINE"',
        '"配送地址","STEP45VILLE, ZZ"',
        f'"箱子数量","{count}"',
        "",
        f'"多件商品（{count} 箱）"' if count > 1 else '"单件商品（1 箱）"',
        *body,
        "",
        ",".join(f'"{cell}"' for cell in marks),
        ",".join(f'"{cell}"' for cell in weight),
        ",".join(f'"{cell}"' for cell in length),
        ",".join(f'"{cell}"' for cell in width),
        ",".join(f'"{cell}"' for cell in height),
    ]).encode("utf-8")


def _poison_jiufang(raw: bytes) -> bytes:
    wb = _book(raw)
    inv = wb["Invoice"]
    for ref in ("M26:N26", "M27:N27"):
        if ref not in {str(merged) for merged in inv.merged_cells.ranges}:
            inv.merge_cells(ref)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _body_mn(ws) -> list[str]:
    return sorted(
        str(merged)
        for merged in ws.merged_cells.ranges
        if merged.min_col == 13 and merged.max_col == 14 and merged.min_row >= 25
    )


def _english_name_csv(title: str) -> bytes:
    return "\n".join([
        '"货件编号","FBAENNAME"',
        '"货件名称","PLAIN"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","1"',
        "",
        '"单件商品（1 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        f'"B5-EN","{title}","B0B5ENNAME","X00TESTEN1","新品","无","4","4"',
        "",
        '"","","","","","","箱号","FBAENNAMEU000001"',
        '"","","","","","","包装箱重量（磅）：","10"',
        '"","","","","","","包装箱长度（英寸）：","10"',
        '"","","","","","","包装箱宽度（英寸）：","10"',
        '"","","","","","","包装箱高度（英寸）：","10"',
    ]).encode("utf-8")


def test_qiluo_and_send_english_name_prefer_catalog():
    """qiluo 与 SEND 都用产品库英文名。库里没有时才用货件标题，并且不写回产品库。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku, get_product_by_sku, upsert_product
    from main import app

    sku = "B5-EN"
    title = "Lingxing Listing Title"
    catalog = "Catalog Rod"
    delete_product_by_sku(sku)
    _seed(sku, "X00TESTEN1", "B0B5ENNAME", "型号乙", english_name=catalog)
    raw = _english_name_csv(title)
    client = TestClient(app)

    def qiluo_english() -> str:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBAENNAME",
                "forwarder_id": "qiluo",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 200, generated.text[:400]
        cells = generated.json()["cells"]
        named = [item.get("value") for item in cells if item.get("field_key") == "english_name"]
        assert named, cells
        return str(named[0])

    def send_english() -> str:
        exported = client.post(
            "/api/invoice-export/export",
            data={"forwarder_id": "send"},
            files={"csv_file": ("en.csv", raw, "text/csv")},
        )
        assert exported.status_code == 200, exported.text[:400]
        book = _book(exported.content)
        try:
            sheet = next(ws for ws in book.worksheets if "Invoice" in ws.title)
            return str(sheet["G15"].value)
        finally:
            book.close()

    try:
        assert qiluo_english() == catalog
        assert send_english() == catalog
        stored = get_product_by_sku(sku) or {}
        assert stored.get("english_name") == catalog

        upsert_product({"sku": sku, "english_name": ""}, overwrite=True)
        assert not str((get_product_by_sku(sku) or {}).get("english_name") or "").strip()
        assert qiluo_english() == title
        assert send_english() == title
        assert not str((get_product_by_sku(sku) or {}).get("english_name") or "").strip()
    finally:
        delete_product_by_sku(sku)


def _insert_warehouse(code: str):
    from fba_tool.master_db import _connect

    now = datetime.now(timezone.utc).isoformat()
    with _connect() as conn:
        conn.execute("DELETE FROM fba_warehouse WHERE code = ?", (code,))
        conn.execute(
            """
            INSERT INTO fba_warehouse
                (code, name, company, address1, city, state, zip, country, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (code, code, "Amazon", "1 Test Road", "STEP45VILLE", "ZZ", "00000", "US", now),
        )
        conn.commit()


def test_generate_multi_warehouse_422_names_candidates_in_chinese():
    """同城多仓时生成接口返回可选仓库，文案是中文，指定仓库后不再 422。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import _connect, delete_product_by_sku
    from main import app

    sku = "B4-WH"
    _seed(sku, "X00TESTB4W", "B0B4WH0001", "型号乙")
    _insert_warehouse("TST2")
    raw = _csv("FBAB4WH", sku, "X00TESTB4W", "B0B4WH0001").decode("utf-8").replace(
        '"货件名称","B13-DESENS"',
        '"货件名称","PLAIN"',
    ).encode("utf-8")
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBAB4WH",
                "forwarder_id": "jiufang",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert generated.status_code == 422, generated.text[:500]
        assert "Unprocessable" not in generated.text
        detail = generated.json()["detail"]
        assert detail["code"] == "warehouse_undecided"
        assert "选择仓库" in detail["message"]
        assert detail["address"] == "STEP45VILLE, ZZ"
        assert "STEP45VILLE, ZZ" in detail["message"]
        assert "TST1" in detail["warehouse_candidates"]
        assert "TST2" in detail["warehouse_candidates"]

        chosen = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBAB4WH",
                "forwarder_id": "jiufang",
                "warehouse": "TST2",
                "csv_base64": base64.b64encode(raw).decode(),
            },
        )
        assert chosen.status_code != 422, chosen.text[:500]
        assert chosen.status_code == 200, chosen.text[:500]
    finally:
        delete_product_by_sku(sku)
        with _connect() as conn:
            conn.execute("DELETE FROM fba_warehouse WHERE code = ?", ("TST2",))
            conn.commit()


def test_jiufang_total_merge_follows_the_total_row(monkeypatch):
    from openpyxl.cell.cell import MergedCell

    from fba_tool.invoice_export import _template_bytes, export_invoice, get_export_forwarder
    from fba_tool.shipment_parser import parse_csv_shipment

    fw = get_export_forwarder("jiufang")
    poisoned = _poison_jiufang(_template_bytes(fw))
    header_before = _book(poisoned)["Invoice"]
    kept = sorted(
        str(merged)
        for merged in header_before.merged_cells.ranges
        if merged.min_row < 25
    )
    header_before.parent.close()
    monkeypatch.setattr("fba_tool.invoice_export._template_bytes", lambda _fw: poisoned)

    def products(count: int) -> dict:
        out = {}
        for index in range(1, count + 1):
            fnsku = f"X00TESTL{index}"
            out[fnsku] = {
                "sku": f"B13-L{index}",
                "fnsku": fnsku,
                "asin": f"B0LINE{index:02d}",
                "listing_price": 10.0,
                "english_name": "Item",
                "chinese_name": "物品",
                "brand": "B",
                "material": "铁",
                "purpose": "用",
                "hs_code": "7326909000",
                "model_name": "无",
                "invoice": {
                    "unit": "个",
                    "purchase_price": 5,
                    "material_en": "iron",
                    "purpose_en": "use",
                    "pcs_per_set": 1,
                },
            }
        return out

    warehouse = {"address1": "1 Test Road", "city": "STEP45VILLE", "state": "ZZ", "zip": "00000"}
    origin = _book(poisoned)
    origin_font = origin["Invoice"]["M26"].font
    saved_font = (origin_font.bold, origin_font.name, origin_font.size)
    origin.close()
    for count in (1, 2, 3, 6):
        ship = parse_csv_shipment(_lines_csv(count), file_name="lines.csv")
        result = export_invoice(
            "jiufang",
            shipment=ship,
            products=products(count),
            reference_id="R",
            warehouse="TST1",
            warehouse_info=warehouse,
            order_no="LCLX",
        )
        wb = _book(result["bytes"])
        try:
            inv = wb["Invoice"]
            trow = 25 + count
            assert _body_mn(inv) == []
            assert "M27:N27" not in {str(merged) for merged in inv.merged_cells.ranges}
            assert "M31:N31" not in {str(merged) for merged in inv.merged_cells.ranges}
            assert inv[f"M{trow}"].value == "*Total（汇总）:"
            total_font = inv[f"M{trow}"].font
            assert (total_font.bold, total_font.name, total_font.size) == saved_font
            data_row = 24 + count
            assert not isinstance(inv[f"N{data_row}"], MergedCell)
            assert str(inv[f"N{data_row}"].value) == "7326909000"
            assert not inv[f"M{data_row}"].font.bold
            header_now = sorted(str(merged) for merged in inv.merged_cells.ranges if merged.min_row < 25)
            assert header_now == kept
        finally:
            wb.close()


def _ui_text(item: dict) -> str:
    """界面把格子读成字符串再导出。整数数量在这一步会变成 \"4.0\"。"""
    value = item.get("value")
    if value in (None, ""):
        return ""
    key = str(item.get("field_key") or "")
    if key in {"qty_per_box", "total_qty", "box_count", "pcs_per_set"}:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return str(value)
        if number == int(number):
            return f"{int(number)}.0"
    return str(value)


def _ui_export_body(body: dict) -> dict:
    from openpyxl.utils import coordinate_to_tuple

    fields = []
    cell_data: dict = {}
    for item in body["cells"]:
        text = _ui_text(item)
        fields.append({**item, "value": text, "key": item.get("field_key") or item.get("key") or ""})
        if not text:
            continue
        row, col = coordinate_to_tuple(str(item["cell"]))
        cell_data.setdefault(str(row - 1), {})[str(col - 1)] = {"v": text}
    snapshot = {
        "sheetOrder": ["s0"],
        "sheets": {"s0": {"name": body["template_sheet"], "cellData": cell_data}},
    }
    return {
        "template_hash": body["template_hash"],
        "forwarder_id": body.get("forwarder_id"),
        "shipment_id": body["shipment_id"],
        "warehouse_code": body["warehouse_code"],
        "template_sheet": body["template_sheet"],
        "fields": fields,
        "snapshot": snapshot,
    }


def _assert_numeric(cell, *, integer: bool = False, value=None) -> None:
    assert cell.data_type == "n", (cell.coordinate, repr(cell.value), cell.data_type)
    assert not isinstance(cell.value, str)
    if integer:
        assert type(cell.value) is int, (cell.coordinate, type(cell.value).__name__, cell.value)
    if value is not None:
        assert float(cell.value) == float(value)


@pytest.mark.parametrize("forwarder_id", FORWARDERS)
def test_numeric_columns_stay_numbers_after_ui_export(forwarder_id):
    """Step 4.5 把预览读成文本再导出时，数量和价格仍是数值格。整数数量是 int。"""
    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku
    from main import app

    sku, fnsku, asin = "B13-NUM", "X00TESTNUM", "B0B13NUM01"
    delete_product_by_sku(sku)
    _seed(sku, fnsku, asin, "型号乙", listing_price=39.99)
    raw = _csv("FBA13DS01", sku, fnsku, asin)
    client = TestClient(app)
    try:
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13DS01",
                "forwarder_id": forwarder_id,
                "csv_base64": base64.b64encode(raw).decode(),
                "po_number": "PO-13",
            },
        )
        assert generated.status_code == 200, generated.text[:400]
        body = generated.json()
        if forwarder_id == "qiluo":
            qty = next(item["value"] for item in body["cells"] if item.get("field_key") == "qty_per_box")
            price = next(item["value"] for item in body["cells"] if item.get("field_key") == "declared_price")
            assert type(qty) is int and qty == 4
            assert type(price) is float and price == 39.99
        exported = client.post("/api/export_invoice", json={**_ui_export_body(body), "forwarder_id": forwarder_id})
        assert exported.status_code == 200, exported.text[:400]
        wb = _book(exported.content)
        try:
            if forwarder_id == "qiluo":
                ws = wb["发票"]
                _assert_numeric(ws["F2"], integer=True, value=4)
                _assert_numeric(ws["G2"], value=39.99)
                for coord in ("M2", "N2", "O2", "P2"):
                    _assert_numeric(ws[coord])
                assert isinstance(ws["A2"].value, str)
                assert isinstance(ws["L2"].value, str)
            elif forwarder_id == "send":
                ws = next(sheet for sheet in wb.worksheets if "Invoice" in sheet.title)
                _assert_numeric(ws["F15"], integer=True, value=4)
                _assert_numeric(ws["R15"], value=39.99)
                _assert_numeric(ws["Q15"], integer=True, value=5)
                _assert_numeric(ws["S15"], value=39.99)
                _assert_numeric(ws["C15"])
                assert isinstance(ws["M15"].value, str)
            elif forwarder_id == "jiufang":
                inv, pk = wb["Invoice"], wb["Packing list"]
                _assert_numeric(inv["O25"], integer=True, value=4)
                _assert_numeric(inv["R25"], value=39.99)
                _assert_numeric(inv["Q25"], integer=True)
                _assert_numeric(inv["S25"], value=159.96)
                _assert_numeric(pk["H6"], integer=True, value=4)
                for coord in ("I6", "J6", "K6", "L6", "M6"):
                    _assert_numeric(pk[coord])
                assert isinstance(inv["N25"].value, str)
            elif forwarder_id == "qiluo_t6":
                ws = wb["发票"]
                _assert_numeric(ws["F14"], integer=True, value=4)
                _assert_numeric(ws["G14"], value=39.99)
                _assert_numeric(ws["B6"], integer=True, value=1)
                for coord in ("N14", "O14", "P14"):
                    _assert_numeric(ws[coord])
                assert isinstance(ws["L14"].value, str)
                assert isinstance(ws["A14"].value, str)
            elif forwarder_id == "qiluo_t5":
                ws = wb.worksheets[0]
                _assert_numeric(ws["H10"], integer=True, value=4)
                _assert_numeric(ws["I10"], value=39.99)
                _assert_numeric(ws["J10"], integer=True, value=1)
                for coord in ("K10", "L10", "M10", "N10"):
                    _assert_numeric(ws[coord])
                assert isinstance(ws["A10"].value, str)
            elif forwarder_id == "deshu":
                ws = wb["发票填写模板"]
                _assert_numeric(ws["N21"], integer=True, value=4)
                _assert_numeric(ws["R21"], value=39.99)
                _assert_numeric(ws["P21"], integer=True)
                _assert_numeric(ws["Q21"], integer=True, value=5)
                _assert_numeric(ws["G4"], integer=True, value=4)
                for coord in ("U21", "V21", "W21"):
                    _assert_numeric(ws[coord])
                assert isinstance(ws["G21"].value, str)
                assert isinstance(wb["填写示例"]["N21"].value, str)
        finally:
            wb.close()
    finally:
        delete_product_by_sku(sku)


def _cell_color(color) -> tuple | None:
    if color is None:
        return None
    if color.type == "rgb":
        return ("rgb", color.rgb)
    if color.type == "theme":
        return ("theme", color.theme, color.tint)
    if color.type == "indexed":
        return ("indexed", color.indexed)
    return (color.type, str(color.value))


def _border_side(side) -> tuple | None:
    if side is None or side.style is None:
        return None
    return (side.style, _cell_color(side.color))


def _cell_style(cell) -> tuple:
    from openpyxl.cell.cell import MergedCell

    if isinstance(cell, MergedCell):
        return ("MERGED",)
    font, fill, border, align = cell.font, cell.fill, cell.border, cell.alignment
    return (
        font.name,
        font.size,
        bool(font.bold),
        bool(font.italic),
        font.underline,
        _cell_color(font.color),
        fill.patternType,
        _cell_color(fill.fgColor),
        _cell_color(fill.bgColor),
        _border_side(border.left),
        _border_side(border.right),
        _border_side(border.top),
        _border_side(border.bottom),
        align.horizontal,
        align.vertical,
        bool(align.wrap_text),
        cell.number_format,
    )


def _style_diffs(blank, filled, sheet_name: str, data_start: int, prototype: int, total_proto: int, count: int) -> list[str]:
    src, dst = blank[sheet_name], filled[sheet_name]
    max_col = max(min(src.max_column or 1, 30), min(dst.max_column or 1, 30))
    diffs = []
    last = data_start + count
    for row in range(1, last + 1):
        if data_start <= row < data_start + count:
            template_row = prototype
        elif row == last:
            template_row = total_proto
        else:
            template_row = row
        if template_row < 1 or template_row > (src.max_row or 1):
            continue
        for col in range(1, max_col + 1):
            got = _cell_style(dst.cell(row, col))
            want = _cell_style(src.cell(template_row, col))
            if got != want:
                diffs.append(f"{sheet_name}!{dst.cell(row, col).coordinate} vs row {template_row}: {want} != {got}")
    return diffs


def test_jiufang_step45_styles_match_the_template():
    """九方 Step 4.5 逐格对照空白模板：数据行、合计行的字体、边框、对齐和数字格式不变。"""
    from fastapi.testclient import TestClient

    from fba_tool.invoice_export import _template_bytes, export_invoice, get_export_forwarder
    from fba_tool.master_db import delete_product_by_sku
    from fba_tool.shipment_parser import parse_csv_shipment
    from main import app

    fw = get_export_forwarder("jiufang")
    blank = _book(_template_bytes(fw))
    warehouse = {"address1": "1 Test Road", "city": "STEP45VILLE", "state": "ZZ", "zip": "00000"}

    def products(count: int) -> dict:
        out = {}
        for index in range(1, count + 1):
            fnsku = f"X00TESTY{index}"
            out[fnsku] = {
                "sku": f"B13-Y{index}",
                "fnsku": fnsku,
                "asin": f"B0STYL{index:02d}",
                "listing_price": 10.0,
                "english_name": "Item",
                "chinese_name": "物品",
                "brand": "B",
                "material": "铁",
                "purpose": "用",
                "hs_code": "7326909000",
                "model_name": "无",
                "invoice": {
                    "unit": "个",
                    "purchase_price": 5,
                    "material_en": "iron",
                    "purpose_en": "use",
                    "pcs_per_set": 1,
                },
            }
        return out

    try:
        for count in (1, 3):
            ship = parse_csv_shipment(_lines_csv(count).replace(b"B13-L", b"B13-Y").replace(b"X00TESTL", b"X00TESTY"), file_name="lines.csv")
            result = export_invoice(
                "jiufang",
                shipment=ship,
                products=products(count),
                reference_id="R",
                warehouse="TST1",
                warehouse_info=warehouse,
                order_no="LCLX",
            )
            filled = _book(result["bytes"])
            try:
                diffs = _style_diffs(blank, filled, "Invoice", 25, 25, 26, count)
                diffs += _style_diffs(blank, filled, "Packing list", 6, 6, 7, count)
                assert diffs == [], "\n".join(diffs[:12])
                assert filled["Packing list"]["H6"].number_format == blank["Packing list"]["H6"].number_format
                assert filled["Packing list"]["I6"].number_format == blank["Packing list"]["I6"].number_format
            finally:
                filled.close()

        sku = "B13-Y1"
        delete_product_by_sku(sku)
        _seed(sku, "X00TESTY1", "B0STYL01", "无", material="铁", purpose="用", listing_price=10)
        client = TestClient(app)
        generated = client.post(
            "/api/invoice/generate",
            json={
                "shipment_id": "FBA13LINE",
                "forwarder_id": "jiufang",
                "csv_base64": base64.b64encode(
                    _lines_csv(1).replace(b"B13-L", b"B13-Y").replace(b"X00TESTL", b"X00TESTY").replace(b"FBA13LINE", b"FBA13LINE")
                ).decode(),
                "po_number": "PO-13",
            },
        )
        assert generated.status_code == 200, generated.text[:400]
        body = generated.json()
        exported = client.post("/api/export_invoice", json={**_ui_export_body(body), "forwarder_id": "jiufang"})
        assert exported.status_code == 200, exported.text[:400]
        stepped = _book(exported.content)
        try:
            diffs = _style_diffs(blank, stepped, "Invoice", 25, 25, 26, 1)
            diffs += _style_diffs(blank, stepped, "Packing list", 6, 6, 7, 1)
            assert diffs == [], "\n".join(diffs[:12])
        finally:
            stepped.close()
            delete_product_by_sku(sku)
    finally:
        blank.close()
