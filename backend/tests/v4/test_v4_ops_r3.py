"""代码方 r2：第 14 批包装箱名称、填写失败禁导出、分步进度、V4、端口 8010。"""
from __future__ import annotations

import csv
import io
import re
import threading
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
FORWARDERS = ("send", "jiufang", "qiluo", "qiluo_t5", "qiluo_t6", "deshu")


def _sta_batch14_csv() -> bytes:
    """第 14 批 Send to Amazon 的形状：6 个 SKU，箱号行是「包装箱名称 / P1 - Bn」，没有「箱号」或 FBA…U。"""
    rows: list[list[object]] = [
        ["货件编号", "FBATEST14B"],
        ["货件名称", "BATCH14-TST1"],
        ["配送地址", "STEP45VILLE, ZZ"],
        ["箱子数量", "2"],
        [],
        ["SKU", "商品名称", "ASIN", "FNSKU", "状况", "预处理类型", "商品预处理方是谁？", "商品贴标方是谁？", "预计数量", "装箱数量", "", "包装箱 1 数量", "包装箱 2 数量"],
    ]
    for index in range(1, 7):
        rows.append([
            f"STA14-SKU-{index}",
            "item",
            f"B0STA14{index:02d}",
            f"X00TEST14{index}",
            "NewItem",
            "贴标",
            "",
            "由卖家提供",
            4,
            4,
            "",
            4 if index <= 3 else "",
            2 if index > 3 else "",
        ])
    rows.append([])

    def foot(label: str, left: object, right: object) -> list[object]:
        return [label, *( [""] * 10 ), left, right]

    rows.extend([
        foot("包装箱名称", "P1 - B1", "P1 - B2"),
        foot("包装箱重量（磅）：", "10", "11"),
        foot("包装箱宽度（英寸）：", "17.70", "10"),
        foot("包装箱长度（英寸）：", "17.70", "12"),
        foot("包装箱高度（英寸）：", "16.10", "8"),
    ])
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8")


def _seed_sta():
    import json
    from datetime import datetime, timezone

    from fba_tool.master_db import _connect, init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields

    init_master_db()
    now = datetime.now(timezone.utc).isoformat()
    with _connect() as conn:
        conn.execute("DELETE FROM fba_warehouse WHERE code = ?", ("TST1",))
        conn.execute(
            """
            INSERT INTO fba_warehouse
                (code, name, company, address1, city, state, zip, country, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("TST1", "TST1", "Amazon", "1 Test Road", "STEP45VILLE", "ZZ", "00000", "US", now),
        )
        conn.commit()
    for index in range(1, 7):
        sku = f"STA14-SKU-{index}"
        upsert_product(
            {
                "sku": sku,
                "fnsku": f"X00TEST14{index}",
                "asin": f"B0STA14{index:02d}",
                "english_name": "Fishing rod",
                "chinese_name": "鱼竿",
                "brand": "Acme",
                "material": "棉",
                "purpose": "宠物",
                "hs_code": "9507100090",
                "model_name": "型号甲",
                "extra_data": json.dumps(
                    {"invoice": {"unit": "个", "purchase_price": 5, "material_en": "cotton", "purpose_en": "pet", "pcs_per_set": 1}},
                    ensure_ascii=False,
                ),
            }
        )
        set_product_fields(sku, {"listing_price": 12.5, "listing_price_currency": "USD"}, SOURCE_MANUAL)


def test_batch14_packaging_name_row_maps_p1_boxes_and_keeps_inches():
    from fba_tool.shipment_parser import parse_csv_shipment

    ship = parse_csv_shipment(_sta_batch14_csv(), file_name="第14批.csv", allow_unassigned_fba=True)
    assert ship.shipment_id == "FBATEST14B"
    assert ship.column_to_box
    assert set(ship.column_to_box.values()) == {1, 2}
    assert ship.box_count == 2
    assert not ship.box_items.empty
    assert int(ship.box_items["数量"].sum()) == 4 * 3 + 2 * 3
    marks = ship.box_specs["FBA箱号"].fillna("").astype(str).str.strip()
    assert list(marks) == ["", ""]
    assert set(ship.box_specs["箱名"].astype(str)) == {"P1 - B1", "P1 - B2"}
    assert ship.box_names == {1: "P1 - B1", 2: "P1 - B2"}
    assert ship.box_fba_ids == {}
    width = float(ship.box_specs.loc[ship.box_specs["箱号"] == 1, "宽(cm)"].iloc[0])
    length = float(ship.box_specs.loc[ship.box_specs["箱号"] == 1, "长(cm)"].iloc[0])
    assert width == 44.96
    assert length == 44.96


def test_send_quantity_is_sku_total_not_per_box():
    """同一 SKU 的 SEND 数量是各箱数量加总，不是单箱数量。重量差 0.003 kg 仍合并成一行。"""
    import io

    from openpyxl import load_workbook

    from fba_tool.invoice_export import SEND_SHEET, export_invoice
    from fba_tool.invoice_rules import Box, box_ranges, sku_totals
    from fba_tool.shipment_parser import parse_csv_shipment

    per_box = 4
    box_count = 8
    dims = (45.0, 45.0, 41.0)
    near = [
        Box(i, f"FBATEST8QU{i:06d}", [("X00TEST8Q", per_box)], 10.004 if i % 2 else 10.001, dims)
        for i in range(1, box_count + 1)
    ]
    sku_total = sum(qty for _sku, qty in sku_totals(near))
    assert sku_total == per_box * box_count
    assert sku_total != per_box
    assert [(row.label, row.qty) for row in box_ranges(near)] == [(f"1-{box_count}", sku_total)]
    apart = box_ranges([
        Box(1, "", [("A", per_box)], 7.5, dims),
        Box(2, "", [("A", per_box)], 7.6, dims),
    ])
    assert [(row.label, row.weight) for row in apart] == [("1", 7.5), ("2", 7.6)]

    def quoted(cells: list[object]) -> str:
        return ",".join(f'"{cell}"' for cell in cells)

    headers = ["SKU", "商品名称", "ASIN", "FNSKU", "状况", "预处理类型", "商品总数"]
    headers += [f"箱子 {i} 的商品数量" for i in range(1, box_count + 1)]
    line = ["QTY-SKU", "rod", "B0QTY00001", "X00TEST8Q", "新品", "无", str(sku_total), *[str(per_box)] * box_count]
    marks = ["", "", "", "", "", "", "箱号"] + [f"FBATEST8QU{i:06d}" for i in range(1, box_count + 1)]
    weight = ["", "", "", "", "", "", "包装箱重量（kg）："] + ["10"] * box_count
    length = ["", "", "", "", "", "", "包装箱长度（cm）："] + ["40"] * box_count
    width = ["", "", "", "", "", "", "包装箱宽度（cm）："] + ["40"] * box_count
    height = ["", "", "", "", "", "", "包装箱高度（cm）："] + ["40"] * box_count
    raw = "\n".join([
        '"货件编号","FBATEST8Q"',
        '"货件名称","QTY"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","8"',
        "",
        '"单件商品（8 箱）"',
        quoted(headers),
        quoted(line),
        "",
        quoted(marks),
        quoted(weight),
        quoted(length),
        quoted(width),
        quoted(height),
    ]).encode()
    ship = parse_csv_shipment(raw, allow_unassigned_fba=True)
    parsed_total = int(ship.box_items["数量"].sum())
    assert parsed_total == sku_total
    result = export_invoice(
        "send",
        shipment=ship,
        products={
            "X00TEST8Q": {
                "sku": "QTY-SKU",
                "fnsku": "X00TEST8Q",
                "asin": "B0QTY00001",
                "listing_price": 10.0,
                "english_name": "Item",
                "chinese_name": "物品",
                "brand": "B",
                "material": "铁",
                "purpose": "用",
                "hs_code": "7326909000",
                "model_name": "无",
                "invoice": {
                    "unit": "个",
                    "purchase_price": 5,
                    "material_en": "iron",
                    "purpose_en": "use",
                    "pcs_per_set": 1,
                },
            }
        },
        reference_id="R",
    )
    wb = load_workbook(io.BytesIO(result["bytes"]))
    try:
        ws = wb[SEND_SHEET]
        assert ws["A15"].value == f"1-{box_count}"
        assert ws["F15"].value == parsed_total
        assert ws["F15"].value != per_box
        assert ws["F16"].value in (None, "")
    finally:
        wb.close()


def test_classic_fba_box_row_still_maps():
    from fba_tool.shipment_parser import parse_csv_shipment

    raw = "\n".join([
        '"货件编号","FBATESTBOX1"',
        '"货件名称","BOX-TST1"',
        '"配送地址","STEP45VILLE, ZZ"',
        '"箱子数量","1"',
        "",
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        '"BOX-SKU","rod","B0BOX00001","X00TESTBOX","新品","无","4","4"',
        "",
        '"","","","","","","箱号","FBATESTBOX1U000001"',
    ]).encode()
    ship = parse_csv_shipment(raw, allow_unassigned_fba=True)
    assert ship.column_to_box
    assert 1 in ship.column_to_box.values()
    assert ship.box_fba_ids.get(1, "").upper().endswith("U000001")


def test_workflow_fill_defines_rec_with_and_without_catalog():
    from fba_tool.invoice_workflow_fill import build_invoice_line_items_workflow
    from fba_tool.shipment_parser import parse_csv_shipment

    raw = _sta_batch14_csv()
    ship = parse_csv_shipment(raw, allow_unassigned_fba=True)

    class Rec:
        unit_cost = 3.5
        asin = "B0FROMREC"

    class Cat:
        def lookup(self, sku):
            return Rec() if str(sku).startswith("STA14-SKU-") else None

    filled = build_invoice_line_items_workflow(ship, raw, Cat(), None)
    assert filled
    assert all(item.values["purchase_price"] == 3.5 for item in filled)
    bare = build_invoice_line_items_workflow(ship, raw, None, None)
    assert bare
    assert bare[0].values["purchase_price"] == ""
    assert bare[0].values["asin"]


@pytest.mark.parametrize("forwarder_id", FORWARDERS)
def test_sta_shape_autofill_then_export(forwarder_id):
    import base64

    from fastapi.testclient import TestClient

    from main import app

    _seed_sta()
    raw = _sta_batch14_csv()
    client = TestClient(app)
    generated = client.post(
        "/api/invoice/generate",
        json={
            "shipment_id": "FBATEST14B",
            "forwarder_id": forwarder_id,
            "csv_base64": base64.b64encode(raw).decode(),
            "po_number": "PO-14",
        },
    )
    assert generated.status_code == 200, generated.text
    assert "未找到「箱号」" not in generated.text
    assert "is not defined" not in generated.text
    body = generated.json()
    exported = client.post(
        "/api/export_invoice",
        json={
            "template_hash": body["template_hash"],
            "forwarder_id": forwarder_id,
            "shipment_id": body["shipment_id"],
            "warehouse_code": body["warehouse_code"],
            "template_sheet": body["template_sheet"],
            "fields": body["cells"],
        },
    )
    assert exported.status_code == 200, exported.text
    assert exported.content[:2] == b"PK"
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=True)
    texts = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=min(ws.max_row or 1, 80), max_col=min(ws.max_column or 1, 30)):
            for cell in row:
                if cell.value not in (None, ""):
                    texts.add(str(cell.value).strip())
    wb.close()
    assert "FBATEST14B" in texts


def test_snapshot_progress_steps_before_100():
    import openpyxl

    from fba_tool.univer_snapshot import workbook_to_snapshot

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in range(1, 41):
        for col in range(1, 6):
            ws.cell(row, col, f"{row}-{col}")
    buf = io.BytesIO()
    wb.save(buf)
    seen: list[int] = []
    workbook_to_snapshot(buf.getvalue(), progress=seen.append)
    assert any(0 < value < 100 for value in seen), seen
    assert seen[-1] == 100
    assert seen == sorted(seen)


def test_job_poll_sees_progress_between_0_and_100():
    from fba_tool.job_runner import RUNNER

    started = threading.Event()
    release = threading.Event()

    def fn(job):
        job.progress = 35
        started.set()
        assert release.wait(2)
        job.progress = 75
        return {"ok": True}

    job, _done = RUNNER.submit(fn, sync_timeout=0)
    assert started.wait(2)
    assert RUNNER.get(job.id).progress == 35
    release.set()
    assert job.done.wait(2)
    assert job.status == "done"
    assert job.progress == 100


def test_run_app_candidate_ports_skip_8010():
    text = (ROOT / "run_app.py").read_text(encoding="utf-8")
    match = re.search(r"^PREFERRED_PORTS = (.+)$", text, re.M)
    assert match
    ports = list(eval(match.group(1), {"DEFAULT_PORT": 8002, "range": range}))
    assert ports[0] == 8002
    assert 8003 in ports and 8009 in ports and 8011 in ports and 18002 in ports
    assert 8010 not in ports


def _batch_csvs(root: Path) -> dict[str, list[Path]]:
    found = {"13": [], "14": []}
    for path in root.rglob("*.csv"):
        if path.name.startswith("~$"):
            continue
        text = str(path)
        if "第13批" in text or "/13/" in text.replace("\\", "/"):
            found["13"].append(path)
        elif "第14批" in text or "/14/" in text.replace("\\", "/"):
            found["14"].append(path)
    return found


def _invoice_files(root: Path) -> list[Path]:
    found = []
    for path in root.rglob("*"):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        name = path.name.lower()
        if any(token in path.name for token in ("发票", "箱单", "Invoice", "invoice")):
            found.append(path)
        elif "send" in name and "amazon" not in name:
            found.append(path)
    return found


def _whole_qty(value) -> int | None:
    """52.0、'52.0'、52 都是 52。半件、空值和布尔不算数量。"""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
    else:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
    if number != number or number <= 0 or not number.is_integer():
        return None
    return int(number)


def _csv_sku_totals(box_items) -> tuple[dict[str, int], dict[str, str]]:
    """CSV 按 SKU 把每箱数量加总。FNSKU 和 SKU 都指回同一个 SKU。"""
    totals: dict[str, int] = {}
    alias: dict[str, str] = {}
    if box_items is None or box_items.empty:
        return totals, alias
    for _, rec in box_items.iterrows():
        sku = str(rec.get("SKU") or "").strip()
        fnsku = str(rec.get("FNSKU") or "").strip()
        qty = _whole_qty(rec.get("数量"))
        if not sku or qty is None:
            continue
        totals[sku] = totals.get(sku, 0) + qty
        alias[sku.casefold()] = sku
        if fnsku:
            alias[fnsku.casefold()] = sku
    return totals, alias


def _qty_kind(header: str) -> str | None:
    from fba_tool.invoice_templates import norm

    key = norm(header).casefold()
    if not key or "每套个数" in key or "qtyperunit" in key:
        return None
    if any(token in key for token in ("每箱数量", "单箱内数量", "箱内数量")):
        return "per_box"
    if any(token in key for token in ("sku个数", "totalqty", "总数量", "商品总数", "产品个数", "numberofunits")):
        return "row_total"
    if "个数" in key or key in {"qty", "quantity", "数量", "商品数量"}:
        return "row_total"
    if key.endswith("数量") and "箱" not in key and "重量" not in key:
        return "row_total"
    return None


def _is_range_header(header: str) -> bool:
    from fba_tool.invoice_templates import norm

    key = norm(header).casefold()
    if not key or "数量" in key:
        return False
    return any(token in key for token in ("箱序号", "箱号段", "货箱编号", "箱号box", "起始箱号", "boxno"))


def _is_ident_header(header: str) -> bool:
    from fba_tool.invoice_templates import norm

    return norm(header).casefold() in {"sku", "fnsku", "msku"}


def _is_detail_header(header: str) -> bool:
    from fba_tool.invoice_templates import norm

    key = norm(header).casefold()
    return any(token in key for token in ("品名", "海关编码", "hscode", "fbaid", "shipmentid", "英文", "中文"))


def _canon_ident(value, alias: dict[str, str]) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value or "").strip()
    if not text:
        return None
    return alias.get(text.casefold())


def _qty_parts(value) -> list[int]:
    if isinstance(value, str) and any(sep in value for sep in (",", "，")):
        parts = []
        for piece in re.split(r"[,，]", value):
            qty = _whole_qty(piece)
            if qty is not None:
                parts.append(qty)
        return parts
    qty = _whole_qty(value)
    return [] if qty is None else [qty]


def _sheet_sku_totals(rows, alias: dict[str, str]) -> tuple[dict[str, int], list[int]]:
    """一张表：有 SKU/FNSKU 就按它加总；琪珞这种不写 SKU 的表，返回每行合计。"""
    from fba_tool.invoice_templates import parse_box_range

    header_at = None
    header = None
    fallback = None
    for index, row in enumerate(rows[:40]):
        qty_at = None
        ident_at = None
        range_at = None
        detail = False
        for col, value in enumerate(row):
            if not isinstance(value, str):
                continue
            kind = _qty_kind(value)
            if kind == "row_total":
                qty_at = (col, kind)
            elif kind == "per_box" and (qty_at is None or qty_at[1] != "row_total"):
                qty_at = (col, kind)
            if _is_ident_header(value):
                ident_at = col
            if _is_range_header(value) and range_at is None:
                range_at = col
            if _is_detail_header(value):
                detail = True
        if qty_at and (ident_at is not None or range_at is not None or detail):
            found = (index, qty_at, ident_at, range_at)
            fallback = found
            if ident_at is not None:
                header_at, header = index, found
    chosen = header if header is not None else fallback
    if chosen is None:
        return {}, []
    start, (qty_col, kind), ident_col, range_col = chosen
    keyed: dict[str, int] = {}
    anonymous: list[int] = []
    blank = 0
    for row in rows[start + 1:start + 181]:
        if not any(cell not in (None, "") for cell in row):
            blank += 1
            if blank >= 3:
                break
            continue
        blank = 0
        ident = _canon_ident(row[ident_col], alias) if ident_col is not None and ident_col < len(row) else None
        if ident is None:
            for cell in row:
                ident = _canon_ident(cell, alias)
                if ident:
                    break
        if ident is None and any(isinstance(cell, str) and ("汇总" in cell or "total" in cell.casefold()) for cell in row):
            break
        parts = _qty_parts(row[qty_col] if qty_col < len(row) else None)
        if not parts:
            continue
        if kind == "per_box":
            span = 1
            if range_col is not None and range_col < len(row):
                try:
                    boxes = parse_box_range(row[range_col])
                except ValueError:
                    boxes = []
                if boxes:
                    span = len(boxes)
            parts = [part * span for part in parts]
        if ident:
            keyed[ident] = keyed.get(ident, 0) + sum(parts)
        elif ident_col is None:
            anonymous.extend(parts)
    return keyed, anonymous


def _qty_compare(data: bytes, expected: dict[str, int], alias: dict[str, str]) -> tuple[bool, str]:
    """发票或导出里的数量先变成整数，再按每个 SKU 的合计和 CSV 比。"""
    import openpyxl

    if not expected:
        return False, "CSV 没有正的 SKU 数量"
    target = sorted(expected.values())
    details: list[str] = []
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    try:
        for ws in wb.worksheets:
            rows = []
            max_row = min(ws.max_row or 1, 220)
            max_col = min(ws.max_column or 1, 40)
            for row in ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True):
                rows.append(row)
            keyed, anonymous = _sheet_sku_totals(rows, alias)
            if keyed == expected:
                return True, ""
            if not keyed and anonymous and sorted(anonymous) == target:
                return True, ""
            if keyed:
                details.append(f"{ws.title} 按 SKU {keyed}")
            elif anonymous:
                details.append(f"{ws.title} 按行 {anonymous}")
    finally:
        wb.close()
    return False, "读到 " + ("；".join(details) if details else "没有数量")


def _workbook_texts(data: bytes) -> set[str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    texts: set[str] = set()
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=min(ws.max_row or 1, 120), max_col=min(ws.max_column or 1, 40)):
                for cell in row:
                    if cell.value not in (None, ""):
                        texts.add(str(cell.value).strip())
    finally:
        wb.close()
    return texts


def _catalog_seed(rec, meta: dict) -> dict | None:
    """真实批次种子：ASIN 来自 CSV 的 SKU 行，没有再看装箱行。"""
    import json

    from fba_tool.sku_keys import lookup_sku

    sku_text = str(rec.get("SKU") or "").strip()
    if not sku_text:
        return None
    hit = lookup_sku(meta, sku_text, {}) or {}
    asin = str(hit.get("asin") or "").strip()
    if not asin:
        for key in getattr(rec, "index", []):
            if str(key).strip().upper() == "ASIN":
                asin = str(rec.get(key) or "").strip()
                break
    fnsku = str(rec.get("FNSKU") or hit.get("fnsku") or "").strip().upper()
    return {
        "sku": sku_text,
        "fnsku": fnsku,
        "asin": asin,
        "english_name": "Item",
        "chinese_name": "物品",
        "brand": "Brand",
        "material": "棉",
        "purpose": "宠物",
        "hs_code": "9507100090",
        "model_name": "M",
        "extra_data": json.dumps(
            {"invoice": {"unit": "个", "purchase_price": 5, "material_en": "cotton", "purpose_en": "pet", "pcs_per_set": 1}},
            ensure_ascii=False,
        ),
    }


def test_real_batch_seed_copies_asin_from_csv():
    from fba_tool.invoice_table import extract_csv_sku_meta
    from fba_tool.shipment_parser import parse_csv_shipment

    raw = _sta_batch14_csv()
    ship = parse_csv_shipment(raw, file_name="batch14.csv", allow_unassigned_fba=True)
    meta = extract_csv_sku_meta(raw)
    seeds = []
    for _, rec in ship.box_items.iterrows():
        seed = _catalog_seed(rec, meta)
        if seed:
            seeds.append(seed)
    assert seeds
    assert all(item["asin"].startswith("B0STA14") for item in seeds)
    assert all(item["fnsku"].startswith("X00TEST14") for item in seeds)


def _warehouse_xlsx(rows: list[tuple]) -> bytes:
    """琪珞「！」表。列顺序与地址库表头一致，合成夹具用，不读真实文件。"""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "！"
    sheet.append(["FBA仓码", "收件人名称", "收件人公司", "邮编", "地址1", "地址2", "地址3", "国家", "城市", "州", "收件人电话"])
    for row in rows:
        sheet.append(list(row))
    buf = io.BytesIO()
    book.save(buf)
    book.close()
    return buf.getvalue()


def _api_warehouse_code(raw: bytes) -> str:
    """生成接口把文件名换成货代 id，路径里的仓码不会进入解析。"""
    from fba_tool.shipment_parser import parse_csv_shipment

    ship = parse_csv_shipment(raw, file_name="invoice.csv", allow_unassigned_fba=True)
    code = (ship.warehouse_code or "").strip().upper()
    return "" if code in ("", "UNKN") else code


def _shared_address_candidates(forwarder_id: str, ship_to: str, csv_code: str) -> list[str] | None:
    """一个地址对应多个仓库、且 CSV 仓码不在其中时，要写派送地址的货代返回这些候选。SEND 不写地址。"""
    from fba_tool.invoice_export import export_forwarders
    from fba_tool.invoice_templates import _ship_to_codes

    adapter = str((export_forwarders().get(forwarder_id) or {}).get("export_adapter") or "")
    if adapter in ("", "send"):
        return None
    codes = _ship_to_codes(ship_to)
    cc = (csv_code or "").strip().upper()
    if cc == "UNKN":
        cc = ""
    if len(codes) > 1 and cc not in codes:
        return codes
    return None


def test_real_batches_13_and_14_autofill_export_matches_invoices(monkeypatch, tmp_path):
    """真实目录在时先导入仓库地址库，再跑六个货代的自动填写→导出，并和同票发票按 SKU 合计对数量。

    原发票把合计存成 52.0 时按 52 比，不和单箱数量 "4" 或整票 "84" 做字符串交。
    一个地址对应多个仓库、货件名称里的仓码又不在其中（第 14 批 GYR2 / LGB8 同城多仓就是这种）时，
    九方等要写地址的货代返回 422 和候选仓库列表，不按 200 失败。
    """
    from tests.smoke.reference import LayoutMissing, available, realdata_root

    if not available():
        pytest.skip("需要 box realdata")
    monkeypatch.setenv("FBA_APP_DATA_DIR", str(tmp_path / "app"))
    monkeypatch.setenv("FBA_DISABLE_SCHEDULER", "1")
    root = realdata_root()
    warehouse_file = root / "tests" / "fixtures" / "warehouse" / "fba_warehouse_qiluo.xlsx"
    if not warehouse_file.is_file():
        pytest.skip("需要 box realdata: " + str(warehouse_file))
    csvs = _batch_csvs(root)
    missing = [batch for batch, paths in csvs.items() if not paths]
    if missing:
        raise LayoutMissing("缺少第 " + "、".join(missing) + " 批货件 CSV")

    import base64

    from fastapi.testclient import TestClient

    from fba_tool.master_db import init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields
    from fba_tool.shipment_parser import parse_csv_shipment
    from fba_tool.warehouse_db import apply_warehouse_import
    from main import create_app

    init_master_db()
    imported = apply_warehouse_import(warehouse_file.read_bytes(), warehouse_file.name)
    assert imported["written"] > 0
    client = TestClient(create_app())
    invoices = _invoice_files(root)
    compared = []
    pipeline = []
    for batch, paths in csvs.items():
        for path in paths:
            raw = path.read_bytes()
            try:
                ship = parse_csv_shipment(raw, file_name=path.name, allow_unassigned_fba=True)
            except ValueError as exc:
                raise AssertionError(f"{path.name} 解析失败：{exc}") from exc
            if ship.box_items.empty:
                raise AssertionError(f"{path.name} 没有箱内数量")
            from fba_tool.invoice_table import extract_csv_sku_meta

            meta = extract_csv_sku_meta(raw)
            for _, rec in ship.box_items.iterrows():
                seed = _catalog_seed(rec, meta)
                if not seed:
                    continue
                upsert_product(seed)
                set_product_fields(seed["sku"], {"listing_price": 10, "listing_price_currency": "USD"}, SOURCE_MANUAL)
            expected, alias = _csv_sku_totals(ship.box_items)
            assert expected, f"{path.name} 没有正的 SKU 数量"
            api_code = _api_warehouse_code(raw)
            for forwarder_id in FORWARDERS:
                generated = client.post(
                    "/api/invoice/generate",
                    json={
                        "shipment_id": ship.shipment_id,
                        "forwarder_id": forwarder_id,
                        "csv_base64": base64.b64encode(raw).decode(),
                    },
                )
                candidates = _shared_address_candidates(forwarder_id, ship.ship_to, api_code)
                if candidates is not None:
                    assert generated.status_code == 422, (
                        f"{batch} {path.name} {forwarder_id} 配送地址对应多个仓库，应 422，"
                        f"实际 {generated.status_code} {generated.text[:400]}"
                    )
                    detail = generated.json()["detail"]
                    assert detail["code"] == "warehouse_undecided"
                    assert detail["address"] == ship.ship_to
                    assert detail["warehouse_candidates"] == candidates, (
                        f"{batch} {path.name} {forwarder_id} 候选仓库应是该配送地址在地址库里的仓库"
                        f"（第 14 批 GYR2/LGB8 同城多仓即这种），实际 {detail.get('warehouse_candidates')}"
                    )
                    assert "对应多个仓库" in detail["message"]
                    assert "选择仓库" in detail["message"]
                    continue
                assert generated.status_code == 200, f"{batch} {path.name} {forwarder_id} {generated.text}"
                assert "箱号" not in generated.text
                body = generated.json()
                exported = client.post(
                    "/api/export_invoice",
                    json={
                        "template_hash": body["template_hash"],
                        "forwarder_id": forwarder_id,
                        "shipment_id": body["shipment_id"],
                        "warehouse_code": body["warehouse_code"],
                        "template_sheet": body["template_sheet"],
                        "fields": body["cells"],
                    },
                )
                assert exported.status_code == 200, f"{batch} {path.name} {forwarder_id} export {exported.text[:300]}"
                texts = _workbook_texts(exported.content)
                assert ship.shipment_id in texts, f"{forwarder_id} 导出没有货件号 {ship.shipment_id}"
                export_ok, export_detail = _qty_compare(exported.content, expected, alias)
                assert export_ok, (
                    f"{forwarder_id} 导出应按每个 SKU 的合计 {expected} 比对，"
                    f"52.0 记成 52，不用单箱数量或整票合计做字符串匹配。{export_detail}"
                )
                pipeline.append((batch, path.name, forwarder_id, ship.shipment_id))
                for invoice in invoices:
                    if ship.shipment_id and ship.shipment_id in invoice.name.upper().replace("-", ""):
                        invoice_bytes = invoice.read_bytes()
                        invoice_texts = _workbook_texts(invoice_bytes)
                        assert ship.shipment_id in invoice_texts
                        invoice_ok, invoice_detail = _qty_compare(invoice_bytes, expected, alias)
                        assert invoice_ok, (
                            f"{invoice.name} 和 CSV 数量对不上：应按每个 SKU 的合计 {expected}，"
                            f"发票里的 52.0 记成 52，而不是去对字符串 \"4\" 或 \"84\"。{invoice_detail}"
                        )
                        compared.append((forwarder_id, ship.shipment_id, invoice.name))
    assert pipeline
    if invoices and not compared:
        raise AssertionError("找到了发票文件，但没有一份能按货件号和第 13/14 批 CSV 对上：" + ", ".join(path.name for path in invoices[:8]))


def _plain_csv(shipment: str, ship_to: str) -> bytes:
    """货件名称不含仓码。生成接口也看不到文件名里的仓码，仓库只能靠配送地址。"""
    return "\n".join([
        f'"货件编号","{shipment}"',
        '"货件名称","PLAIN"',
        f'"配送地址","{ship_to}"',
        '"箱子数量","1"',
        "",
        '"单件商品（1 箱）"',
        '"SKU","商品名称","ASIN","FNSKU","状况","预处理类型","商品总数","箱子 1 的商品数量"',
        '"N5C-SKU","rod","B0N5C00001","X00TESTN5C","新品","无","4","4"',
        "",
        f'"","","","","","","箱号","{shipment}U000001"',
        '"","","","","","","包装箱重量（磅）：","10"',
        '"","","","","","","包装箱长度（英寸）：","17.70"',
        '"","","","","","","包装箱宽度（英寸）：","17.70"',
        '"","","","","","","包装箱高度（英寸）：","16.10"',
    ]).encode("utf-8")


def test_import_warehouse_library_then_fill(monkeypatch, tmp_path):
    """先导入地址库再填写。唯一地址的九方在导入前 422，导入后写出该仓；同城多仓返回 422 和候选。"""
    import base64
    import json

    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku, init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields
    from fba_tool.warehouse_db import apply_warehouse_import
    from main import create_app

    monkeypatch.setenv("FBA_APP_DATA_DIR", str(tmp_path / "app"))
    monkeypatch.setenv("FBA_DISABLE_SCHEDULER", "1")
    init_master_db()
    sku = "N5C-SKU"
    upsert_product(
        {
            "sku": sku,
            "fnsku": "X00TESTN5C",
            "asin": "B0N5C00001",
            "english_name": "Fishing rod",
            "chinese_name": "鱼竿",
            "brand": "Acme",
            "material": "棉",
            "purpose": "宠物",
            "hs_code": "9507100090",
            "model_name": "型号甲",
            "extra_data": json.dumps(
                {"invoice": {"unit": "个", "purchase_price": 5, "material_en": "cotton", "purpose_en": "pet", "pcs_per_set": 1}},
                ensure_ascii=False,
            ),
        }
    )
    set_product_fields(sku, {"listing_price": 10, "listing_price_currency": "USD"}, SOURCE_MANUAL)
    library = _warehouse_xlsx([
        ("ONL1", "ONL1", "Amazon", "00001", "1 Only Road", "", "", "US", "ONLYVILLE", "ZZ", ""),
        ("GYR2", "GYR2", "Amazon", "00002", "2 Share Road", "", "", "US", "SHAREVILLE", "ZZ", ""),
        ("LGB8", "LGB8", "Amazon", "00003", "3 Share Road", "", "", "US", "SHAREVILLE", "ZZ", ""),
    ])
    unique = _plain_csv("FBAN5CONLY", "ONLYVILLE, ZZ")
    shared = _plain_csv("FBAN5CSHARE", "SHAREVILLE, ZZ")
    client = TestClient(create_app())

    def generate(raw: bytes, shipment: str, warehouse: str = ""):
        body = {
            "shipment_id": shipment,
            "forwarder_id": "jiufang",
            "csv_base64": base64.b64encode(raw).decode(),
        }
        if warehouse:
            body["warehouse"] = warehouse
        return client.post("/api/invoice/generate", json=body)

    try:
        blocked = generate(unique, "FBAN5CONLY")
        assert blocked.status_code == 422, blocked.text[:400]
        assert blocked.json()["detail"]["code"] == "warehouse_undecided"

        imported = apply_warehouse_import(library, "synthetic-warehouse.xlsx")
        assert imported["written"] == 3

        filled = generate(unique, "FBAN5CONLY")
        assert filled.status_code == 200, filled.text[:400]
        body = filled.json()
        exported = client.post(
            "/api/export_invoice",
            json={
                "template_hash": body["template_hash"],
                "forwarder_id": "jiufang",
                "shipment_id": body["shipment_id"],
                "warehouse_code": body["warehouse_code"],
                "template_sheet": body["template_sheet"],
                "fields": body["cells"],
            },
        )
        assert exported.status_code == 200, exported.text[:400]
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=True)
        try:
            sheet = book["Invoice"]
            assert sheet["B12"].value == "ONL1"
            assert sheet["B16"].value == "ONLYVILLE"
        finally:
            book.close()

        ambiguous = generate(shared, "FBAN5CSHARE")
        assert ambiguous.status_code == 422, ambiguous.text[:400]
        detail = ambiguous.json()["detail"]
        assert detail["code"] == "warehouse_undecided"
        assert detail["address"] == "SHAREVILLE, ZZ"
        assert detail["warehouse_candidates"] == ["GYR2", "LGB8"]
        assert "对应多个仓库" in detail["message"]
        assert "GYR2" in detail["message"] and "LGB8" in detail["message"]
        assert "选择仓库" in detail["message"]

        chosen = generate(shared, "FBAN5CSHARE", warehouse="LGB8")
        assert chosen.status_code == 200, chosen.text[:400]
    finally:
        delete_product_by_sku(sku)


def _grid_xlsx(rows: list[list], *, float_cells: dict[int, str] | None = None) -> bytes:
    """float_cells 把整数格写成 52.0 这种 XML。openpyxl 会把 52.0 存成 52。"""
    import zipfile

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(list(row))
    buf = io.BytesIO()
    book.save(buf)
    book.close()
    if not float_cells:
        return buf.getvalue()
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dest:
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                text = blob.decode("utf-8")
                for number, literal in float_cells.items():
                    text = text.replace(f"<v>{number}</v>", f"<v>{literal}</v>")
                blob = text.encode("utf-8")
            dest.writestr(info, blob)
    src.close()
    return out.getvalue()


def test_float_qty_compares_sku_totals_not_per_box_strings():
    """原发票把 SKU 合计存成 52.0 / 32.0。不能再和单箱 "4" 或整票 "84" 做字符串交。"""
    import pandas as pd

    rows = [{"SKU": "SKU-A", "FNSKU": "X00TESTA", "数量": 4.0} for _ in range(13)]
    rows += [{"SKU": "SKU-B", "FNSKU": "X00TESTB", "数量": 4.0} for _ in range(8)]
    expected, alias = _csv_sku_totals(pd.DataFrame(rows))
    assert expected == {"SKU-A": 52, "SKU-B": 32}
    assert sum(expected.values()) == 84

    invoice = _grid_xlsx(
        [
            ["箱号 Box No", "SKU", "SKU个数\nTotal Qty"],
            ["1-13", "X00TESTA", 52],
            ["14-21", "X00TESTB", 32],
        ],
        float_cells={52: "52.0", 32: "32.0"},
    )
    per_box = _grid_xlsx(
        [
            ["箱序号*", "FBAID*", "单箱内数量*"],
            ["1,2,3,4,5,6,7,8,9,10,11,12,13", "FBAFLOAT", 4],
            ["14,15,16,17,18,19,20,21", "FBAFLOAT", 4],
        ],
        float_cells={4: "4.0"},
    )
    per_box_by_sku = _grid_xlsx(
        [
            ["箱号段(起始-结束)", "SKU", "每箱数量"],
            ["1-13", "SKU-A", 4],
            ["14-21", "SKU-B", 4],
        ],
        float_cells={4: "4.0"},
    )
    invoice_texts = _workbook_texts(invoice)
    assert "52.0" in invoice_texts and "32.0" in invoice_texts
    assert not ({"4"} & invoice_texts)
    assert "84" not in invoice_texts
    for blob in (invoice, per_box, per_box_by_sku):
        texts = _workbook_texts(blob)
        assert not ({"4"} & texts)
        assert "84" not in texts
        ok, detail = _qty_compare(blob, expected, alias)
        assert ok, detail

    text_qty = _grid_xlsx([
        ["SKU", "总数量"],
        ["SKU-A", "52.0"],
        ["SKU-B", "32.0"],
    ])
    ok, detail = _qty_compare(text_qty, expected, alias)
    assert ok, detail


def _float_forwarder_csv() -> bytes:
    """两行 SKU，单箱 4，合计 8 和 12。合计不等于单箱，也不等于整票 20。"""
    headers = ["SKU", "商品名称", "ASIN", "FNSKU", "状况", "预处理类型", "商品总数"]
    headers += [f"箱子 {index} 的商品数量" for index in range(1, 6)]
    left = ["FLOAT-A", "rod", "B0FLOATA01", "X00TESTFLA", "新品", "无", "8", "4", "4", "", "", ""]
    right = ["FLOAT-B", "rod", "B0FLOATB01", "X00TESTFLB", "新品", "无", "12", "", "", "4", "4", "4"]
    marks = ["", "", "", "", "", "", "箱号"] + [f"FBAFLOAT1U{index:06d}" for index in range(1, 6)]
    weight = ["", "", "", "", "", "", "包装箱重量（磅）："] + ["10"] * 5
    length = ["", "", "", "", "", "", "包装箱长度（英寸）："] + ["10"] * 5
    width = ["", "", "", "", "", "", "包装箱宽度（英寸）："] + ["10"] * 5
    height = ["", "", "", "", "", "", "包装箱高度（英寸）："] + ["10"] * 5

    def quoted(cells: list[object]) -> str:
        return ",".join(f'"{cell}"' for cell in cells)

    return "\n".join([
        '"货件编号","FBAFLOAT1"',
        '"货件名称","PLAIN"',
        '"配送地址","ONLYVILLE, ZZ"',
        '"箱子数量","5"',
        "",
        '"多件商品（5 箱）"',
        quoted(headers),
        quoted(left),
        quoted(right),
        "",
        quoted(marks),
        quoted(weight),
        quoted(length),
        quoted(width),
        quoted(height),
    ]).encode("utf-8")


def test_each_forwarder_export_matches_float_invoice_sku_totals(monkeypatch, tmp_path):
    """六个货代的导出和一份浮点数量发票，都按 SKU 合计对齐，不看单元格里的 "4"。"""
    import base64

    from fastapi.testclient import TestClient

    from fba_tool.master_db import delete_product_by_sku, init_master_db, upsert_product
    from fba_tool.product_fields import SOURCE_MANUAL, set_product_fields
    from fba_tool.shipment_parser import parse_csv_shipment
    from fba_tool.warehouse_db import apply_warehouse_import
    from main import create_app

    monkeypatch.setenv("FBA_APP_DATA_DIR", str(tmp_path / "app"))
    monkeypatch.setenv("FBA_DISABLE_SCHEDULER", "1")
    init_master_db()
    raw = _float_forwarder_csv()
    ship = parse_csv_shipment(raw, file_name="float.csv", allow_unassigned_fba=True)
    expected, alias = _csv_sku_totals(ship.box_items)
    assert expected == {"FLOAT-A": 8, "FLOAT-B": 12}
    invoice = _grid_xlsx(
        [
            ["箱号 Box No", "SKU", "SKU个数\nTotal Qty"],
            ["1-2", "X00TESTFLA", 8],
            ["3-5", "X00TESTFLB", 12],
        ],
        float_cells={8: "8.0", 12: "12.0"},
    )
    invoice_texts = _workbook_texts(invoice)
    assert "8.0" in invoice_texts and "12.0" in invoice_texts
    assert not ({"4"} & invoice_texts)
    assert "20" not in invoice_texts
    ok, detail = _qty_compare(invoice, expected, alias)
    assert ok, detail

    library = _warehouse_xlsx([
        ("ONL1", "ONL1", "Amazon", "00001", "1 Only Road", "", "", "US", "ONLYVILLE", "ZZ", ""),
    ])
    assert apply_warehouse_import(library, "synthetic-warehouse.xlsx")["written"] == 1
    import json

    for sku, fnsku, asin in (
        ("FLOAT-A", "X00TESTFLA", "B0FLOATA01"),
        ("FLOAT-B", "X00TESTFLB", "B0FLOATB01"),
    ):
        upsert_product(
            {
                "sku": sku,
                "fnsku": fnsku,
                "asin": asin,
                "english_name": "Fishing rod",
                "chinese_name": "鱼竿",
                "brand": "Acme",
                "material": "棉",
                "purpose": "宠物",
                "hs_code": "9507100090",
                "model_name": "型号甲",
                "extra_data": json.dumps(
                    {"invoice": {"unit": "个", "purchase_price": 5, "material_en": "cotton", "purpose_en": "pet", "pcs_per_set": 1}},
                    ensure_ascii=False,
                ),
            }
        )
        set_product_fields(sku, {"listing_price": 10, "listing_price_currency": "USD"}, SOURCE_MANUAL)
    client = TestClient(create_app())
    try:
        for forwarder_id in FORWARDERS:
            generated = client.post(
                "/api/invoice/generate",
                json={
                    "shipment_id": ship.shipment_id,
                    "forwarder_id": forwarder_id,
                    "csv_base64": base64.b64encode(raw).decode(),
                },
            )
            assert generated.status_code == 200, f"{forwarder_id} {generated.text[:400]}"
            body = generated.json()
            exported = client.post(
                "/api/export_invoice",
                json={
                    "template_hash": body["template_hash"],
                    "forwarder_id": forwarder_id,
                    "shipment_id": body["shipment_id"],
                    "warehouse_code": body["warehouse_code"],
                    "template_sheet": body["template_sheet"],
                    "fields": body["cells"],
                },
            )
            assert exported.status_code == 200, f"{forwarder_id} export {exported.text[:300]}"
            export_ok, export_detail = _qty_compare(exported.content, expected, alias)
            assert export_ok, f"{forwarder_id} {export_detail}"
    finally:
        delete_product_by_sku("FLOAT-A")
        delete_product_by_sku("FLOAT-B")
