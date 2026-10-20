"""Test Step 3 exports against user templates."""
from pathlib import Path

import openpyxl

from fba_tool.boxing import dataframe_to_sku_rows, generate_box_plan, plans_to_editor_df
from fba_tool.data_processing import load_from_excel
from fba_tool.excel_export import export_amazon_template, export_internal_archive
from fba_tool.sta_template_export import pick_sta_packing_sheet
from fba_tool.input_template import ensure_template_on_disk

inp = ensure_template_on_disk()
df, fnsku = load_from_excel(inp.read_bytes())
sku_rows = dataframe_to_sku_rows(df)
plans, _ = generate_box_plan(sku_rows)
editor_df = plans_to_editor_df(plans)

internal = export_internal_archive(editor_df, sku_rows, fnsku_overrides=fnsku)
amazon = export_amazon_template(editor_df, sku_rows, fnsku_overrides=fnsku)

out = Path("_export_test")
out.mkdir(exist_ok=True)
(out / "test_internal.xlsx").write_bytes(internal)
(out / "test_amazon.xlsx").write_bytes(amazon)

wb1 = openpyxl.load_workbook(out / "test_internal.xlsx", data_only=False)
ws1 = wb1.active
print("INTERNAL sheet", ws1.title)
print("R2", [ws1.cell(2, c).value for c in range(1, 10)])
print("R9 C3", ws1.cell(9, 3).value, "G", ws1.cell(9, 7).value)
print("R12 G", ws1.cell(12, 7).value)

wb2 = openpyxl.load_workbook(out / "test_amazon.xlsx", data_only=False)
pack = pick_sta_packing_sheet(wb2)
print("\nAMAZON sheets", wb2.sheetnames)
print("K3", pack.cell(3, 11).value, "L3", pack.cell(3, 12).value)
print("R6 SKU", pack.cell(6, 1).value, "ASIN", pack.cell(6, 3).value, "box1", pack.cell(6, 12).value)
print("K5 blank", pack.cell(5, 11).value, "L5", pack.cell(5, 12).value)
for r in range(6, pack.max_row + 1):
    if pack.cell(r, 1).value == "包装箱名称":
        print("footer name row", r, "P1-B1", pack.cell(r, 12).value)
        print("weight row", r + 1, pack.cell(r + 1, 1).value)
        break
print("OK")
