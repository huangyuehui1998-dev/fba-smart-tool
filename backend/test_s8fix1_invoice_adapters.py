"""S8-fix1 N2：德速 T3 / 琪珞 T6 / 琪珞 T5 导出内容测试——以货代原件为期望值（数量、箱号段、重量口径、尺寸、汇总）。
重量口径不再写死单箱重：由货代原件和货件 CSV 的重量核对判定（与琪珞模板解析同一做法），或由用户指定。"""
from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pytest

from fba_tool.invoice_export import decide_weight_basis, export_invoice
from fba_tool.shipment_parser import parse_csv_shipment

from tests._realdata import skip_module_if_missing

HERE = Path(__file__).resolve().parent
S = HERE / "tests" / "golden" / "samples"
skip_module_if_missing(HERE / "tests/golden/product_master_golden.json")
SNAP = {p["fnsku"]: p for p in json.loads((HERE / "tests/golden/product_master_golden.json").read_text(encoding="utf-8"))["products"]}


def _prod(f):
    if f in SNAP:
        p = dict(SNAP[f])
        p.setdefault("hs_code", "6307909000")
        p.setdefault("material", "布")
        p.setdefault("purpose", "宠物用")
        return p
    return {"fnsku": f, "listing_price": 10.0, "english_name": "Item", "chinese_name": "物品", "brand": "B", "material": "铁",
            "purpose": "用", "hs_code": "7326909000", "asin": "B0X", "model_name": "无",
            "invoice": {"unit": "个", "purchase_price": 5, "material_en": "iron", "purpose_en": "use", "pcs_per_set": 1}}


def _ship(csv_name):
    raw = (S / "inputs" / csv_name).read_bytes()
    s = parse_csv_shipment(raw, file_name=csv_name, allow_unassigned_fba=True)
    return raw, s, {str(f).upper(): _prod(str(f).upper()) for f in s.box_items["FNSKU"]}


def _num(v):
    try:
        return round(float(v), 4)
    except (TypeError, ValueError):
        return v


def _rows(ws, first, cols):
    out = []
    r = first
    while ws[f"{cols[0]}{r}"].value not in (None, ""):
        out.append(tuple(_num(ws[f"{c}{r}"].value) for c in cols))
        r += 1
    return out


# ---------------------------------------------------------------- 德速 T3 第9批（单件重 1.7025…）
T3_9 = S / "invoice_templates/T3_第9批第1票_FBA19GQXZW0N.xlsx"


def test_deshu_t3_batch9_basis_decided_per_unit_from_csv():
    raw, s, _ = _ship("第9批-FBA19GQXZW0N.csv")
    basis, why = decide_weight_basis(T3_9.read_bytes(), T3_9.name, raw, s.shipment_id)
    assert basis == "per_unit" and "单件重" in why


def test_deshu_t3_batch9_content_matches_original():
    raw, s, prods = _ship("第9批-FBA19GQXZW0N.csv")
    basis, why = decide_weight_basis(T3_9.read_bytes(), T3_9.name, raw, s.shipment_id)
    res = export_invoice("deshu", shipment=s, products=prods, reference_id="R", weight_basis=basis, weight_basis_source=why)
    out = openpyxl.load_workbook(io.BytesIO(res["bytes"]))["发票填写模板"]
    org = openpyxl.load_workbook(T3_9, data_only=True)["发票填写模板"]
    cols = "CNTUVW"  # 箱号段、每箱件数、重量、长宽高
    got, want = _rows(out, 21, cols), _rows(org, 21, cols)
    assert got == want and len(got) == 6
    assert got[0][2] == 1.7025
    assert out["B2"].value == org["B2"].value == s.shipment_id
    assert any(w["code"] == "weight_basis" and w["basis"] == "per_unit" for w in res["warnings"])


def test_deshu_t3_per_box_when_chosen_and_blank_when_undecided():
    _, s, prods = _ship("第9批-FBA19GQXZW0N.csv")
    ws = openpyxl.load_workbook(io.BytesIO(export_invoice("deshu", shipment=s, products=prods, weight_basis="per_box")["bytes"]))["发票填写模板"]
    assert _num(ws["T21"].value) == 6.81  # 1.7025 × 4
    res = export_invoice("deshu", shipment=s, products=prods)
    ws = openpyxl.load_workbook(io.BytesIO(res["bytes"]))["发票填写模板"]
    assert ws["T21"].value is None  # 口径没定：不猜，留空
    w = [x for x in res["warnings"] if x["code"] == "weight_basis_undecided"]
    assert w and "6.81" in w[0]["message"] and "1.7025" in w[0]["message"] and "德速第2批单箱重" in w[0]["message"]
    with pytest.raises(ValueError):
        export_invoice("deshu", shipment=s, products=prods, weight_basis="per_kg")


# ---------------------------------------------------------------- 琪珞 T6 第6批（单件重）/ 第7批（单箱重）
T6_6 = {
    "第6批-FBA19GDLZ88C.csv": "J-US-第6批9810第1票美国海运ILG1-6箱-PETBARRIERS发票.xlsx",
    "第6批-FBA19GDN35JP.csv": "J-US-第6批9810第2票美国海运SMF6-6箱-PETBARRIERS发票.xlsx",
    "第6批-FBA19GDPP6M8.csv": "J-US-第6批9810第3票美国海运JAX3-6箱-PETBARRIERS发票.xlsx",
    "第6批-FBA19GDL9K3X.csv": "J-US-第6批9810第4票美国海运ACY2-6箱-PETBARRIERS发票.xlsx",
    "第6批-FBA19GDJR25Y.csv": "J-US-第6批9810第5票美国海运MGE3-6箱-PETBARRIERS发票.xlsx",
}


@pytest.mark.parametrize("csv_name,inv", list(T6_6.items()))
def test_qiluo_t6_batch6_per_unit_matches_original(csv_name, inv):
    ref = S / "invoice_legacy" / inv
    raw, s, prods = _ship(csv_name)
    basis, why = decide_weight_basis(ref.read_bytes(), ref.name, raw, s.shipment_id)
    assert basis == "per_unit", why
    res = export_invoice("qiluo_t6", shipment=s, products=prods, weight_basis=basis, weight_basis_source=why)
    out = openpyxl.load_workbook(io.BytesIO(res["bytes"]))["发票"]
    org = openpyxl.load_workbook(ref, data_only=True)["发票"]
    # 每箱件数、单件重逐格对原件。不比：箱号（原件写整批 1-30，见 D-09 第3点）；长宽高（导出取货件 CSV 60×45×36，
    # 第6批原件写的 35×35×42 和 CSV 不符，是复制旧票没改）
    got, want = _rows(out, 14, "FM"), _rows(org, 14, "FM")
    assert sorted(set(got)) == sorted(set(want)) == [(5.0, 4.0)], (got, want)
    assert out["B3"].value == s.shipment_id and out["B6"].value == 6


def test_qiluo_t6_batch7_per_box_matches_original():
    ref = S / "invoice_templates/T6_第7批第1票_VGT2-29箱.xlsx"
    raw, s, prods = _ship("第7批-FBA19GVFX5SC.csv")
    basis, why = decide_weight_basis(ref.read_bytes(), ref.name, raw, s.shipment_id)
    assert basis == "per_box", why
    res = export_invoice("qiluo_t6", shipment=s, products=prods, weight_basis=basis, weight_basis_source=why)
    out = openpyxl.load_workbook(io.BytesIO(res["bytes"]))["发票"]
    org = openpyxl.load_workbook(ref, data_only=True)["发票"]
    assert _rows(out, 14, "AFMNOP") == _rows(org, 14, "AFMNOP")
    assert _num(out["M14"].value) == 6.81
    assert out["B6"].value == org["B6"].value == 29


def test_qiluo_t6_undecided_blank_with_qiluo_example():
    _, s, prods = _ship("第7批-FBA19GVFX5SC.csv")
    res = export_invoice("qiluo_t6", shipment=s, products=prods)
    assert openpyxl.load_workbook(io.BytesIO(res["bytes"]))["发票"]["M14"].value is None
    w = [x for x in res["warnings"] if x["code"] == "weight_basis_undecided"]
    assert w and "琪珞第6批单件重、第7批单箱重" in w[0]["message"] and "德速" not in w[0]["message"]


# ---------------------------------------------------------------- 琪珞 T5（单行汇总；没有带 CSV 的原件，按 CSV 汇总值核对）
def test_qiluo_t5_summary_from_csv():
    raw, s, prods = _ship("第7批-FBA19GVFX5SC.csv")
    wh = {"zip": "89115", "state": "NV", "city": "Las Vegas", "address1": "6401 Howdy Wells Ave"}
    res = export_invoice("qiluo_t5", shipment=s, products=prods, warehouse="VGT2", warehouse_info=wh, reference_id="REF7")
    ws = openpyxl.load_workbook(io.BytesIO(res["bytes"])).worksheets[0]
    m = res["model"]
    assert (ws["B3"].value, ws["D3"].value, ws["B7"].value, ws["F7"].value, ws["H7"].value) == (s.shipment_id, "REF7", "VGT2", "89115", "NV")
    assert ws["J10"].value == len(m.boxes) == 29
    assert _num(ws["N10"].value) == round(sum(b.weight for b in m.boxes), 2)
    qtys = [int(x) for x in str(ws["H10"].value).split(",")]
    assert qtys == [l["qty"] for l in m.lines] and sum(qtys) == int(s.box_items["数量"].sum())
    dims = {tuple(b.dims) for b in m.boxes}
    if len(dims) == 1:
        assert (ws["K10"].value, ws["L10"].value, ws["M10"].value) == next(iter(dims))
    else:
        assert any(w["code"] == "qiluo_t5_mixed_dims" for w in res["warnings"])
