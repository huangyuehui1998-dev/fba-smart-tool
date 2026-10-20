"""从货代原始样例生成空白导出模板（一次性脚本，结果已提交到 templates/forwarders/）。

- send.xlsx    ← tests/golden/samples/invoice/FBA19P4N6MM2发票.xlsx：清空 C4 和第 15 行起的数据，去掉图片；
- jiufang.xlsx ← tests/golden/samples/invoice/LCL2610S97BA-箱单发票.xlsx：清空票级表头、派送地址、发货人，
  Invoice 第 25/26 行、Packing list 第 6/7 行只留格式（数据行样式 / 合计行样式）。
模板只取版式，不取任何票的数据。
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "tests" / "golden" / "samples" / "invoice"
OUT = ROOT / "templates" / "forwarders"


def _clear(ws, refs):
    for r in refs:
        ws[r].value = None


def _drop_rows(ws, first):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first:
            ws.unmerge_cells(str(rng))
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)


def build_send():
    wb = openpyxl.load_workbook(SAMPLES / "FBA19P4N6MM2发票.xlsx")
    ws = wb.worksheets[0]
    ws._images = []
    ws["C4"].value = None
    for row in ws.iter_rows(min_row=15):
        for c in row:
            c.value = None
    _drop_rows(ws, 16)
    wb.save(OUT / "send.xlsx")


def build_jiufang():
    wb = openpyxl.load_workbook(SAMPLES / "LCL2610S97BA-箱单发票.xlsx")
    inv, pk = wb["Invoice"], wb["Packing list"]
    for ws in (inv, pk):
        ws._images = []
    _clear(inv, ["B2", "P2", "P3", "P4", "P5", "B4", "B5", "B7", "B8", "B12", "B13", "B15", "B16", "B17", "B19", "B22"])
    _clear(pk, ["C2", "C3"])
    # S97BA 只有 1 行数据：Invoice 第 25 行=数据行样式、第 26 行=合计行样式；Packing list 第 6 行=数据行、第 11 行=合计行
    _drop_rows(inv, 27)
    for rng in list(inv.merged_cells.ranges):
        if rng.min_row >= 25:
            inv.unmerge_cells(str(rng))
    pk.delete_rows(7, 4)
    _drop_rows(pk, 8)
    for ws, rows in ((inv, (25, 26)), (pk, (6, 7))):
        for row in ws.iter_rows(min_row=rows[0], max_row=rows[1]):
            for c in row:
                c.value = None
    wb.save(OUT / "jiufang.xlsx")


if __name__ == "__main__":
    build_send()
    build_jiufang()
    print("ok")
