"""发货装箱单（Packing List 样式）测试。"""
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from fba_tool.excel_export import export_shipping_sheet
from fba_tool.master_db import init_master_db
from fba_tool.shipment_parser import CsvShipment, parse_csv_shipment
from fba_tool.shipment_validator import ShipmentPair, validate_shipments
from fba_tool.shipping_sheet import (
    SHIPPING_SHEET_COLUMNS,
    SHIPPING_SHEET_LEGEND_COLUMNS,
    build_shipping_sheet_legend,
    build_shipping_sheet_df,
)


@pytest.fixture(autouse=True)
def isolated_app_data(monkeypatch, tmp_path):
    ad = tmp_path / "app_data"
    ad.mkdir()
    monkeypatch.setattr("fba_tool.master_db._master_db_file", lambda: ad / "fba_master.db")
    monkeypatch.setattr("fba_tool.app_paths.app_data_dir", lambda: ad)
    init_master_db()
    return ad


def _pair_from_csv(path: Path) -> ShipmentPair:
    ship = parse_csv_shipment(path.read_bytes())
    return ShipmentPair(ship.shipment_id, ship.warehouse_code, ship, None)  # type: ignore[arg-type]


def _make_pair() -> ShipmentPair:
    csv = CsvShipment(
        shipment_id="FBA123456789",
        warehouse_code="FWA4",
        box_count=2,
        box_data={1: {"SKU-A": 2}, 2: {"SKU-A": 1, "SKU-B": 1}},
        box_fba_ids={1: "FBA123456789U001", 2: "FBA123456789U002"},
    )
    csv.box_specs = pd.DataFrame(
        [
            {"箱号": 1, "长(cm)": 40.0, "宽(cm)": 30.0, "高(cm)": 30.0, "实重(kg)": 10.0},
            {"箱号": 2, "长(cm)": 50.0, "宽(cm)": 30.0, "高(cm)": 30.0, "实重(kg)": 12.0},
        ]
    )
    csv.box_items = pd.DataFrame(
        [
            {"箱号": 1, "SKU": "SKU-A", "数量": 2, "FNSKU": ""},
            {"箱号": 2, "SKU": "SKU-A", "数量": 1, "FNSKU": ""},
            {"箱号": 2, "SKU": "SKU-B", "数量": 1, "FNSKU": ""},
        ]
    )
    return ShipmentPair("FBA123456789", "FWA4", csv, None)  # type: ignore[arg-type]


def test_build_columns_and_fnsku_merge():
    df = build_shipping_sheet_df([_make_pair()])
    assert list(df.columns) == SHIPPING_SHEET_COLUMNS
    assert df.iloc[0]["票数"] == "第一票"
    assert df.iloc[0]["仓库代码"] == "FWA4"
    assert "英文品名" not in df.columns
    # SKU-A 在箱 1、2 -> 合并为一行；SKU-B 仅箱 2 -> 单独一行
    assert len(df) == 2
    sku_a = df[df["中文品名"].astype(str).str.contains("SKU-A", na=False)]
    assert len(sku_a) == 1
    assert sku_a.iloc[0]["箱号"] == "1,2"
    assert int(sku_a.iloc[0]["总数量"]) == 3
    sku_b = df[df["中文品名"].astype(str).str.contains("SKU-B", na=False)]
    assert len(sku_b) == 1
    assert sku_b.iloc[0]["箱号"] == "2"


def test_ticket_from_pdf_filename():
    pairs = [
        ShipmentPair(
            "FBA19GQW64C7",
            "SCK4",
            _make_pair().csv,
            None,
            "第二票外箱贴US-7TH-9810-PETBEDMAT-FBA19GQW64C7-SCK4.pdf",
        ),  # type: ignore[arg-type]
        ShipmentPair(
            "FBA1",
            "FWA4",
            _make_pair().csv,
            None,
            "第一票外箱贴Y-FBA1-FWA4.pdf",
        ),  # type: ignore[arg-type]
    ]
    df = build_shipping_sheet_df(pairs)
    tickets = df["票数"].unique().tolist()
    assert "第二票" in tickets
    assert "第一票" in tickets


def test_export_total_row():
    df = build_shipping_sheet_df([_make_pair()])
    data = export_shipping_sheet(df)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert "发货装箱单" in str(ws.cell(1, 1).value)
    total_row = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value or "").startswith("合计"):
                total_row = r
                break
        if total_row:
            break
    assert total_row is not None
    vals = [ws.cell(total_row, c).value for c in range(1, ws.max_column + 1)]
    assert 4 in vals
    # 毛重：箱1(10) + 箱2(12) 各计一次；SKU-A 行 22，SKU-B 行 12，合计列求和
    assert 34 in vals or 34.0 in vals


def test_export_bytes_from_sample():
    csv_path = Path("templates/亚马逊FBA表1.csv")
    if not csv_path.exists():
        return
    result = validate_shipments([(csv_path.name, csv_path.read_bytes())], [])
    pairs = result.pairs or [_pair_from_csv(csv_path)]
    df = build_shipping_sheet_df(pairs)
    data = export_shipping_sheet(df)
    assert len(data) > 500


def test_build_legend_from_step1():
    pair = _make_pair()
    step1 = [
        {
            "SKU简称": "SKU-A",
            "FNSKU": "",
            "发货总数": 90,
            "单箱装量": 6,
            "箱长(cm)": 51.0,
            "箱宽(cm)": 29.5,
            "箱高(cm)": 38.0,
        },
        {
            "SKU简称": "SKU-B",
            "FNSKU": "",
            "发货总数": 40,
            "单箱装量": 4,
            "箱长(cm)": 50.0,
            "箱宽(cm)": 30.0,
            "箱高(cm)": 30.0,
        },
        {
            "SKU简称": "SKU-OTHER",
            "FNSKU": "",
            "发货总数": 10,
            "单箱装量": 1,
            "箱长(cm)": 10.0,
            "箱宽(cm)": 10.0,
            "箱高(cm)": 10.0,
        },
    ]
    legend = build_shipping_sheet_legend(step1, [pair])
    assert len(legend) == 2
    assert legend[0]["SKU简称"] == "SKU-A"
    assert int(legend[0]["发货总数"]) == 90
    assert list(legend[0].keys()) == SHIPPING_SHEET_LEGEND_COLUMNS


def test_export_with_legend():
    pair = _make_pair()
    df = build_shipping_sheet_df([pair])
    legend = build_shipping_sheet_legend(
        [
            {
                "SKU简称": "SKU-A",
                "发货总数": 90,
                "单箱装量": 6,
                "箱长(cm)": 51,
                "箱宽(cm)": 29.5,
                "箱高(cm)": 38,
            }
        ],
        [pair],
    )
    data = export_shipping_sheet(df, legend)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.cell(2, 1).value == "SKU简称*"
    assert ws.cell(2, 4).value == "发货总数*"
    assert ws.cell(3, 1).value == "SKU-A"
    assert ws.cell(3, 4).value == 90


def test_shipping_sheet_uses_csv_fnsku():
    pair = _make_pair()
    pair.csv.box_items = pd.DataFrame(
        [
            {"箱号": 1, "SKU": "SKU-A", "数量": 2, "FNSKU": "X0054S6RG5"},
            {"箱号": 2, "SKU": "SKU-A", "数量": 1, "FNSKU": "X0054S6RG5"},
            {"箱号": 2, "SKU": "SKU-B", "数量": 1, "FNSKU": "X0054S6R65"},
        ]
    )
    pair.csv.sku_fnsku = {"SKU-A": "X0054S6RG5", "SKU-B": "X0054S6R65"}
    df = build_shipping_sheet_df([pair])
    codes = set(df["亚马逊条码(FNSKU)"].astype(str))
    assert "X0054S6RG5" in codes
    assert "X0054S6R65" in codes


def test_parse_sample_csv_keeps_fnsku():
    csv_path = Path("templates/亚马逊FBA表.csv")
    if not csv_path.exists():
        return
    from fba_tool.shipment_parser import parse_csv_shipment

    ship = parse_csv_shipment(csv_path.read_bytes())
    assert ship.sku_fnsku.get("J-PETWOODSTAIRS-HUI-4PCS-FBA") == "X0054S6RG5"
    df = build_shipping_sheet_df(
        [
            ShipmentPair(ship.shipment_id, ship.warehouse_code, ship, None)  # type: ignore[arg-type]
        ]
    )
    assert "X0054S6RG5" in set(df["亚马逊条码(FNSKU)"].astype(str))
    assert "X0054S6R65" in set(df["亚马逊条码(FNSKU)"].astype(str))


def test_step1_fnsku_fills_when_csv_blank():
    pair = _make_pair()
    df = build_shipping_sheet_df(
        [pair],
        step1_rows=[
            {"SKU简称": "SKU-A", "FNSKU": "X0STEP1AAA"},
            {"SKU简称": "SKU-B", "FNSKU": "X0STEP1BBB"},
        ],
    )
    codes = set(df["亚马逊条码(FNSKU)"].astype(str))
    assert "X0STEP1AAA" in codes
    assert "X0STEP1BBB" in codes


if __name__ == "__main__":
    test_build_columns_and_fnsku_merge()
    test_ticket_from_pdf_filename()
    test_build_legend_from_step1()
    test_export_with_legend()
    test_export_total_row()
    test_export_bytes_from_sample()
    test_shipping_sheet_uses_csv_fnsku()
    test_parse_sample_csv_keeps_fnsku()
    test_step1_fnsku_fills_when_csv_blank()
    print("ALL OK")
